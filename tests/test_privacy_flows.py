"""用真实服务和隐私出口验证模板引文、源码证据及证书附件的完整链路"""

import json

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from test_honors import certificate_bytes, wait_honor
from test_jobs import wait_job
from test_privacy import provider_at, reply
from test_template_analysis import completed, simple_document

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.infrastructure.database import uid
from resume_maker.integrations.privacy import TOKEN
from resume_maker.services.jobs import Jobs
from resume_maker.services.templates.tasks import Templates


def test_masked_source_paths_and_quotes_are_restored_before_evidence_validation(tmp_path, catalog):
    """模型只引用脱敏路径和正文，最终源码证据仍核对原始文件"""
    root = tmp_path / "original-private-source"
    root.mkdir()
    for index in range(205):
        (root / f"part-{index:03}.py").write_text("# ordinary source\n" * 120, encoding="utf-8")
    original = "# source\n" * 16000 + "print('测试甲', 'late@example.invalid')\n"
    (root / "测试甲.py").write_text(original, encoding="utf-8")
    project = catalog.create_project("例子", [str(root)])
    catalog.db.set_setting("privacy_terms", ["测试甲"])
    seen = []

    def handle(request, access):
        """依据收到的安全片段构造真实行号的源码证据"""
        body = request
        seen.append(body)
        context = json.loads(body["input"].split("本轮上下文数据：\n")[1])
        rows, args = [], {}
        while True:
            page = access.call("list_source_files", args)
            seen.append(page)
            rows.extend(page["results"])
            if page["complete"]:
                break
            args["cursor"] = page["next_cursor"]
        assert len(rows) == 206
        row = next(row for row in rows if "[[RM_" in row["path"])
        file = access.call("read_source", {**row, "start_line": 16001})
        seen.append(file)
        evidence = {
            "source": file["source"],
            "path": file["path"],
            "line_start": 16001,
            "line_end": 16001,
            "quote": file["lines"][0]["text"],
            "status": "code",
        }
        result = {
            **reply(),
            "experience": {
                **context["current_experience"],
                "highlights": [
                    {"id": "one", "title": "处理", "text": "完成处理", "evidence": [evidence]}
                ],
            },
        }
        return result

    provider = provider_at(tmp_path, None, catalog.db, source_handler=handle)
    jobs = Jobs(catalog.db, catalog, tmp_path / "data", provider)
    conversation = catalog.db.one(
        "SELECT id FROM conversations WHERE project_id=?", (project["id"],)
    )
    jobs.start()
    try:
        job = jobs.submit(
            conversation["id"], "分析", "analysis", project["head_revision"], "all", uid()
        )
        result = wait_job(catalog, job["id"])
        assert result["status"] == "completed", result["error"]
    finally:
        jobs.stop()
    serialized = json.dumps(seen, ensure_ascii=False)
    assert "测试甲" not in serialized and "original-private-source" not in serialized
    assert "late@example.invalid" not in serialized
    proposal = catalog.db.one("SELECT after_json FROM proposals WHERE job_id=?", (job["id"],))
    evidence = proposal["after"]["highlights"][0]["evidence"][0]
    assert evidence["status"] == "code" and evidence["path"] == "测试甲.py"
    assert evidence["quote"] == "print('测试甲', 'late@example.invalid')"
    audit = catalog.db.setting("privacy_audit")[0]
    assert audit["status"] == "completed" and audit["payload"]["source_tool_calls"] == 4
    assert "测试甲" not in json.dumps(audit, ensure_ascii=False)
    assert (root / "测试甲.py").read_text(encoding="utf-8") == original


def test_template_quote_restored_and_original_images_never_sent(tmp_path, catalog):
    """脱敏引文仍可准确定位原模板，图片只在本机保留"""
    source = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("姓名：测试甲")
    document.save(source)
    seen = []

    def handle(request):
        """只根据安全节点清单返回姓名映射，不读取任何本机文件"""
        body = request
        seen.append(body)
        context = json.loads(body["input"].splitlines()[-1])
        rows = [row for part in context["template"]["parts"].values() for row in part]
        row = next(row for row in rows if row[1] == "p")
        plan = {
            "summary": "姓名映射",
            "fields": [{"node": row[0], "quote": row[4], "target": "personal.name"}],
            "repeats": [],
            "photos": [],
            "keep": [],
            "remove": [],
            "warnings": [],
        }
        return plan

    provider = provider_at(tmp_path, handle, catalog.db)
    service = Templates(catalog, tmp_path / "data", provider)
    task = service.analyze(source, simple_document())
    result = completed(service, task["id"])
    assert result["status"] == "completed", result["error"]
    assert result["plan"]["fields"][0]["quote"] == "姓名：测试甲"
    serialized = json.dumps(seen, ensure_ascii=False)
    assert "测试甲" not in serialized and "新的用户资料" not in serialized
    assert "input_image" not in serialized and "data:image" not in serialized
    assert Document(source).paragraphs[0].text == "姓名：测试甲"


def test_blank_certificate_stays_local_and_pdf_text_is_masked(tmp_path):
    """荣誉上传不把原件发给模型，PDF 文字中的获奖人本地还原"""
    config = Config(data_dir=tmp_path / "data", token="test")
    app = create_app(config)
    seen = []

    def handle(request):
        """供应商只接收 PDF 的脱敏文字"""
        body = request
        seen.append(body)
        context = json.loads(body["input"].split("\n本地 OCR")[0].splitlines()[-1])
        token = TOKEN.search(context["pdf_text"])[0]
        return {
            "fields": {"name": "Synthetic Award", "recipient": token},
            "text": token,
            "warnings": [],
        }

    provider = provider_at(tmp_path, handle, app.state.services.db)
    app.state.services.honors.provider = provider
    with TestClient(app, headers={"x-resume-token": "test"}) as client:
        image = client.post(
            "/api/honors/upload?filename=certificate.png", content=certificate_bytes()
        ).json()
        blocked = wait_honor(client, image["id"], "failed")
        assert "本地 OCR" in blocked["error"] and not seen
        assert client.get(f"/api/honors/{image['id']}/original").status_code == 200
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((30, 30), "Recipient: Synthetic Person\nSynthetic Award")
            raw = pdf.tobytes()
        uploaded = client.post("/api/honors/upload?filename=certificate.pdf", content=raw).json()
        result = wait_honor(client, uploaded["id"])
        assert result["fields"]["recipient"] == "Synthetic Person"
        assert len(seen) == 1 and "Synthetic Person" not in json.dumps(seen)
        assert "input_image" not in json.dumps(seen)
