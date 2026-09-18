"""模板浏览器组织信息持久化、原模板隔离和缩略图缓存契约。"""

from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient
from test_resume_previews import register_template

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.core.errors import Problem
from resume_maker.services.template_library import TemplateLibrary


def test_library_persists_and_category_deletion_keeps_likes(tmp_path):
    """收藏和分类跨实例保留，删除分类不删除模板或取消收藏。"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library").status_code == 401
        assert client.get("/api/template-library", headers=headers).json() == {
            "categories": [],
            "items": {},
            "templates": [],
        }
        register_template(client.app.state.services.catalog, config.data_dir)
        result = client.post(
            "/api/template-library/categories", headers=headers, json={"name": "技术"}
        )
        category = result.json()["categories"][0]["id"]
        client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"category_id": category}
        )
        state = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"liked": True}
        ).json()
        assert state["items"]["mapped"] == {"category_id": category, "liked": True}
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library", headers=headers).json() == state
        state = client.delete(
            f"/api/template-library/categories/{category}", headers=headers
        ).json()
        assert state["items"]["mapped"] == {"category_id": "", "liked": True}
        assert client.app.state.services.catalog.template("mapped")["name"] == "测试模板"
        assert not client.app.state.services.db.all("SELECT * FROM resumes")


def test_library_validation_and_independent_updates(catalog, tmp_path):
    """新分类校验空白与重名，并发修改分类和收藏互不覆盖。"""
    service = TemplateLibrary(catalog, tmp_path / "data")
    category = service.create_category(" 技术 ")["categories"][0]["id"]
    for name in (" ", "技术", "未分类", "全部模板", "我的喜欢", "回收站"):
        with pytest.raises(Problem):
            service.create_category(name)
    with pytest.raises(Problem):
        service.update("missing", {"liked": True})
    with pytest.raises(Problem):
        service.update("builtin", {"category_id": "missing"})
    with ThreadPoolExecutor(max_workers=2) as pool:
        operations = [
            pool.submit(service.update, "builtin", changes)
            for changes in ({"category_id": category}, {"liked": True})
        ]
        for operation in operations:
            operation.result()
    assert service.state()["items"]["builtin"] == {"category_id": category, "liked": True}


def test_template_rename_persists_without_changing_references(tmp_path):
    """已使用模板可改名，名称跨重启保留，模板文件、映射与所有简历数据不变。"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        services = client.app.state.services
        source = register_template(services.catalog, config.data_dir)
        original = source.read_bytes()
        record = services.catalog.template("mapped")
        category = services.template_library.create_category("技术")["categories"][0]["id"]
        services.template_library.update("mapped", {"category_id": category, "liked": True})
        for name in ("第一份简历", "第二份简历"):
            assert (
                client.post(
                    "/api/resumes",
                    headers=headers,
                    json={"name": name, "template_id": "mapped", "items": []},
                ).status_code
                == 200
            )
        resumes = services.db.all("SELECT * FROM resumes")
        assert (
            client.patch(
                "/api/template-library/items/mapped", json={"name": "未授权修改"}
            ).status_code
            == 401
        )
        response = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"name": "  新名称  "}
        )
        assert response.status_code == 200
        state = response.json()
        assert state["templates"][0]["name"] == "新名称"
        assert state["templates"][0]["usage_count"] == 2
        assert state["items"]["mapped"] == {"category_id": category, "liked": True}
        assert services.catalog.template("mapped") == {**record, "name": "新名称"}
        assert source.read_bytes() == original
        assert services.db.all("SELECT * FROM resumes") == resumes
        assert (
            client.delete("/api/template-library/items/mapped", headers=headers).status_code == 409
        )
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library", headers=headers).json() == state
        assert client.get("/api/state", headers=headers).json()["templates"][0]["name"] == "新名称"


def test_template_rename_validation_is_atomic(tmp_path):
    """非法名称、内置模板及回收站拒绝改名；联合修改失败时不部分保存。"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        services = client.app.state.services
        register_template(services.catalog, config.data_dir)
        before = services.template_library.state()
        for identifier, changes, status in (
            ("mapped", {"name": ""}, 422),
            ("mapped", {"name": None}, 422),
            ("mapped", {"name": " "}, 400),
            ("mapped", {"name": "字" * 201}, 422),
            ("builtin", {"name": "新内置名称"}, 409),
            ("missing", {"name": "不存在"}, 404),
            ("mapped", {"name": "不应保存", "category_id": "missing", "liked": True}, 409),
        ):
            result = client.patch(
                f"/api/template-library/items/{identifier}", headers=headers, json=changes
            )
            assert result.status_code == status, result.text
            assert services.template_library.state() == before
        recycled = services.template_library.delete("mapped")
        result = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"name": "不应保存"}
        )
        assert result.status_code == 409
        assert services.template_library.state() == recycled


def test_thumbnail_cache_and_original_are_isolated(catalog, tmp_path, monkeypatch):
    """重复及并发读取只渲染一次，原始文件和数据库模板映射不变。"""
    data_dir = tmp_path / "data"
    source = register_template(catalog, data_dir)
    original = source.read_bytes()
    record = catalog.template("mapped")
    calls = []

    def render(path, pdf):
        """只模拟排版输出，同时检查传入的是独立源模板快照。"""
        assert path != source and path.read_bytes() == original
        calls.append(path)
        (pdf.parent / "page-1.png").write_bytes(b"thumbnail")
        return 1, None

    monkeypatch.setattr("resume_maker.services.template_library.render_word", render)
    service = TemplateLibrary(catalog, data_dir)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(service.thumbnail, ["mapped", "mapped"]))
    assert results[0] == results[1]
    assert results[0].read_bytes() == b"thumbnail"
    assert len(calls) == 1
    assert TemplateLibrary(catalog, data_dir).thumbnail("mapped") == results[0]
    assert source.read_bytes() == original and catalog.template("mapped") == record
    source.write_bytes(b"changed outside app")
    with pytest.raises(Problem, match="程序外变化"):
        service.thumbnail("mapped")


def test_thumbnail_failure_can_retry_and_builtin_is_real_docx(catalog, tmp_path, monkeypatch):
    """渲染失败允许重试；内置模板也用实际 DOCX 排版器生成示例预览。"""
    from docx import Document

    calls = []

    def render(path, pdf):
        """首次模拟失败，重试验证真实示例文档后提供图片。"""
        calls.append(path)
        assert "你的姓名" in "".join(
            p.text
            for table in Document(path).tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        )
        if len(calls) == 1:
            return None, "Word 暂不可用"
        (pdf.parent / "page-1.png").write_bytes(b"png")
        return 1, None

    monkeypatch.setattr("resume_maker.services.template_library.render_word", render)
    service = TemplateLibrary(catalog, tmp_path / "data")
    with pytest.raises(Problem, match="Word 暂不可用"):
        service.thumbnail("builtin")
    assert service.thumbnail("builtin").is_file()
    assert len(calls) == 2
