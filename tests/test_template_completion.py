"""不同模板容器和任意栏目名共用补位流程，检查、保存及输出不会相互矛盾"""

import base64
import threading
from copy import deepcopy

import pytest
from docx import Document
from fastapi.testclient import TestClient
from test_template_analysis import TemplateProvider, completed
from test_template_mapping import photo_bytes

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.core.errors import Problem
from resume_maker.domain.models import ProjectVisibility, ProviderSettings
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.infrastructure.database import dump, now
from resume_maker.integrations.sources import digest
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.completion import complete_template
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text
from resume_maker.integrations.word.templates.values import missing_targets
from resume_maker.services.templates.analysis import analyze_plan, assess_plan
from resume_maker.services.templates.cache import cache_path, cached_plan, remember_plan


def generic_template(path, layout="body", padding=0):
    """构造正文、表格行或双栏模板并按原文定位节点"""
    document = Document()
    for _ in range(padding):
        document.add_paragraph()
    document.add_paragraph("Seed Person")
    labels = ["Portfolio Ω", "Recognition · β"]
    samples = [["Old project", "Old period", "Old project body"], ["Old entry body"]]
    table = (
        document.add_table(rows=0 if layout == "rows" else 1, cols=2) if layout != "body" else None
    )
    for index, (label, content) in enumerate(zip(labels, samples, strict=True)):
        if layout == "body":
            document.add_paragraph(label)
            for literal in content:
                document.add_paragraph(literal)
        elif layout == "rows":
            table.add_row().cells[0].text = label
            cell = table.add_row().cells[0]
            cell.text = content[0]
            for literal in content[1:]:
                cell.add_paragraph(literal)
        else:
            cell = table.cell(0, index)
            cell.text = label
            for literal in content:
                cell.add_paragraph(literal)
    document.add_paragraph("Fixed ending")
    document.save(path)
    package = TemplatePackage(path)
    ids = {paragraph_text(node): key for key, node in package.nodes.items() if node.tag == w("p")}

    def binding(literal, target):
        """以清单中的完整原文建立字段绑定"""
        return TextBinding(node=ids[literal], quote=literal, target=target)

    def block(literal):
        """表格行重复整行，其他版式只重复所属容器中的段落"""
        node = package.node(ids[literal])
        return package.ids[next(node.iterancestors(w("tr")))] if layout == "rows" else ids[literal]

    plan = TemplatePlan(
        summary="Independent generic fixture",
        photos=[],
        remove=[],
        warnings=[],
        keep=[ids["Fixed ending"]],
        fields=[
            binding("Seed Person", "personal.name"),
            *[binding(label, "section-title:" + label) for label in labels],
        ],
        repeats=[
            RepeatBinding(
                section="projects",
                start=block(samples[0][0]),
                end=block(samples[0][-1]),
                sample_start=block(samples[0][0]),
                sample_end=block(samples[0][-1]),
                fields=[
                    binding(literal, target)
                    for literal, target in zip(
                        samples[0], ["title", "period", "description"], strict=True
                    )
                ],
            ),
            RepeatBinding(
                section=labels[1],
                start=block(samples[1][0]),
                end=block(samples[1][0]),
                sample_start=block(samples[1][0]),
                sample_end=block(samples[1][0]),
                fields=[binding(samples[1][0], "details")],
            ),
        ],
    )
    assert package.review(plan)["ready"]
    return package, plan


def generic_content():
    """构造两个不同项目和两条普通经历，包含未来新增、可隐藏的自定义信息"""
    document = ResumeDocument.model_validate(
        {
            "personal": {"name": "Current Person", "website": "https://profile.example.test"},
            "sections": [
                {"id": "portfolio", "kind": "projects", "title": "Portfolio Ω"},
                {
                    "id": "recognition",
                    "title": "Recognition · β",
                    "entries": [
                        {
                            "id": f"entry-{i}",
                            "title": f"Award {i}",
                            "period": f"Date {i}",
                            "details": f"Entry body {i}",
                            "custom_fields": [
                                {"id": "issuer", "label": "Issuer", "value": f"Institute {i}"}
                            ],
                        }
                        for i in range(2)
                    ],
                },
            ],
        }
    )
    projects = [
        {
            "project_id": f"project-{i}",
            "highlight_ids": [],
            "content": {
                "title": f"Project {i}",
                "period": f"Project date {i}",
                "role": "",
                "stack": [],
                "description": f"Project body {i}",
                "highlights": [],
                "custom_fields": [
                    {"id": "link", "label": "Repository", "value": f"https://repo{i}.example.test"}
                ],
            },
        }
        for i in range(2)
    ]
    return document, projects


def visible_text(path):
    """提取所有容器的段落文字，检查重复记录、隐藏资料和占位符"""
    return "\n".join(
        row["text"] for row in TemplatePackage(path).inventory()["nodes"] if row["kind"] == "p"
    )


@pytest.mark.parametrize("layout", ["body", "rows", "cells"])
@pytest.mark.parametrize("padding", [0, 11])
def test_mixed_fields_complete_in_unrelated_layouts(tmp_path, layout, padding):
    """不同节点编号和容器均补齐混合缺项，反复运行不增行，显隐变化不泄露资料"""
    source, snapshot, output = [
        tmp_path / name for name in ("source.docx", "snapshot.docx", "output.docx")
    ]
    package, plan = generic_template(source, layout, padding)
    original, original_plan = source.read_bytes(), deepcopy(plan)
    document, projects = generic_content()
    assert set(missing_targets(document, plan, projects)) == {
        "personal.website",
        "Portfolio Ω · custom_fields",
        "Recognition · β · title",
        "Recognition · β · period",
        "Recognition · β · custom_fields",
    }
    repaired, mapping, notices = complete_template(package, plan, document, projects, snapshot)
    assert notices and assess_plan(repaired, mapping, document, projects)["ready"]
    assert any("Portfolio Ω" in notice for notice in notices)
    assert not any("projects" in notice for notice in notices)
    assert TemplatePackage(snapshot).review(mapping)["ready"]
    before = snapshot.read_bytes()
    _, same, notices = complete_template(repaired, mapping, document, projects, snapshot)
    assert same == mapping and not notices and before == snapshot.read_bytes()
    fill_template(snapshot, output, mapping, document.model_dump(), projects)
    text = visible_text(output)
    assert text.count(document.personal.website) == 1
    for index in range(2):
        for literal in (
            f"Project {index}",
            f"Repository：https://repo{index}.example.test",
            f"Award {index}",
            f"Date {index}",
            f"Entry body {index}",
            f"Issuer：Institute {index}",
        ):
            assert text.count(literal) == 1
    assert "Old " not in text and "〔" not in text and text.endswith("Fixed ending")
    document.personal.hidden_fields = ["website"]
    document.sections[1].entries[0].hidden_fields = ["period"]
    document.project_visibility = {"project-0": ProjectVisibility(custom_fields={"link": False})}
    fill_template(snapshot, output, mapping, document.model_dump(), projects)
    text = visible_text(output)
    assert "profile.example.test" not in text and "repo0.example.test" not in text
    assert "Date 0" not in text and "Date 1" in text and "repo1.example.test" in text
    assert source.read_bytes() == original and plan == original_plan


def test_missing_photo_and_unknown_original_content_still_block(tmp_path):
    """可补位字段不会掩盖未识别照片或原文，无法完成时仍阻止丢失信息的输出"""
    source = tmp_path / "source.docx"
    package, plan = generic_template(source)
    document, projects = generic_content()
    document.personal.photo = "data:image/png;base64," + base64.b64encode(photo_bytes(50)).decode()
    repaired, mapping, _ = complete_template(package, plan, document, projects)
    assert missing_targets(document, mapping, projects) == ["照片"]
    assert not assess_plan(repaired, mapping, document, projects)["ready"]
    with pytest.raises(Problem, match="照片"):
        fill_template(source, tmp_path / "output.docx", plan, document.model_dump(), projects)
    plan.fields = [field for field in plan.fields if field.target != "personal.name"]
    unchanged, same, notices = complete_template(package, plan, document, projects)
    assert unchanged is package and same == plan and not notices
    assert not unchanged.review(same)["ready"]


def test_cache_cannot_hide_new_field_requirements(tmp_path):
    """相同字段的新值可复用缓存，新增可见字段会失效，旧缓存不能掩盖新增内容"""
    package, plan = generic_template(tmp_path / "source.docx")
    document, projects = generic_content()
    document.personal.website = ""
    for entry in document.sections[1].entries:
        entry.title = entry.period = ""
        entry.custom_fields = []
    for project in projects:
        project["content"]["custom_fields"] = []
    settings = ProviderSettings()
    original_path = cache_path(tmp_path, package, document, projects, settings)
    remember_plan(original_path, plan)
    assert cached_plan(original_path, package, document, projects) is not None
    document.personal.name = "Another Person"
    assert cache_path(tmp_path, package, document, projects, settings) == original_path
    document.personal.website = "https://new.example.test"
    assert cache_path(tmp_path, package, document, projects, settings) != original_path
    assert cached_plan(original_path, package, document, projects) is None
    document.personal.hidden_fields = ["website"]
    assert cache_path(tmp_path, package, document, projects, settings) == original_path
    assert cached_plan(original_path, package, document, projects) is not None


def test_renamed_project_heading_can_supply_new_section_style(tmp_path):
    """项目栏目改名后仍可复用大标题样式"""
    source = tmp_path / "source.docx"
    package, plan = generic_template(source)
    # 移除普通栏目示例且只保留一个以任意名称命名的项目栏目
    ordinary = plan.repeats.pop()
    heading = next(field for field in plan.fields if field.target.endswith("Recognition · β"))
    plan.fields.remove(heading)
    plan.remove.extend([heading.node, ordinary.start])
    document, projects = generic_content()
    repaired, mapping, _ = complete_template(package, plan, document, projects)
    assert assess_plan(repaired, mapping, document, projects)["ready"]
    snapshot = tmp_path / "snapshot.docx"
    repaired.write(snapshot)
    output = tmp_path / "output.docx"
    fill_template(snapshot, output, mapping, document.model_dump(), projects)
    text = visible_text(output)
    assert text.count("Recognition · β") == 1 and "Award 0" in text


def test_first_analysis_completes_all_supported_missing_fields(tmp_path, monkeypatch):
    """首轮识别完整原文后即补齐混合缺项，最终清单、源快照和方案使用同一套编号"""
    source = tmp_path / "original.docx"
    package, plan = generic_template(source, "cells", 8)
    document, projects = generic_content()
    monkeypatch.setattr(
        "resume_maker.services.templates.analysis.source_pages", lambda *_: ([], {}, [])
    )

    class MappedProvider(TemplateProvider):
        """只识别已有原文，缺少的位置必须由通用补位器处理"""

        def run_structured(self, **kwargs):
            """返回预先核对的模板映射"""
            self.calls.append(kwargs)
            return plan.model_copy(deep=True)

    def emit(*_):
        """测试忽略公开进度且只核验结果和实际文件"""

    provider = MappedProvider()
    mapping, review, attempts, error = analyze_plan(
        package, provider, tmp_path, document, projects, ProviderSettings(), threading.Event(), emit
    )
    assert review["ready"] and not review["missing"] and attempts == 1 and error is None
    assert len(provider.calls) == 1 and TemplatePackage(source).review(mapping)["ready"]
    output = tmp_path / "output.docx"
    fill_template(source, output, mapping, document.model_dump(), projects)
    assert "Award 1" in visible_text(
        output
    ) and "Repository：https://repo0.example.test" in visible_text(output)


def test_library_review_save_and_reopen_share_completion(tmp_path, monkeypatch):
    """真实接口打开旧模板、资料新增后的检查、试填及保存均补齐，原版本和任务源文件不变"""
    config = Config(data_dir=tmp_path / "data", token="test")
    provider = TemplateProvider(failure=True)
    app = create_app(config, provider)
    monkeypatch.setattr("resume_maker.services.templates.tasks.render_word", lambda *_: (1, None))
    with TestClient(app) as client:
        catalog = app.state.services.catalog
        folder = config.data_dir / "templates" / "generic"
        folder.mkdir(parents=True)
        _, plan = generic_template(folder / "template.docx", "rows", 5)
        original = (folder / "template.docx").read_bytes()
        with catalog.db.transaction() as conn:
            conn.execute(
                "INSERT INTO templates VALUES (?,?,?,?,?)",
                ("generic", "任意模板", digest(original), dump({"plan": plan.model_dump()}), now()),
            )
        document, projects = generic_content()
        items = []
        for data in projects:
            project_root = tmp_path / data["project_id"]
            project_root.mkdir()
            project = catalog.create_project(data["project_id"], [str(project_root)])
            base = project["head_revision"]
            catalog.put_draft(project["id"], base, "experience", data["content"], 0)
            revision = catalog.save_revision(project["id"], base, base)
            items.append(
                {"project_id": project["id"], "revision_id": revision["id"], "highlight_ids": []}
            )
        headers = {"x-resume-token": "test"}
        body = {"document": document.model_dump(), "items": items}
        legacy = client.post("/api/templates/generic/edit", headers=headers).json()
        repair = client.post(
            f"/api/templates/analyses/{legacy['id']}/repair",
            json={**body, "plan": legacy["plan"]},
            headers=headers,
        )
        assert repair.status_code == 200
        repaired = completed(app.state.services.templates, repair.json()["id"])
        assert repaired["review"]["ready"] and repaired["attempts"] == 0
        assert not provider.calls
        response = client.post("/api/templates/generic/edit", json=body, headers=headers)
        assert response.status_code == 200
        opened = response.json()
        assert opened["review"]["ready"] and not opened["review"]["missing"]
        assert not provider.calls
        task_source = app.state.services.templates.source(opened["id"])
        assert opened["inventory"]["nodes"] == TemplatePackage(task_source).inventory()["nodes"]
        # 打开后继续新增字段，检查和保存必须和预览使用相同的自动补位规则
        body["document"]["personal"]["phone"] = "123456789"
        body["document"]["sections"][1]["entries"][0]["subtitle"] = "New subtitle"
        body["plan"] = opened["plan"]
        task_bytes = task_source.read_bytes()
        prefix = f"/api/templates/analyses/{opened['id']}"
        reviewed = client.post(prefix + "/review", json=body, headers=headers).json()
        assert reviewed["ready"] and reviewed["missing"] == [] and reviewed["notices"]
        preview = client.post(prefix + "/preview", json=body, headers=headers)
        assert preview.status_code == 200
        output = task_source.parent / preview.json()["id"] / "resume.docx"
        assert "电话：123456789" in visible_text(output) and "New subtitle" in visible_text(output)
        saved = client.post(prefix + "/save", json={**body, "name": "新版"}, headers=headers)
        assert saved.status_code == 200
        reopened = client.post(
            f"/api/templates/{saved.json()['id']}/edit",
            json={"document": body["document"], "items": items},
            headers=headers,
        )
        assert reopened.status_code == 200 and reopened.json()["review"]["ready"]
        assert (
            folder / "template.docx"
        ).read_bytes() == original and task_source.read_bytes() == task_bytes
        assert catalog.template("generic")["mapping"]["plan"] == plan.model_dump()
