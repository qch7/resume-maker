"""程序补齐模板缺项时保留姓名位置"""

from copy import deepcopy

import pytest
from docx import Document
from pydantic import ValidationError
from test_template_completion import generic_content, generic_template, visible_text
from test_template_order import body_text

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.completion import complete_template
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text
from resume_maker.services.templates.analysis import assess_plan, complete_labels
from resume_maker.services.templates.schema import plan_schema


def test_model_schema_only_accepts_actual_nodes_and_known_sections(tmp_path):
    """模型不能把原文写入节点字段，也不能发明一个不存在的栏目"""
    package, plan = generic_template(tmp_path / "source.docx")
    content, _ = generic_content()
    model = plan_schema(package, content)
    assert model.model_validate(plan.model_dump()).model_dump() == plan.model_dump()
    wrong = plan.model_dump()
    wrong["remove"] = ["Seed Person"]
    with pytest.raises(ValidationError):
        model.model_validate(wrong)
    wrong = plan.model_dump()
    wrong["repeats"][0]["section"] = "Invented section"
    with pytest.raises(ValidationError):
        model.model_validate(wrong)


@pytest.mark.parametrize("conflicting", [False, True])
def test_redundant_nested_deletions_are_reduced_without_hiding_real_conflicts(
    tmp_path, conflicting
):
    """父块删除自动覆盖其子段落，父块包含实际填写字段时仍拒绝"""
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Old person"
    doc.save(path)
    package = TemplatePackage(path)
    nodes = package.inventory()["nodes"]
    paragraph = next(row["id"] for row in nodes if row["kind"] == "p")
    table = next(row["id"] for row in nodes if row["kind"] == "tbl")
    plan = TemplatePlan(
        summary="Nested removal",
        fields=[TextBinding(node=paragraph, quote="Old person", target="personal.name")]
        if conflicting
        else [],
        repeats=[],
        photos=[],
        keep=[paragraph],
        remove=[paragraph, table],
        warnings=[],
    )
    updated = complete_labels(package, plan)
    assert updated.remove == [table]
    assert not updated.keep
    assert package.review(updated)["ready"] is not conflicting


@pytest.mark.parametrize("layout", ["body", "rows", "columns"])
def test_missing_child_section_is_inserted_after_parent_and_can_be_hidden(tmp_path, layout):
    """原稿没有子栏目时程序补齐其标题和记录，父级隐藏也隐藏子栏目"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    package, plan = generic_template(source, layout)
    content, projects = generic_content()
    from resume_maker.domain.resume import ResumeSection

    content.sections.append(
        ResumeSection(
            id="child",
            parent_id="recognition",
            title="Additional study",
            entries=[
                {"id": "one", "details": "Course Alpha"},
                {"id": "two", "details": "Course Beta"},
            ],
        )
    )
    fill_template(source, output, plan, content.model_dump(), projects)
    text = body_text(output)
    assert text.index("Award 1") < text.index("Additional study") < text.index("Course Alpha")
    assert text.count("Additional study") == 1 and text.count("Course Beta") == 1
    result = TemplatePackage(output)
    assert all(
        child.tag != w("p")
        for node in result.nodes.values()
        if node.tag == w("tbl")
        for child in node
    )
    if layout == "columns":
        child = next(n for n in result.nodes.values() if paragraph_text(n) == "Additional study")
        cell = next(child.iterancestors(w("tc")))
        assert "Award 1" in "".join(cell.itertext()) and "Project 0" not in "".join(cell.itertext())
    content.sections[1].visible = False
    fill_template(source, output, plan, content.model_dump(), projects)
    assert "Course Alpha" not in body_text(output) and "Additional study" not in body_text(output)


@pytest.mark.parametrize("blank", [True, False])
@pytest.mark.parametrize("table", [True, False])
def test_section_before_personal_header_cannot_move_header(tmp_path, blank, table):
    """模拟无标题荣誉位于姓名前，真实荣誉也不能把其后的独立个人区吸入"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("" if blank else "Old award")
    doc.add_paragraph("Old name")
    if table:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "old@example.test"
    else:
        doc.add_paragraph("old@example.test")
    doc.add_paragraph("Skills", "Heading 1")
    doc.add_paragraph("Old skill")
    doc.save(source)
    package = TemplatePackage(source)
    ids = {paragraph_text(node): key for key, node in package.nodes.items() if node.tag == w("p")}
    first = package.ids[package.parts["word/document.xml"].find(w("body"))[0]]
    plan = TemplatePlan(
        summary="Header boundary fixture",
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
        fields=[
            TextBinding(node=ids["Old name"], quote="Old name", target="personal.name"),
            TextBinding(
                node=ids["old@example.test"], quote="old@example.test", target="personal.email"
            ),
            TextBinding(node=ids["Skills"], quote="Skills", target="section-title:Skills"),
        ],
        repeats=[
            RepeatBinding(
                section="Awards",
                start=first,
                end=first,
                sample_start=first,
                sample_end=first,
                fields=[
                    TextBinding(node=first, quote="" if blank else "Old award", target="title")
                ],
            ),
            RepeatBinding(
                section="Skills",
                start=ids["Old skill"],
                end=ids["Old skill"],
                sample_start=ids["Old skill"],
                sample_end=ids["Old skill"],
                fields=[TextBinding(node=ids["Old skill"], quote="Old skill", target="details")],
            ),
        ],
    )
    content = ResumeDocument(
        personal={"name": "New name", "email": "new@example.test"},
        sections=[
            {
                "id": "skills",
                "title": "Skills",
                "entries": [{"id": "skill", "details": "New skill"}],
            },
            {
                "id": "awards",
                "title": "Awards",
                "entries": [{"id": "award", "title": "New award", "period": "2026"}],
            },
            {"id": "projects", "title": "Projects", "kind": "projects"},
        ],
    )
    original, original_plan = source.read_bytes(), deepcopy(plan)
    fill_template(source, output, plan, content.model_dump(), [])
    text = body_text(output)
    assert text.index("New name") < text.index("new@example.test") < text.index("Skills")
    assert text.index("New skill") < text.index("New award") < text.index("2026")
    assert "Old" not in text and "〔" not in text
    if blank:
        assert text.count("Awards") == 1
    assert source.read_bytes() == original and plan == original_plan


def test_missing_project_body_is_completed_without_model_slots(tmp_path):
    """只有项目名称的映射也可补入全部真实正文，并保留每个项目的不同值"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    package, plan = generic_template(source)
    region = plan.repeats[0]
    removed = [field for field in region.fields if field.target != "title"]
    # 旧示例仍须明确清理，缺少新字段不能成为保留旧正文的借口
    for field in removed:
        package.node(field.node).getparent().remove(package.node(field.node))
    previous = dict(package.nodes)
    region.fields = [field for field in region.fields if field.target == "title"]
    region.end = region.sample_end = region.start
    package.reindex()
    from resume_maker.integrations.word.templates.supplement import remap_plan

    plan = remap_plan(plan, previous, package)
    package.write(source)
    content, projects = generic_content()
    for i, project in enumerate(projects):
        project["content"].update(
            role=f"Role {i}",
            stack=[f"Stack {i}"],
            highlights=[{"id": "one", "title": f"Point {i}", "text": f"Evidence {i}"}],
        )
        project["highlight_ids"] = ["one"]
    completed, mapping, notices = complete_template(package, plan, content, projects)
    assert notices and assess_plan(completed, mapping, content, projects)["ready"]
    fill_template(source, output, plan, content.model_dump(), projects)
    text = visible_text(output)
    for i in range(2):
        for value in (
            f"Project {i}",
            f"Role {i}",
            f"Stack {i}",
            f"Evidence {i}",
            f"Project body {i}",
        ):
            assert text.count(value) == 1
    assert "〔" not in text
