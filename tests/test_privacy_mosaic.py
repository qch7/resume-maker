"""验证模板图片打码、模型附件和无原图回退的隐私边界"""

import base64
import hashlib
import json
import threading
from io import BytesIO

import pytest
from docx import Document
from PIL import Image, ImageDraw, PngImagePlugin
from test_template_analysis import completed, simple_document

from resume_maker.domain.models import ProviderSettings
from resume_maker.domain.templates import TemplatePlan
from resume_maker.integrations.privacy_store import PrivacyStore
from resume_maker.integrations.providers.base import Cancelled, MosaicImage, ProviderError
from resume_maker.integrations.providers.codex import CodexProvider
from resume_maker.integrations.providers.mosaic import mosaic_sheets, mosaic_tile
from resume_maker.services.templates.tasks import Templates


def synthetic_image():
    """生成含细小身份文字和元数据的合成人像，不读取任何真实照片"""
    image = Image.new("RGB", (240, 320), "#d8edf6")
    draw = ImageDraw.Draw(image)
    draw.ellipse((65, 25, 175, 185), fill="#e8b997")
    draw.ellipse((65, 15, 175, 70), fill="#303030")
    draw.rectangle((50, 190, 190, 320), fill="#284568")
    draw.text((80, 100), "PRIVATE-ID", fill="black")
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("private", "ORIGINAL-METADATA-CANARY")
    output = BytesIO()
    image.save(output, format="PNG", pnginfo=metadata)
    return output.getvalue()


def test_mosaic_removes_detail_metadata_and_preserves_node_labels():
    """原图降成有限色块后重新编码，附件标签独立绘制且可审计"""
    raw = synthetic_image()
    tile = mosaic_tile(raw)
    assert len(tile.getcolors(tile.width * tile.height)) <= 64
    with Image.open(BytesIO(raw)) as original:
        assert tile.tobytes() != original.resize(tile.size).tobytes()
    sheets, records = mosaic_sheets([MosaicImage("n71", raw)], threading.Event())
    assert len(sheets) == 1 and raw not in sheets[0]
    assert b"ORIGINAL-METADATA-CANARY" not in sheets[0]
    with Image.open(BytesIO(sheets[0])) as sheet:
        assert sheet.size == (640, 720) and sheet.mode == "RGB"
        assert not sheet.info and not sheet.getexif()
        assert sheet.crop((16, 10, 90, 32)).getbbox()
        assert sheet.crop((16, 10, 90, 32)).getextrema()[0][0] < 255
    assert records[0]["nodes"] == ["n71"]
    assert records[0]["sha256"] == hashlib.sha256(sheets[0]).hexdigest()
    assert records[0]["grid_long_edge"] == 8


def test_mosaic_batches_labels_and_cancellation():
    """多图片分组时标签完整保留，取消后不生成剩余附件"""
    images = [MosaicImage(f"n{i + 1}", synthetic_image()) for i in range(7)]
    sheets, records = mosaic_sheets(images, threading.Event())
    assert len(sheets) == 2
    assert [node for record in records for node in record["nodes"]] == [i.node for i in images]
    flag = threading.Event()
    flag.set()
    with pytest.raises(Cancelled):
        mosaic_sheets(images, flag)


@pytest.mark.parametrize("failure", ["decode", "node", "duplicate", "count", "bytes", "pixels"])
def test_unsafe_mosaics_never_reach_runner(tmp_path, failure):
    """解码、边界或编号校验失败时停止发送，不降级为原图附件"""
    images = [MosaicImage("n1", synthetic_image())]
    if failure == "decode":
        images = [MosaicImage("n1", b"not an image")]
    elif failure == "node":
        images = [MosaicImage("PRIVATE-NAME", images[0].data)]
    elif failure == "duplicate":
        images *= 2
    elif failure == "count":
        images = [MosaicImage(f"n{i + 1}", images[0].data) for i in range(25)]
    elif failure == "bytes":
        images = [MosaicImage("n1", b"a" * (20 * 1024 * 1024 + 1))]
    else:
        buffer = BytesIO()
        Image.new("1", (5000, 5000)).save(buffer, format="PNG")
        images = [MosaicImage("n1", buffer.getvalue())]
    calls = []
    provider = CodexProvider(runner=lambda *args, **kwargs: calls.append(args))
    with pytest.raises(ProviderError):
        provider.run_structured(
            result_model=TemplatePlan,
            workspace=tmp_path,
            prompt="检查模板",
            thread_id=None,
            settings=ProviderSettings(),
            cancelled=threading.Event(),
            emit=lambda *_: None,
            images=images,
        )
    assert not calls


def test_template_mosaic_repair_covers_photo_without_original_upload(tmp_path, catalog):
    """从遗漏照片的旧方案修复到可试填，首次和每轮修复均只附带马赛克"""
    raw = synthetic_image()
    source = tmp_path / "template.docx"
    template = Document()
    template.add_paragraph("姓名：合成测试甲")
    template.add_picture(BytesIO(raw))
    template.save(source)
    original = source.read_bytes()
    seen = []

    def runner(payload, settings, environment, cancelled, emit, *, safe_images):
        """替身核对真实隐私出口的附件，再按节点编号返回照片位置"""
        seen.append(payload)
        context = json.loads(payload["input"].splitlines()[-1])
        assert not context["source_pages"]["available"]
        assert len(safe_images) == len(payload["images"]) == 1
        assert all(
            raw != image and b"ORIGINAL-METADATA-CANARY" not in image for image in safe_images
        )
        assert "合成测试甲" not in payload["input"]
        assert base64.b64encode(raw).decode() not in json.dumps(payload)
        rows = [row for part in context["template"]["parts"].values() for row in part]
        paragraph = next(row for row in rows if row[1] == "p" and row[4])
        photo = next(row[0] for row in rows if row[1] == "image")
        assert photo in context["visible_images"]
        assert payload["images"][0]["nodes"] == [photo]
        return json.dumps(
            {
                "summary": "根据马赛克人像轮廓映射照片",
                "fields": [
                    {"node": paragraph[0], "quote": paragraph[4], "target": "personal.name"}
                ],
                "photos": [photo],
                "repeats": [],
                "keep": [],
                "remove": [],
                "warnings": [],
            }
        )

    provider = CodexProvider(runner=runner, privacy=PrivacyStore(catalog.db))
    service = Templates(catalog, tmp_path / "data", provider)
    document = simple_document()
    document.personal.name = "合成测试甲"
    document.personal.photo = "data:image/png;base64," + base64.b64encode(raw).decode()
    task = completed(service, service.analyze(source, document)["id"])
    assert task["status"] == "completed", task["error"]
    assert task["review"]["ready"] and task["attempts"] == 1
    old_plan = TemplatePlan.model_validate(task["plan"])
    old_plan.keep = old_plan.photos
    old_plan.photos = []
    repaired = completed(service, service.repair(task["id"], old_plan, document, [])["id"])
    assert repaired["review"]["ready"] and repaired["attempts"] == 1
    assert repaired["plan"]["photos"] == task["plan"]["photos"]
    assert len(seen) == 2 and source.read_bytes() == original
    assert catalog.db.setting("privacy_audit")[0]["payload"]["images"] == seen[-1]["images"]
