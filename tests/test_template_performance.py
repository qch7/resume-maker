"""验证缓存、增量会话和活动查询保留覆盖校验、任务隔离及取消处理"""

import json

from docx import Document
from fastapi.testclient import TestClient
from test_template_analysis import TemplateProvider, completed, simple_document, simple_template
from test_template_repair import RepairProvider

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.services.templates.analysis import compact_inventory
from resume_maker.services.templates.tasks import Templates


class SessionProvider(RepairProvider):
    """返回可用于下一轮续聊的会话事件"""

    def run_structured(self, **kwargs):
        """每轮上报同一会话标识，首轮遗漏会触发增量修正"""
        kwargs["emit"]("thread", {"id": "template-session"})
        kwargs["emit"]("usage", {"input_tokens": 100 * (len(self.calls) + 1), "cumulative": True})
        return super().run_structured(**kwargs)


def test_repair_reuses_only_its_own_session(catalog, tmp_path):
    """新任务从独立会话开始，自动修正省去清单、图片和未变化的旧方案"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = SessionProvider()
    service = Templates(catalog, tmp_path, provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    first, second = provider.calls
    assert first["settings"].reasoning_effort == second["settings"].reasoning_effort == ""
    assert task["review"]["ready"] and task["attempts"] == 2
    assert task["usage"]["input_tokens"] == 200
    assert first["thread_id"] is None and second["thread_id"] == "template-session"
    request = json.loads(second["prompt"].split("\n")[-1])
    assert "template" not in request and "previous_plan" not in request
    assert second["images"] == [] and len(second["prompt"]) < len(first["prompt"]) / 2
    changed = service.repair(
        task["id"],
        first["result_model"].model_validate(task["plan"]),
        simple_document(),
        [],
        "重新核对",
    )
    completed(service, changed["id"])
    assert provider.calls[2]["thread_id"] is None


def test_cache_survives_restart_but_rechecks_changed_requirements(catalog, tmp_path):
    """资料值变化可复用，新增待填字段必须重识别，人工要求不能被缓存吞掉"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    first = completed(service, service.analyze(source, simple_document())["id"])
    service.stop()
    service = Templates(catalog, tmp_path, provider)
    doc = simple_document()
    doc.personal.name = "另一位使用者"
    second = completed(service, service.analyze(source, doc)["id"])
    assert second["reused"] and second["attempts"] == 0 and len(provider.calls) == 1
    assert second["plan"] == first["plan"] and second["review"]["ready"]
    doc.personal.phone = "10000000000"
    third = completed(service, service.analyze(source, doc)["id"])
    assert not third["reused"] and third["review"]["ready"]
    assert not third["review"]["missing"] and len(provider.calls) == 2
    assert len(list((tmp_path / "template-cache").glob("*.json"))) == 1


def test_corrupted_or_invalid_cache_falls_back_to_analysis(catalog, tmp_path):
    """损坏或遗漏字段的缓存都不能伪装成完成结果"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    first = completed(service, service.analyze(source, simple_document())["id"])
    path = next((tmp_path / "template-cache").glob("*.json"))
    for content in ("broken json", json.dumps({**first["plan"], "fields": []})):
        path.write_text(content, encoding="utf-8")
        result = completed(service, service.analyze(source, simple_document())["id"])
        assert not result["reused"] and result["review"]["ready"]
    assert len(provider.calls) == 3
    changed = Document(source)
    changed.add_paragraph("另一段正文")
    changed.save(source)
    result = completed(service, service.analyze(source, simple_document())["id"])
    assert not result["reused"] and not result["review"]["ready"]


def test_compact_inventory_preserves_exact_text_and_boundaries(tmp_path):
    """清单压缩只消除重复结构键和容器正文，保留所有段落、空位和结构约束"""
    path = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for cell in table.rows[0].cells:
        cell.text = "长文本  精确空格\t及联系方式：10000000000" * 10
    doc.save(path)
    inventory = TemplatePackage(path).inventory()
    compact = compact_inventory(inventory)
    rebuilt = []
    for part, rows in compact["parts"].items():
        rebuilt.extend(
            {**dict(zip(compact["columns"], row, strict=True)), "part": part} for row in rows
        )
    assert len(rebuilt) == len(inventory["nodes"])
    for original, node in zip(inventory["nodes"], rebuilt, strict=True):
        assert node == {
            **original,
            "text": original["text"] if original["kind"] in {"p", "image"} else "",
        }
    assert len(json.dumps(compact)) < len(json.dumps(inventory))


def test_progress_is_incremental_bounded_and_freezes_on_cancel(tmp_path, monkeypatch):
    """进度鉴权和结果相同，游标增量不携带模板原文，取消后事件和耗时冻结"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider(block=True)
    app = create_app(Config(data_dir=tmp_path / "data", token="test"), provider)
    with TestClient(app) as client:
        service = app.state.services.templates
        task = service.analyze(source, simple_document())
        assert provider.started.wait(2)
        emit = provider.calls[0]["emit"]
        for index in range(100):
            emit("activity", {"text": f"活动 {index} api_key=secret-value"})
        emit(
            "usage",
            {
                "input_tokens": 123,
                "output_tokens": 45,
                "cached_input_tokens": 100,
                "private": "hidden",
            },
        )
        endpoint = f"/api/templates/analyses/{task['id']}/progress"
        headers = {"x-resume-token": "test"}
        assert client.get(endpoint).status_code == 401
        progress = client.get(endpoint, headers=headers).json()
        assert len(progress["events"]) == 80 and "secret-value" not in json.dumps(progress)
        assert not {"inventory", "plan", "review"} & progress.keys()
        assert progress["usage"] == {
            "input_tokens": 123,
            "output_tokens": 45,
            "cached_input_tokens": 100,
        }
        assert not client.get(endpoint + f"?after={progress['cursor']}", headers=headers).json()[
            "events"
        ]
        cancelled = service.cancel(task["id"])
        emit("activity", {"text": "迟到的事件"})
        provider.release.set()
        completed(service, task["id"])
        assert service.progress(task["id"])["elapsed_ms"] == cancelled["elapsed_ms"]
        assert service.progress(task["id"])["cursor"] == progress["cursor"]
        assert client.get(endpoint + "?after=-1", headers=headers).status_code == 422
        assert (
            client.get("/api/templates/analyses/unknown/progress", headers=headers).status_code
            == 404
        )
