"""原生模板的组合图形、浮动坐标、照片和横向分栏不能因识别或填充丢失"""

import base64
import threading
from io import BytesIO

import pytest
from docx import Document
from lxml import etree
from test_template_mapping import photo_bytes

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.recovery import prepare_template
from resume_maker.integrations.word.templates.fill import fill_fields, fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WPG = "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def grouped_header(path):
    """构造照片、蓝色背景和白色姓名共用组合图形；锚点借用课程正文的模板"""
    document = Document()
    document.add_paragraph("固定说明")
    document.add_paragraph("教育背景")
    document.add_paragraph("旧学校")
    course = document.add_paragraph("旧课程")
    picture = course.add_run().add_picture(BytesIO(photo_bytes(30)))
    inline = picture._inline
    inline.tag = f"{{{WP}}}anchor"
    inline.set("behindDoc", "0")
    for axis, offset in (("H", "12000"), ("V", "0")):
        position = etree.Element(f"{{{WP}}}position{axis}", relativeFrom="page")
        etree.SubElement(position, f"{{{WP}}}posOffset").text = offset
        inline.insert(0, position)
    graphic = inline.find(f"{{{A}}}graphic/{{{A}}}graphicData")
    graphic.set("uri", WPG)
    photo = graphic[0]
    group = etree.SubElement(graphic, f"{{{WPG}}}wgp")
    group.append(photo)
    shape = etree.SubElement(group, f"{{{WPS}}}wsp")
    properties = etree.SubElement(shape, f"{{{WPS}}}spPr")
    fill = etree.SubElement(properties, f"{{{A}}}solidFill")
    etree.SubElement(fill, f"{{{A}}}srgbClr", val="718DB5")
    geometry = etree.SubElement(properties, f"{{{A}}}prstGeom", prst="rect")
    etree.SubElement(geometry, f"{{{A}}}avLst")
    textbox = etree.SubElement(etree.SubElement(shape, f"{{{WPS}}}txbx"), w("txbxContent"))
    paragraph = etree.SubElement(textbox, w("p"))
    run = etree.SubElement(paragraph, w("r"))
    etree.SubElement(etree.SubElement(run, w("rPr")), w("color")).set(w("val"), "FFFFFF")
    etree.SubElement(run, w("t")).text = "旧姓名"
    document.add_paragraph("项目经历")
    document.add_paragraph("旧项目")
    document.save(path)
    package = TemplatePackage(path)
    ids = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    photo_id = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "image")
    plan = TemplatePlan(
        summary="保留组合设计",
        fields=[
            TextBinding(node=ids[text], quote=text, target=target)
            for text, target in (
                ("旧姓名", "personal.name"),
                ("教育背景", "section-title:教育背景"),
                ("项目经历", "section-title:项目经历"),
            )
        ],
        repeats=[
            RepeatBinding(
                section=section,
                start=ids[text],
                end=ids[text],
                sample_start=ids[text],
                sample_end=ids[text],
                fields=[TextBinding(node=ids[text], quote=text, target="title")],
            )
            for section, text in (
                ("教育背景", "旧学校"),
                ("主修课程", "旧课程"),
                ("projects", "旧项目"),
            )
        ],
        photos=[photo_id],
        keep=[ids["固定说明"]],
        remove=[],
        warnings=[],
    )
    return package, plan


@pytest.mark.parametrize("photo", [True, False])
@pytest.mark.parametrize("education", [True, False])
def test_grouped_photo_changes_leave_background_text_and_page_anchor(tmp_path, photo, education):
    """更换或隐藏照片、隐藏其原锚点栏目时；页首背景、姓名、坐标和其他图形仍保留"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    package, plan = grouped_header(source)
    before = source.read_bytes()
    assert package.review(plan)["ready"]
    content = ResumeDocument.model_validate(
        {
            "personal": {
                "name": "新姓名",
                "photo": "data:image/png;base64," + base64.b64encode(photo_bytes(180)).decode()
                if photo
                else "",
            },
            "sections": [
                {"id": "projects", "title": "项目经历", "kind": "projects"},
                {
                    "id": "edu",
                    "title": "教育背景",
                    "kind": "education",
                    "visible": education,
                    "entries": [{"id": "school", "title": "新学校"}],
                },
                {
                    "id": "course",
                    "title": "主修课程",
                    "kind": "text",
                    "parent_id": "edu",
                    "entries": [{"id": "course", "title": "新课程"}],
                },
            ],
        }
    )
    fill_template(source, output, plan, content.model_dump(), [])
    result = TemplatePackage(output)
    body = result.parts["word/document.xml"]
    drawing_ids = body.xpath(".//*[local-name()='docPr' or local-name()='cNvPr']/@id")
    assert len(drawing_ids) == len(set(drawing_ids))
    assert body.find(f".//{{{A}}}srgbClr").get("val") == "718DB5"
    assert body.find(f".//{{{A}}}prstGeom").get("prst") == "rect"
    assert body.find(f".//{{{WP}}}positionH/{{{WP}}}posOffset").text == "12000"
    assert body.find(".//w:txbxContent/w:p/w:r/w:rPr/w:color", NS).get(w("val")) == "FFFFFF"
    assert "新姓名" in "".join(body.itertext()) and "旧姓名" not in "".join(body.itertext())
    assert ("新学校" in "".join(body.itertext())) == education
    images = [row for row in result.inventory()["nodes"] if row["kind"] == "image"]
    assert len(images) == int(photo)
    if photo:
        assert result.image(images[0]["id"]) == photo_bytes(180)
    assert source.read_bytes() == before


def test_native_normalization_is_idempotent_and_keeps_separate_drawing_geometry(tmp_path):
    """共用锚点分离后再读取编号稳定；绘图 XML 完整保留且不混入文字重复区"""
    source, output = tmp_path / "source.docx", tmp_path / "normalized.docx"
    package, plan = grouped_header(source)
    package.write(output)
    repeated = TemplatePackage(output)
    assert repeated.inventory() == package.inventory()
    assert repeated.review(plan)["ready"]
    course = next(region for region in plan.repeats if region.section == "主修课程")
    assert not package.node(course.start).findall(".//w:drawing", NS)
    assert len(package.parts["word/document.xml"].findall(".//w:drawing", NS)) == 1


def test_import_separates_paragraph_relative_drawing_without_reindexing_saved_sources(tmp_path):
    """首次识别才分离相对段落锚点；既有模板按旧编号加载且写回后编号稳定。"""
    source, output = tmp_path / "source.docx", tmp_path / "prepared.docx"
    grouped_header(source)
    doc = Document(source)
    anchor = doc.element.find(f".//{{{WP}}}anchor")
    anchor.find(f"{{{WP}}}positionV").set("relativeFrom", "paragraph")
    before_geometry = etree.tostring(anchor)
    doc.save(source)
    original = source.read_bytes()
    existing = TemplatePackage(source)
    course = next(
        node
        for node in existing.nodes.values()
        if node.tag == w("p") and paragraph_text(node) == "旧课程"
    )
    assert course.findall(".//w:drawing", NS)
    package, notices = prepare_template(
        source, output, None, None, threading.Event(), lambda *_: None, None, []
    )
    course = next(
        node
        for node in package.nodes.values()
        if node.tag == w("p") and paragraph_text(node) == "旧课程"
    )
    assert not course.findall(".//w:drawing", NS)
    actual = package.parts["word/document.xml"].find(f".//{{{WP}}}anchor")
    # 序列化后命名空间声明可能不同；比较几何和图形内容而非前缀。
    old = etree.fromstring(before_geometry)
    assert actual.attrib == old.attrib
    assert actual.find(f"{{{WP}}}positionV/{{{WP}}}posOffset").text == "0"
    assert actual.find(f".//{{{A}}}srgbClr").get("val") == "718DB5"
    assert TemplatePackage(output).inventory() == package.inventory()
    assert notices and source.read_bytes() == original


def test_repeated_education_preserves_three_column_layout_and_fonts(tmp_path):
    """多条教育记录仍复制原三栏节属性、字体和横向段落且不改建通用表格"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    document = Document()
    for text in ("旧日期", "旧学校", "旧专业"):
        paragraph = document.add_paragraph(text)
        paragraph.runs[0].font.name = "楷体"
    closing = document.add_paragraph()
    section = etree.SubElement(closing._p.get_or_add_pPr(), w("sectPr"))
    etree.SubElement(section, w("type")).set(w("val"), "continuous")
    etree.SubElement(section, w("cols")).set(w("num"), "3")
    document.save(source)
    package = TemplatePackage(source)
    rows = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="横向三栏",
        fields=[],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=rows[0]["id"],
                end=rows[-1]["id"],
                sample_start=rows[0]["id"],
                sample_end=rows[-1]["id"],
                fields=[
                    TextBinding(node=row["id"], quote=row["text"], target=target)
                    for row, target in zip(rows[:3], ("period", "title", "subtitle"), strict=True)
                ],
            )
        ],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    content = ResumeDocument.model_validate(
        {
            "sections": [
                {
                    "id": "edu",
                    "title": "教育背景",
                    "kind": "education",
                    "entries": [
                        {
                            "id": str(i),
                            "title": f"学校{i}",
                            "subtitle": f"专业{i}",
                            "period": str(i),
                        }
                        for i in (1, 2)
                    ],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
            ]
        }
    )
    fill_template(source, output, plan, content.model_dump(), [])
    result = Document(output)
    assert not result.tables
    assert len(result.element.body.xpath(".//w:cols[@w:num='3']")) == 2
    assert [p.text for p in result.paragraphs if p.text] == [
        "1",
        "学校1",
        "专业1",
        "2",
        "学校2",
        "专业2",
    ]
    assert all(p.runs[0].font.name == "楷体" for p in result.paragraphs if p.text)


def test_highlight_keeps_bold_title_and_regular_body_across_styled_spaces():
    """替换整条亮点仍区分标题和正文且不把旧标题末尾的粗体空格样式扩散到全文"""
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("旧标题： ").bold = True
    paragraph.add_run("旧正文").bold = False
    fill_fields(
        {"p": paragraph._p},
        [TextBinding(node="p", quote="旧标题： 旧正文", target="highlights")],
        {"highlights": "新的标题：这是一段更长的新正文"},
    )
    assert paragraph.text == "新的标题：这是一段更长的新正文"
    assert next(run for run in paragraph.runs if "新的标题" in run.text).bold
    assert next(run for run in paragraph.runs if "新正文" in run.text).bold is False
