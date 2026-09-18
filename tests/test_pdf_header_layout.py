"""以独立合成 PDF 验证实际资料填充后的顶部布局，不使用用户姓名或固定节点编号。"""

from copy import deepcopy
from threading import Event

import pymupdf
import pytest
from docx import Document
from lxml import etree
from test_template_mapping import photo_bytes, resume_content

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.pdf_geometry import PDF, SOURCE, TEXT, WP
from resume_maker.integrations.word.pdf_header_layout import PDFHeaderLayout
from resume_maker.integrations.word.pdf_recovery import rebuild_pdf
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text


def header_source(directory, left=False):
    """生成带同排联系方式、独立图标、头像和正文边界的 PDF，再按原文建立完整映射。"""
    pdf, source = directory / "source.pdf", directory / "source.docx"
    with pymupdf.open() as document:
        page = document.new_page(width=540, height=700)
        page.insert_text((200, 45), "OLD NAME", fontsize=20, fontname="hebo")
        for x, y, text in [
            (90, 80, "OLD PHONE"),
            (255, 80, "OLD EMAIL"),
            (90, 103, "OLD ROLE"),
            (255, 103, "OLD CUSTOM"),
        ]:
            page.draw_circle((x - 10, y - 3), 4, color=(0.1, 0.2, 0.3))
            page.insert_text((x, y), text, fontsize=11)
        page.insert_image(
            (25 if left else 455, 22, 73 if left else 503, 86), stream=photo_bytes(40)
        )
        page.draw_line((25, 125), (510, 125), color=(0.7, 0.7, 0.7))
        page.insert_text((25, 165), "FIXED BODY", fontsize=12)
        document.save(pdf)

    def fallback(*args):
        """有可靠文字层的合成 PDF 不允许回退视觉 OCR。"""
        raise AssertionError("unexpected OCR fallback")

    def quiet(*args):
        """忽略测试进度。"""

    rebuild_pdf(pdf, source, Event(), quiet, fallback)
    package = TemplatePackage(source)
    inventory = package.inventory()["nodes"]
    pairs = [
        ("OLD NAME", "name"),
        ("OLD PHONE", "phone"),
        ("OLD EMAIL", "email"),
        ("OLD ROLE", "job_title"),
        ("OLD CUSTOM", "custom:Source item"),
    ]
    fields = [
        TextBinding(
            node=next(
                row["id"] for row in inventory if row["kind"] == "p" and quote in row["text"]
            ),
            quote=quote,
            target="personal." + key,
        )
        for quote, key in pairs
    ]
    photos = []
    for row in inventory:
        if row["kind"] == "image":
            node = package.node(row["id"])
            anchor = next(node.iterancestors(f"{{{WP}}}anchor"), None)
            if (
                anchor is not None
                and anchor.find(f"{{{WP}}}docPr").get("descr") == "PDF 原图或照片"
            ):
                photos.append(row["id"])
    assert len(photos) == 1
    claimed = {field.node for field in fields} | set(photos)
    plan = TemplatePlan(
        summary="PDF 顶部资料",
        fields=fields,
        repeats=[],
        photos=photos,
        keep=[
            row["id"]
            for row in inventory
            if row["kind"] in {"p", "image"} and row["id"] not in claimed
        ],
        remove=[],
        warnings=[],
    )
    assert package.review(plan)["ready"]
    return source, plan


def header_content(**personal):
    """构造有新增字段、长网址和照片的真实填充资料。"""
    return ResumeDocument.model_validate(
        {
            "personal": {
                "name": "NEW NAME",
                "phone": "10000000000",
                "email": "long.address@example.test",
                "job_title": "Software engineer",
                "gender": "女",
                "age": "28",
                "gpa": "4.2/5.0",
                "location": "Example city",
                "website": "https://example.test/portfolio",
                "photo": resume_content().personal.photo,
                **personal,
            },
            "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
        }
    )


@pytest.mark.parametrize("left", [False, True])
@pytest.mark.parametrize(
    "hidden", [[], ["email", "job_title"], ["photo"], ["phone", "email", "job_title", "photo"]]
)
def test_fields_icons_and_photo_reflow_together(tmp_path, left, hidden):
    """左右照片、隐藏联系方式和新增资料都在独立容器中流动，不残留空图标或旧头像偏移。"""
    source, plan = header_source(tmp_path, left)
    original, original_plan = source.read_bytes(), deepcopy(plan)
    content = header_content(hidden_fields=hidden)
    output = tmp_path / "filled.docx"
    notices = fill_template(source, output, plan, content.model_dump(), [])
    document = Document(output)
    text = "".join(node.text or "" for node in document._element.iter(w("t")))
    assert "NEW NAME" in text and "FIXED BODY" in text
    name = next(
        node
        for node in document.tables[0]._element.iter(w("p"))
        if paragraph_text(node) == "NEW NAME"
    )
    assert name.find(f"{w('pPr')}/{w('jc')}").get(w("val")) == "center"
    assert not any(value in text for value in ("OLD", "〔待填写〕", "Source item"))
    assert not any(
        key.startswith(f"{{{PDF}}}") for node in document._element.iter() for key in node.attrib
    )
    assert not list(document.tables[0]._element.iter(f"{{{WP}}}anchor"))
    # 电话、邮箱、角色图标各属于自己的字段；缺值的来源自定义字段连同图标消失。
    expected = sum(key not in hidden for key in ("phone", "email", "job_title", "photo"))
    assert len(list(document.tables[0]._element.iter(f"{{{WP}}}inline"))) == expected
    if "photo" not in hidden:
        cells = document.tables[0].rows[0].cells
        photo_cell = cells[0 if left else -1]
        assert not photo_cell.text.strip()
        assert len(list(photo_cell._tc.iter(f"{{{WP}}}inline"))) == 1
    for key in ("phone", "email", "job_title"):
        assert (getattr(content.personal, key) in text) == (key not in hidden)
    assert any("重排 PDF 顶部" in note for note in notices)
    assert source.read_bytes() == original and plan == original_plan


def test_long_values_use_wrapping_cells_and_never_fixed_height(tmp_path):
    """超长姓名、意向、邮箱、网址和新增自定义字段都不超出分配宽度或挤入照片列。"""
    source, plan = header_source(tmp_path)
    content = header_content(
        name="A much longer bilingual name 中英文姓名",
        job_title="Senior software engineer and document layout specialist " * 2,
        email="a.very.long.contact.address.for.testing@example.test",
        website="https://example.test/" + "portfolio/" * 35,
        custom_fields=[
            {"id": str(i), "label": f"Custom {i}", "value": "Editable value " * (i + 1)}
            for i in range(6)
        ],
    )
    output = tmp_path / "long.docx"
    fill_template(source, output, plan, content.model_dump(), [])
    doc = Document(output)
    table = doc.tables[0]
    assert len(table.rows[0].cells) == 2
    assert not table._element.xpath(".//w:trHeight[@w:hRule='exact'] | .//w:noWrap")
    available = int(table.rows[0].cells[0]._tc.tcPr.tcW.w)
    for nested in table.rows[0].cells[0].tables:
        assert sum(int(x.get(w("w"))) for x in nested._tbl.tblGrid) == available
    text = "".join(node.text or "" for node in doc._element.iter(w("t")))
    assert all(
        value in text
        for value in (
            content.personal.name,
            content.personal.email,
            content.personal.website,
            "Custom 5",
        )
    )


def test_same_source_reacts_to_value_and_visibility_changes_without_persisting_layout(tmp_path):
    """重复试填重新按当前值计算，长短字段和显隐切换不会复用旧行宽或永久删除图标。"""
    source, plan = header_source(tmp_path)
    output = tmp_path / "result.docx"
    fill_template(
        source, output, plan, header_content(hidden_fields=["email", "photo"]).model_dump(), []
    )
    first = Document(output)
    fill_template(source, output, plan, header_content(email="short@test.dev").model_dump(), [])
    second = Document(output)
    assert len(list(first.tables[0]._element.iter(f"{{{WP}}}inline"))) == 2
    assert len(list(second.tables[0]._element.iter(f"{{{WP}}}inline"))) == 4
    assert "short@test.dev" in "".join(second._element.itertext())


def test_unrelated_word_layout_is_not_rewritten(tmp_path):
    """普通 Word 即使同样有姓名、电话和表格，也不触发任何 PDF 顶部重排。"""
    path = tmp_path / "native.docx"
    doc = Document()
    doc.add_paragraph("Old name")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Old phone"
    doc.save(path)
    package = TemplatePackage(path)
    before = etree.tostring(package.parts["word/document.xml"])
    layout = PDFHeaderLayout(package, None, {})
    layout.apply()
    assert not layout.fields and not layout.nodes and not layout.notices
    assert etree.tostring(package.parts["word/document.xml"]) == before


def test_mixed_pdf_does_not_reflow_unmarked_scanned_header(tmp_path):
    """后续原生页的文档标记不能误触发首页扫描恢复区的重排。"""
    path = tmp_path / "mixed.docx"
    doc = Document()
    doc._element.set(SOURCE, "1")
    doc.add_paragraph("Old name")
    doc.add_paragraph("Old phone")
    doc.save(path)
    package = TemplatePackage(path)
    nodes = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="扫描首页",
        fields=[
            TextBinding(node=nodes[0]["id"], quote="Old name", target="personal.name"),
            TextBinding(node=nodes[1]["id"], quote="Old phone", target="personal.phone"),
        ],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    before = etree.tostring(package.parts["word/document.xml"])
    layout = PDFHeaderLayout(
        package, plan, {"personal.name": "New name", "personal.phone": "10000000000"}
    )
    layout.apply()
    assert not layout.nodes
    assert etree.tostring(package.parts["word/document.xml"]) == before


@pytest.mark.parametrize("source_kind", ["1", "image-v1"])
def test_long_pdf_section_title_expands_background_and_reserves_spacing(tmp_path, source_kind):
    """复制的长栏目标题底块随文字扩展，并为下方没有段前空白的正文预留高度。"""
    from docx.shared import Pt

    from resume_maker.integrations.word.pdf_assets import attach_asset
    from resume_maker.integrations.word.pdf_titles import fit_pdf_titles

    path = tmp_path / "title.docx"
    doc = Document()
    doc._element.set(SOURCE, source_kind)
    heading = doc.add_paragraph("A longer editable section title")
    heading.runs[0].font.size = Pt(12)
    heading.paragraph_format.left_indent = Pt(8)
    heading.paragraph_format.line_spacing = Pt(16)
    heading.paragraph_format.space_before = Pt(10)
    attach_asset(
        heading,
        photo_bytes(35),
        pymupdf.Rect(0, 20, 65, 43),
        20,
        0,
        "PDF 固定装饰（图标、底色或线条）",
        1,
    )
    doc.add_paragraph("First body paragraph has no reserved space")
    doc.save(path)
    package = TemplatePackage(path)
    node = next(
        node
        for node in package.nodes.values()
        if node.tag == w("p") and paragraph_text(node).startswith("A longer")
    )
    field = TextBinding(
        node=package.ids[node], quote=paragraph_text(node), target="section-title:Long title"
    )
    fit_pdf_titles(package, [field])
    anchor = next(node.iter(f"{{{WP}}}anchor"))
    assert int(anchor.find(f"{{{WP}}}extent").get("cx")) > 65 * 12700
    spacing = node.find(f"{w('pPr')}/{w('spacing')}")
    assert int(spacing.get(w("after"))) >= 10 * 20
    assert spacing.get(w("lineRule")) == "atLeast"


def test_ambiguous_legacy_icon_binding_reports_problem_instead_of_guessing(tmp_path):
    """同段多个字段没有来源几何时报告可操作问题，不能默默把图标全挂到第一个字段。"""
    source, plan = header_source(tmp_path)
    package = TemplatePackage(source)
    # 验证新 PDF 已持久化来源标记；后续保存不会仅靠文字内容猜测模板种类。
    assert package.parts["word/document.xml"].get(SOURCE) == "1"
    assert any(node.get(TEXT) for node in package.nodes.values())
    field = next(f for f in plan.fields if f.target == "personal.phone")
    node = package.node(field.node)
    other = next(f for f in plan.fields if f.target == "personal.email")
    if other.node != field.node:
        pytest.skip("当前转换器为两个字段建立了独立段落，无歧义需要模拟。")
    from resume_maker.integrations.word.pdf_geometry import BOX

    for item in package.parts["word/document.xml"].iter():
        item.attrib.pop(BOX, None)
        item.attrib.pop(TEXT, None)
    with pytest.raises(Problem, match="位置证据"):
        PDFHeaderLayout(package, plan, {f.target: "visible" for f in plan.fields})
    assert paragraph_text(node)
