"""项目基本信息的显隐、自定义条目、版本隔离及两种 Word 版式回归。"""

from copy import deepcopy

import pytest
from conftest import experience
from docx import Document
from pydantic import ValidationError
from test_jobs import FakeProvider, wait_job
from test_template_project_slots import metadata_template, plan_for, project_document

from resume_maker.domain.experience import replace_field
from resume_maker.domain.models import Experience, ResumeItem
from resume_maker.infrastructure.database import Database, uid
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import paragraph_text
from resume_maker.services.catalog import Catalog
from resume_maker.services.jobs import Jobs


def body_text(path):
    """按段落拼接文字，标签与内容允许分属不同字重的文字运行。"""
    return "\n".join(paragraph_text(node) for node in Document(path).element.body.iter(w("p")))


def project_info():
    """固定项和自定义项各含可见与隐藏内容，便于验证恢复不丢原值。"""
    return {
        **experience("项目标题原值"),
        "period": "2025.01–2026.02",
        "role": "角色原值",
        "hidden_fields": ["role", "stack"],
        "custom_fields": [
            {
                "id": "link",
                "label": "项目链接",
                "value": "https://example.test/project",
                "visible": True,
            },
            {"id": "team", "label": "团队", "value": "隐藏团队原值", "visible": False},
            {"id": "blank", "label": "空白条目", "value": " ", "visible": True},
        ],
    }


def test_project_info_drafts_publish_and_preserve_pinned_versions(catalog, project, populated):
    """资料随版本发布并在重启后保留，简历引用的旧版本及隐藏原文不改变。"""
    base, identifier = populated["id"], project["id"]
    pinned = catalog.save_resume(
        "旧引用", None, [ResumeItem(project_id=identifier, revision_id=base, highlight_ids=["one"])]
    )
    value = project_info()
    meta = {key: item for key, item in value.items() if key != "highlights"}
    working = catalog.put_draft(identifier, base, "meta", meta, 0)
    assert working["content"] == value
    saved = catalog.save_revision(identifier, base, base)
    reopened = Catalog(Database(catalog.db.path))
    assert reopened.working(identifier, saved["id"])["content"] == value
    assert reopened.revision(base)["content"] == populated["content"]
    assert (
        reopened.db.one("SELECT * FROM resumes WHERE id=?", (pinned["id"],))["items"][0][
            "revision_id"
        ]
        == base
    )
    restored = replace_field(value, "meta", {"hidden_fields": [], "custom_fields": []})
    assert restored["role"] == "角色原值" and restored["stack"] == ["Python"]
    assert restored["custom_fields"] == []
    legacy = replace_field(value, "experience", experience("旧格式建议"))
    assert legacy["custom_fields"] == value["custom_fields"]
    assert legacy["hidden_fields"] == value["hidden_fields"]


def test_project_info_validation_and_old_revision_defaults():
    """旧版本自动获得空设置，未知显隐字段和重复自定义标识必须拒绝。"""
    old = Experience.model_validate(experience())
    assert old.custom_fields == old.hidden_fields == []
    with pytest.raises(ValidationError):
        Experience.model_validate({**experience(), "hidden_fields": ["highlights"]})
    field = {"id": "same", "label": "网址", "value": "示例"}
    with pytest.raises(ValidationError, match="不能重复"):
        Experience.model_validate({**experience(), "custom_fields": [field, field]})


@pytest.mark.parametrize("layout", ["builtin", "body", "cell", "row", "details"])
@pytest.mark.parametrize("hide_all", [False, True])
def test_project_info_export_visibility_and_custom_fallback(tmp_path, layout, hide_all):
    """内置、综合正文和独立字段模板都保留新增项；隐藏字段不泄漏且原始资料不变。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    value = project_info()
    if hide_all:
        value["hidden_fields"] = ["title", "period", "role", "stack", "description"]
    projects = [{"content": value, "highlight_ids": ["one"]}]
    original = deepcopy(projects)
    document = project_document().model_dump()
    if layout == "builtin":
        write_full_resume(output, document, projects)
    else:
        if layout == "details":
            doc = Document()
            for text in ("Title", "Period", "Details"):
                doc.add_paragraph(text)
            doc.save(source)
            plan = plan_for(
                source,
                [
                    ("Title", "Title", "title"),
                    ("Period", "Period", "period"),
                    ("Details", "Details", "details"),
                ],
            )
        else:
            plan = metadata_template(source, layout, 7)
        before, plan_before = source.read_bytes(), plan.model_dump()
        fill_template(source, output, plan, document, projects)
        assert source.read_bytes() == before and plan.model_dump() == plan_before
    text = body_text(output)
    assert text.count("项目链接：https://example.test/project") == 1
    assert "角色原值" not in text and "Python" not in text and "隐藏团队原值" not in text
    assert "空白条目" not in text and "Export Word" not in text
    assert "Parse documents" in text
    for field in ("title", "period", "description"):
        assert (value[field] in text) is not hide_all
    assert projects == original
    value["hidden_fields"] = []
    value["custom_fields"][1]["visible"] = True
    if layout == "builtin":
        write_full_resume(output, document, projects)
    else:
        fill_template(source, output, plan, document, projects)
    text = body_text(output)
    assert all(term in text for term in ("角色原值", "Python", "团队：隐藏团队原值"))


class ResettingProvider(FakeProvider):
    """模拟重新分析时未返回用户新增字段或显隐设置的模型。"""

    def run(self, **kwargs):
        """正常产生分析建议后清空扩展资料，验证服务会保留用户设置。"""
        result = super().run(**kwargs)
        result.experience.hidden_fields = []
        result.experience.custom_fields = []
        return result


def test_reanalysis_preserves_user_defined_project_info(catalog, project, populated, tmp_path):
    """采用 AI 整段建议时新增条目与显隐设置不被模型默认值覆盖。"""
    base, identifier = populated["id"], project["id"]
    value = project_info()
    value["body_order"] = ["custom:link", "highlights", "role"]
    catalog.put_draft(identifier, base, "experience", value, 0)
    jobs = Jobs(catalog.db, catalog, tmp_path / "data", ResettingProvider())
    conversation = catalog.db.all("SELECT * FROM conversations")[0]
    jobs.start()
    try:
        job = jobs.submit(conversation["id"], "重新分析", "analysis", base, "all", uid())
        assert wait_job(catalog, job["id"])["status"] == "completed"
        proposal = catalog.db.all("SELECT * FROM proposals")[0]
        result = catalog.adopt(proposal["id"])["content"]
        assert result["description"] == "New analysis"
        assert result["custom_fields"] == value["custom_fields"]
        assert result["hidden_fields"] == value["hidden_fields"]
        assert result["body_order"] == value["body_order"]
    finally:
        jobs.stop()
