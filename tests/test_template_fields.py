"""超链接域的真实 DOCX 归一化、格式保留、重复读取及其他动态域边界"""

from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from test_template_analysis import TemplateProvider, completed, simple_document

from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.services.templates.tasks import Templates


def add_complex(paragraph, codes, text, *, separate=True, close=True):
    """构造被 Word 拆分的域指令和带粗体样式的显示结果"""
    etree.SubElement(paragraph.add_run()._r, w("fldChar")).set(w("fldCharType"), "begin")
    for code in codes:
        etree.SubElement(paragraph.add_run()._r, w("instrText")).text = code
    if separate:
        etree.SubElement(paragraph.add_run()._r, w("fldChar")).set(w("fldCharType"), "separate")
    if text:
        paragraph.add_run(text).bold = True
    if close:
        etree.SubElement(paragraph.add_run()._r, w("fldChar")).set(w("fldCharType"), "end")


@pytest.mark.parametrize("simple", [True, False])
@pytest.mark.parametrize("part", ["body", "header", "footer"])
def test_hyperlink_fields_preserve_display_style_and_allow_filling(tmp_path, simple, part):
    """邮箱域在各文字部件可直接识别填充；原文件、相邻文字、字体和页码仍保留"""
    source, normalized, output = [
        tmp_path / name for name in ("source.docx", "copy.docx", "out.docx")
    ]
    doc = Document()
    paragraph = (
        doc.add_paragraph() if part == "body" else getattr(doc.sections[0], part).paragraphs[0]
    )
    paragraph.add_run("邮箱：")
    code = ' HYPERLINK "mailto:old@example.test" \\h '
    if simple:
        field = etree.SubElement(paragraph._p, w("fldSimple"), {w("instr"): code})
        run = etree.SubElement(field, w("r"))
        etree.SubElement(etree.SubElement(run, w("rPr")), w("b"))
        etree.SubElement(run, w("t")).text = "旧邮箱"
    else:
        add_complex(paragraph, [" HY", 'PERLINK "mailto:old@example.test"', " \\h "], "旧邮箱")
    paragraph.add_run(" · 联系方式")
    page = etree.SubElement(doc.add_paragraph()._p, w("fldSimple"))
    page.set(w("instr"), " PAGE \\* MERGEFORMAT ")
    doc.save(source)
    source_bytes = source.read_bytes()
    package = TemplatePackage(source)
    assert not package.inventory()["warnings"]
    assert any("HYPERLINK" in message for message in package.notices)
    row = next(row for row in package.inventory()["nodes"] if "旧邮箱" in row["text"])
    plan = TemplatePlan(
        summary="替换邮箱",
        fields=[TextBinding(node=row["id"], quote="旧邮箱", target="personal.email")],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    content = simple_document()
    content.personal.name = ""
    content.personal.email = "new@example.test"
    assert package.review(plan)["ready"]
    package.write(normalized)
    assert TemplatePackage(normalized).inventory()["nodes"] == package.inventory()["nodes"]
    fill_template(source, output, plan, content.model_dump(), [])
    assert source.read_bytes() == source_bytes
    actual = TemplatePackage(output)
    root = actual.parts[row["part"]]
    assert (
        "".join(root.xpath(".//w:t/text()", namespaces=NS)) == "邮箱：new@example.test · 联系方式"
    )
    bold = root.xpath(".//w:r[w:rPr/w:b]/w:t/text()", namespaces=NS)
    assert bold == ["new@example.test"]
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        assert all(b"old@example.test" not in after.read(name) for name in after.namelist())
    assert actual.parts["word/document.xml"].xpath(".//w:fldSimple/@w:instr", namespaces=NS) == [
        " PAGE \\* MERGEFORMAT "
    ]


def test_hyperlink_import_reaches_analysis_and_keeps_snapshot_ids(catalog, tmp_path):
    """链接域无需人工修改即可进入识别；保存的归一化副本与映射节点编号一致"""
    source = tmp_path / "hyperlink.docx"
    doc = Document()
    add_complex(doc.add_paragraph(), [' HYPERLINK "https://example.test" '], "原姓名")
    doc.save(source)
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = service.analyze(source, simple_document())
    result = completed(service, task["id"])
    assert result["status"] == "completed" and result["review"]["ready"]
    assert len(provider.calls) == 1
    assert "HYPERLINK" in "".join(result["inventory"]["notices"])
    assert (
        TemplatePackage(service.source(task["id"])).inventory()["nodes"]
        == result["inventory"]["nodes"]
    )
    service.stop()


@pytest.mark.parametrize(
    "text,separate,close", [("", True, True), ("旧邮箱", False, True), ("旧邮箱", True, False)]
)
def test_incomplete_fields_keep_text_and_become_editable(tmp_path, text, separate, close):
    """不完整域保留可见文字并变为可编辑内容且不再因缺少指令边界阻止识别"""
    source = tmp_path / "incomplete.docx"
    doc = Document()
    add_complex(
        doc.add_paragraph(),
        ['HYPERLINK "mailto:old@example.test"'],
        text,
        separate=separate,
        close=close,
    )
    doc.save(source)
    package = TemplatePackage(source)
    root = package.parts["word/document.xml"]
    assert not root.xpath(".//w:instrText | .//w:fldChar", namespaces=NS)
    assert "".join(root.xpath(".//w:t/text()", namespaces=NS)) == text
    assert "动态域" in "".join(package.notices)


def test_split_page_fields_and_nested_dynamic_fields_are_distinguished(tmp_path):
    """拼接页码指令并接受格式开关；嵌套的其他域都冻结为原有显示结果"""
    source = tmp_path / "nested.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    add_complex(paragraph, [" PA", "GE \\* MERGEFORMAT "], "1")
    add_complex(paragraph, [' HYPERLINK "https://example.test" '], "链接", close=False)
    add_complex(paragraph, ["MERGE", "FIELD OldName"], "旧姓名")
    etree.SubElement(paragraph.add_run()._r, w("fldChar")).set(w("fldCharType"), "end")
    doc.save(source)
    package = TemplatePackage(source)
    warnings = "".join(package.inventory()["warnings"])
    assert not warnings
    assert "MERGEFIELD" in "".join(package.notices)
    root = package.parts["word/document.xml"]
    assert "".join(root.xpath(".//w:t/text()", namespaces=NS)) == "1链接旧姓名"
    assert len(root.xpath(".//w:fldChar", namespaces=NS)) == 3
