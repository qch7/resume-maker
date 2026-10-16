"""验证已规范化模板的内存克隆、副本隔离及补位解析次数"""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from resume_maker.core.errors import Problem
from resume_maker.integrations import document_limits
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates import mapping
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text
from tests.support.documents import add_annotation_part, photo_bytes
from tests.support.layouts import generic_content, generic_template, template_body_text


def test_clone_preserves_current_parts_resources_and_independent_indexes():
    """内存修改、规范化结果、包资源在克隆中完整保留，各副本编辑互不影响"""
    source = BytesIO()
    document = Document()
    paragraph = document.add_paragraph("Seed")
    for tag, text in (("ins", "Inserted"), ("del", "Deleted")):
        revision = etree.SubElement(paragraph._p, w(tag))
        etree.SubElement(etree.SubElement(revision, w("r")), w("t")).text = text
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    document.add_picture(BytesIO(photo_bytes(50)))
    add_annotation_part(document, "comment", "")
    add_annotation_part(
        document,
        "footnote",
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p>'
        '</w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>Note</w:t></w:r></w:p></w:footnote>',
    )
    document.save(source)
    package = TemplatePackage(source)
    assert "word/comments.xml" not in package.files
    root = package.parts["word/document.xml"]
    assert not list(root.iter(w("ins"), w("del")))
    assert "Inserted" in paragraph_text(next(root.iter(w("p"))))
    assert package.notices

    next(root.iter(w("t"))).text = "Edited in memory"
    added = etree.Element(w("p"))
    etree.SubElement(etree.SubElement(added, w("r")), w("t")).text = "New paragraph"
    root.find(w("body")).insert(1, added)
    original_nodes = dict(package.nodes)
    original_files = dict(package.files)
    original_xml = {name: etree.tostring(part) for name, part in package.parts.items()}
    original_notices = list(package.notices)
    buffer = BytesIO()
    package.write(buffer)
    restored = TemplatePackage(buffer)
    cloned, sibling = package.clone(), package.clone()

    with ZipFile(buffer) as archive:
        assert cloned.files == {name: archive.read(name) for name in archive.namelist()}
    assert cloned.locations == restored.locations
    assert cloned.inventory()["nodes"] == restored.inventory()["nodes"]
    assert cloned.inventory()["warnings"] == restored.inventory()["warnings"]
    assert cloned.notices == original_notices
    for identifier, node in cloned.nodes.items():
        assert cloned.ids[node] == identifier
        assert etree.tostring(node) == etree.tostring(restored.node(identifier))
        assert node is not sibling.node(identifier)
    media = next(name for name in package.files if name.startswith("word/media/"))
    assert cloned.files[media] is package.files[media]
    assert cloned.files["word/styles.xml"] is package.files["word/styles.xml"]

    cloned.parts["word/document.xml"].find(".//" + w("t")).text = "Clone only"
    cloned.files[media] = b"replacement in clone"
    cloned.notices.append("克隆独立提示")
    cloned.nodes.clear()
    assert package.nodes == original_nodes
    assert package.files == original_files
    assert {name: etree.tostring(part) for name, part in package.parts.items()} == original_xml
    assert package.notices == original_notices
    assert sibling.notices == original_notices
    assert sibling.inventory()["nodes"] == restored.inventory()["nodes"]
    package.files[media] = b"replacement in original"
    assert sibling.files[media] == original_files[media]


def test_clone_keeps_package_size_limit(monkeypatch):
    """内存修改后的包仍受原解压大小边界约束"""
    source = BytesIO()
    Document().save(source)
    package = TemplatePackage(source)
    monkeypatch.setattr(document_limits, "DOCX_MAX_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(Problem, match="100 MB"):
        package.clone()


@pytest.mark.parametrize("layout", ["body", "rows", "cells"])
def test_mixed_completion_normalizes_only_the_external_package(monkeypatch, layout):
    """同时补栏目、个人资料只解析外部输入一次，补位不生成中间 ZIP"""
    source, output = BytesIO(), BytesIO()
    _, plan = generic_template(source, layout)
    document, projects = generic_content()
    original_source, original_plan = source.getvalue(), deepcopy(plan)
    prepare, write = mapping.prepare_parts, TemplatePackage.write
    preparations, outputs = [], []

    def counted_prepare(files):
        """保留外部输入的真实规范化并统计执行次数"""
        preparations.append(1)
        return prepare(files)

    def counted_write(package, target):
        """记录真正的 ZIP 输出，补位副本不应触发此操作"""
        outputs.append(target)
        return write(package, target)

    monkeypatch.setattr(mapping, "prepare_parts", counted_prepare)
    monkeypatch.setattr(TemplatePackage, "write", counted_write)
    notices = fill_template(source, output, plan, document.model_dump(), projects)
    assert len(preparations) == 1
    assert outputs == [output]
    assert len(notices) >= 3
    text = template_body_text(output)
    assert document.personal.website in text
    assert "Repository：https://repo0.example.test" in text
    assert "Award 1" in text and "Date 1" in text and "Issuer：Institute 1" in text
    assert "Old " not in text and "〔" not in text
    assert source.getvalue() == original_source and plan == original_plan


@pytest.mark.parametrize("layout", ["body", "rows", "cells"])
def test_mixed_completion_matches_zip_roundtrip(monkeypatch, layout):
    """不同容器经过新旧克隆路径补位后，生成包的每个文件保持一致"""
    source, result, baseline = BytesIO(), BytesIO(), BytesIO()
    _, plan = generic_template(source, layout)
    document, projects = generic_content()
    fill_template(source, result, plan, document.model_dump(), projects)

    def serialized_clone(package):
        """还原优化前的内存 ZIP 往返作为行为基线"""
        buffer = BytesIO()
        package.write(buffer)
        return TemplatePackage(buffer)

    monkeypatch.setattr(TemplatePackage, "clone", serialized_clone)
    fill_template(source, baseline, plan, document.model_dump(), projects)
    with ZipFile(result) as actual, ZipFile(baseline) as expected:
        assert {name: actual.read(name) for name in actual.namelist()} == {
            name: expected.read(name) for name in expected.namelist()
        }
