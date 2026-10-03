"""陌生模板的完整填充、样式保留、旧信息清理和错误映射边界"""

from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import NS, TemplatePackage, paragraph_text
from resume_maker.integrations.word.templates.values import missing_targets
from tests.support.documents import make_template, photo_bytes, project_content, resume_content


def test_complete_fill_keeps_layout_and_removes_old_data(tmp_path):
    """整份替换所有容器、重复记录和照片，同时保留样式、布局和包资源"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    package, plan = make_template(source)
    assert package.review(plan)["ready"]
    fill_template(source, output, plan, resume_content().model_dump(), project_content())
    actual = TemplatePackage(output)
    text = "".join(paragraph_text(node) for node in actual.nodes.values() if node.tag == w("p"))
    for expected in [
        "测试新姓名",
        "新大学一",
        "新大学二",
        "new@example.test",
        "10000000000",
        "测试城市",
        "文档开发工程师",
        "新文档项目",
        "覆盖姓名和教育",
        "语言：中文",
    ]:
        assert expected in text
    assert "旧" not in text and "不应输出" not in text
    new_name = next(
        node
        for node in actual.nodes.values()
        if node.tag == w("p") and paragraph_text(node).startswith("测试新姓名")
    )
    assert new_name.find("w:r/w:rPr/w:b", NS) is not None
    assert len(actual.parts["word/document.xml"].xpath(".//w:tr", namespaces=NS)) == 2
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        assert before.read("word/theme/theme1.xml") == after.read("word/theme/theme1.xml")
        assert "word/media/image1.png" not in after.namelist()
        assert photo_bytes(20) not in [after.read(name) for name in after.namelist()]
        assert photo_bytes(180) in [after.read(name) for name in after.namelist()]
    assert package.review(plan)["ready"]


def test_hidden_rows_fields_and_photos_do_not_reappear(tmp_path):
    """隐藏资料时清空原位置并安全移除空重复区的全部表格行"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    _, plan = make_template(source)
    document = resume_content()
    document.personal.hidden_fields = ["phone", "photo"]
    document.sections[0].visible = False
    document.sections[1].visible = False
    fill_template(source, output, plan, document.model_dump(), project_content())
    actual = TemplatePackage(output)
    all_text = "".join(
        root.xpath(".//w:t/text()", namespaces=NS)[0] for root in actual.parts.values()
    )
    assert "旧" not in all_text and "10000000000" not in all_text
    assert not actual.parts["word/document.xml"].xpath(".//w:tbl", namespaces=NS)
    assert not any(name.startswith("word/media/") for name in actual.files)
    Document(output)


@pytest.mark.parametrize(
    "case", ["quote", "overlap", "sample", "unknown", "keep_remove", "photo_duplicate"]
)
def test_invalid_mapping_is_rejected(tmp_path, case):
    """错误位置、引文、范围和相互矛盾的分类不能产生导出文件"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source)
    if case == "quote":
        plan.fields[0].quote = "不存在的文字"
    elif case == "overlap":
        plan.fields.append(plan.repeats[0].fields[0].model_copy(update={"target": "personal.name"}))
    elif case == "sample":
        plan.repeats[0].sample_start = plan.repeats[0].sample_end = plan.repeats[1].start
    elif case == "unknown":
        plan.fields[0].target = "personal.password"
    elif case == "keep_remove":
        plan.remove.append(plan.keep[0])
    else:
        plan.photos *= 2
    assert package.review(plan)["errors"]
    with pytest.raises(Problem):
        fill_template(source, tmp_path / "bad.docx", plan, resume_content().model_dump(), [])
    assert not (tmp_path / "bad.docx").exists()


def test_unresolved_and_missing_current_information_block_export(tmp_path):
    """原文未处理或当前新填资料没有位置时，明确阻止遗漏信息的导出"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source)
    plan.remove = []
    assert package.review(plan)["unresolved"][0]["text"] == "旧多余说明"
    document = resume_content()
    document.personal.website = "https://example.test"
    assert "personal.website" in missing_targets(document, plan, project_content())
    document.sections[0].title = "新的教育名称"
    assert "栏目：新的教育名称" in missing_targets(document, plan, project_content())


@pytest.mark.parametrize("encoded_path", [False, True])
def test_sample_decoration_requires_explicit_confirmation(tmp_path, encoded_path):
    """重复样本图片需要确认，仍被装饰引用的同一图片资源必须保留"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source, decoration=True)
    media = "word/media/image1.png"
    if encoded_path:
        media = "word/media/image 1.png"
        package.files[media] = package.files.pop("word/media/image1.png")
        path = "word/_rels/document.xml.rels"
        package.files[path] = package.files[path].replace(
            b'Target="media/image1.png"', b'Target="/word/media/image%201.png"'
        )
        package.write(source)
    assert package.image(plan.photos[0]) == photo_bytes(20)
    assert "图片" in "".join(package.review(plan)["errors"])
    decoration = next(
        row["id"]
        for row in package.inventory()["nodes"]
        if row["kind"] == "image" and row["id"] not in plan.photos
    )
    plan.keep.append(decoration)
    assert package.review(plan)["ready"]
    output = tmp_path / "filled.docx"
    fill_template(source, output, plan, resume_content().model_dump(), project_content())
    with ZipFile(output) as archive:
        assert archive.read(media) == photo_bytes(20)
    actual = TemplatePackage(output)
    assert len([row for row in actual.inventory()["nodes"] if row["kind"] == "image"]) == 3


def test_empty_insertion_and_occurrence_rules(tmp_path):
    """空引文不能覆盖非空原文，同段多处相同值按指定次数独立替换"""
    path = tmp_path / "repeat.docx"
    doc = Document()
    doc.add_paragraph("示例 / 示例")
    doc.add_paragraph()
    doc.save(path)
    package = TemplatePackage(path)
    node = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p")
    plan = TemplatePlan(
        summary="测试",
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
        fields=[
            TextBinding(node=node, quote="示例", target="personal.name", occurrence=1),
            TextBinding(node=node, quote="示例", target="personal.phone", occurrence=2),
        ],
    )
    document = ResumeDocument(
        sections=[{"id": "p", "kind": "projects", "title": "项目"}],
        personal={"name": "甲\n乙", "phone": "电话值"},
    )
    output = tmp_path / "output.docx"
    fill_template(path, output, plan, document.model_dump(), [])
    assert Document(output).paragraphs[0].text == "甲\n乙 / 电话值"
    from resume_maker.domain.resume import CustomInfoField

    document.personal.custom_fields = [CustomInfoField(id="lang", label="语言", value="中文")]
    empty = next(row["id"] for row in package.inventory()["nodes"] if row["can_insert"])
    plan.fields.append(TextBinding(node=empty, quote="", target="personal.custom:语言"))
    fill_template(path, output, plan, document.model_dump(), [])
    assert Document(output).paragraphs[1].text == "语言：中文"
    plan.fields[0].quote = ""
    assert "空引文" in "".join(package.review(plan)["errors"])


def test_empty_text_does_not_make_a_photo_paragraph_an_insertion_slot(tmp_path):
    """含图片或文本框的无文字段落禁止用作字段空位"""
    package, plan = make_template(tmp_path / "source.docx")
    image = next(row for row in package.inventory()["nodes"] if row["kind"] == "image")
    owner = next(row for row in package.inventory()["nodes"] if row["id"] == image["ancestors"][0])
    assert not owner["text"] and not owner["can_insert"]
    plan.fields[-1].node = owner["id"]
    assert "空白的段落" in "".join(package.review(plan)["errors"])


def test_dynamic_fields_are_frozen_without_restoring_old_values(tmp_path):
    """无缓存的合并域自动变为可编辑空位，指令不会恢复旧值"""
    path = tmp_path / "dynamic.docx"
    doc = Document()
    paragraph = doc.add_paragraph("姓名")
    field = OxmlElement("w:fldSimple")
    field.set(w("instr"), "MERGEFIELD OldName")
    paragraph._p.append(field)
    doc.save(path)
    package = TemplatePackage(path)
    assert not package.inventory()["warnings"]
    assert "动态域" in "".join(package.notices)
    assert not package.parts["word/document.xml"].xpath(".//w:fldSimple", namespaces=NS)


def test_only_replaced_hyperlinks_are_removed(tmp_path):
    """替换个人主页后不再跳转旧地址，同段明确保留的固定链接仍可使用"""
    source, output = tmp_path / "links.docx", tmp_path / "result.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    for text, target in [
        ("旧个人主页", "https://old.example.test"),
        ("固定说明", "https://fixed.example.test"),
    ]:
        link = OxmlElement("w:hyperlink")
        link.set(
            f"{{{NS['r']}}}id",
            doc.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
        )
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
    doc.save(source)
    package = TemplatePackage(source)
    node_id = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p")
    plan = TemplatePlan(
        summary="主页",
        warnings=[],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        fields=[
            TextBinding(node=node_id, quote="旧个人主页", target="personal.website"),
        ],
    )
    document = ResumeDocument(
        personal={"website": "https://new.example.test"},
        sections=[
            {"id": "p", "title": "项目", "kind": "projects"},
        ],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    actual = TemplatePackage(output)
    links = actual.parts["word/document.xml"].xpath(".//w:hyperlink", namespaces=NS)
    assert len(links) == 1 and "".join(links[0].itertext()) == "固定说明"
    assert b"old.example.test" not in actual.files["word/_rels/document.xml.rels"]
    assert b"fixed.example.test" in actual.files["word/_rels/document.xml.rels"]
