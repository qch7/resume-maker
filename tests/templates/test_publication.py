"""模板补位在短暂文件占用时有限重试，发布失败保留原件及可恢复副本"""

from pathlib import Path

import pytest
from docx import Document

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.infrastructure import filesystem
from resume_maker.integrations.word.templates.completion import complete_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.integrations.word.templates.supplement import supplement_personal_fields

PUBLISHERS = [complete_template, supplement_personal_fields]


def personal_template(tmp_path):
    """构造只缺少所在地位置的合成模板，两个补位入口都需要发布副本"""
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Old Name")
    document.save(source)
    package = TemplatePackage(source)
    node = next(row for row in package.inventory()["nodes"] if row["text"] == "Old Name")
    plan = TemplatePlan(
        summary="合成补位模板",
        fields=[TextBinding(node=node["id"], quote="Old Name", target="personal.name")],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    content = ResumeDocument(
        personal={"name": "New Name", "location": "Synthetic City"},
        sections=[{"id": "projects", "kind": "projects", "title": "projects"}],
    )
    return source, package, plan, content


@pytest.mark.parametrize("publisher", PUBLISHERS)
def test_template_publication_survives_temporary_windows_denial(tmp_path, monkeypatch, publisher):
    """短暂拒绝期间保留原模板，释放后一次发布可重新读取及核对的补位结果"""
    source, package, plan, content = personal_template(tmp_path)
    before = source.read_bytes()
    original = Path.replace
    attempts = []

    def occupied(temporary, target):
        """模拟 Windows 对完整 DOCX 副本的前两次替换暂时拒绝"""
        assert target == source and source.read_bytes() == before
        assert TemplatePackage(temporary).inventory()["nodes"]
        attempts.append(temporary)
        if len(attempts) <= 2:
            error = PermissionError("合成的临时文件占用")
            error.winerror = 5
            raise error
        return original(temporary, target)

    monkeypatch.setattr(Path, "replace", occupied)
    updated, mapping, notices = publisher(package, plan, content, [], source)
    assert notices and updated.review(mapping)["ready"]
    assert TemplatePackage(source).review(mapping)["ready"]
    assert any(field.target == "personal.location" for field in mapping.fields)
    assert source.read_bytes() != before
    assert 3 <= len(attempts) <= 7
    assert not attempts[-1].exists()
    assert len(plan.fields) == 1 and package.review(plan)["ready"]


@pytest.mark.parametrize("publisher", PUBLISHERS)
@pytest.mark.parametrize("code", [5, 13, None])
def test_failed_template_publication_preserves_original(tmp_path, monkeypatch, publisher, code):
    """持续 Windows 占用有限重试，其他权限错误立即抛出并保留旧模板、暂存副本"""
    source, package, plan, content = personal_template(tmp_path)
    before = source.read_bytes()
    attempts, waits = [], []
    error = PermissionError("合成的持久文件访问拒绝")
    if code is not None:
        error.winerror = code

    def denied(temporary, target):
        """持续阻止发布，核对每次尝试都没有破坏原模板"""
        assert target == source and source.read_bytes() == before
        attempts.append(temporary)
        raise error

    monkeypatch.setattr(Path, "replace", denied)
    monkeypatch.setattr(filesystem.time, "sleep", waits.append)
    with pytest.raises(PermissionError) as result:
        publisher(package, plan, content, [], source)
    assert result.value is error
    assert source.read_bytes() == before
    assert len(attempts) == (7 if code == 5 else 1)
    assert len(waits) == len(attempts) - 1
    assert TemplatePackage(attempts[-1]).inventory()["nodes"]
    assert len(plan.fields) == 1 and package.review(plan)["ready"]
