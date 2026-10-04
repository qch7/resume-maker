"""模板库的完整映射、分类、缩略图和持久化"""

from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.core.errors import Problem
from resume_maker.domain.templates import TemplatePlan
from resume_maker.infrastructure.database import dump, now
from resume_maker.plugins.queries import templates as query_templates
from resume_maker.runtime.host import Contribution
from resume_maker.services.resumes import Resumes
from resume_maker.services.template_records import TemplateRecords
from resume_maker.services.templates.library import TemplateLibrary
from tests.support.data import workspace
from tests.support.document_services import Documents, ResumePreviews
from tests.support.documents import resume_content
from tests.support.templates import (
    TemplateProvider,
    completed,
    register_template,
    simple_document,
    simple_template,
)


def test_library_persists_and_category_deletion_keeps_likes(tmp_path):
    """收藏和分类跨实例保留，删除分类不删除模板或取消收藏"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library").status_code == 401
        assert client.get("/api/template-library", headers=headers).json() == {
            "categories": [],
            "items": {},
            "templates": [],
        }
        register_template(client.app.state.services.catalog, config.data_dir)
        result = client.post(
            "/api/template-library/categories", headers=headers, json={"name": "技术"}
        )
        category = result.json()["categories"][0]["id"]
        client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"category_id": category}
        )
        state = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"liked": True}
        ).json()
        assert state["items"]["mapped"] == {"category_id": category, "liked": True}
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library", headers=headers).json() == state
        state = client.delete(
            f"/api/template-library/categories/{category}", headers=headers
        ).json()
        assert state["items"]["mapped"] == {"category_id": "", "liked": True}
        assert client.app.state.services.resume.template("mapped")["name"] == "测试模板"
        assert not client.app.state.services.db.all("SELECT * FROM resumes")


def test_library_validation_and_independent_updates(catalog, tmp_path):
    """新分类校验空白和重名并发修改分类和收藏互不覆盖"""
    service = TemplateLibrary(
        Resumes(catalog, storage=catalog.db),
        tmp_path / "data",
        storage=catalog.db,
        records=TemplateRecords(),
    )
    category = service.create_category(" 技术 ")["categories"][0]["id"]
    for name in (" ", "技术", "未分类", "全部模板", "我的喜欢", "回收站"):
        with pytest.raises(Problem):
            service.create_category(name)
    with pytest.raises(Problem):
        service.update("missing", {"liked": True})
    with pytest.raises(Problem):
        service.update("builtin", {"category_id": "missing"})
    with ThreadPoolExecutor(max_workers=2) as pool:
        operations = [
            pool.submit(service.update, "builtin", changes)
            for changes in ({"category_id": category}, {"liked": True})
        ]
        for operation in operations:
            operation.result()
    assert service.state()["items"]["builtin"] == {"category_id": category, "liked": True}


def test_template_rename_persists_without_changing_references(tmp_path):
    """已使用模板可改名，名称跨重启保留，模板文件、映射和所有简历数据不变"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        services = client.app.state.services
        source = register_template(services.catalog, config.data_dir)
        original = source.read_bytes()
        record = services.resume.template("mapped")
        category = services.template_library.create_category("技术")["categories"][0]["id"]
        services.template_library.update("mapped", {"category_id": category, "liked": True})
        for name in ("第一份简历", "第二份简历"):
            assert (
                client.post(
                    "/api/resumes",
                    headers=headers,
                    json={"name": name, "template_id": "mapped", "items": []},
                ).status_code
                == 200
            )
        resumes = services.db.all("SELECT * FROM resumes")
        assert (
            client.patch(
                "/api/template-library/items/mapped", json={"name": "未授权修改"}
            ).status_code
            == 401
        )
        response = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"name": "  新名称  "}
        )
        assert response.status_code == 200
        state = response.json()
        assert state["templates"][0]["name"] == "新名称"
        assert state["templates"][0]["usage_count"] == 2
        assert state["items"]["mapped"] == {"category_id": category, "liked": True}
        assert services.resume.template("mapped") == {**record, "name": "新名称"}
        assert source.read_bytes() == original
        assert services.db.all("SELECT * FROM resumes") == resumes
        assert (
            client.delete("/api/template-library/items/mapped", headers=headers).status_code == 409
        )
    with TestClient(create_app(config)) as client:
        assert client.get("/api/template-library", headers=headers).json() == state
        assert client.get("/api/state", headers=headers).json()["templates"][0]["name"] == "新名称"


def test_template_rename_validation_is_atomic(tmp_path):
    """非法名称、内置模板及回收站拒绝改名，联合修改失败时不部分保存"""
    config = Config(data_dir=tmp_path / "data", token="test")
    headers = {"x-resume-token": "test"}
    with TestClient(create_app(config)) as client:
        services = client.app.state.services
        register_template(services.catalog, config.data_dir)
        before = services.template_library.state()
        for identifier, changes, status in (
            ("mapped", {"name": ""}, 422),
            ("mapped", {"name": None}, 422),
            ("mapped", {"name": " "}, 400),
            ("mapped", {"name": "字" * 201}, 422),
            ("builtin", {"name": "新内置名称"}, 409),
            ("missing", {"name": "不存在"}, 404),
            ("mapped", {"name": "不应保存", "category_id": "missing", "liked": True}, 409),
        ):
            result = client.patch(
                f"/api/template-library/items/{identifier}", headers=headers, json=changes
            )
            assert result.status_code == status, result.text
            assert services.template_library.state() == before
        recycled = services.template_library.delete("mapped")
        result = client.patch(
            "/api/template-library/items/mapped", headers=headers, json={"name": "不应保存"}
        )
        assert result.status_code == 409
        assert services.template_library.state() == recycled


def test_thumbnail_cache_and_original_are_isolated(catalog, tmp_path, monkeypatch):
    """重复及并发读取只渲染一次，原始文件和数据库模板映射不变"""
    data_dir = tmp_path / "data"
    source = register_template(catalog, data_dir)
    original = source.read_bytes()
    record = Resumes(catalog, storage=catalog.db).template("mapped")
    calls = []

    def render(path, pdf):
        """只模拟排版输出，同时检查传入的是独立源模板快照"""
        assert path != source and path.read_bytes() == original
        calls.append(path)
        (pdf.parent / "page-1.png").write_bytes(b"thumbnail")
        return 1, None

    monkeypatch.setattr("resume_maker.services.templates.library.render_word", render)
    service = TemplateLibrary(
        Resumes(catalog, storage=catalog.db),
        data_dir,
        storage=catalog.db,
        records=TemplateRecords(),
    )
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(service.thumbnail, ["mapped", "mapped"]))
    assert results[0] == results[1]
    assert results[0].read_bytes() == b"thumbnail"
    assert len(calls) == 1
    assert (
        TemplateLibrary(
            Resumes(catalog, storage=catalog.db),
            data_dir,
            storage=catalog.db,
            records=TemplateRecords(),
        ).thumbnail("mapped")
        == results[0]
    )
    assert (
        source.read_bytes() == original
        and Resumes(catalog, storage=catalog.db).template("mapped") == record
    )
    source.write_bytes(b"changed outside app")
    with pytest.raises(Problem, match="程序外变化"):
        service.thumbnail("mapped")


def test_thumbnail_failure_can_retry_and_builtin_is_real_docx(catalog, tmp_path, monkeypatch):
    """渲染失败允许重试，内置模板也用实际 DOCX 排版器生成示例预览"""
    from docx import Document

    calls = []

    def render(path, pdf):
        """首次模拟失败，重试验证真实示例文档后提供图片"""
        calls.append(path)
        assert "你的姓名" in "".join(
            p.text
            for table in Document(path).tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        )
        if len(calls) == 1:
            return None, "Word 暂不可用"
        (pdf.parent / "page-1.png").write_bytes(b"png")
        return 1, None

    monkeypatch.setattr("resume_maker.services.templates.library.render_word", render)
    service = TemplateLibrary(
        Resumes(catalog, storage=catalog.db),
        tmp_path / "data",
        storage=catalog.db,
        records=TemplateRecords(),
    )
    with pytest.raises(Problem, match="Word 暂不可用"):
        service.thumbnail("builtin")
    assert service.thumbnail("builtin").is_file()
    assert len(calls) == 2


def test_only_complete_templates_can_be_selected(catalog, tmp_path):
    """缺少完整映射的记录保留原数据并提示重新识别"""
    data = tmp_path / "data"
    register_template(catalog, data)
    with catalog.db.transaction() as conn:
        conn.execute(
            "INSERT INTO templates VALUES (?,?,?,?,?)",
            ("unavailable", "无完整映射", "unused", dump({}), now()),
        )
    library = workspace(
        catalog,
        contributors=lambda: [
            Contribution("ext.template-library", "workspace.queries", "templates", query_templates)
        ],
    ).state()["templates"]
    assert [item["id"] for item in library] == ["mapped"]
    assert all(set(item) == {"id", "name", "created_at"} for item in library)
    document = resume_content()
    with pytest.raises(Problem, match="AI 识别"):
        Resumes(catalog, storage=catalog.db).save_resume(
            "缺失映射", "unavailable", [], document=document
        )
    previews = ResumePreviews(Resumes(catalog, storage=catalog.db), data)
    with pytest.raises(Problem, match="AI 识别"):
        previews.render("unavailable", document.model_dump(), [])
    resume = Resumes(catalog, storage=catalog.db).save_resume(
        "完整方案", "mapped", [], document=document
    )
    with catalog.db.transaction() as conn:
        conn.execute("UPDATE resumes SET template_id='unavailable' WHERE id=?", (resume["id"],))
    with pytest.raises(Problem, match="AI 识别"):
        Documents(Resumes(catalog, storage=catalog.db), data, storage=catalog.db).export(
            resume["id"]
        )
    assert catalog.db.one("SELECT mapping_json FROM templates WHERE id='unavailable'") == {
        "mapping": {}
    }
    assert not catalog.db.all("SELECT * FROM exports")
    assert previews.directory is None
    previews.stop()


def test_saved_recognition_history_survives_restart_and_resave(tmp_path):
    """保存并重启后保留真实活动、轮次、耗时和用量，人工另存不会清空历史或再调用 AI"""

    class RecordedProvider(TemplateProvider):
        """为真实分析流程提供确定的公开活动和统计"""

        def run_structured(self, **kwargs):
            """发出可识别的活动，再沿用脱敏模板的结构化建议"""
            kwargs["emit"]("activity", {"text": "已定位姓名与联系方式 api_key=private"})
            kwargs["emit"](
                "usage", {"input_tokens": 321, "cached_input_tokens": 120, "output_tokens": 65}
            )
            return super().run_structured(**kwargs)

    config = Config(data_dir=tmp_path / "data", token="test")
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = RecordedProvider()
    keys = (
        "events",
        "cursor",
        "round",
        "elapsed_ms",
        "usage",
        "metrics",
        "attempts",
        "repair_error",
        "reused",
        "phase",
    )
    with TestClient(create_app(config, provider)) as client:
        service = client.app.state.services.templates
        task = completed(service, service.analyze(source, simple_document())["id"])
        assert task["events"] and task["round"] > 0
        history = {key: task[key] for key in keys}
        assert history["usage"]["input_tokens"] == 321
        assert "private" not in dump(history)
        saved = service.save(
            task["id"], "保存记录", TemplatePlan.model_validate(task["plan"]), simple_document(), []
        )
        assert saved["mapping"]["analysis"] == history
        assert saved["mapping"]["importer"] == task["importer"]

    restarted_provider = TemplateProvider(failure=True)
    with TestClient(create_app(config, restarted_provider)) as client:
        service = client.app.state.services.templates
        opened = service.open(saved["id"])
        assert opened["from_library"] and opened["status"] == "completed"
        assert opened["id"] != task["id"]
        assert {key: opened[key] for key in keys} == history
        assert opened["importer"] == task["importer"]
        plan = TemplatePlan.model_validate(opened["plan"])
        plan.summary = "已人工核对"
        resaved = service.save(opened["id"], "人工另存", plan, simple_document(), [])
        assert resaved["mapping"]["analysis"] == history
        assert resaved["mapping"]["importer"] == task["importer"]
        opened["events"].clear()
        assert service.get(opened["id"])["events"] == history["events"]
        assert not restarted_provider.calls

    with TestClient(create_app(config, restarted_provider)) as client:
        opened = client.post(
            f"/api/templates/{resaved['id']}/edit", headers={"x-resume-token": "test"}
        ).json()
        assert {key: opened[key] for key in keys} == history
        assert opened["plan"]["summary"] == "已人工核对"
