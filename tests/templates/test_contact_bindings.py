"""验证同排联系方式误合并时保留原布局，歧义映射继续交给用户确认"""

import pytest
from docx import Document

from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.word.templates.contact_bindings import separate_contact_bindings
from resume_maker.integrations.word.templates.contact_style import SLOT_VERSION
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage


def contact_plan(fields):
    """构造仅包含联系方式的完整模板映射"""
    return TemplatePlan(
        summary="联系信息", fields=fields, repeats=[], photos=[], keep=[], remove=[], warnings=[]
    )


@pytest.mark.parametrize("supplemented", [False, True])
@pytest.mark.parametrize("reverse", [False, True])
def test_split_combined_contacts_preserves_one_line_and_source(tmp_path, supplemented, reverse):
    """双标签误合并后原位拆分，旧自动补位可回收，原件及输入映射保持不变"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    values = ["电话：13800001234", "邮箱：old@example.test"]
    if reverse:
        values.reverse()
    doc.add_paragraph("\t".join(values))
    if supplemented:
        doc.add_paragraph("邮箱：〔待填写〕")._p.set(SLOT_VERSION, "1")
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [node for node in package.inventory()["nodes"] if node["kind"] == "p"]
    quote = "".join(values)
    plan = contact_plan(
        [TextBinding(node=paragraphs[0]["id"], quote=quote, target="personal.phone")]
        + (
            [TextBinding(node=paragraphs[1]["id"], quote="〔待填写〕", target="personal.email")]
            if supplemented
            else []
        ),
    )
    before, mapping = source.read_bytes(), plan.model_dump()
    document = {
        "personal": {"phone": "10000000000", "email": "new@example.test"},
        "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
    }
    fill_template(source, output, plan, document, [])
    lines = [paragraph.text for paragraph in Document(output).paragraphs]
    expected = ["电话：10000000000", "邮箱：new@example.test"]
    if reverse:
        expected.reverse()
    assert lines == ["\t".join(expected)]
    assert source.read_bytes() == before
    assert plan.model_dump() == mapping


def test_ambiguous_contact_values_and_existing_real_binding_are_not_rewritten(tmp_path):
    """未能验证的联系方式、用户独立绑定位置均不能由规则猜测替换"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("电话：13800001234\t邮箱：old@example.test")
    doc.add_paragraph("独立邮箱：other@example.test")
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [node for node in package.inventory()["nodes"] if node["kind"] == "p"]
    plan = contact_plan(
        [
            TextBinding(
                node=paragraphs[0]["id"],
                quote="13800001234邮箱：old@example.test",
                target="personal.phone",
            ),
            TextBinding(
                node=paragraphs[1]["id"], quote="other@example.test", target="personal.email"
            ),
        ]
    )
    result, updated, notices = separate_contact_bindings(package, plan)
    assert result is package and updated is plan and notices == []
    plan.fields = plan.fields[:1]
    plan.fields[0].quote = "13800001234"
    result, updated, notices = separate_contact_bindings(package, plan)
    assert result is package and updated is plan and notices == []
