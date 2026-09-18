"""新增大栏目与普通条目字段自动沿用模板，预览不依赖再次识别。"""

from copy import deepcopy

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from lxml import etree
from test_template_order import body_text, layout_content, simple_template

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import ResumeDocument, ResumeSection, SectionEntry
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_flow import effective_section
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text


def new_section(title="实习经历"):
    """生成含多条记录、空元数据及自定义字段的新栏目。"""
    return ResumeSection(
        id=title,
        title=title,
        entries=[
            SectionEntry(
                id="first",
                title="示例公司",
                subtitle="研发实习生",
                period="2026.06—2026.08",
                details="完成接口开发\n参与性能优化",
                custom_fields=[{"id": "team", "label": "团队", "value": "平台组"}],
            ),
            SectionEntry(id="second", details="第二条经历"),
        ],
    )


@pytest.mark.parametrize("table", [False, True])
@pytest.mark.parametrize("kept", [False, True])
def test_new_sections_reuse_headings_and_follow_current_order(tmp_path, table, kept):
    """正文和表格均复制大标题样式，多个新增栏目随顺序、隐藏、恢复及清空同步变化。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = simple_template(source, table=table)
    if kept:
        plan.keep.extend(field.node for field in plan.fields)
        plan.fields = []
    original, original_plan = source.read_bytes(), deepcopy(plan)
    content = layout_content()
    internship, volunteer = new_section(), new_section("志愿服务")
    volunteer.entries = [SectionEntry(id="volunteer", title="社区志愿者")]
    content.sections.insert(1, internship)
    content.sections.append(volunteer)
    fill_template(source, output, plan, content.model_dump(), [])
    text = body_text(output)
    assert text.splitlines() == [
        "固定开头",
        "实习经历",
        "示例公司",
        "研发实习生",
        "2026.06—2026.08",
        "完成接口开发",
        "参与性能优化",
        "团队：平台组",
        "第二条经历",
        "教育背景",
        "教育条目",
        "主修课程",
        "课程条目",
        "志愿服务",
        "社区志愿者",
        "固定结尾",
    ]
    paragraphs = list(Document(output).element.body.iter(w("p")))
    for title in ("教育背景", "实习经历", "志愿服务"):
        heading = next(node for node in paragraphs if paragraph_text(node) == title)
        style = heading.find("w:pPr/w:pStyle", NS)
        assert (style.get(w("val")) if style is not None else None) == (
            None if table else "Heading1"
        )
    assert all(paragraph_text(node) for node in paragraphs)
    if table:
        assert len(Document(output).tables[0].rows) == 9
    internship.visible = False
    volunteer.entries = []
    fill_template(source, output, plan, content.model_dump(), [])
    assert all(value not in body_text(output) for value in ("实习经历", "示例公司", "志愿服务"))
    internship.visible = True
    content.sections.remove(internship)
    content.sections.append(internship)
    fill_template(source, output, plan, content.model_dump(), [])
    assert body_text(output).index("课程条目") < body_text(output).index("实习经历")
    assert body_text(output).endswith("第二条经历\n固定结尾")
    assert source.read_bytes() == original and plan == original_plan


def graphic_template(path):
    """建立含蓝色文本框标题、底纹与独立联系方式的脱敏图形模板。"""
    doc = Document()
    doc.add_paragraph("邮箱：旧邮箱")
    heading = doc.add_paragraph()
    heading._p.append(
        etree.fromstring(
            f'<w:r xmlns:w="{NS["w"]}" xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:pict><v:shape id="heading-shape" fillcolor="#718fb9" '
            'style="width:420pt;height:24pt"><v:textbox><w:txbxContent>'
            '<w:p><w:pPr><w:shd w:fill="ECECEC"/></w:pPr><w:r><w:rPr>'
            '<w:rFonts w:ascii="Arial" w:eastAsia="黑体"/><w:b/><w:color w:val="FFFFFF"/>'
            '<w:sz w:val="28"/></w:rPr><w:t>荣誉证书</w:t></w:r></w:p>'
            "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
        )
    )
    for text in ("旧荣誉一", "旧荣誉二"):
        run = doc.add_paragraph(text).runs[0]
        run.font.name = "宋体"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string("333333")
    doc.add_paragraph("固定结尾")
    doc.save(path)
    package = TemplatePackage(path)
    nodes = {paragraph_text(node): key for key, node in package.nodes.items() if node.tag == w("p")}
    return TemplatePlan(
        summary="图形标题与普通条目扩展",
        fields=[
            TextBinding(node=nodes["邮箱：旧邮箱"], quote="旧邮箱", target="personal.email"),
            TextBinding(node=nodes["荣誉证书"], quote="荣誉证书", target="section-title:荣誉证书"),
        ],
        repeats=[
            RepeatBinding(
                section="荣誉证书",
                start=nodes["旧荣誉一"],
                end=nodes["旧荣誉二"],
                sample_start=nodes["旧荣誉一"],
                sample_end=nodes["旧荣誉一"],
                fields=[TextBinding(node=nodes["旧荣誉一"], quote="旧荣誉一", target="details")],
            )
        ],
        photos=[],
        keep=[nodes["固定结尾"]],
        remove=[],
        warnings=[],
    )


def test_graphic_heading_and_new_honor_fields_expand_together(tmp_path):
    """一次补齐主页、荣誉元数据及实习大栏目，完整复用图形，空字段不产生占位或空行。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = graphic_template(source)
    original, original_plan = source.read_bytes(), deepcopy(plan)
    honors = new_section("荣誉证书")
    content = ResumeDocument(
        personal={"email": "test@example.test", "website": "https://example.test"},
        sections=[
            new_section(),
            honors,
            ResumeSection(id="projects", title="项目经历", kind="projects"),
        ],
    )
    fill_template(source, output, plan, content.model_dump(), [])
    body = Document(output).element.body
    # 标签和值可具有不同字重；按可见段落连接，不能把运行边界误当成换行。
    text = "\n".join(paragraph_text(node) for node in body.iter(w("p")))
    assert (
        text.index("个人主页：https://example.test")
        < text.index("实习经历")
        < text.index("荣誉证书")
    )
    for value in ("示例公司", "研发实习生", "2026.06—2026.08", "第二条经历", "团队：平台组"):
        assert text.count(value) == 2
    assert "旧荣誉" not in text and "〔自动条目占位〕" not in text
    for node in body.iter(w("p")):
        assert paragraph_text(node) or node.find(".//w:pict", NS) is not None
    shapes = body.xpath(".//*[local-name()='shape']")
    assert len(shapes) == 2 and len({shape.get("id") for shape in shapes}) == 2
    for shape in shapes:
        assert shape.get("fillcolor") == "#718fb9"
        assert shape.get("style") == "width:420pt;height:24pt"
        assert shape.find(".//w:pPr/w:shd", NS).get(w("fill")) == "ECECEC"
        assert shape.find(".//w:rPr/w:sz", NS).get(w("val")) == "28"
        assert shape.find(".//w:rPr/w:color", NS).get(w("val")) == "FFFFFF"
    title = next(node for node in body.iter(w("p")) if paragraph_text(node) == "示例公司")
    assert title.find(".//w:rPr/w:sz", NS).get(w("val")) == "22"
    assert source.read_bytes() == original and plan == original_plan


def test_new_section_still_requires_a_reusable_heading_and_reviewed_template(tmp_path):
    """没有大标题样本或原映射不完整时仍给出明确错误，不能静默丢弃资料。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = simple_template(source)
    plan.keep.extend(field.node for field in plan.fields)
    plan.fields = []
    plan.repeats = []
    package = TemplatePackage(source)
    plan.keep = [key for key, node in package.nodes.items() if node.tag == w("p")]
    content = ResumeDocument(
        sections=[new_section(), ResumeSection(id="projects", title="项目经历", kind="projects")]
    )
    with pytest.raises(Problem, match="栏目：实习经历"):
        fill_template(source, output, plan, content.model_dump(), [])
    plan.keep = []
    with pytest.raises(Problem, match="映射尚未完成"):
        fill_template(source, output, plan, content.model_dump(), [])
    assert not output.exists()


def test_empty_added_fields_do_not_leave_blank_table_rows(tmp_path):
    """表格内已有栏目补出元数据和正文后，记录未填写的新增字段不残留空表格行。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = simple_template(source, table=True)
    content = layout_content()
    education = content.sections[-1]
    education.entries = new_section().entries
    education.entries[1].title = "第二家公司"
    education.entries[1].details = ""
    fill_template(source, output, plan, content.model_dump(), [])
    rows = Document(output).tables[0].rows
    assert [row.cells[0].text for row in rows] == [
        "教育背景",
        "示例公司",
        "研发实习生\n2026.06—2026.08",
        "完成接口开发\n参与性能优化\n团队：平台组",
        "第二家公司",
        "主修课程",
        "课程条目",
    ]


@pytest.mark.parametrize("table", [False, True])
def test_major_sections_continue_without_inherited_page_breaks(tmp_path, table):
    """已有和新增栏目连续排版，消除硬分页、样式分页和末节默认分页，保留列宽及区外设置。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    opening = doc.add_paragraph("固定开头")
    opening.paragraph_format.page_break_before = True
    container = doc.add_table(rows=0, cols=1) if table else doc
    for title in ("教育背景", "专业技能"):
        heading = container.add_row().cells[0].paragraphs[0] if table else doc.add_paragraph()
        heading.style = "Heading 1"
        heading.paragraph_format.page_break_before = True
        heading.add_run().add_break(WD_BREAK.PAGE)
        heading.add_run(title)
        body = container.add_row().cells[0].paragraphs[0] if table else doc.add_paragraph()
        body.add_run("旧" + title).add_break(WD_BREAK.LINE)
        if not table and title == "教育背景":
            doc.add_section(WD_SECTION_START.NEW_PAGE)
    # 最终节缺少 type 也意味着另起一页，不能只处理标题段落里的显式分页。
    final = doc.sections[-1]._sectPr
    kind = final.find(w("type"))
    if kind is not None:
        final.remove(kind)
    final.find(w("cols")).set(w("num"), "2")
    doc.add_paragraph("固定结尾").paragraph_format.page_break_before = True
    doc.save(source)
    package = TemplatePackage(source)
    nodes = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    regions = []
    for title in ("教育背景", "专业技能"):
        node = package.node(nodes["旧" + title])
        root = next(node.iterancestors(w("tr"))) if table else node
        identifier = package.ids[root]
        regions.append(
            RepeatBinding(
                section=title,
                start=identifier,
                end=identifier,
                sample_start=identifier,
                sample_end=identifier,
                fields=[TextBinding(node=nodes["旧" + title], quote="旧" + title, target="title")],
            )
        )
    plan = TemplatePlan(
        summary="带强制分页的栏目样本",
        fields=[
            TextBinding(node=nodes[title], quote=title, target=f"section-title:{title}")
            for title in ("教育背景", "专业技能")
        ],
        repeats=regions,
        photos=[],
        keep=[nodes["固定开头"], nodes["固定结尾"]],
        remove=[],
        warnings=[],
    )
    original, original_plan = source.read_bytes(), deepcopy(plan)
    content = ResumeDocument(
        sections=[
            ResumeSection(
                id="education",
                title="教育背景",
                entries=[SectionEntry(id="school", title="示例大学")],
            ),
            ResumeSection(
                id="skills", title="专业技能", entries=[SectionEntry(id="skill", title="Python")]
            ),
            new_section(),
            ResumeSection(id="projects", title="项目经历", kind="projects"),
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    filled = Document(output)
    text = body_text(output)
    assert text.index("教育背景") < text.index("专业技能") < text.index("实习经历")
    for node in etree.fromstring(etree.tostring(filled.element.body)).iter(w("p")):
        before = node.find("w:pPr/w:pageBreakBefore", NS)
        if paragraph_text(node) in {"固定开头", "固定结尾"}:
            assert before is not None and before.get(w("val")) != "0"
        elif paragraph_text(node):
            assert before is not None and before.get(w("val")) == "0"
            assert effective_section(node).find(w("type")).get(w("val")) == "continuous"
        if paragraph_text(node) in {"教育背景", "专业技能", "实习经历"} and not table:
            assert node.find("w:pPr/w:keepNext", NS) is not None
    assert not filled.element.body.xpath(".//w:br[@w:type='page']")
    assert filled.element.body.xpath(".//w:br[not(@w:type) or @w:type='textWrapping']")
    assert filled.sections[-1]._sectPr.find(w("cols")).get(w("num")) == "2"
    assert source.read_bytes() == original and plan == original_plan
