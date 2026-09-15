"""注释文字、备用绘图及连续分节的真实 DOCX 填充验证。"""

from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from test_template_annotations import add_annotation_part

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage


def empty_plan(**kwargs):
    """提供无预设分类的完整方案，测试只声明需要的替换。"""
    return TemplatePlan(
        **(
            {
                "summary": "测试",
                "fields": [],
                "repeats": [],
                "photos": [],
                "keep": [],
                "remove": [],
                "warnings": [],
            }
            | kwargs
        )
    )


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_note_text_is_replaced_and_comments_are_removed_from_copy(tmp_path, kind):
    """正文与脚注尾注同时替换，批注及其引用清理，原文件完全保留。"""
    doc = Document()
    paragraph = doc.add_paragraph("原姓名")
    reference = etree.SubElement(paragraph.add_run()._r, w(f"{kind}Reference"))
    reference.set(w("id"), "1")
    etree.SubElement(paragraph._p, w("commentRangeStart")).set(w("id"), "0")
    add_annotation_part(
        doc, kind, f'<w:{kind} w:id="1"><w:p><w:r><w:t>原电话</w:t></w:r></w:p></w:{kind}>'
    )
    add_annotation_part(
        doc, "comment", '<w:comment w:id="0"><w:p><w:r><w:t>编辑批注</w:t></w:r></w:p></w:comment>'
    )
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc.save(source)
    original = source.read_bytes()
    package = TemplatePackage(source)
    fields = [
        TextBinding(
            node=n["id"],
            quote=n["text"],
            target="personal.name" if n["text"] == "原姓名" else "personal.phone",
        )
        for n in package.inventory()["nodes"]
        if n["text"] in {"原姓名", "原电话"}
    ]
    document = ResumeDocument(
        personal={"name": "新姓名", "phone": "10000000000"},
        sections=[{"id": "p", "kind": "projects", "title": "项目"}],
    )
    fill_template(source, output, empty_plan(fields=fields), document.model_dump(), [])
    assert source.read_bytes() == original
    with ZipFile(output) as archive:
        assert "word/comments.xml" not in archive.namelist()
        assert b"commentRange" not in archive.read("word/document.xml")
        note = etree.fromstring(archive.read(f"word/{kind}s.xml"))
        assert note.xpath(".//w:t/text()", namespaces=NS) == ["10000000000"]
        assert "comments" not in archive.read("word/_rels/document.xml.rels").decode()
    assert Document(output).paragraphs[0].text == "新姓名"


@pytest.mark.parametrize("requires,expected", [("wps", "现代表示"), ("future", "备用表示")])
def test_word_alternate_representation_is_not_counted_twice(tmp_path, requires, expected):
    """Word 的同一内容只选择一种有效表示，避免 AI 和待处理清单重复计算。"""
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph._p.append(
        etree.fromstring(
            (
                "<mc:AlternateContent "
                'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                f'xmlns:w="{NS["w"]}"><mc:Choice Requires="{requires}">'
                "<w:r><w:t>现代表示</w:t></w:r></mc:Choice>"
                "<mc:Fallback><w:r><w:t>备用表示</w:t></w:r></mc:Fallback>"
                "</mc:AlternateContent>"
            ).encode()
        )
    )
    source = tmp_path / "source.docx"
    doc.save(source)
    package = TemplatePackage(source)
    assert [n["text"] for n in package.inventory()["nodes"] if n["kind"] == "p"] == [expected]
    assert not package.inventory()["notices"]


@pytest.mark.parametrize("count", [0, 2])
def test_repeated_experience_preserves_continuous_column_boundaries(tmp_path, count):
    """多栏标题加单栏正文按条目重复，清空或增加经历不会改变后续正文分栏。"""
    doc = Document()
    title = doc.add_paragraph("原项目")
    section = etree.SubElement(title._p.get_or_add_pPr(), w("sectPr"))
    etree.SubElement(section, w("type")).set(w("val"), "continuous")
    etree.SubElement(section, w("cols")).set(w("num"), "4")
    doc.add_paragraph("原正文")
    doc.add_paragraph("后续固定内容")
    final = doc.sections[-1]._sectPr
    etree.SubElement(final, w("type")).set(w("val"), "continuous")
    final.find(w("cols")).set(w("num"), "1")
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc.save(source)
    package = TemplatePackage(source)
    rows = [n for n in package.inventory()["nodes"] if n["kind"] == "p"]
    plan = empty_plan(
        keep=[rows[2]["id"]],
        repeats=[
            RepeatBinding(
                section="projects",
                start=rows[0]["id"],
                end=rows[1]["id"],
                sample_start=rows[0]["id"],
                sample_end=rows[1]["id"],
                fields=[
                    TextBinding(node=rows[0]["id"], quote="原项目", target="title"),
                    TextBinding(node=rows[1]["id"], quote="原正文", target="description"),
                ],
            )
        ],
    )
    projects = [
        {
            "highlight_ids": [],
            "content": {
                "title": f"新项目{i}",
                "description": f"新正文{i}",
                "period": "",
                "role": "",
                "stack": [],
                "highlights": [],
            },
        }
        for i in range(count)
    ]
    document = ResumeDocument(sections=[{"id": "p", "kind": "projects", "title": "项目"}])
    assert package.review(plan)["ready"]
    fill_template(source, output, plan, document.model_dump(), projects)
    result = Document(output)
    texts = result.element.xpath(".//w:t/text()")
    assert "原项目" not in texts and "原正文" not in texts
    assert texts[-1] == "后续固定内容"
    assert sum(text.startswith("新项目") for text in texts) == count
    columns = [s.find(w("cols")).get(w("num")) for s in result.element.iter(w("sectPr"))]
    assert columns == (["4", "1"] * count + ["1"] if count else ["4", "1"])
    assert not result.tables
