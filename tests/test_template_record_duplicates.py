"""重复样本须表示一条记录，标量字段多位置不能悄悄复制整条内容"""

import json
import threading

import pytest
from docx import Document
from provider_stub import ProviderStub

from resume_maker.core.errors import Problem
from resume_maker.domain.models import ProviderSettings
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.services.templates.analysis import analyze_plan


def repeated_sample(path, layout, target):
    """构造两条独立样例被误选为一条的段落、单元格或整行模板"""
    doc = Document()
    doc.add_paragraph("Research capabilities")
    if layout == "rows":
        table = doc.add_table(rows=2, cols=1)
        for row, text in zip(table.rows, ["Archived entry A", "Archived entry B"], strict=True):
            row.cells[0].text = text
    else:
        area = doc.add_table(rows=1, cols=1).cell(0, 0) if layout == "cell" else doc
        for text in ["Archived entry A", "Archived entry B"]:
            area.add_paragraph(text)
    doc.save(path)
    package = TemplatePackage(path)
    nodes = package.inventory()["nodes"]
    header = next(row for row in nodes if row["text"] == "Research capabilities")
    rows = [row for row in nodes if row["kind"] == "p" and row["text"].startswith("Archived")]
    boundaries = [row["ancestors"][0] if layout == "rows" else row["id"] for row in rows]
    plan = TemplatePlan(
        summary="两条样例",
        fields=[
            TextBinding(
                node=header["id"],
                quote=header["text"],
                target="section-title:Research capabilities",
            )
        ],
        repeats=[
            RepeatBinding(
                section="Research capabilities",
                start=boundaries[0],
                end=boundaries[-1],
                sample_start=boundaries[0],
                sample_end=boundaries[-1],
                fields=[
                    TextBinding(node=row["id"], quote=row["text"], target=target) for row in rows
                ],
            )
        ],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    document = ResumeDocument(
        sections=[
            {"id": "projects", "kind": "projects", "title": "Projects"},
            {
                "id": "research",
                "title": "Research capabilities",
                "entries": [{"id": str(i), target: f"Unique new entry {i}"} for i in range(2)],
            },
        ]
    )
    return package, plan, document


@pytest.mark.parametrize("layout", ["body", "cell", "rows"])
@pytest.mark.parametrize("target", ["details", "title"])
def test_duplicate_scalar_record_fields_block_export(tmp_path, layout, target):
    """多条样例合并后会让每条新记录出现两次，须在试填和导出前拒绝"""
    source, output = tmp_path / "original.docx", tmp_path / "filled.docx"
    package, plan, document = repeated_sample(source, layout, target)
    review = package.review(plan)
    assert not review["ready"]
    assert any("同一条记录" in error and target in error for error in review["errors"])
    with pytest.raises(Problem, match="同一条记录"):
        fill_template(source, output, plan, document.model_dump(), [])
    assert not output.exists()


def test_multiple_paragraphs_with_distinct_fields_remain_valid(tmp_path):
    """一条记录本来可含标题和正文，不能仅因样本占多个段落就拒绝"""
    package, plan, _ = repeated_sample(tmp_path / "original.docx", "body", "title")
    plan.repeats[0].fields[1].target = "details"
    assert package.review(plan)["ready"]


def test_repeated_personal_fields_are_not_record_slots(tmp_path):
    """个人信息在正文和页眉重复显示不属于重复记录的错误边界"""
    source = tmp_path / "original.docx"
    doc = Document()
    doc.add_paragraph("Old name")
    doc.sections[0].header.paragraphs[0].text = "Old name"
    doc.save(source)
    package = TemplatePackage(source)
    bindings = [
        TextBinding(node=r["id"], quote=r["text"], target="personal.name")
        for r in package.inventory()["nodes"]
        if r["kind"] == "p" and r["text"] == "Old name"
    ]
    assert len(package.validate_fields(bindings)) == 2


@pytest.mark.parametrize("repair", [True, False])
@pytest.mark.parametrize("invalid", ["duplicate", "quote"])
def test_duplicate_record_feedback_is_bounded_and_reusable(tmp_path, monkeypatch, repair, invalid):
    """反馈明确后可缩小样本，持续返回错误方案时不得接受或无限重试"""
    source = tmp_path / "original.docx"
    package, bad, document = repeated_sample(source, "body", "details")
    good = bad.model_copy(deep=True)
    good.repeats[0].sample_end = good.repeats[0].sample_start
    good.repeats[0].fields = good.repeats[0].fields[:1]
    if invalid == "quote":
        bad = good.model_copy(deep=True)
        bad.repeats[0].fields[0].quote = "Incorrectly remembered source wording"
    monkeypatch.setattr(
        "resume_maker.services.templates.analysis.source_pages", lambda *_: ([], [], [])
    )

    class Provider(ProviderStub):
        """模拟先错误选择样本，再根据同一契约返回映射"""

        def __init__(self):
            """记录模拟请求"""
            self.calls = []

        def run_structured(self, **kwargs):
            """验证后续请求携带具体错误，其他映射保持原样"""
            request = json.loads(kwargs["prompt"].splitlines()[-1])
            if self.calls:
                expected_error = "同一条记录" if invalid == "duplicate" else "找不到引文"
                assert any(expected_error in e for e in request["validation"]["errors"])
                context = request["validation"]["node_context"]
                assert any(
                    row[4] == "Archived entry A"
                    for rows in context["parts"].values()
                    for row in rows
                )
            self.calls.append(request)
            return good if repair and len(self.calls) > 1 else bad

    provider = Provider()
    plan, review, attempts, _ = analyze_plan(
        package,
        provider,
        tmp_path,
        document,
        [],
        ProviderSettings(),
        threading.Event(),
        lambda *_: None,
    )
    assert 1 < attempts <= 3
    assert review["ready"] is repair
    if repair:
        output = tmp_path / "filled.docx"
        fill_template(source, output, plan, document.model_dump(), [])
        text = "\n".join(
            r["text"] for r in TemplatePackage(output).inventory()["nodes"] if r["kind"] == "p"
        )
        assert text.count("Unique new entry 0") == text.count("Unique new entry 1") == 1
        assert "Archived" not in text
