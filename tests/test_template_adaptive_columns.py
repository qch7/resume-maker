"""使用合成模板验证新增标签样式和长短记录混排"""

from copy import deepcopy

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt
from lxml import etree
from test_template_entry_layout import record_content, record_plan

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text
from resume_maker.integrations.word.templates.supplement import supplement_personal_fields


@pytest.mark.parametrize("gap", ["", " ", "\t", "\n"])
@pytest.mark.parametrize("period_first", [False, True])
def test_adjacent_metadata_gets_spacing_without_replacing_existing_separators(
    tmp_path, gap, period_first
):
    """源样本日期和名称相连时增加间距，保留已有换行、制表位和单空格"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    left, right = ("Old period", "Old title") if period_first else ("Old title", "Old period")
    paragraph = doc.add_paragraph(left + gap + right, "List Bullet")
    if gap == "\t":
        paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(360), WD_TAB_ALIGNMENT.RIGHT)
    doc.save(source)
    literal = left + (gap if gap == " " else "") + right
    plan = record_plan(source, [(literal, "Old title", "title"), (literal, "Old period", "period")])
    content = record_content([{"id": "one", "title": "Result", "period": "2026-09-20"}])
    fill_template(source, output, plan, content.model_dump(), [])
    result = Document(output).paragraphs[0].text
    expected = ("2026-09-20", "Result") if period_first else ("Result", "2026-09-20")
    assert result == expected[0] + (gap or "\u2002") + expected[1]


@pytest.mark.parametrize("gap", [" " * 70, "\u3000" * 12, "\t"])
@pytest.mark.parametrize("cell", [False, True])
def test_long_repeat_titles_wrap_independently_of_dates(tmp_path, gap, cell):
    """不同空白排版及容器中，长名称和完整日期占独立列，所有新增条目共用列宽"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    area = doc.add_table(rows=1, cols=1).cell(0, 0) if cell else doc
    header = area.add_paragraph()
    header.add_run("成果：").bold = True
    header.add_run("旧标题").bold = False
    header.add_run(gap)
    header.add_run("日期：").bold = True
    header.add_run("旧日期").bold = False
    if gap == "\t":
        header.paragraph_format.tab_stops.add_tab_stop(Pt(380), WD_TAB_ALIGNMENT.RIGHT)
    if cell:
        area._tc.remove(area.paragraphs[0]._p)
    doc.save(source)
    literal = "成果：旧标题" + ("" if gap == "\t" else gap) + "日期：旧日期"
    plan = record_plan(source, [(literal, "旧标题", "title"), (literal, "旧日期", "period")], cell)
    content = record_content(
        [
            {
                "id": "a",
                "title": "跨院校联合研究与创新成果展示评选活动一等奖" * 3,
                "period": "2026-06",
            },
            {"id": "b", "title": "短名称", "period": "2025-11-03"},
            {"id": "c", "title": "包含  双空格的名称", "period": "2026-09"},
        ]
    )
    before, previous = source.read_bytes(), deepcopy(plan)
    fill_template(source, output, plan, content.model_dump(), [])
    package = TemplatePackage(output)
    tables = [
        node
        for node in package.nodes.values()
        if node.tag == w("tbl") and node.find("w:tblPr/w:tblLayout", NS) is not None
    ]
    assert len(tables) == 3
    assert (
        len({tuple(node.xpath("w:tblGrid/w:gridCol/@w:w", namespaces=NS)) for node in tables}) == 1
    )
    for table, entry in zip(tables, content.sections[0].entries, strict=True):
        cells = table.findall("w:tr/w:tc", NS)
        assert len(cells) == 2
        assert "".join(cells[0].xpath(".//w:t/text()", namespaces=NS)) == "成果：" + entry.title
        assert "".join(cells[1].xpath(".//w:t/text()", namespaces=NS)) == "日期：" + entry.period
        assert not table.xpath(".//w:trHeight", namespaces=NS)
        for paragraph in table.iter(w("p")):
            runs = [
                run for run in paragraph.findall(w("r")) if run.xpath("w:t/text()", namespaces=NS)
            ]
            assert runs[0].find("w:rPr/w:b", NS) is not None
            assert runs[-1].find("w:rPr/w:b", NS).get(w("val")) == "0"
    assert source.read_bytes() == before and plan == previous


def test_narrow_container_stacks_metadata_without_losing_labels(tmp_path):
    """窄栏容不下完整时间范围时，每条记录统一纵排，固定说明和标签保留"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.width = Pt(120)
    cell.paragraphs[0].text = "Title: Old title        Period: Old period"
    doc.save(source)
    text = cell.paragraphs[0].text
    plan = record_plan(source, [(text, "Old title", "title"), (text, "Old period", "period")], True)
    content = record_content(
        [{"id": "a", "title": "Result", "period": "September 2025 – September 2026"}]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    cell = Document(output).tables[0].cell(0, 0)
    assert not cell.tables
    assert [p.text for p in cell.paragraphs] == [
        "Title: Result",
        "Period: September 2025 – September 2026",
    ]


def test_repeat_sample_uses_its_own_section_width(tmp_path):
    """范围起点在多栏旧记录而样本在后续单栏时，复制过程不能误判样本的可用栏宽"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.sections[0]._sectPr.find(w("cols")).set(w("num"), "2")
    section = doc.add_section(WD_SECTION_START.CONTINUOUS)
    section._sectPr.find(w("cols")).set(w("num"), "1")
    etree.SubElement(doc.paragraphs[0]._p.pPr.sectPr, w("type")).set(w("val"), "continuous")
    literal = "Title: Old title" + " " * 40 + "Period: Old period"
    doc.add_paragraph(literal)
    doc.save(source)
    plan = record_plan(source, [(literal, "Old title", "title"), (literal, "Old period", "period")])
    package = TemplatePackage(source)
    sample = package.node(plan.repeats[0].sample_start)
    plan.repeats[0].start = package.ids[sample.getprevious()]
    content = record_content([{"id": "a", "title": "A new result", "period": "2026-09-17"}])
    fill_template(source, output, plan, content.model_dump(), [])
    table = Document(output).tables[0]
    assert table.cell(0, 1).text == "Period: 2026-09-17"
    assert table.cell(0, 0).text == "Title: A new result"


def test_out_of_bounds_right_tab_is_replaced_even_for_short_values(tmp_path):
    """右制表位超出页面时即使内容很短也要修复"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Old title\tOld period")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(800), WD_TAB_ALIGNMENT.RIGHT)
    doc.save(source)
    plan = record_plan(
        source,
        [
            ("Old titleOld period", "Old title", "title"),
            ("Old titleOld period", "Old period", "period"),
        ],
    )
    content = record_content([{"id": "a", "title": "Short", "period": "2026"}])
    fill_template(source, output, plan, content.model_dump(), [])
    table = Document(output).tables[0]
    assert table.cell(0, 0).text == "Short" and table.cell(0, 1).text == "2026"
    assert sum(column.width for column in table.columns) <= Pt(432)


def test_generated_columns_respect_the_aligned_body_indent(tmp_path):
    """标题按正文缩进重新对齐后，列总宽同时收缩，右侧日期不会超出正文页面边界"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    literal = "Old title" + " " * 30 + "Old period"
    doc.add_paragraph(literal).paragraph_format.left_indent = Pt(3)
    doc.add_paragraph("Old body").paragraph_format.left_indent = Pt(37)
    doc.save(source)
    plan = record_plan(
        source,
        [
            (literal, "Old title", "title"),
            (literal, "Old period", "period"),
            ("Old body", "Old body", "details"),
        ],
    )
    content = record_content(
        [{"id": "a", "title": "Result", "period": "2026", "details": "Details"}]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    table = Document(output).tables[0]
    assert table._tbl.find("w:tblPr/w:tblInd", NS).get(w("w")) == "740"
    assert sum(column.width for column in table.columns) == Pt(432 - 37)


@pytest.mark.parametrize("after", [False, True])
@pytest.mark.parametrize("section_break", [False, True])
def test_unlabelled_contact_blank_is_repaired_on_either_side(tmp_path, after, section_break):
    """位于联系方式前后留白中的主页补齐标签并继承完整条目样式"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    if not after:
        doc.add_paragraph()
    contact = doc.add_paragraph()
    contact.paragraph_format.left_indent = Pt(10)
    contact.add_run("电话：").bold = True
    contact.add_run("旧电话").bold = False
    if section_break:
        doc.add_section(WD_SECTION_START.CONTINUOUS)
    if after:
        doc.add_paragraph()
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [n for n in package.nodes.values() if n.tag == w("p")]
    contact_node = next(n for n in paragraphs if paragraph_text(n) == "电话：旧电话")
    blank_node = paragraphs[-1 if after else 0]
    plan = TemplatePlan(
        summary="空位测试",
        fields=[
            TextBinding(node=package.ids[contact_node], quote="旧电话", target="personal.phone"),
            TextBinding(node=package.ids[blank_node], quote="", target="personal.website"),
        ],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    content = ResumeDocument.model_validate(
        {
            "personal": {"phone": "12345", "website": "https://example.test"},
            "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
        }
    )
    original, previous = source.read_bytes(), deepcopy(plan)
    fill_template(source, output, plan, content.model_dump(), [])
    result = Document(output)
    homepage = next(p for p in result.paragraphs if "https://example.test" in p.text)
    assert homepage.text == "个人主页：https://example.test"
    assert homepage.paragraph_format.left_indent == Pt(10)
    assert homepage.runs[0].bold and homepage.runs[1].bold is False
    assert source.read_bytes() == original and plan == previous


def test_contact_blank_with_an_explicit_label_after_contacts_is_preserved(tmp_path):
    """主页已有独立标签时保留原位置"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("电话：旧电话")
    doc.add_paragraph("Website:")
    doc.add_paragraph()
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [n for n in package.nodes.values() if n.tag == w("p")]
    plan = TemplatePlan(
        summary="明确标签",
        fields=[
            TextBinding(node=package.ids[paragraphs[0]], quote="旧电话", target="personal.phone"),
            TextBinding(node=package.ids[paragraphs[2]], quote="", target="personal.website"),
        ],
        repeats=[],
        photos=[],
        keep=[package.ids[paragraphs[1]]],
        remove=[],
        warnings=[],
    )
    content = ResumeDocument.model_validate(
        {
            "personal": {"phone": "12345", "website": "https://example.test"},
            "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
        }
    )
    fill_template(source, output, plan, content.model_dump(), [])
    assert [p.text for p in Document(output).paragraphs] == [
        "电话：12345",
        "Website:",
        "https://example.test",
    ]


@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("hyperlink", [False, True])
def test_supplement_uses_label_and_value_styles_instead_of_isolated_city(
    tmp_path, legacy, hyperlink
):
    """新增字段和旧生成字段均沿用完整联系方式样式"""
    source, completed, output = (
        tmp_path / name for name in ("source.docx", "completed.docx", "result.docx")
    )
    doc = Document()
    phone = doc.add_paragraph()
    phone.paragraph_format.left_indent = Pt(10)
    label, value = phone.add_run("电话： "), phone.add_run("Old phone")
    label.bold, value.bold = True, False
    for run in (label, value):
        run.font.size, run.font.name = Pt(12), "Arial"
    if hyperlink:
        link = etree.SubElement(phone._p, w("hyperlink"), {w("anchor"): "contact"})
        link.append(value._r)
    city = doc.add_paragraph("Old city")
    city.paragraph_format.right_indent = Pt(320)
    city.paragraph_format.first_line_indent = Pt(12)
    if legacy:
        added = doc.add_paragraph("个人主页：〔待填写〕")
        added.paragraph_format.right_indent = Pt(320)
    doc.save(source)
    package = TemplatePackage(source)
    nodes = {paragraph_text(node): key for key, node in package.nodes.items() if node.tag == w("p")}
    fields = [
        TextBinding(node=nodes["电话： Old phone"], quote="Old phone", target="personal.phone"),
        TextBinding(node=nodes["Old city"], quote="Old city", target="personal.location"),
    ]
    if legacy:
        fields.append(
            TextBinding(
                node=nodes["个人主页：〔待填写〕"], quote="〔待填写〕", target="personal.website"
            )
        )
    plan = TemplatePlan(
        summary="测试", fields=fields, repeats=[], photos=[], keep=[], remove=[], warnings=[]
    )
    content = ResumeDocument.model_validate(
        {
            "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
            "personal": {
                "phone": "12345",
                "location": "City",
                "website": "https://example.test/profile",
            },
        }
    )
    before = source.read_bytes()
    package, mapped, _ = supplement_personal_fields(package, plan, content, [], completed)
    website = next(field for field in mapped.fields if field.target == "personal.website")
    paragraph = package.node(website.node)
    assert paragraph.find("w:pPr/w:ind", NS).get(w("left")) == "200"
    assert paragraph.find("w:pPr/w:ind", NS).get(w("right")) is None
    runs = paragraph.findall(w("r"))
    assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
    assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"
    assert all(run.find("w:rPr/w:rFonts", NS).get(w("ascii")) == "Arial" for run in runs)
    again, same, notices = supplement_personal_fields(package, mapped, content, [])
    assert not notices and same == mapped and again is package
    fill_template(completed, output, mapped, content.model_dump(), [])
    assert sum("个人主页" in p.text for p in Document(output).paragraphs) == 1
    assert source.read_bytes() == before
