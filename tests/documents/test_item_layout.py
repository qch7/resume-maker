"""逐项样式按稳定身份填充，重复文字、重排及模板聚合不会串改"""

from copy import deepcopy

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree

from resume_maker.domain.page_layout import LAYOUT_KEY, item_key
from resume_maker.domain.resume import ResumeDocument
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.integrations.word.item_layout import ITEM
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.page_layout import apply_page_layout
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from tests.support.documents import make_template, project_content, resume_content
from tests.support.layouts import order_template, project_document


def configure(document, *, items=None, types=None, **extra):
    """通过实际简历扩展保存逐项设置，保留其他插件数据"""
    document.extensions[LAYOUT_KEY] = {
        "display": {"hidden": True},
        "options": {"item_styles": {"types": types or {}, "items": items or {}}, **extra},
    }
    return document.model_dump()


def runs(package, literal):
    """从真正生成的全部 Word 部件中读取可见运行，避免把文本框外层运行重复计算"""
    return [
        run
        for root in package.parts.values()
        for run in root.iter(w("r"))
        if "".join(text.text or "" for text in run.findall(w("t"))).strip() == literal
    ]


def sizes(package, literal):
    """读取最终 DOCX 的字号，验证生成行为而非内部配置形状"""
    return [float(run.find("w:rPr/w:sz", NS).get(w("val"))) / 2 for run in runs(package, literal)]


def test_linked_items_detach_reconnect_and_keep_identity_after_rename(tmp_path):
    """同类默认联动，单项独立覆盖、复链、改名及跨栏目同标识保持可预测"""
    document = ResumeDocument.model_validate(
        {
            "sections": [
                {"id": "s:A", "title": "栏目甲", "entries": [{"id": "same", "title": "共同标题"}]},
                {"id": "s:B", "title": "栏目乙", "entries": [{"id": "same", "title": "共同标题"}]},
                {"id": "p", "title": "项目经历", "kind": "projects"},
            ],
        }
    )
    key = item_key("entry", "s:A", "same", "title")
    content = configure(
        document, items={key: {"font_size": 17}}, types={"title": {"font_size": 13}}
    )
    output = tmp_path / "styled.docx"
    write_full_resume(output, content, [])
    assert sizes(TemplatePackage(output), "共同标题") == [17, 13]
    content["sections"].reverse()
    write_full_resume(output, content, [])
    assert sizes(TemplatePackage(output), "共同标题") == [13, 17]
    content["sections"][-1]["entries"][0]["title"] = "改名后仍独立"
    write_full_resume(output, content, [])
    assert sizes(TemplatePackage(output), "改名后仍独立") == [17]
    del content["extensions"][LAYOUT_KEY]["options"]["item_styles"]["items"][key]
    write_full_resume(output, content, [])
    assert sizes(TemplatePackage(output), "改名后仍独立") == [13]


def test_same_paragraph_uses_distinct_sizes_and_maximum_spacing(tmp_path):
    """教育三列可分别设字号，共段间距取显式最大值且不受替换顺序影响"""
    document = resume_content()
    output = tmp_path / "header.docx"
    content = configure(
        document,
        items={
            item_key("entry", "education", "school1", "title"): {"font_size": 14, "space_after": 2},
            item_key("entry", "education", "school1", "subtitle"): {
                "font_size": 9,
                "space_after": 7,
            },
            item_key("entry", "education", "school1", "period"): {
                "font_size": 8,
                "space_before": 4,
            },
            item_key("personal", "name"): {"font_size": 20},
        },
        types={"section_title": {"font_size": 16}},
    )
    write_full_resume(output, content, [])
    package = TemplatePackage(output)
    assert sizes(package, "新大学一") == [14]
    assert sizes(package, "计算机 · 本科") == [9]
    assert sizes(package, "2020–2024") == [8]
    assert sizes(package, "测试新姓名") == [20]
    assert sizes(package, "教育背景") == [16]
    paragraph = next(run.iterancestors(w("p")) for run in runs(package, "新大学一"))
    paragraph = next(paragraph)
    spacing = paragraph.find("w:pPr/w:spacing", NS)
    assert spacing.get(w("after")) == "140"
    assert spacing.get(w("before")) == "80"


@pytest.mark.parametrize("layout", ["details", "split_rows", "highlight_slots"])
@pytest.mark.parametrize("ordered", [False, True])
def test_template_aggregates_keep_highlight_and_custom_identity(tmp_path, layout, ordered):
    """综合正文、亮点槽位、表格及手动重排均保留每项字号、独立段距"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    plan = order_template(source, layout)
    projects = project_content()
    projects[0]["project_id"] = "p"
    value = projects[0]["content"]
    value["highlights"] = [
        {"id": "h1", "title": "同名亮点", "text": "共同内容"},
        {"id": "h2", "title": "同名亮点", "text": "共同内容"},
    ]
    value["custom_fields"] = [
        {"id": "custom", "label": "补充", "value": "自定义内容", "visible": True}
    ]
    projects[0]["highlight_ids"] = ["h1", "h2"]
    if ordered:
        value["body_order"] = ["highlights", "custom:custom", "role", "stack", "description"]
    document = configure(
        project_document(),
        items={
            item_key("project", "p", "highlight", "h1"): {"font_size": 12, "space_after": 6},
            item_key("project", "p", "highlight", "h2"): {"font_size": 9, "space_after": 2},
            item_key("project", "p", "custom", "custom"): {"font_size": 8, "space_after": 1},
        },
    )
    original = source.read_bytes()
    fill_template(source, output, plan, document, projects)
    package = TemplatePackage(output)
    assert sizes(package, "共同内容") == [12, 9]
    assert sizes(package, "自定义内容") == [8]
    paragraphs = [next(run.iterancestors(w("p"))) for run in runs(package, "共同内容")]
    assert paragraphs[0] is not paragraphs[1]
    assert [paragraph.find("w:pPr/w:spacing", NS).get(w("after")) for paragraph in paragraphs] == [
        "120",
        "40",
    ]
    assert source.read_bytes() == original
    assert all(ITEM not in node.attrib for root in package.parts.values() for node in root.iter())


def test_imported_headers_textboxes_and_exact_quotes_keep_item_styles(tmp_path):
    """页眉联系方式、文本框职位、栏目标题及重复区通过原映射精确套样式"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _, plan = make_template(source)
    content = configure(
        resume_content(),
        types={"name": {"font_size": 18}, "job_title": {"font_size": 15}},
        items={
            item_key("personal", "phone"): {"font_size": 8},
            item_key("personal", "email"): {"font_size": 12},
        },
    )
    fill_template(source, output, plan, content, project_content())
    package = TemplatePackage(output)
    assert sizes(package, "测试新姓名") == [18]
    assert sizes(package, "文档开发工程师") == [15]
    assert sizes(package, "10000000000") == [8]
    assert sizes(package, "new@example.test") == [12]
    box = next(
        node
        for root in package.parts.values()
        for node in root.iter("{urn:schemas-microsoft-com:vml}textbox")
    )
    assert "mso-fit-shape-to-text:t" in box.get("style")


def test_auto_fit_preserves_hierarchy_small_manual_sizes_and_fixed_spacing(tmp_path):
    """自动压页基于实际继承字号收紧，不放大手工小字或裁短原固定行高"""
    output = tmp_path / "auto.docx"
    document = Document()
    document.styles["Normal"].font.size = Pt(18)
    paragraph = document.add_paragraph("继承段落直接字号")
    props = paragraph._p.get_or_add_pPr()
    etree.SubElement(etree.SubElement(props, w("rPr")), w("sz")).set(w("val"), "28")
    paragraph.paragraph_format.line_spacing = Pt(22)
    small = document.add_paragraph("已有小字")
    small.runs[0].font.size = Pt(8)
    small.paragraph_format.line_spacing = 0.8
    small.paragraph_format.space_after = Pt(1)
    document.save(output)
    apply_page_layout(
        output, {LAYOUT_KEY: {"options": {"auto_fit": {"spacing_factor": 0, "font_delta": 2}}}}
    )
    actual = Document(output)
    assert actual.paragraphs[0].runs[0].font.size.pt == 12
    assert actual.paragraphs[0].paragraph_format.line_spacing.pt == 14.4
    assert actual.paragraphs[1].runs[0].font.size.pt == 8
    assert actual.paragraphs[1].paragraph_format.line_spacing == 0.8
    assert actual.paragraphs[1].paragraph_format.space_after.pt == 0


def test_layout_roundtrip_preserves_unknown_options_and_tuple_keys():
    """保存再恢复保留扩展资料及特殊标识，独立空声明继续屏蔽同类设置"""
    document = project_document()
    key = item_key("project", '中文/"项目', "custom", "a:b")
    content = configure(document, items={key: {}}, future={"owned": True})
    content["extensions"]["example.future"] = {"value": [1, 2]}
    restored = ResumeDocument.model_validate(deepcopy(content)).model_dump()
    assert restored["extensions"] == content["extensions"]
    assert key == '["project","中文/\\"项目","custom","a:b"]'
