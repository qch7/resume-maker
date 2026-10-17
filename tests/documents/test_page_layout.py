"""版面设置保留完整事实、标题样式、原文件及其他方案的默认排版"""

from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from resume_maker.domain.page_layout import LAYOUT_KEY, layout_options
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.page_layout import apply_page_layout
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage


def content():
    """构造有隐藏字段、来源身份、完整荣誉事实的独立方案"""
    return {
        "personal": {"name": "合成测试同学", "phone": "10000000000"},
        "sections": [
            {"id": "projects", "title": "项目经历", "kind": "projects"},
            {
                "id": "renamed",
                "title": "成果",
                "entries": [
                    {
                        "id": "entry",
                        "source": {"provider": "ext.honors/library", "id": "honor", "version": "1"},
                        "title": "合成竞赛",
                        "period": "2026-10",
                        "details": "经过核实的成果",
                        "subtitle": "隐藏单位",
                        "hidden_fields": ["subtitle"],
                        "custom_fields": [{"id": "award", "label": "奖项", "value": "一等奖"}],
                    }
                ],
            },
        ],
    }


def compact(document):
    """只在当前副本启用有界紧凑参数"""
    result = deepcopy(document)
    result["extensions"] = {
        LAYOUT_KEY: {
            "display": {"hidden": True},
            "options": {
                "target_pages": 1,
                "font_size": 9.5,
                "space_after": 1,
                "line_spacing": 1.05,
            },
        }
    }
    return result


def test_builtin_compact_keeps_visible_facts_and_large_header(tmp_path):
    """通用正文参数保留每条事实及字段结构，姓名字号及另一份方案保持原样"""
    document = content()
    original = deepcopy(document)
    output = tmp_path / "compact.docx"
    write_full_resume(output, compact(document), [])
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    for value in ("合成竞赛", "2026-10", "一等奖", "经过核实的成果"):
        assert text.count(value) == 1
    assert "隐藏单位" not in text
    honor = next(paragraph for paragraph in rendered.paragraphs if "合成竞赛" in paragraph.text)
    assert "一等奖" not in honor.text and "经过核实的成果" not in honor.text
    assert any(paragraph.text == "经过核实的成果" for paragraph in rendered.paragraphs)
    assert any(paragraph.text == "奖项：一等奖" for paragraph in rendered.paragraphs)
    assert honor.runs[0].font.size.pt == 9.5
    assert honor.paragraph_format.space_after.pt == 1
    name = next(
        run
        for table in rendered.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if "合成测试同学" in run.text
    )
    assert name.font.size.pt == 23
    assert document == original


@pytest.mark.parametrize("retired", [True, False, "ignored"])
@pytest.mark.parametrize("minimal", [False, True])
def test_adaptive_retired_option_keeps_normal_fields_and_template(tmp_path, retired, minimal):
    """旧开关不再合并模板字段，原映射、源文件及逐字段内容保持不变"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    samples = ("成果", "旧标题") if minimal else ("成果", "旧标题", "旧时间", "旧正文", "旧奖项")
    for value in samples:
        doc.add_paragraph(value)
    doc.save(source)
    package = TemplatePackage(source)
    rows = [node for node in package.inventory()["nodes"] if node["kind"] == "p"]
    fields = [
        TextBinding(node=node["id"], quote=node["text"], target=target)
        for node, target in zip(
            rows[1:],
            ("title",) if minimal else ("title", "period", "details", "custom_fields"),
            strict=True,
        )
    ]
    plan = TemplatePlan(
        summary="成果模板",
        fields=[],
        repeats=[
            RepeatBinding(
                section="成果",
                start=rows[1]["id"],
                end=rows[-1]["id"],
                sample_start=rows[1]["id"],
                sample_end=rows[-1]["id"],
                fields=fields,
            )
        ],
        photos=[],
        keep=[rows[0]["id"]],
        remove=[],
        warnings=[],
    )
    original, mapping = source.read_bytes(), plan.model_dump()
    document = compact(content())
    baseline = tmp_path / "baseline.docx"
    fill_template(source, baseline, plan, document, [])
    document["extensions"][LAYOUT_KEY]["options"]["honors_inline"] = retired
    fill_template(source, output, plan, document, [])
    paragraphs = [paragraph.text for paragraph in Document(output).paragraphs if paragraph.text]
    for value in ("合成竞赛", "2026-10", "经过核实的成果", "奖项：一等奖"):
        assert "\n".join(paragraphs).count(value) == 1
    if not minimal:
        assert paragraphs[-5:] == ["成果", "合成竞赛", "2026-10", "经过核实的成果", "奖项：一等奖"]
    assert "隐藏单位" not in "\n".join(paragraphs)
    with ZipFile(baseline) as before, ZipFile(output) as after:
        assert before.read("word/document.xml") == after.read("word/document.xml")
    assert source.read_bytes() == original and plan.model_dump() == mapping


@pytest.mark.parametrize("retired", [True, False, "ignored"])
def test_builtin_retired_option_is_ignored_without_losing_other_extensions(tmp_path, retired):
    """旧数据可以继续恢复及导出，废弃开关无效，未知选项及用户内容不变"""
    document = compact(content())
    baseline, output = tmp_path / "baseline.docx", tmp_path / "legacy.docx"
    write_full_resume(baseline, document, [])
    options = document["extensions"][LAYOUT_KEY]["options"]
    options.update({"honors_inline": retired, "future": {"retained": True}})
    before = deepcopy(document)
    restored = ResumeDocument.model_validate(document).model_dump()
    assert restored["extensions"] == before["extensions"]
    effective = layout_options(restored["extensions"]).model_dump()
    assert "honors_inline" not in effective
    assert effective["future"] == {"retained": True}
    write_full_resume(output, restored, [])
    with ZipFile(baseline) as original, ZipFile(output) as actual:
        assert original.read("word/document.xml") == actual.read("word/document.xml")
    assert document == before


def test_postprocess_retains_unrelated_parts_and_default_is_noop(tmp_path):
    """正文调整保留压缩包其他资源，无版面参数时不重写文件"""
    output = tmp_path / "document.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.save(output)
    with ZipFile(output, "a") as archive:
        archive.writestr("custom/unchanged.bin", b"preserved")
    original = output.read_bytes()
    apply_page_layout(output, {})
    assert output.read_bytes() == original
    apply_page_layout(output, compact(content())["extensions"])
    with ZipFile(output) as archive:
        assert archive.read("custom/unchanged.bin") == b"preserved"
    assert Document(output).paragraphs[0]._p.find("w:r/w:rPr/w:sz", NS).get(w("val")) == "19"


def test_custom_heading_inherits_large_size_without_body_override(tmp_path):
    """自定义标题通过父样式继承大字号时仍不受正文字号调整影响"""
    output = tmp_path / "styles.docx"
    doc = Document()
    parent = doc.styles.add_style("LargeSource", WD_STYLE_TYPE.PARAGRAPH)
    parent.font.size = Pt(20)
    child = doc.styles.add_style("InheritedHeading", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    doc.add_paragraph("标题", style=child)
    doc.add_paragraph("正文")
    doc.save(output)
    apply_page_layout(output, compact(content())["extensions"])
    rendered = Document(output)
    assert rendered.paragraphs[0]._p.find("w:r/w:rPr/w:sz", NS) is None
    assert rendered.paragraphs[1]._p.find("w:r/w:rPr/w:sz", NS).get(w("val")) == "19"


@pytest.mark.parametrize(
    "options",
    [
        {"font_size": 8},
        {"font_size": 13},
        {"space_after": -1},
        {"line_spacing": 2},
        {"target_pages": 4},
    ],
)
def test_invalid_layout_is_rejected_before_saving(options):
    """接口不能保存超出可读范围或无效预算的版面参数"""
    document = content()
    document["extensions"] = {LAYOUT_KEY: {"display": {"hidden": True}, "options": options}}
    with pytest.raises(ValueError):
        ResumeDocument.model_validate(document)
