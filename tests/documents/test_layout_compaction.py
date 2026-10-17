"""通用压页覆盖隐藏格式、分页约束及图形保护，不依赖资料名称或来源"""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Pt
from lxml import etree

from resume_maker.domain.page_layout import LAYOUT_KEY, AutoFit
from resume_maker.integrations.word.item_layout import scaled_size
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.page_layout import apply_page_layout
from resume_maker.integrations.word.templates.mapping import TemplatePackage


def compact(doc, tmp_path, **options):
    """从真实 DOCX 原件执行正式后处理，返回保留全部部件的解析结果"""
    path = tmp_path / "compact.docx"
    doc.save(path)
    apply_page_layout(path, {LAYOUT_KEY: {"options": {"auto_fit": options}}})
    return TemplatePackage(path)


def sz(parent, path="w:rPr/w:sz"):
    """按 Word 半磅存储读取明确字号"""
    return float(parent.find(path, NS).get(w("val"))) / 2


def test_text_markers_empty_runs_and_empty_paragraphs_share_actual_font_limits(tmp_path):
    """段标记及纯控制运行不能继续按十一磅撑高已经缩小的正文"""
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("完整正文").font.size = Pt(10.5)
    paragraph.add_run().add_tab()
    paragraph.add_run().add_break()
    doc.add_paragraph()
    package = compact(doc, tmp_path, spacing_factor=0, font_scale=0.7)
    paragraphs = package.parts["word/document.xml"].findall("w:body/w:p", NS)
    assert [sz(run) for run in paragraphs[0].findall(w("r"))] == [7, 7, 7]
    assert sz(paragraphs[0], "w:pPr/w:rPr/w:sz") == 7
    assert sz(paragraphs[1], "w:pPr/w:rPr/w:sz") == 1
    assert len(paragraphs[0].findall(".//w:tab", NS)) == 1
    assert len(paragraphs[0].findall(".//w:br", NS)) == 1


def test_paragraph_mark_size_does_not_replace_inherited_visible_text_size(tmp_path):
    """段落标记只控制结束标记，不能误当正文继承字号并放大实际文字"""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    paragraph = doc.add_paragraph("继承十磅正文")
    marker = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(w("val"), "60")
    marker.append(size)
    paragraph._p.get_or_add_pPr().append(marker)
    package = compact(doc, tmp_path, font_scale=0.8)
    paragraph = package.parts["word/document.xml"].find("w:body/w:p", NS)
    assert sz(paragraph.find(w("r"))) == 8
    assert sz(paragraph, "w:pPr/w:rPr/w:sz") == 8


@pytest.mark.parametrize(
    "size,delta,scale,expected",
    [
        (20, 0, 0.75, 15),
        (10.5, 0, 0.75, 7.5),
        (10, 3, 1, 9),
        (10, 3, 0.8, 7),
        (5, 0, 0, 5),
        (72, 0, 0, 6),
        (6, 0, 0.1, 6),
    ],
)
def test_font_scaling_preserves_legacy_size_hierarchy_and_half_point_floor(
    size, delta, scale, expected
):
    """旧减字号继续按九磅解释，新比例统一向下半磅且不放大已有小字"""
    assert scaled_size(size, AutoFit(font_delta=delta, font_scale=scale)) == expected


def test_inherited_paragraph_spacing_and_fixed_lines_shrink_without_clipping(tmp_path):
    """样式继承的前后间距及固定行框均进入压缩，固定行框仍容纳最大文字"""
    doc = Document()
    parent = doc.styles.add_style("CustomParent", WD_STYLE_TYPE.PARAGRAPH)
    parent.font.size = Pt(10)
    parent.paragraph_format.space_before = Pt(8)
    parent.paragraph_format.space_after = Pt(6)
    parent.paragraph_format.line_spacing = Pt(28)
    style = doc.styles.add_style("CustomChild", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = parent
    doc.add_paragraph("固定行框中的完整文字", style=style)
    package = compact(doc, tmp_path, spacing_factor=0, font_scale=0.8)
    paragraph = package.parts["word/document.xml"].find("w:body/w:p", NS)
    spacing = paragraph.find("w:pPr/w:spacing", NS)
    assert spacing.get(w("before")) == spacing.get(w("after")) == "0"
    assert spacing.get(w("line")) == "192"
    assert spacing.get(w("lineRule")) == "exact"
    assert sz(paragraph.find(w("r"))) == 8


def test_plain_text_hard_pages_and_long_keep_chain_can_flow(tmp_path):
    """普通文字硬分页改成换行，长链只保留短首段跟随首段正文"""
    doc = Document()
    style = doc.styles.add_style("AlwaysNext", WD_STYLE_TYPE.PARAGRAPH)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.page_break_before = True
    style.paragraph_format.keep_together = True
    for index in range(6):
        paragraph = doc.add_paragraph(f"完整条目 {index}", style=style)
        if index == 3:
            paragraph.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("尾部完整文字")
    package = compact(doc, tmp_path, spacing_factor=0, relax_pagination=True)
    root = package.parts["word/document.xml"]
    paragraphs = root.findall("w:body/w:p", NS)
    assert paragraphs[0].find("w:pPr/w:keepNext", NS) is None
    assert all(p.find("w:pPr/w:keepNext", NS).get(w("val")) == "0" for p in paragraphs[1:6])
    assert all(p.find("w:pPr/w:pageBreakBefore", NS).get(w("val")) == "0" for p in paragraphs[:6])
    assert all(p.find("w:pPr/w:keepLines", NS).get(w("val")) == "0" for p in paragraphs[:6])
    assert not root.findall(".//w:br[@w:type='page']", NS)
    assert len(root.findall(".//w:br", NS)) == 1
    assert "尾部完整文字" in "".join(root.itertext())


def test_fixed_text_table_rows_release_height_without_removing_cells(tmp_path):
    """纯文字表格解除固定高及不许跨页，表格内容和编辑结构完整保留"""
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    for index, row in enumerate(table.rows):
        row.height = Pt(90)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column, cell in enumerate(row.cells):
            cell.text = f"完整单元格 {index}-{column}"
    package = compact(doc, tmp_path, spacing_factor=0, relax_pagination=True)
    root = package.parts["word/document.xml"]
    assert len(list(root.iter(w("tr")))) == 4 and len(list(root.iter(w("tc")))) == 8
    assert not list(root.iter(w("trHeight")))
    assert all(node.get(w("val")) == "0" for node in root.iter(w("cantSplit")))
    assert all(
        f"完整单元格 {i}-{j}" in "".join(root.itertext()) for i in range(4) for j in range(2)
    )


def test_positioned_objects_preserve_neighbor_spacing_pages_sections_and_extent(tmp_path):
    """页定位对象可依赖后段留白，保守保留整篇定位约束以免合页后叠图"""
    doc = Document()
    drawing = OxmlElement("w:drawing")
    anchor = etree.SubElement(
        drawing, "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor"
    )
    etree.SubElement(
        anchor,
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent",
        cx="1200000",
        cy="700000",
    )
    doc.add_paragraph().add_run()._r.append(drawing)
    paragraph = doc.add_paragraph("图形后的完整正文")
    paragraph.paragraph_format.space_before = Pt(70)
    paragraph.paragraph_format.page_break_before = True
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    original = etree.tostring(drawing)
    package = compact(doc, tmp_path, spacing_factor=0, font_scale=0.7, relax_pagination=True)
    root = package.parts["word/document.xml"]
    paragraph = root.findall("w:body/w:p", NS)[1]
    assert paragraph.find("w:pPr/w:spacing", NS).get(w("before")) == "1400"
    assert paragraph.find("w:pPr/w:pageBreakBefore", NS).get(w("val"), "1") == "1"
    assert len(root.findall(".//w:br[@w:type='page']", NS)) == 1
    assert etree.tostring(root.find(".//w:drawing", NS)) == original


def test_image_row_keeps_original_fixed_height_and_dimensions(tmp_path):
    """有图片的行不缩到图片之外，行高及绘图尺寸保持原值"""
    doc = Document()
    row = doc.add_table(rows=1, cols=1).rows[0]
    row.height = Pt(90)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    drawing = OxmlElement("w:drawing")
    inline = etree.SubElement(
        drawing, "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
    )
    etree.SubElement(
        inline,
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent",
        cx="500000",
        cy="800000",
    )
    row.cells[0].paragraphs[0].add_run()._r.append(drawing)
    package = compact(doc, tmp_path, spacing_factor=0, relax_pagination=True)
    root = package.parts["word/document.xml"]
    height = next(root.iter(w("trHeight")))
    assert height.get(w("val")) == "1800" and height.get(w("hRule")) == "exact"
    assert root.xpath(".//*[local-name()='extent']/@cy") == ["800000"]


def test_same_geometry_sections_can_flow_but_different_page_size_stays_separate(tmp_path):
    """只有页面配置一致的相邻节可以连续，纸张或方向变化保持分页"""
    doc = Document()
    doc.add_paragraph("第一节完整文字")
    doc.add_section()
    doc.add_paragraph("第二节完整文字")
    doc.add_section().page_width = Pt(800)
    doc.add_paragraph("第三节完整文字")
    package = compact(doc, tmp_path, relax_pagination=True)
    sections = list(package.parts["word/document.xml"].iter(w("sectPr")))
    assert sections[1].find(w("type")).get(w("val")) == "continuous"
    assert (
        sections[2].find(w("type")) is None
        or sections[2].find(w("type")).get(w("val")) != "continuous"
    )
