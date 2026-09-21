"""验收脚本按真实页面核验内容，区分浮动标题导致的抽取乱序和实际漏文。"""

from pathlib import Path
from runpy import run_path

import pymupdf
import pytest
from docx import Document

from resume_maker.domain.resume import ResumeDocument


@pytest.mark.parametrize("missing", [False, True])
def test_evaluation_finds_cross_page_text_but_still_detects_missing_content(tmp_path, missing):
    """晚写入的浮动标题不能打断跨页正文；实际少掉后半句仍必须报告。"""
    evaluation = run_path(
        str(Path(__file__).resolve().parents[1] / "scripts/evaluate_templates.py")
    )
    source = tmp_path / "result.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 700), "First half")
        page.insert_text((50, 50), "Ada")
        page.insert_text((50, 100), "Awards")
        page = pdf.new_page()
        page.insert_text((50, 50), "omitted" if missing else "second half")
        pdf.save(source)
    document = ResumeDocument(
        personal={"name": "Ada"},
        sections=[
            {"id": "projects", "kind": "projects", "title": "Projects"},
            {
                "id": "awards",
                "title": "Awards",
                "entries": [{"id": "one", "details": "First half second half"}],
            },
        ],
    )
    report = evaluation["inspect_output"](source, document, [])
    assert report["missing_values"] == (["First half second half"] if missing else [])
    assert report["name_on_first_page"] and report["missing_headings"] == []
    assert evaluation["accepted_output"]({"ready": True, "content_check": report}) is not missing
    assert not evaluation["accepted_output"]({"ready": True})


@pytest.mark.parametrize("missing", [False, True])
def test_evaluation_checks_actual_project_values_without_requiring_generated_labels(
    tmp_path, missing
):
    """模板可只显示原始字段值；不强制派生正文的标签，但真实角色缺失仍失败。"""
    evaluation = run_path(
        str(Path(__file__).resolve().parents[1] / "scripts/evaluate_templates.py")
    )
    source = tmp_path / "result.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 50), "Projects\nArchive\n2026\nPython\nA reproducible dataset.")
        if not missing:
            page.insert_text((50, 150), "Maintainer")
        pdf.save(source)
    document = ResumeDocument(
        sections=[{"id": "projects", "kind": "projects", "title": "Projects"}]
    )
    projects = [
        {
            "project_id": "p",
            "highlight_ids": [],
            "content": {
                "title": "Archive",
                "period": "2026",
                "role": "Maintainer",
                "stack": ["Python"],
                "description": "A reproducible dataset.",
                "highlights": [],
                "custom_fields": [],
            },
        }
    ]
    report = evaluation["inspect_output"](source, document, projects)
    assert report["missing_values"] == (["Maintainer"] if missing else [])


@pytest.mark.parametrize("extra", [False, True])
def test_record_count_audit_allows_shared_values_but_detects_added_copies(tmp_path, extra):
    """用户确实填写两条同名记录时应保留两次；第三次复制才属于多填。"""
    evaluation = run_path(
        str(Path(__file__).resolve().parents[1] / "scripts/evaluate_templates.py")
    )
    source = tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("Credentials")
    for _ in range(3 if extra else 2):
        doc.add_paragraph("Shared research credential")
    doc.save(source)
    document = ResumeDocument(
        sections=[
            {"id": "projects", "kind": "projects", "title": "Projects"},
            {
                "id": "credentials",
                "title": "Credentials",
                "entries": [
                    {"id": str(i), "details": "Shared research credential"} for i in range(2)
                ],
            },
        ]
    )
    check = evaluation["inspect_record_counts"](source, document, [])
    assert bool(check["duplicates"]) is extra
    assert check["counts"] == [
        {"identity": "Shared research credential", "expected": 2, "actual": 3 if extra else 2}
    ]
    report = {
        "ready": True,
        "content_check": {"pages": 1, "name_on_first_page": True},
        "record_count_check": check,
    }
    assert evaluation["accepted_output"](report) is not extra
