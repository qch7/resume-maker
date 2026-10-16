"""自动排版基于冻结输入及实际页数，失败取消不发布候选或修改资料"""

import json
import threading
import time
from copy import deepcopy
from itertools import pairwise
from types import SimpleNamespace

import anyio
import pytest
from docx import Document

from resume_maker.api.schemas import ResumePreviewInput
from resume_maker.domain.layout_fit import (
    FIT_TIMEOUT_SECONDS,
    MAX_FIT_ATTEMPTS,
    fit_candidates,
    saved_options,
    with_fit_options,
)
from resume_maker.domain.page_layout import LAYOUT_KEY, item_key
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.plugin_packages.sys_documents.routes.resumes import fit_resume
from resume_maker.plugin_packages.sys_documents.services.resume_previews import ResumePreviews
from resume_maker.sdk import documents
from resume_maker.sdk.documents import DocumentInput, current_document_work, document_work
from resume_maker.sdk.model import Cancelled, ProviderError
from tests.support.documents import resume_content


class Catalog:
    """只允许一次读取的合成资料源，渲染候选不得重新读取可变内容"""

    def __init__(self):
        """记录读取次数供冻结行为验证"""
        self.reads = 0

    def freeze_preview(self, directory, template_id, document, items):
        """模拟同一事务冻结的独立输入"""
        self.reads += 1
        return DocumentInput(json.dumps({"document": document}), json.dumps(items), "{}", None)


@pytest.fixture
def make_preview(tmp_path):
    """构造计数分页器及可检查的临时文件，不启动 Word"""
    services = []

    def create(pages, *, settings=None):
        """页数回调可以依据候选参数或顺序产生确定结果"""
        calls = []

        def engine(output, document, projects):
            """用 JSON 代替 DOCX 以直接检查实际送给分页器的输入"""
            output.write_text(json.dumps(document), encoding="utf-8")

        def renderer(source, output):
            """每个候选写独立分页文件，模拟真实渲染器契约"""
            document = json.loads(source.read_text(encoding="utf-8"))
            calls.append(document)
            count, error = pages(document, len(calls))
            if count:
                output.write_bytes(b"pdf")
                for index in range(1, count + 1):
                    for extension in ("svg", "png"):
                        (output.parent / f"page-{index}.{extension}").write_bytes(b"page")
            return count, error

        service = ResumePreviews(
            Catalog(),
            tmp_path / str(len(services)),
            engine=engine,
            render=renderer,
            settings=settings,
        )
        services.append(service)
        return service, calls

    yield create
    for service in services:
        service.stop()


def test_budget_met_preserves_exact_options_and_reuses_normal_preview(make_preview):
    """预算内不改写原版式，普通预览直接复用本次真实结果"""
    service, calls = make_preview(lambda *_: (1, None))
    original = resume_content().model_dump()
    before = deepcopy(original)
    result = service.fit(None, original, [])
    assert result["status"] == "unchanged" and result["options"] == {}
    assert result["attempts"] == 1 and original == before and calls == [before]
    assert service.render(None, original, []) == result["preview"]
    assert len(calls) == 1 and service.catalog.reads == 2


def test_multiple_candidates_freeze_once_preserve_styles_and_publish_only_final(make_preview):
    """逐步收紧到首次实测达标，其他候选不占用发布缓存"""
    service, calls = make_preview(lambda _, number: (1 if number == 7 else 2, None))
    document = with_fit_options(
        resume_content().model_dump(),
        {"target_pages": 1, "item_styles": {"types": {"body": {"font_size": 12}}}, "future": [1]},
    )
    document["extensions"][LAYOUT_KEY]["display"]["future_display"] = "保留显示元数据"
    original = deepcopy(document)
    result = service.fit(None, document, [])
    assert result["status"] == "fit" and result["attempts"] == 7
    assert result["options"]["auto_fit"] == {"spacing_factor": 0, "font_delta": 2}
    assert result["options"]["item_styles"] == saved_options(original)["item_styles"]
    assert result["options"]["future"] == [1]
    assert document == original and service.catalog.reads == 1
    assert all(
        value["extensions"][LAYOUT_KEY]["display"]["future_display"] == "保留显示元数据"
        for value in calls
    )
    assert len(service.pool.entries) == 1 and len(service.cache) == 1
    assert len(list(service.pool.workspace.rglob("resume.docx"))) == 1
    final = with_fit_options(document, result["options"])
    assert service.render(None, final, []) == result["preview"]
    assert len(calls) == 7
    with service.lease(result["preview"]["id"], "resume.docx") as path:
        assert json.loads(path.read_text(encoding="utf-8")) == final


def test_unfit_keeps_best_observed_candidate_without_monotonic_assumption(make_preview):
    """后续更紧的候选页数变多时仍返回前面真实测得的最佳结果"""
    service, calls = make_preview(lambda _, number: (3 if number in (3, 5) else 4, None))
    result = service.fit(None, resume_content().model_dump(), [])
    assert result["status"] == "unfit" and result["pages"] == 3
    assert result["attempts"] == MAX_FIT_ATTEMPTS
    assert result["options"]["auto_fit"] == {"spacing_factor": 0, "font_delta": 0}
    stored = json.loads(service.file(result["preview"]["id"], "resume.docx").read_text("utf-8"))
    assert stored == calls[2]
    assert len(service.pool.entries) == 1
    assert len(list(service.pool.workspace.rglob("resume.docx"))) == 1


def test_search_keeps_trying_after_temporary_page_increase(make_preview):
    """中间候选回流增页不能提前推断后续候选不可行"""
    service, _ = make_preview(lambda _, number: ({1: 2, 2: 3, 3: 1}[number], None))
    result = service.fit(None, resume_content().model_dump(), [])
    assert result["status"] == "fit" and result["attempts"] == 3


def test_large_custom_font_reaches_final_readable_candidate(make_preview):
    """手动大字号也测到可读下限，不能只减三磅就判断无法压页"""
    service, _ = make_preview(
        lambda document, _: (
            1 if (saved_options(document).get("auto_fit") or {}).get("font_delta") == 63 else 2,
            None,
        )
    )
    document = with_fit_options(
        resume_content().model_dump(), {"item_styles": {"types": {"body": {"font_size": 72}}}}
    )
    result = service.fit(None, document, [])
    assert result["status"] == "fit" and result["attempts"] == MAX_FIT_ATTEMPTS
    assert result["options"]["auto_fit"]["font_delta"] == 63


def honor_document():
    """构造确有多字段可合并的可见荣誉，身份不依赖栏目中文名"""
    document = resume_content().model_dump()
    document["sections"].append(
        {
            "id": "awards",
            "title": "成果",
            "kind": "text",
            "visible": True,
            "entries": [
                {
                    "id": "award1",
                    "title": "合成竞赛",
                    "period": "2026-10",
                    "details": "核实成果",
                    "visible": True,
                    "source": {"provider": "ext.honors/library", "id": "award1", "version": "1"},
                    "custom_fields": [{"id": "award", "label": "奖项", "value": "合成一等奖"}],
                }
            ],
        }
    )
    return document


def test_honor_compaction_precedes_smaller_fonts_and_preserves_field_styles(tmp_path):
    """完整 DOCX 流程先合并荣誉再测页数，紧凑行保留各字段独立字号及事实"""
    calls = []

    def renderer(source, output):
        """用真实 DOCX 段落确认合并效果，只替代外部 Word 分页"""
        paragraphs = Document(source).paragraphs
        compact = any(
            all(value in paragraph.text for value in ("合成竞赛", "核实成果", "合成一等奖"))
            for paragraph in paragraphs
        )
        calls.append(compact)
        pages = 1 if compact else 2
        output.write_bytes(b"pdf")
        for index in range(1, pages + 1):
            for extension in ("png", "svg"):
                (output.parent / f"page-{index}.{extension}").write_bytes(b"page")
        return pages, None

    styles = {
        item_key("entry", "awards", "award1", "title"): {"font_size": 8},
        item_key("entry", "awards", "award1", "period"): {"font_size": 12},
        item_key("entry", "awards", "award1", "details"): {"font_size": 9},
        item_key("entry", "awards", "award1", "custom", "award"): {"font_size": 10},
    }
    document = with_fit_options(honor_document(), {"item_styles": {"items": styles}})
    service = ResumePreviews(Catalog(), tmp_path, engine=write_full_resume, render=renderer)
    try:
        result = service.fit(None, document, [])
        assert result["status"] == "fit" and result["attempts"] == 4
        assert calls == [False, False, False, True]
        assert result["options"]["honors_inline"] is True
        assert result["options"]["auto_fit"]["font_delta"] == 0
        assert service.catalog.reads == 1
        final = Document(service.file(result["preview"]["id"], "resume.docx"))
        paragraph = next(value for value in final.paragraphs if "合成竞赛" in value.text)
        for text, size in (("合成竞赛", 8), ("2026-10", 12), ("核实成果", 9), ("合成一等奖", 10)):
            assert next(run for run in paragraph.runs if text in run.text).font.size.pt == size
    finally:
        service.stop()


@pytest.mark.parametrize("case", ["hidden", "parent_hidden", "single", "already_inline", "visible"])
def test_honor_extra_attempt_requires_visible_mergeable_content(make_preview, case):
    """隐藏、仅有单字段或已经紧凑的荣誉不额外消耗排版次数"""
    document = honor_document()
    section = document["sections"][-1]
    if case == "hidden":
        section["visible"] = False
    elif case == "parent_hidden":
        section["parent_id"] = "parent"
        document["sections"].append({"id": "parent", "title": "隐藏父栏目", "visible": False})
    elif case == "single":
        section["entries"][0].update({"period": "", "details": "", "custom_fields": []})
    elif case == "already_inline":
        document = with_fit_options(document, {"honors_inline": True})
    service, calls = make_preview(lambda *_: (2, None))
    result = service.fit(None, document, [])
    assert result["status"] == "unfit"
    assert len(calls) == MAX_FIT_ATTEMPTS + (case == "visible")
    assert result["attempts"] == len(calls)


def test_existing_fit_only_tightens_and_respects_half_point_floor():
    """继续压页不会放大字号或恢复已经收起的空白"""
    original = {"auto_fit": {"spacing_factor": 0.4, "font_delta": 2.5}, "unknown": {"kept": 1}}
    candidates = list(fit_candidates(original))
    assert candidates[0] == original and len(candidates) <= MAX_FIT_ATTEMPTS
    assert all(candidate["unknown"] == original["unknown"] for candidate in candidates)
    for previous, current in pairwise(candidates):
        assert current["auto_fit"]["spacing_factor"] <= previous["auto_fit"]["spacing_factor"]
        assert current["auto_fit"]["font_delta"] >= previous["auto_fit"]["font_delta"]
    assert list(fit_candidates({"auto_fit": {"spacing_factor": 0, "font_delta": 63}})) == [
        {"auto_fit": {"spacing_factor": 0, "font_delta": 63}}
    ]


@pytest.mark.parametrize("failed_at", [1, 3])
def test_renderer_failure_preserves_original_and_does_not_adopt_partial_candidate(
    make_preview, failed_at
):
    """真实分页不可用时不能采用尚未确认的压缩参数"""
    service, _ = make_preview(
        lambda _, number: (None, "排版暂不可用") if number == failed_at else (2, None)
    )
    document = with_fit_options(resume_content().model_dump(), {"target_pages": 1})
    result = service.fit(None, document, [])
    assert result["status"] == "unavailable" and result["options"] == {"target_pages": 1}
    assert result["attempts"] == failed_at and result["message"] == "排版暂不可用"
    assert len(service.pool.entries) == 1
    stored = json.loads(service.file(result["preview"]["id"], "resume.docx").read_text("utf-8"))
    assert stored == document


def test_missing_renderer_reports_unavailable_without_searching(make_preview):
    """无可用精确分页能力时只保留当前排版及可下载的文档"""
    service, calls = make_preview(lambda *_: (1, None))
    service.renderer = None
    result = service.fit(None, resume_content().model_dump(), [])
    assert result["status"] == "unavailable" and result["pages"] is None
    assert result["attempts"] == 1 and calls == []


def test_boolean_page_count_is_not_a_successful_measurement(make_preview):
    """渲染器返回布尔值时不能把成功标志当作真实一页"""
    service, _ = make_preview(lambda *_: (True, None))
    result = service.fit(None, resume_content().model_dump(), [])
    assert result["status"] == "unavailable" and result["options"] == {}


@pytest.mark.parametrize("ending", ["cancel", "deadline"])
def test_late_cancellation_or_deadline_removes_all_unpublished_candidates(make_preview, ending):
    """最终候选完成后也核验取消及总截止，不能返回过期成功"""
    cancelled = threading.Event()

    def pages(_, number):
        """第三轮达到预算的同时结束请求，验证发布屏障"""
        if number == 3:
            if ending == "cancel":
                cancelled.set()
            else:
                current_document_work().deadline = time.monotonic() - 1
        return (1 if number == 3 else 2), None

    service, calls = make_preview(pages)
    with pytest.raises(Cancelled if ending == "cancel" else ProviderError):
        service.fit(None, resume_content().model_dump(), [], cancelled=cancelled)
    assert len(calls) == 3 and service.requests == {} and service.pool.entries == {}
    assert service.cache == {} and not list(service.pool.workspace.rglob("resume.docx"))


def test_fit_waiting_for_render_lock_obeys_original_deadline(make_preview):
    """排队计入同一次总预算，超时前尚未冻结资料或启动分页"""
    service, calls = make_preview(lambda *_: (1, None))
    service.lock.acquire()
    try:
        with pytest.raises(ProviderError, match="预算"):
            service.fit(None, resume_content().model_dump(), [], timeout=0.01)
    finally:
        service.lock.release()
    assert calls == [] and service.catalog.reads == 0 and service.requests == {}


@pytest.mark.parametrize("initial_budget", [10, 660])
def test_existing_request_scope_keeps_original_start_and_shorter_deadline(
    make_preview, monkeypatch, initial_budget
):
    """外部调用复用较长请求时仍从准入起限制压页时间，较短截止保持不变"""
    clock = [0.0]
    monkeypatch.setattr(documents, "time", SimpleNamespace(monotonic=lambda: clock[0]))
    service, calls = make_preview(lambda *_: (1, None))
    with service.request_work(timeout=initial_budget) as work:
        limit = min(initial_budget, FIT_TIMEOUT_SECONDS)
        clock[0] = limit + 1
        with pytest.raises(ProviderError, match="预算"):
            service.fit(None, resume_content().model_dump(), [])
        assert work.deadline == limit
    assert calls == [] and service.catalog.reads == 0 and service.requests == {}


def test_fit_route_disconnect_waits_for_actual_cleanup():
    """压页路由复用取消屏障，断开后等待真正的子进程清理"""
    entered, cancelled_seen, cleanup = threading.Event(), threading.Event(), threading.Event()
    results, timeouts = [], []

    class Service:
        """受控同步压页替身"""

        def request_work(self, cancelled, *, timeout):
            """记录线程池排队前设置的压页总预算"""
            timeouts.append(timeout)
            return document_work(cancelled, timeout=timeout, max_bytes=100)

        def fit(self, *args, cancelled, **kwargs):
            """收到断开仍必须等待实际资源回收"""
            entered.set()
            assert cancelled.wait(2)
            cancelled_seen.set()
            assert cleanup.wait(2)
            raise Cancelled("合成断开")

    async def scenario():
        """验证路由未在同步清理结束前返回"""

        async def disconnected():
            """开始执行后模拟页面离开"""
            return entered.is_set()

        async def invoke():
            """保存路由结果的异常类型"""
            try:
                await fit_resume(
                    Service(),
                    ResumePreviewInput(document=resume_content(), items=[]),
                    SimpleNamespace(headers={}, is_disconnected=disconnected),
                )
            except Exception as exc:
                results.append(type(exc))

        with anyio.fail_after(3):
            async with anyio.create_task_group() as group:
                group.start_soon(invoke)
                while not cancelled_seen.is_set():
                    await anyio.sleep(0.01)
                assert results == []
                cleanup.set()

    anyio.run(scenario)
    assert results == [Cancelled] and timeouts == [FIT_TIMEOUT_SECONDS]
