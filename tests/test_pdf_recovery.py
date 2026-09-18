"""使用脱敏合成 PDF 验证原生版面、混合页、图文分离及 Word 路径隔离。"""

from io import BytesIO
from threading import Event

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops
from test_template_analysis import simple_document
from test_template_entry_layout import record_content, record_plan
from test_template_recovery import RecoveryProvider

from resume_maker.core.errors import Problem
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.providers.base import Cancelled
from resume_maker.integrations.word.pdf_assets import extract_assets, separate_bullets
from resume_maker.integrations.word.pdf_flow import text_counter
from resume_maker.integrations.word.pdf_recovery import native_text, normalized_page, rebuild_pdf
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage
from resume_maker.integrations.word.template_recovery import prepare_template


def synthetic_pdf(path, *, rotate=0):
    """创建文字、局部图标、白字底块、图片和三列资料，不使用用户文件或固定节点号。"""
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        page.insert_text((190, 45), "EXAMPLE NAME", fontsize=19, fontname="hebo")
        page.draw_rect((25, 75, 145, 100), fill=(0.1, 0.2, 0.3), color=None)
        page.insert_text((31, 92), "EXPERIENCE", color=(1, 1, 1), fontsize=13)
        page.draw_circle((30, 55), 4, color=(0, 0, 0), width=1)
        page.draw_line((26, 55), (34, 55), color=(0, 0, 0))
        page.insert_text((41, 59), "contact@example.test", fontsize=10)
        for x, text in [(25, "Example Company"), (230, "Engineer"), (385, "2024 - 2026")]:
            page.insert_text((x, 124), text, fontsize=11)
        for index in range(3):
            y = 150 + index * 20
            page.draw_circle((30, y - 3), 1.7, fill=(0, 0, 0), color=None)
            page.insert_text(
                (40, y), f"Item {index + 1}: editable project description.", fontsize=11
            )
        raw = BytesIO()
        Image.new("RGB", (60, 75), (190, 40, 70)).save(raw, format="PNG")
        page.insert_image((420, 22, 468, 82), stream=raw.getvalue())
        page.set_rotation(rotate)
        pdf.save(path)


def forbidden_fallback(*args):
    """原生夹具不得意外进入有损视觉恢复。"""
    raise AssertionError("原生 PDF 不应触发视觉恢复")


def quiet(*args):
    """忽略测试进度，避免测试输出包含原文资料。"""


def test_native_preserves_text_assets_and_editable_fields(tmp_path):
    """原生文字不经 OCR，装饰不含旧标题，改写字段后原素材仍保留且源文件不变。"""
    source, output = tmp_path / "source.pdf", tmp_path / "result.docx"
    synthetic_pdf(source)
    original = source.read_bytes()
    notices = rebuild_pdf(source, output, Event(), quiet, forbidden_fallback)
    package = TemplatePackage(output)
    inventory = package.inventory()
    assert not inventory["warnings"]
    rows = [row for row in inventory["nodes"] if row["kind"] == "p" and row["text"]]
    text = "".join(row["text"] for row in rows)
    with pymupdf.open(source) as pdf:
        assert not (text_counter(pdf[0].get_text()) - text_counter(text))
    assert "1 页保留" in notices[0]
    assert "Item 1" in text and "Item 3" in text
    contact = next(p for p in Document(output).paragraphs if "contact@example.test" in p.text)
    assert contact.text.startswith("contact@example.test")
    pictures = [row for row in inventory["nodes"] if row["kind"] == "image"]
    assert len(pictures) >= 3
    root = package.parts["word/document.xml"]
    for anchor in root.iter(qn("wp:anchor")):
        assert anchor.get("behindDoc") == "1"
        assert anchor.find(qn("wp:positionV")).get("relativeFrom") == "paragraph"
        assert anchor.find(qn("wp:positionH")).get("relativeFrom") == "column"
    name = next(row for row in rows if "EXAMPLE NAME" in row["text"])
    plan = TemplatePlan(
        summary="原生字段",
        fields=[TextBinding(node=name["id"], quote="EXAMPLE NAME", target="personal.name")],
        repeats=[],
        photos=[],
        remove=[],
        keep=[
            row["id"]
            for row in inventory["nodes"]
            if row["id"] != name["id"] and row["kind"] in {"p", "image"}
        ],
        warnings=[],
    )
    filled = tmp_path / "filled.docx"
    fill_template(output, filled, plan, simple_document().model_dump(), [])
    filled_package = TemplatePackage(filled)
    filled_text = "".join(
        row["text"] for row in filled_package.inventory()["nodes"] if row["kind"] == "p"
    )
    assert "EXAMPLE NAME" not in filled_text and "新的用户资料" in filled_text
    assert len(
        [row for row in filled_package.inventory()["nodes"] if row["kind"] == "image"]
    ) == len(pictures)
    assert source.read_bytes() == original


def test_visual_layer_has_no_baked_in_heading(tmp_path):
    """深色标题底图中不残留白色旧文字；重复项目符号转换为文字而非图片。"""
    source = tmp_path / "assets.pdf"
    synthetic_pdf(source)
    with pymupdf.open(source) as pdf:
        assets, bullets = separate_bullets(pdf[0], extract_assets(pdf[0]))
    assert len(bullets) == 3
    background = next(raw for box, photo, raw in assets if not photo and box.width > 100)
    image = Image.open(BytesIO(background)).convert("RGB")
    inset = image.crop((12, 12, image.width - 12, image.height - 12))
    assert max(high for _, high in inset.getextrema()) < 100


def test_mixed_pages_keep_order_sizes_and_media(tmp_path):
    """原生、扫描和旋转页面逐页分流，合并后文字顺序、分节尺寸及素材关系都有效。"""
    native = tmp_path / "native.pdf"
    synthetic_pdf(native)
    source = tmp_path / "mixed.pdf"
    with pymupdf.open(native) as page_source, pymupdf.open() as pdf:
        pdf.insert_pdf(page_source)
        scan = pdf.new_page(width=400, height=600)
        scan.insert_image(scan.rect, stream=page_source[0].get_pixmap().tobytes("png"))
        page = pdf.new_page(width=300, height=500)
        page.insert_text((70, 450), "FINAL NATIVE PAGE", rotate=90)
        page.set_rotation(90)
        pdf.save(source)
    recovered = []

    def fallback(document, page, number):
        """扫描页只识别一次，并保留它在混合文档中的真实页号。"""
        recovered.append(number)
        document.add_paragraph("SCANNED PAGE TWO")
        return []

    output = tmp_path / "mixed.docx"
    rebuild_pdf(source, output, Event(), quiet, fallback)
    assert recovered == [2]
    document = Document(output)
    text = "".join(item.text or "" for item in document._element.iter(qn("w:t")))
    assert (
        text.index("EXAMPLE NAME")
        < text.index("SCANNED PAGE TWO")
        < text.index("FINAL NATIVE PAGE")
    )
    sizes = [(round(s.page_width.pt), round(s.page_height.pt)) for s in document.sections]
    assert (500, 700) in sizes and (400, 600) in sizes and (500, 300) in sizes
    package = TemplatePackage(output)
    for row in package.inventory()["nodes"]:
        if row["kind"] == "image":
            assert package.image(row["id"])


def test_late_cancel_and_failure_do_not_replace_output(tmp_path, monkeypatch):
    """本地转换结束时取消或后续页失败都不覆盖已有输出，也不留下半份可用模板。"""
    source, output = tmp_path / "source.pdf", tmp_path / "result.docx"
    synthetic_pdf(source)
    output.write_bytes(b"previous result")
    flag = Event()

    def cancel(page, bullets, symbols):
        """模拟取消发生在底层转换内部，返回结果后仍必须拒绝发布。"""
        flag.set()
        return Document(), []

    monkeypatch.setattr("resume_maker.integrations.word.pdf_recovery.convert_flow", cancel)
    with pytest.raises(Cancelled):
        rebuild_pdf(source, output, flag, quiet, forbidden_fallback)
    assert output.read_bytes() == b"previous result"


def test_converter_failure_falls_back_with_notice(tmp_path, monkeypatch):
    """版面引擎拒绝当前页时回退识别并记录原因，不能静默丢页。"""
    source, output = tmp_path / "source.pdf", tmp_path / "result.docx"
    synthetic_pdf(source)

    def broken(page, bullets, symbols):
        """模拟底层引擎不能转换当前格式。"""
        raise ValueError("unsupported geometry")

    def fallback(document, page, number):
        """视觉恢复产生可编辑内容，便于核验最终输出。"""
        document.add_paragraph("Recovered visually")
        return []

    monkeypatch.setattr("resume_maker.integrations.word.pdf_recovery.convert_flow", broken)
    notes = rebuild_pdf(source, output, Event(), quiet, fallback)
    assert "unsupported geometry" in "".join(notes)
    assert Document(output).paragraphs[0].text == "Recovered visually"


def test_scan_and_hidden_ocr_layer_are_not_native(tmp_path):
    """纯扫描及带隐藏 OCR 层的扫描页进入视觉恢复，不把整页旧简历变成底图。"""
    source = tmp_path / "source.pdf"
    synthetic_pdf(source)
    with pymupdf.open(source) as original, pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        page.insert_image(page.rect, stream=original[0].get_pixmap().tobytes("png"))
        assert not native_text(page)
        page.insert_text((50, 50), "HIDDEN OCR TEXT", render_mode=3)
        assert not native_text(page)


def test_native_word_never_enters_pdf_converter(tmp_path, monkeypatch):
    """PDF 新入口即使失效，原生 Word 的表格、图片和文字仍走原路径。"""
    source, output = tmp_path / "word.docx", tmp_path / "result.docx"
    document = Document()
    document.add_paragraph("原姓名")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Native table"
    document.save(source)
    monkeypatch.setattr(
        "resume_maker.integrations.word.pdf_recovery.rebuild_pdf", forbidden_fallback
    )
    package, _ = prepare_template(
        source, output, RecoveryProvider(), None, Event(), quiet, simple_document(), []
    )
    assert any(row["kind"] == "tbl" for row in package.inventory()["nodes"])
    assert Document(output).tables[0].cell(0, 0).text == "Native table"


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_rotation_and_crop_keep_visible_page(rotation):
    """四个方向的旋转与非零裁切框都保持逐像素一致，避免原文和裁图坐标错位。"""
    with pymupdf.open() as source:
        page = source.new_page(width=300, height=500)
        page.insert_text((40, 70), "ROTATION")
        page.draw_rect((50, 100, 200, 150), fill=(1, 0, 0))
        page.set_cropbox(pymupdf.Rect(20, 20, 280, 460))
        page.set_rotation(rotation)
        original = page.get_pixmap()
        with normalized_page(source, 0) as normalized:
            copied = normalized[0].get_pixmap()
        left = Image.frombytes("RGB", (original.width, original.height), original.samples)
        right = Image.frombytes("RGB", (copied.width, copied.height), copied.samples)
        assert left.size == right.size
        assert ImageChops.difference(left, right).getbbox() is None


def test_section_background_and_entry_decoration_follow_repeats(tmp_path):
    """PDF 标题底图经过栏目填充仍在文字后，条目局部图标随记录复制并可随空栏目收起。"""
    pdf, source, output = [tmp_path / name for name in ("source.pdf", "source.docx", "out.docx")]
    with pymupdf.open() as document:
        page = document.new_page(width=500, height=700)
        page.draw_rect((25, 25, 175, 55), fill=(0.1, 0.2, 0.3), color=None)
        page.insert_text((30, 46), "PROJECTS", fontsize=14, color=(1, 1, 1))
        page.draw_circle((20, 86), 5, color=(0.1, 0.2, 0.3))
        page.insert_text((32, 90), "Old entry", fontsize=12)
        page.insert_text((32, 115), "Old details", fontsize=11)
        document.save(pdf)
    rebuild_pdf(pdf, source, Event(), quiet, forbidden_fallback)
    package = TemplatePackage(source)
    rows = package.inventory()["nodes"]
    plan = record_plan(
        source,
        [
            (
                next(row["text"] for row in rows if row["kind"] == "p" and quote in row["text"]),
                quote,
                target,
            )
            for quote, target in [("Old entry", "title"), ("Old details", "details")]
        ],
    )
    heading = next(row for row in rows if row["kind"] == "p" and row["text"].strip() == "PROJECTS")
    plan.keep.remove(heading["id"])
    plan.keep.extend(row["id"] for row in rows if row["kind"] == "image")
    plan.fields.append(
        TextBinding(node=heading["id"], quote="PROJECTS", target="section-title:独立栏目")
    )
    content = record_content(
        [
            {"id": str(i), "title": f"Record {i}", "details": "Longer editable body " * 12}
            for i in range(3)
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    filled = TemplatePackage(output)
    paragraphs = [node for node in filled.nodes.values() if node.tag == qn("w:p")]
    for index in range(3):
        paragraph = next(
            node
            for node in paragraphs
            if "".join(t.text or "" for t in node.iter(qn("w:t"))).startswith(f"Record {index}")
        )
        assert list(paragraph.iter(qn("wp:anchor")))
    title = next(
        node
        for node in paragraphs
        if "".join(t.text or "" for t in node.iter(qn("w:t"))).strip() == "独立栏目"
    )
    assert list(title.iter(qn("wp:anchor"))) and not list(title.iter(qn("wp:inline")))
    fill_template(source, output, plan, record_content([]).model_dump(), [])
    assert not any(row["kind"] == "image" for row in TemplatePackage(output).inventory()["nodes"])


def test_blank_middle_page_and_hyperlink_are_not_lost(tmp_path):
    """空白中间页仍有分节，原生文字超链接在多页合并后仍引用正确关系。"""
    source, output = tmp_path / "source.pdf", tmp_path / "source.docx"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 50), "Example website")
        page.insert_link(
            {
                "kind": pymupdf.LINK_URI,
                "from": pymupdf.Rect(30, 35, 120, 55),
                "uri": "https://example.test",
            }
        )
        pdf.new_page()
        pdf.new_page().insert_text((30, 50), "LAST PAGE")
        pdf.save(source)

    def blank(document, page, number):
        """空白页没有识别内容，也不能因此省略该页。"""
        assert number == 2
        return []

    rebuild_pdf(source, output, Event(), quiet, blank)
    document = Document(output)
    assert len(document.sections) == 3
    links = list(document._element.iter(qn("w:hyperlink")))
    assert (
        links and document.part.rels[links[0].get(qn("r:id"))].target_ref == "https://example.test"
    )


def test_password_protected_pdf_reports_reason(tmp_path):
    """加密 PDF 直接报告密码问题，不误送 Word 转换或报告识别成功。"""
    source = tmp_path / "encrypted.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((30, 50), "Private")
        pdf.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret")
    with pytest.raises(Problem, match="加密"):
        prepare_template(
            source,
            tmp_path / "output.docx",
            RecoveryProvider(),
            None,
            Event(),
            quiet,
            simple_document(),
            [],
        )


def test_font_icon_is_kept_as_independent_asset(tmp_path):
    """字体中的电话图标按实际外观保留，不依赖导出机器安装同一字体。"""
    source, output = tmp_path / "symbols.pdf", tmp_path / "symbols.docx"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 40), "Example applicant")
        page.insert_font(fontname="Icons", fontbuffer=pymupdf.Font("cjk").buffer)
        page.insert_text((30, 80), "☎", fontname="Icons", fontsize=14)
        page.insert_text((52, 80), "123456789")
        assert "☎" in page.get_text()
        pdf.save(source)
    rebuild_pdf(source, output, Event(), quiet, forbidden_fallback)
    rows = TemplatePackage(output).inventory()["nodes"]
    assert len([row for row in rows if row["kind"] == "image"]) == 1
    text = "".join(row["text"] for row in rows if row["kind"] == "p")
    assert "123456789" in text and "Example applicant" in text and "☎" not in text


def test_colored_background_does_not_absorb_photo(tmp_path):
    """页面底色单独生成，照片仍是可独立替换的局部图片。"""
    source = tmp_path / "color.pdf"
    synthetic_pdf(source)
    with pymupdf.open(source) as pdf:
        pdf[0].draw_rect(pdf[0].rect, fill=(0.8, 0.9, 1), color=None, overlay=False)
        assets = extract_assets(pdf[0])
    backgrounds = [asset for asset in assets if asset[1] is None]
    photos = [asset for asset in assets if asset[1] is True]
    assert len(backgrounds) == 1 and len(photos) == 1
    assert photos[0][0].width < 60 and photos[0][0].height < 80
    picture = Image.open(BytesIO(backgrounds[0][2])).convert("RGB")
    assert all(low == high for low, high in picture.getextrema())


def test_stroked_horizontal_and_vertical_lines_are_preserved():
    """零面积路径按可见笔画裁图，水平分隔线和竖线不会被当作空素材丢弃。"""
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        page.draw_line((25, 65), (475, 65), color=(0.1, 0.2, 0.3), width=1)
        page.draw_line((20, 100), (20, 600), color=(0.1, 0.2, 0.3), width=1)
        assets = extract_assets(page)
    assert len(assets) == 2
    assert any(box.width >= 450 and box.height >= 1 for box, _, _ in assets)
    assert any(box.width >= 1 and box.height >= 500 for box, _, _ in assets)
    assert all(Image.open(BytesIO(raw)).getchannel("A").getbbox() for _, _, raw in assets)


def test_wrapped_list_keeps_each_marker_with_its_editable_body(tmp_path):
    """首项换行的列表仍逐项生成段落，不把圆点拆成独立列而在 Word 中留下孤立符号。"""
    source, output = tmp_path / "list.pdf", tmp_path / "list.docx"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=500, height=700)
        for baseline in (100, 133, 150):
            page.draw_circle((31, baseline - 3), 1.7, fill=(0, 0, 0), color=None)
        page.insert_text(
            (41, 100), "First item with a wrapped description and editable content.", fontsize=11
        )
        page.insert_text((41, 115), "Continuation of the first item.", fontsize=11)
        page.insert_text((41, 133), "Second item remains associated with its marker.", fontsize=11)
        page.insert_text((41, 150), "Third item remains associated with its marker.", fontsize=11)
        pdf.save(source)
    rebuild_pdf(source, output, Event(), quiet, forbidden_fallback)
    paragraphs = [
        "".join(t.text or "" for t in node.iter(qn("w:t")))
        for node in Document(output)._element.iter(qn("w:p"))
    ]
    assert not any(text.strip() == "•" for text in paragraphs)
    for first in ("First", "Second", "Third"):
        assert any(text.startswith("•") and first in text for text in paragraphs)


def test_two_columns_remain_independent_containers(tmp_path):
    """左右栏正文不能被逐行交错抄成一个段落，应保留分栏或独立表格容器。"""
    source, output = tmp_path / "columns.pdf", tmp_path / "columns.docx"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=520, height=700)
        for x, name in [(30, "LEFT"), (280, "RIGHT")]:
            page.insert_text((x, 45), name + " COLUMN", fontsize=14)
            for row in range(6):
                page.insert_text(
                    (x, 75 + row * 22), f"{name} independent paragraph {row}.", fontsize=10
                )
        pdf.save(source)
    rebuild_pdf(source, output, Event(), quiet, forbidden_fallback)
    document = Document(output)
    columns = list(document._element.iter(qn("w:cols")))
    assert any(int(node.get(qn("w:num"), "1")) >= 2 for node in columns) or any(
        len(table.columns) >= 2 for table in document.tables
    )
    paragraphs = [
        "".join(t.text or "" for t in node.iter(qn("w:t")))
        for node in document._element.iter(qn("w:p"))
    ]
    assert all(not ("LEFT" in text and "RIGHT" in text) for text in paragraphs)
    assert sum("independent paragraph" in text for text in paragraphs) >= 2
