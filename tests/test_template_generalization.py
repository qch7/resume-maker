"""用独立构造的不同语言、列数和容器验证适配规则，不使用用户简历或固定节点编号。"""

import json
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt
from lxml import etree
from test_template_analysis import TemplateProvider, completed, simple_document, simple_template
from test_template_mapping import photo_bytes

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text
from resume_maker.integrations.word.template_supplement import supplement_personal_fields
from resume_maker.integrations.word.template_visuals import layout_context
from resume_maker.services.templates import Templates


def mapping(package, fields=(), repeats=(), keep=()):
    """按真实原文查找节点，额外空段落或不同字体造成的编号变化不影响测试方案。"""
    ids = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p")
    }
    return TemplatePlan(
        summary="独立测试模板",
        fields=[
            TextBinding(node=ids[text], quote=quote, target=target)
            for text, quote, target in fields
        ],
        repeats=list(repeats),
        keep=list(keep),
        photos=[],
        remove=[],
        warnings=[],
    )


@pytest.mark.parametrize("columns", [2, 3, 4])
@pytest.mark.parametrize("padding,language", [(0, "中文"), (9, "English")])
def test_explicit_columns_do_not_depend_on_email_order_or_node_numbers(
    tmp_path, columns, padding, language
):
    """不同语言与列数只复用明确空列，邮箱可以在任意一行，字段顺序和编号均不是规则。"""
    source, snapshot, output = [
        tmp_path / name for name in ("source.docx", "snapshot.docx", "output.docx")
    ]
    doc = Document()
    for _ in range(padding):
        doc.add_paragraph()
    paragraph = doc.add_paragraph()
    for column in range(1, columns):
        paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(column * 110))
    keys = ["email", "website", "phone", "name", "gpa"][: columns + 1]
    labels = [f"{language}{index}：OLD{index}" for index in range(columns + 1)]
    paragraph.add_run("\t".join(labels[:columns]) + "\n" + labels[-1]).font.name = "Arial"
    doc.save(source)
    original = source.read_bytes()
    package = TemplatePackage(source)
    text = paragraph.text.replace("\t", "").replace("\n", "")
    plan = mapping(
        package, [(text, f"OLD{index}", f"personal.{key}") for index, key in enumerate(keys)]
    )
    document = simple_document()
    document.personal.name = ""
    for index, key in enumerate(keys):
        setattr(document.personal, key, f"NEW{index}")
    document.personal.location = "New City"
    package, plan, notices = supplement_personal_fields(package, plan, document, [], snapshot)
    assert notices and package.review(plan)["ready"]
    fill_template(snapshot, output, plan, document.model_dump(), [])
    filled = Document(output)
    assert len(filled.paragraphs) == padding + 1
    assert filled.paragraphs[-1].text.endswith(
        "\n" + labels[-1].replace(f"OLD{columns}", f"NEW{columns}") + "\t所在地：New City"
    )
    assert len(filled.paragraphs[-1].paragraph_format.tab_stops) == columns - 1
    assert source.read_bytes() == original


@pytest.mark.parametrize("fixed_height", [False, True])
def test_contact_supplement_respects_table_growth_constraints(tmp_path, fixed_height):
    """可扩展单元格内补齐资料，固定行高时移到表格后，既不裁切新字段也不改表格宽高。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name: Old Name"
    table.cell(0, 1).text = "Tel: Old Phone"
    if fixed_height:
        table.rows[0].height = Pt(18)
        table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    doc.save(source)
    package = TemplatePackage(source)
    plan = mapping(
        package,
        [
            ("Name: Old Name", "Old Name", "personal.name"),
            ("Tel: Old Phone", "Old Phone", "personal.phone"),
        ],
    )
    document = simple_document()
    document.personal.phone, document.personal.location = "12345", "Somewhere"
    package, plan, notices = supplement_personal_fields(package, plan, document, [], source)
    assert notices and package.review(plan)["ready"]
    fill_template(source, output, plan, document.model_dump(), [])
    filled = Document(output)
    location = "所在地：Somewhere"
    assert (location in filled.tables[0].cell(0, 1).text) == (not fixed_height)
    assert (location in "\n".join(p.text for p in filled.paragraphs)) == fixed_height
    if fixed_height:
        assert filled.tables[0].rows[0].height_rule == WD_ROW_HEIGHT_RULE.EXACTLY
        assert filled.tables[0].rows[0].height == Pt(18)


def test_icon_labelled_blank_contact_is_not_relocated(tmp_path):
    """图标标识的空白城市位置是模板设计的一部分，不能仅凭缺少文字标签移动它。"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_picture(BytesIO(photo_bytes(40)))
    doc.add_paragraph()
    doc.add_paragraph("E-mail: old@example.test")
    doc.save(source)
    package = TemplatePackage(source)
    plan = mapping(
        package,
        [
            ("", "", "personal.location"),
            ("E-mail: old@example.test", "old@example.test", "personal.email"),
        ],
        keep=[node["id"] for node in package.inventory()["nodes"] if node["kind"] == "image"],
    )
    document = simple_document()
    document.personal.name = ""
    document.personal.email, document.personal.location = "new@example.test", "Somewhere"
    before = source.read_bytes()
    _, same, notices = supplement_personal_fields(package, plan, document, [], source)
    assert same == plan and not notices and source.read_bytes() == before


@pytest.mark.parametrize("columns", [2, 3])
@pytest.mark.parametrize("hide_first", [False, True])
def test_independent_sidebar_sections_preserve_columns_and_clear_only_their_own_content(
    tmp_path, columns, hide_first
):
    """左右栏各自重复和隐藏，跨栏排序不拆散容器，长正文、装饰与不同列宽继续保留。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=columns)
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = Pt(120 + index * 25)
        etree.SubElement(cell._tc.get_or_add_tcPr(), w("shd")).set(w("fill"), "E3EEF7")
        cell.text = f"Area {index}"
        cell.add_paragraph(f"Old title {index}")
        cell.add_paragraph(f"Old details {index}")
    doc.save(source)
    original = source.read_bytes()
    package = TemplatePackage(source)
    ids = {row["text"]: row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p"}
    plan = mapping(
        package,
        [
            (f"Area {index}", f"Area {index}", f"section-title:栏目 {index}")
            for index in range(columns)
        ],
        repeats=[
            RepeatBinding(
                section=f"栏目 {index}",
                start=ids[f"Old title {index}"],
                end=ids[f"Old details {index}"],
                sample_start=ids[f"Old title {index}"],
                sample_end=ids[f"Old details {index}"],
                fields=[
                    TextBinding(
                        node=ids[f"Old title {index}"], quote=f"Old title {index}", target="title"
                    ),
                    TextBinding(
                        node=ids[f"Old details {index}"],
                        quote=f"Old details {index}",
                        target="details",
                    ),
                ],
            )
            for index in range(columns)
        ],
    )
    document = ResumeDocument(
        sections=[
            {
                "id": str(index),
                "title": f"栏目 {index}",
                "kind": "text",
                "visible": not (hide_first and index == 0),
                "entries": [
                    {
                        "id": str(record),
                        "title": f"New {index}.{record}",
                        "details": "Long text " * 15 + "\nSecond line",
                    }
                    for record in range(3)
                ],
            }
            for index in reversed(range(columns))
        ]
        + [{"id": "projects", "title": "Projects", "kind": "projects"}]
    )
    notices = fill_template(source, output, plan, document.model_dump(), [])
    filled = Document(output)
    assert any("各容器内部" in notice for notice in notices)
    assert len(filled.tables) == 1 and len(filled.tables[0].columns) == columns
    for index, cell in enumerate(filled.tables[0].rows[0].cells):
        assert cell.width == Pt(120 + index * 25)
        assert cell._tc.find("w:tcPr/w:shd", NS).get(w("fill")) == "E3EEF7"
        assert "Old " not in cell.text
        if hide_first and index == 0:
            assert not cell.text.strip()
        else:
            assert cell.text.startswith(f"栏目 {index}") and cell.text.count("Second line") == 3
            for record in range(3):
                assert cell.text.count(f"New {index}.{record}") == 1
    assert source.read_bytes() == original


@pytest.mark.parametrize("relative", ["paragraph", "line"])
def test_relative_drawing_stays_with_its_text_anchor(tmp_path, relative):
    """相对段落或行的图形不能迁到新空段落，否则同样的坐标会产生不同的实际位置。"""
    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Old Name")
    picture = paragraph.add_run().add_picture(BytesIO(photo_bytes(30)))
    anchor = picture._inline
    anchor.tag = f"{{{wp}}}anchor"
    for axis, reference in (("H", "column"), ("V", relative)):
        position = etree.SubElement(anchor, f"{{{wp}}}position{axis}", relativeFrom=reference)
        etree.SubElement(position, f"{{{wp}}}posOffset").text = "182880"
    doc.save(source)
    package = TemplatePackage(source)
    plan = mapping(
        package,
        [("Old Name", "Old Name", "personal.name")],
        keep=[node["id"] for node in package.inventory()["nodes"] if node["kind"] == "image"],
    )
    fill_template(source, output, plan, simple_document().model_dump(), [])
    filled = Document(output)
    assert len(filled.paragraphs) == 1
    assert filled.paragraphs[0].text == simple_document().personal.name
    assert filled.paragraphs[0]._p.find(f".//{{{wp}}}positionV").get("relativeFrom") == relative


def test_source_pages_and_structural_controls_reach_the_model_without_current_values(
    catalog, tmp_path, monkeypatch
):
    """整页版式、换行与容器证据一起送入识别，附件数量有界，当前资料值不被发送。"""
    source = tmp_path / "source.docx"
    simple_template(source)

    def render(document, pdf):
        """生成八页独立模板图，用来核验附件顺序与未展示页说明。"""
        with pymupdf.open() as output:
            for number in range(8):
                output.new_page().insert_text((40, 40), f"Template page {number + 1}")
            output.save(pdf)
        return None

    monkeypatch.setattr("resume_maker.integrations.word.template_visuals.word_process", render)
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    request = json.loads(provider.calls[0]["prompt"].split("\n")[-1])
    assert request["source_pages"]["total"] == 8 and request["source_pages"]["omitted"] == 2
    assert len(provider.calls[0]["images"]) == 6
    assert all(image.read_bytes().startswith(b"\x89PNG") for image in provider.calls[0]["images"])
    assert simple_document().personal.name not in provider.calls[0]["prompt"]
    assert task["review"]["ready"] and any("后续页" in text for text in task["review"]["notices"])
    doc = Document()
    paragraph = doc.add_paragraph("Email: old\tTel: old\nOther")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(130))
    doc.save(source)
    package = TemplatePackage(source)
    evidence = next(iter(layout_context(package)["paragraphs"].values()))
    assert evidence["controls"] == [{"kind": "tab", "offset": 10}, {"kind": "br", "offset": 18}]
    assert evidence["tabs"][0]["pos"] == "2600"


def test_failed_structural_repair_restores_matching_best_snapshot_and_resets_model_context(
    catalog, tmp_path
):
    """补齐位置改变编号后重发清单；后续失败时将最佳方案及其源快照一起恢复。"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("原姓名")
    doc.add_paragraph("教育背景 / 项目经历")
    doc.save(source)
    original = source.read_bytes()
    original_inventory = TemplatePackage(source).inventory()["nodes"]

    class ChangingProvider(TemplateProvider):
        """先遗漏整段标题，再补齐但留下真实栏目冲突，最后一次修正模拟服务失败。"""

        def run_structured(self, **kwargs):
            """记录新会话收到的映射和节点关系，确保模型不会混用旧节点编号。"""
            plan = super().run_structured(**kwargs)
            kwargs["emit"]("thread", {"id": f"session-{len(self.calls)}"})
            package = TemplatePackage(kwargs["workspace"] / "original.docx")
            if len(self.calls) == 3:
                request = json.loads(kwargs["prompt"].split("\n")[-1])
                assert kwargs["thread_id"] is None
                assert request["template"]["parts"] and request["layout"]["paragraphs"]
                assert package.review(TemplatePlan.model_validate(request["previous_plan"]))[
                    "ready"
                ]
                assert any("栏目边界" in error for error in request["validation"]["errors"])
                raise RuntimeError("第三轮模拟失败")
            if len(self.calls) == 2:
                heading = next(
                    row
                    for row in package.inventory()["nodes"]
                    if row["text"] == "教育背景 / 项目经历"
                )
                plan.fields.extend(
                    TextBinding(node=heading["id"], quote=title, target="section-title:" + title)
                    for title in ("教育背景", "项目经历")
                )
            return plan

    document = simple_document()
    document.personal.location = "New City"
    provider = ChangingProvider()
    service = Templates(catalog, tmp_path, provider)
    task = completed(service, service.analyze(source, document)["id"])
    assert task["status"] == "completed" and not task["review"]["ready"]
    assert task["repair_error"] == "第三轮模拟失败" and len(provider.calls) == 3
    assert "personal.location" in task["review"]["missing"]
    assert task["inventory"]["nodes"] == original_inventory
    assert "〔待填写〕" not in "".join(row["text"] for row in task["inventory"]["nodes"])
    assert len(task["plan"]["fields"]) == 1
    assert source.read_bytes() == original


def test_cached_mapping_reenters_repair_when_current_trial_fails(catalog, tmp_path, monkeypatch):
    """缓存结构有效但当前试填失败时继续自动识别，不能直接返回不可用的缓存结果。"""
    from resume_maker.services.template_analysis import check_trial

    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    first = completed(service, service.analyze(source, simple_document())["id"])
    assert first["review"]["ready"]
    calls = []

    def fail_cached_trial_once(source, plan, review, document, projects):
        """模拟只有缓存试填遇到的新限制，下一轮按正常填充器验证修正结果。"""
        calls.append(plan)
        if len(calls) == 1:
            review["errors"].append("当前资料触发新的排版约束")
            review["ready"] = False
            return review
        return check_trial(source, plan, review, document, projects)

    monkeypatch.setattr("resume_maker.services.templates.check_trial", fail_cached_trial_once)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["review"]["ready"] and not task["reused"] and len(provider.calls) == 2
