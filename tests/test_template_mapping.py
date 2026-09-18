"""陌生模板的完整填充、样式保留、旧信息清理与错误映射边界"""

import base64
from io import BytesIO
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from lxml import etree

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import NS, TemplatePackage, paragraph_text
from resume_maker.integrations.word.templates.values import missing_targets


def photo_bytes(color):
    """生成无个人信息的纯色照片；验证包内图片是否真正更换"""
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 30, 40), False)
    pixmap.clear_with(color)
    return pixmap.tobytes("png")


def resume_content():
    """构造包含个人资料、教育、项目和隐藏字段的完整脱敏资料"""
    return ResumeDocument.model_validate(
        {
            "personal": {
                "name": "测试新姓名",
                "phone": "10000000000",
                "email": "new@example.test",
                "job_title": "文档开发工程师",
                "location": "测试城市",
                "photo": "data:image/png;base64," + base64.b64encode(photo_bytes(180)).decode(),
                "custom_fields": [{"id": "custom", "label": "语言", "value": "中文"}],
            },
            "sections": [
                {
                    "id": "education",
                    "title": "教育背景",
                    "kind": "education",
                    "entries": [
                        {
                            "id": "school1",
                            "title": "新大学一",
                            "subtitle": "计算机 · 本科",
                            "period": "2020–2024",
                            "details": "研究文档结构\n完成排版系统",
                        },
                        {
                            "id": "school2",
                            "title": "新大学二",
                            "subtitle": "计算机 · 硕士",
                            "period": "2024–2026",
                            "details": "构建自动填充工具",
                        },
                    ],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
            ],
        }
    )


def project_content():
    """只选择一条固定项目亮点以防把未选择的说明写入 Word"""
    return [
        {
            "highlight_ids": ["chosen"],
            "content": {
                "title": "新文档项目",
                "period": "2025",
                "role": "开发",
                "stack": ["Python"],
                "description": "生成可编辑文档",
                "highlights": [
                    {"id": "chosen", "title": "自动映射", "text": "覆盖姓名和教育"},
                    {"id": "skipped", "title": "不应输出", "text": "未选中的亮点"},
                ],
            },
        }
    ]


def make_template(path, *, decoration=False):
    """建立含表格、跨样式文字、页眉页脚、文本框和照片的陌生 DOCX"""
    doc = Document()
    doc.sections[0].page_width = Cm(21)
    doc.sections[0].page_height = Cm(29.7)
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(10)
    name = doc.add_paragraph()
    name.add_run("旧").bold = True
    name.add_run("姓名").italic = True
    name.add_run(" · 个人简历")
    doc.add_picture(BytesIO(photo_bytes(20)), width=Cm(1.5), height=Cm(2))
    doc.sections[0].header.paragraphs[0].text = "电话：旧电话 | 邮箱：旧邮箱"
    doc.sections[0].footer.paragraphs[0].text = "旧城市"
    box = etree.fromstring(
        f'<w:p xmlns:w="{NS["w"]}" xmlns:v="{NS["v"]}"><w:r><w:pict>'
        '<v:shape id="TextBox1" style="width:220pt;height:22pt" stroked="f">'
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>旧意向</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    doc._element.body.insert(2, box)
    doc.add_paragraph("教育背景", "Heading 1")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for row in table.rows:
        for cell, text in zip(row.cells, ["旧大学", "旧专业", "旧时间"], strict=True):
            cell.text = text
        row.cells[0].add_paragraph("旧教育正文")
    if decoration:
        table.rows[0].cells[0].paragraphs[0].add_run().add_picture(
            BytesIO(photo_bytes(20)), width=Cm(0.3)
        )
    doc.add_paragraph("项目经历", "Heading 1")
    doc.add_paragraph("旧项目 · 旧项目时间").runs[0].bold = True
    doc.add_paragraph("旧项目正文\n旧第二行")
    doc.add_paragraph("旧项目二")
    doc.add_paragraph("旧多余说明")
    doc.add_paragraph("")
    doc.save(path)
    package = TemplatePackage(path)

    def locate(text):
        """用测试原文定位首次出现的段落且不依赖 XML 节点编号的具体数值"""
        return next(
            key
            for key, node in package.nodes.items()
            if node.tag == w("p") and paragraph_text(node) == text
        )

    def bind(text, target, quote=None):
        """从明确样本位置创建可审查映射"""
        return TextBinding(node=locate(text), quote=text if quote is None else quote, target=target)

    rows = [key for key, node in package.nodes.items() if node.tag == w("tr")]
    empty = [
        key
        for key, node in package.nodes.items()
        if node.tag == w("p") and not len(node) and package.locations[key] == "word/document.xml"
    ][-1]
    plan = TemplatePlan(
        summary="测试完整模板",
        warnings=[],
        fields=[
            bind("旧姓名 · 个人简历", "personal.name", "旧姓名"),
            bind("电话：旧电话 | 邮箱：旧邮箱", "personal.phone", "旧电话"),
            bind("电话：旧电话 | 邮箱：旧邮箱", "personal.email", "旧邮箱"),
            bind("旧城市", "personal.location"),
            bind("旧意向", "personal.job_title"),
            TextBinding(node=empty, quote="", target="personal.custom_fields"),
        ],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=rows[0],
                end=rows[1],
                sample_start=rows[0],
                sample_end=rows[0],
                fields=[
                    bind("旧大学", "title"),
                    bind("旧专业", "subtitle"),
                    bind("旧时间", "period"),
                    bind("旧教育正文", "details"),
                ],
            ),
            RepeatBinding(
                section="projects",
                start=locate("旧项目 · 旧项目时间"),
                end=locate("旧项目二"),
                sample_start=locate("旧项目 · 旧项目时间"),
                sample_end=locate("旧项目正文旧第二行"),
                fields=[
                    bind("旧项目 · 旧项目时间", "title", "旧项目"),
                    bind("旧项目 · 旧项目时间", "period", "旧项目时间"),
                    bind("旧项目正文旧第二行", "details"),
                ],
            ),
        ],
        photos=[
            node["id"]
            for node in package.inventory()["nodes"]
            if node["kind"] == "image" and not list(package.node(node["id"]).iterancestors(w("tr")))
        ],
        keep=[locate("教育背景"), locate("项目经历")],
        remove=[locate("旧多余说明")],
    )
    return package, plan


def test_complete_fill_keeps_layout_and_removes_old_data(tmp_path):
    """整份替换所有容器、重复记录和照片；同时保留样式、布局与包资源"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    package, plan = make_template(source)
    assert package.review(plan)["ready"]
    fill_template(source, output, plan, resume_content().model_dump(), project_content())
    actual = TemplatePackage(output)
    text = "".join(paragraph_text(node) for node in actual.nodes.values() if node.tag == w("p"))
    for expected in [
        "测试新姓名",
        "新大学一",
        "新大学二",
        "new@example.test",
        "10000000000",
        "测试城市",
        "文档开发工程师",
        "新文档项目",
        "覆盖姓名和教育",
        "语言：中文",
    ]:
        assert expected in text
    assert "旧" not in text and "不应输出" not in text
    new_name = next(
        node
        for node in actual.nodes.values()
        if node.tag == w("p") and paragraph_text(node).startswith("测试新姓名")
    )
    assert new_name.find("w:r/w:rPr/w:b", NS) is not None
    assert len(actual.parts["word/document.xml"].xpath(".//w:tr", namespaces=NS)) == 2
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        assert before.read("word/theme/theme1.xml") == after.read("word/theme/theme1.xml")
        assert "word/media/image1.png" not in after.namelist()
        assert photo_bytes(20) not in [after.read(name) for name in after.namelist()]
        assert photo_bytes(180) in [after.read(name) for name in after.namelist()]
    assert package.review(plan)["ready"]  # 导出不能修改模板与映射


def test_hidden_rows_fields_and_photos_do_not_reappear(tmp_path):
    """隐藏资料清空原位置；空重复区可移除全部表格行而不损坏文档"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    _, plan = make_template(source)
    document = resume_content()
    document.personal.hidden_fields = ["phone", "photo"]
    document.sections[0].visible = False
    document.sections[1].visible = False
    fill_template(source, output, plan, document.model_dump(), project_content())
    actual = TemplatePackage(output)
    all_text = "".join(
        root.xpath(".//w:t/text()", namespaces=NS)[0] for root in actual.parts.values()
    )
    assert "旧" not in all_text and "10000000000" not in all_text
    assert not actual.parts["word/document.xml"].xpath(".//w:tbl", namespaces=NS)
    assert not any(name.startswith("word/media/") for name in actual.files)
    Document(output)


@pytest.mark.parametrize(
    "case", ["quote", "overlap", "sample", "unknown", "keep_remove", "photo_duplicate"]
)
def test_invalid_mapping_is_rejected(tmp_path, case):
    """错误位置、引文、范围与相互矛盾的分类不能产生导出文件"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source)
    if case == "quote":
        plan.fields[0].quote = "不存在的文字"
    elif case == "overlap":
        plan.fields.append(plan.repeats[0].fields[0].model_copy(update={"target": "personal.name"}))
    elif case == "sample":
        plan.repeats[0].sample_start = plan.repeats[0].sample_end = plan.repeats[1].start
    elif case == "unknown":
        plan.fields[0].target = "personal.password"
    elif case == "keep_remove":
        plan.remove.append(plan.keep[0])
    else:
        plan.photos *= 2
    assert package.review(plan)["errors"]
    with pytest.raises(Problem):
        fill_template(source, tmp_path / "bad.docx", plan, resume_content().model_dump(), [])
    assert not (tmp_path / "bad.docx").exists()


def test_unresolved_and_missing_current_information_block_export(tmp_path):
    """原文未处理或当前新填资料没有位置时；明确阻止遗漏信息的导出"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source)
    plan.remove = []
    assert package.review(plan)["unresolved"][0]["text"] == "旧多余说明"
    document = resume_content()
    document.personal.website = "https://example.test"
    assert "personal.website" in missing_targets(document, plan, project_content())
    document.sections[0].title = "新的教育名称"
    assert "栏目：新的教育名称" in missing_targets(document, plan, project_content())


@pytest.mark.parametrize("encoded_path", [False, True])
def test_sample_decoration_requires_explicit_confirmation(tmp_path, encoded_path):
    """重复样本图片需要确认；仍被装饰引用的同一图片资源必须保留"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source, decoration=True)
    media = "word/media/image1.png"
    if encoded_path:
        media = "word/media/image 1.png"
        package.files[media] = package.files.pop("word/media/image1.png")
        path = "word/_rels/document.xml.rels"
        package.files[path] = package.files[path].replace(
            b'Target="media/image1.png"', b'Target="/word/media/image%201.png"'
        )
        package.write(source)
    assert package.image(plan.photos[0]) == photo_bytes(20)
    assert "图片" in "".join(package.review(plan)["errors"])
    decoration = next(
        row["id"]
        for row in package.inventory()["nodes"]
        if row["kind"] == "image" and row["id"] not in plan.photos
    )
    plan.keep.append(decoration)
    assert package.review(plan)["ready"]
    output = tmp_path / "filled.docx"
    fill_template(source, output, plan, resume_content().model_dump(), project_content())
    with ZipFile(output) as archive:
        assert archive.read(media) == photo_bytes(20)
    actual = TemplatePackage(output)
    assert len([row for row in actual.inventory()["nodes"] if row["kind"] == "image"]) == 3


def test_empty_insertion_and_occurrence_rules(tmp_path):
    """空引文不能覆盖非空原文；同段多处相同值按指定次数独立替换"""
    path = tmp_path / "repeat.docx"
    doc = Document()
    doc.add_paragraph("示例 / 示例")
    doc.add_paragraph()
    doc.save(path)
    package = TemplatePackage(path)
    node = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p")
    plan = TemplatePlan(
        summary="测试",
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
        fields=[
            TextBinding(node=node, quote="示例", target="personal.name", occurrence=1),
            TextBinding(node=node, quote="示例", target="personal.phone", occurrence=2),
        ],
    )
    document = ResumeDocument(
        sections=[{"id": "p", "kind": "projects", "title": "项目"}],
        personal={"name": "甲\n乙", "phone": "电话值"},
    )
    output = tmp_path / "output.docx"
    fill_template(path, output, plan, document.model_dump(), [])
    assert Document(output).paragraphs[0].text == "甲\n乙 / 电话值"
    from resume_maker.domain.resume import CustomInfoField

    document.personal.custom_fields = [CustomInfoField(id="lang", label="语言", value="中文")]
    empty = next(row["id"] for row in package.inventory()["nodes"] if row["can_insert"])
    plan.fields.append(TextBinding(node=empty, quote="", target="personal.custom:语言"))
    fill_template(path, output, plan, document.model_dump(), [])
    assert Document(output).paragraphs[1].text == "语言：中文"
    plan.fields[0].quote = ""
    assert "空引文" in "".join(package.review(plan)["errors"])


def test_empty_text_does_not_make_a_photo_paragraph_an_insertion_slot(tmp_path):
    """无文字的图片与文本框容器不是空白段落；AI 和人工映射都不能占用它"""
    package, plan = make_template(tmp_path / "source.docx")
    image = next(row for row in package.inventory()["nodes"] if row["kind"] == "image")
    owner = next(row for row in package.inventory()["nodes"] if row["id"] == image["ancestors"][0])
    assert not owner["text"] and not owner["can_insert"]
    plan.fields[-1].node = owner["id"]
    assert "空白的段落" in "".join(package.review(plan)["errors"])


def test_dynamic_fields_are_frozen_without_restoring_old_values(tmp_path):
    """无缓存的合并域自动变为可编辑空位；指令不会恢复旧值"""
    path = tmp_path / "dynamic.docx"
    doc = Document()
    paragraph = doc.add_paragraph("姓名")
    field = OxmlElement("w:fldSimple")
    field.set(w("instr"), "MERGEFIELD OldName")
    paragraph._p.append(field)
    doc.save(path)
    package = TemplatePackage(path)
    assert not package.inventory()["warnings"]
    assert "动态域" in "".join(package.notices)
    assert not package.parts["word/document.xml"].xpath(".//w:fldSimple", namespaces=NS)


def test_only_replaced_hyperlinks_are_removed(tmp_path):
    """替换个人主页后不再跳转旧地址；同段明确保留的固定链接仍可使用"""
    source, output = tmp_path / "links.docx", tmp_path / "result.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    for text, target in [
        ("旧个人主页", "https://old.example.test"),
        ("固定说明", "https://fixed.example.test"),
    ]:
        link = OxmlElement("w:hyperlink")
        link.set(
            f"{{{NS['r']}}}id",
            doc.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
        )
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
    doc.save(source)
    package = TemplatePackage(source)
    node_id = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p")
    plan = TemplatePlan(
        summary="主页",
        warnings=[],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        fields=[
            TextBinding(node=node_id, quote="旧个人主页", target="personal.website"),
        ],
    )
    document = ResumeDocument(
        personal={"website": "https://new.example.test"},
        sections=[
            {"id": "p", "title": "项目", "kind": "projects"},
        ],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    actual = TemplatePackage(output)
    links = actual.parts["word/document.xml"].xpath(".//w:hyperlink", namespaces=NS)
    assert len(links) == 1 and "".join(links[0].itertext()) == "固定说明"
    assert b"old.example.test" not in actual.files["word/_rels/document.xml.rels"]
    assert b"fixed.example.test" in actual.files["word/_rels/document.xml.rels"]
