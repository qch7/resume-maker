"""验证整页脱敏、原文还原及真实图片恢复入口的隐私边界"""

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from threading import Event

import pytest
from docx import Document
from PIL import Image, ImageDraw, PngImagePlugin
from test_image_recovery import image_fixture
from test_privacy_mosaic import synthetic_image

from resume_maker.domain.image_layout import ImagePage
from resume_maker.domain.models import ProviderSettings
from resume_maker.domain.templates import TemplatePlan
from resume_maker.integrations.privacy import Redactor
from resume_maker.integrations.providers import page_images
from resume_maker.integrations.providers.base import Cancelled, PageImage, ProviderError
from resume_maker.integrations.providers.codex import CodexProvider
from resume_maker.integrations.word.image.layout import text_layer
from resume_maker.integrations.word.image.recovery import rebuild_image
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.services.templates.analysis import visual_evidence


def page_fixture(path, monkeypatch):
    """创建带敏感文字、白字底块和独立人像的合成页面及本地 OCR 结果"""
    layout = image_fixture(path)
    with Image.open(path) as source, Image.open(BytesIO(synthetic_image())) as photo:
        image = source.copy()
        asset = next(asset for asset in layout.assets if asset.kind == "photo")
        box = page_images.pixel_box(asset.box, image.size)
        image.paste(photo.resize((box[2] - box[0], box[3] - box[1])), box[:2])
        # 未识别的小字也不能作为原始像素从文字块以外直接外发
        ImageDraw.Draw(image).text((700, 800), "MISSED-PRIVATE-TEXT", fill="black")
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("private", "PAGE-PRIVATE-METADATA")
        image.save(path, pnginfo=metadata)
    local = {
        "pages": [
            {
                "blocks": [
                    {"text": row.text, "box": row.box, "confidence": 0.99} for row in layout.texts
                ]
            }
        ],
        "text": "\n".join(row.text for row in layout.texts),
    }
    monkeypatch.setattr(page_images, "read_document", lambda *_: local)
    return layout, local


def test_page_pixels_metadata_and_sensitive_rows_are_removed(tmp_path, monkeypatch):
    """敏感行不含原始文字像素，头像只剩粗色块，坐标和安全文字仍可用于版面恢复"""
    source = tmp_path / "private-name.png"
    layout, _ = page_fixture(source, monkeypatch)
    original = source.read_bytes()
    redactor = Redactor(["SAMPLE NAME", "OLD PHONE", "OLD EMAIL"])
    safe, context, record = page_images.sanitized_page(source, redactor, Event())
    assert source.read_bytes() == original and safe != original
    assert b"PAGE-PRIVATE-METADATA" not in safe
    assert all(value not in json.dumps(context) for value in redactor.values)
    assert record["covered_rows"] >= 3
    assert record["sha256"] == hashlib.sha256(safe).hexdigest()
    assert redactor.restore(context["blocks"][0]["text"]) == "SAMPLE NAME"
    with Image.open(BytesIO(safe)) as image:
        assert not image.info and not image.getexif()
        first = page_images.pixel_box(layout.texts[0].box, image.size)
        # 编号只占行首，其余位置由不透明遮盖替换
        assert image.getpixel((first[2] - 3, first[3] - 3)) == (83, 89, 101)
        photo = next(a for a in layout.assets if a.kind == "photo")
        tile = image.crop(page_images.pixel_box(photo.box, image.size))
        assert len(tile.getcolors(tile.width * tile.height)) <= 65
        assert context["graphics"]


def test_private_page_rebuild_restores_text_assets_and_preserves_source(tmp_path, monkeypatch):
    """生产 Provider 发送脱敏整页，返回后本机恢复原文及原图照片，覆盖首次失败重试"""
    source, output = tmp_path / "source.png", tmp_path / "recovered.docx"
    layout, _ = page_fixture(source, monkeypatch)
    original, calls = source.read_bytes(), []

    def runner(payload, *args, safe_images):
        """验证真实网关的图片和上下文，仅返回合成坐标和受保护文字"""
        context = json.loads(payload["input"].splitlines()[-1])
        calls.append(context)
        assert len(safe_images) == 1
        assert payload["images"][0]["kind"] == "sanitized-page"
        assert "SAMPLE NAME" not in payload["input"]
        assert source.name not in payload["input"]
        assert b"PAGE-PRIVATE-METADATA" not in safe_images[0]
        result = layout.model_dump()
        for row, safe in zip(result["texts"], context["blocks"], strict=True):
            row["text"] = safe["text"]
        if len(calls) == 1:
            result["texts"].pop()
        return json.dumps(result)

    provider = CodexProvider(runner=runner).with_private_data({"personal": {"name": "SAMPLE NAME"}})
    notes = rebuild_image(source, output, provider, ProviderSettings(), Event(), lambda *_: None)
    assert len(calls) == 2
    assert "OLD ROLE" in provider.sensitive_values
    assert calls[0]["blocks"][0]["text"] != calls[1]["blocks"][0]["text"]
    assert source.read_bytes() == original
    document = Document(output)
    assert "SAMPLE NAME" in "".join(p.text for p in document.paragraphs)
    assert "[[RM_" not in document._element.xml
    assert any("脱敏重绘" in note for note in notes)
    # Word 中的头像使用本机原件，不能把外发的马赛克当成最终照片
    with Image.open(source) as image:
        photo = next(a for a in layout.assets if a.kind == "photo")
        box = tuple(round(v * image.size[i % 2]) for i, v in enumerate(photo.box))
        pixels = image.crop(box).tobytes()
    assert any(
        Image.open(BytesIO(rel.target_part.blob)).tobytes() == pixels
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/image")
    )
    seen = []

    def reopened_runner(payload, *args, **kwargs):
        """新的修复任务没有旧 Provider 内存，仍须保护模板中的未知身份"""
        seen.append(payload)
        assert "SAMPLE NAME" not in payload["input"]
        assert "OLD ROLE" not in payload["input"]
        return TemplatePlan(
            summary="checked", fields=[], repeats=[], photos=[], keep=[], remove=[], warnings=[]
        ).model_dump_json()

    fresh = CodexProvider(runner=reopened_runner)
    package = TemplatePackage(output)
    visual_evidence(fresh, package, output, tmp_path, Event())
    fresh.run_structured(
        result_model=TemplatePlan,
        workspace=tmp_path,
        prompt="\n".join(row.text for row in layout.texts),
        thread_id=None,
        settings=ProviderSettings(),
        cancelled=Event(),
        emit=lambda *_: None,
    )
    assert len(seen) == 1


@pytest.mark.parametrize("failure", ["ocr", "graphics", "cancel", "empty", "multiple"])
def test_page_sanitization_failure_never_sends_original(tmp_path, monkeypatch, failure):
    """识别、打码、数量或取消校验失败时不调用供应商，也不会回退原图"""
    source = tmp_path / "source.png"
    _, local = page_fixture(source, monkeypatch)
    flag = Event()
    if failure == "empty":
        local["pages"][0]["blocks"] = []
    elif failure == "cancel":
        flag.set()
    elif failure in {"ocr", "graphics"}:

        def fail(*args):
            """模拟本地处理失败，验证请求在生成附件以前终止"""
            raise ProviderError("本地处理失败")

        monkeypatch.setattr(
            page_images, "read_document" if failure == "ocr" else "graphic_regions", fail
        )
    calls = []
    with pytest.raises((ProviderError, Cancelled)):
        CodexProvider(runner=lambda *args, **kwargs: calls.append(args)).run_structured(
            result_model=ImagePage,
            workspace=tmp_path,
            prompt="恢复图片模板",
            thread_id=None,
            settings=ProviderSettings(),
            cancelled=flag,
            emit=lambda *_: None,
            images=[PageImage(source)] * (2 if failure == "multiple" else 1),
        )
    assert not calls


def test_low_confidence_rows_and_secrets_use_same_local_redactor(tmp_path, monkeypatch):
    """不确定文字整行遮盖，凭据即使随隐私占位符返回也不能还原"""
    source = tmp_path / "source.png"
    _, local = page_fixture(source, monkeypatch)
    local["pages"][0]["blocks"][0].update(text="Uncertain Person", confidence=0.4)
    local["pages"][0]["blocks"][1]["text"] = "api_key=private-secret-canary"
    local["text"] = "\n".join(row["text"] for row in local["pages"][0]["blocks"])
    redactor = Redactor()
    _, context, _ = page_images.sanitized_page(source, redactor, Event())
    assert "Uncertain Person" not in json.dumps(context)
    assert "private-secret-canary" not in json.dumps(context)
    restored = redactor.restore(context)
    assert restored["blocks"][0]["text"] == "Uncertain Person"
    assert "private-secret-canary" not in json.dumps(restored)


def test_hyphens_keep_original_codepoints_in_pdf_and_word(tmp_path):
    """等线共享横线字形仍保留原始半角连字符，中文横线也不被统一替换"""
    from resume_maker.integrations.word.image.layout import build_image_document, font_for

    text = "2022-2026 foo-bar -- A\u2010B"
    layout = ImagePage(texts=[{"text": text, "box": [0.1, 0.1, 0.8, 0.12]}])
    font = font_for(layout.texts[0])
    layout.texts[0].box[2] = 0.1 + font.text_length(text, fontsize=12) / 595.276
    with text_layer(layout, 595.276, 842) as pdf:
        recovered = pdf[0].get_text()
        assert recovered.count("-") == text.count("-")
        assert recovered.count("\u2010") == text.count("\u2010")
    source = tmp_path / "source.png"
    Image.new("RGB", (595, 842), "white").save(source)
    document = build_image_document(source, layout, Event())
    recovered = "".join(p.text for p in document.paragraphs)
    assert recovered.count("-") == text.count("-")
    assert recovered.count("\u2010") == text.count("\u2010")


@pytest.mark.parametrize("change", ["text", "box"])
def test_model_cannot_replace_text_or_move_original_ocr_boxes(change):
    """模型不能猜测被遮盖的文字，也不能移动原框让裁图绕过文字重叠检查"""
    context = {"blocks": [{"text": "[[RM_123456_1]]", "box": [0.1, 0.1, 0.4, 0.2]}]}
    result = {"texts": deepcopy(context["blocks"])}
    if change == "text":
        result["texts"][0]["text"] = "Guessed identity"
    else:
        result["texts"][0]["box"] = [0.1, 0.5, 0.4, 0.6]
    with pytest.raises(ProviderError):
        page_images.validate_page_text(result, context)


def test_actual_character_loss_still_fails_font_verification(monkeypatch):
    """连字符修复不能放宽真实字符缺失的校验"""
    from resume_maker.core.errors import Problem
    from resume_maker.integrations.word.image import layout as image_layout

    original = ImagePage(texts=[{"text": "alpha-beta", "box": [0.1, 0.1, 0.3, 0.12]}])

    def missing_character(page, layout):
        """模拟字体或排版过程丢失了普通字母"""
        page.insert_text((50, 50), "alpha-bet", fontname="helv")

    monkeypatch.setattr(image_layout, "insert_texts", missing_character)
    with pytest.raises(Problem, match="无法完整恢复"):
        with image_layout.text_layer(original, 595.276, 842):
            pass
