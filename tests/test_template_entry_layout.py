"""用不同陌生模板验证条目对齐和隐藏字段收起且不依赖个人资料或固定坐标"""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.shared import Pt
from lxml import etree
from test_template_mapping import photo_bytes

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.entry_layout import ParagraphStyles
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text


def record_plan(path, targets, table=False):
    """按当前原文建立单条重复范围；任意空段落和样式变化都不会依赖节点编号"""
    package = TemplatePackage(path)
    nodes = {
        paragraph_text(node): key
        for key, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    fields = [
        TextBinding(node=nodes[text], quote=quote, target=target) for text, quote, target in targets
    ]
    paragraphs = list(dict.fromkeys(package.node(field.node) for field in fields))
    roots = list(
        dict.fromkeys(next(node.iterancestors(w("tr"))) if table else node for node in paragraphs)
    )
    return TemplatePlan(
        summary="独立条目布局测试",
        fields=[],
        repeats=[
            RepeatBinding(
                section="独立栏目",
                start=package.ids[roots[0]],
                end=package.ids[roots[-1]],
                sample_start=package.ids[roots[0]],
                sample_end=package.ids[roots[-1]],
                fields=fields,
            )
        ],
        keep=[key for text, key in nodes.items() if text not in {item[0] for item in targets}],
        photos=[],
        remove=[],
        warnings=[],
    )


def record_content(entries):
    """构造任意普通栏目和必需的项目区；隐藏规则与实际个人信息界面相同"""
    return ResumeDocument.model_validate(
        {
            "sections": [
                {"id": "records", "title": "独立栏目", "entries": entries},
                {"id": "projects", "kind": "projects", "title": "项目经历"},
            ]
        }
    )


@pytest.mark.parametrize("indent", [11, 37])
@pytest.mark.parametrize("table", [False, True])
def test_repeat_titles_align_to_effective_body_indent_and_keep_right_tabs(tmp_path, indent, table):
    """不同继承样式及容器中的多条记录沿用本列正文起点；日期右制表位和字体保留"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    base = doc.styles.add_style("Base body", WD_STYLE_TYPE.PARAGRAPH)
    base.paragraph_format.left_indent = Pt(indent)
    body_style = doc.styles.add_style("Inherited body", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = base
    area = doc.add_table(rows=1, cols=1).cell(0, 0) if table else doc
    header = area.add_paragraph("Old title\tOld period")
    header.paragraph_format.left_indent = Pt(3)
    header.paragraph_format.first_line_indent = Pt(4)
    header.paragraph_format.tab_stops.add_tab_stop(Pt(360), WD_TAB_ALIGNMENT.RIGHT)
    header.runs[0].bold = True
    header.runs[0].font.name = "Arial"
    area.add_paragraph("Old body", "Inherited body")
    if table:
        area._tc.remove(area.paragraphs[0]._p)
    doc.save(source)
    plan = record_plan(
        source,
        [
            ("Old titleOld period", "Old title", "title"),
            ("Old titleOld period", "Old period", "period"),
            ("Old body", "Old body", "details"),
        ],
        table,
    )
    original, original_plan = source.read_bytes(), deepcopy(plan)
    content = record_content(
        [
            {"id": str(i), "title": f"Record {i}", "period": "2026", "details": "Body content"}
            for i in range(3)
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    package = TemplatePackage(output)
    styles = ParagraphStyles(package)
    headers = [
        node
        for node in package.nodes.values()
        if node.tag == w("p") and paragraph_text(node).startswith("Record")
    ]
    assert len(headers) == 3
    for node in headers:
        assert styles.indent(node)[w("left")] == str(indent * 20)
        assert node.find("w:pPr/w:ind", NS).get(w("firstLine")) == "0"
        assert node.find("w:pPr/w:tabs/w:tab", NS).get(w("pos")) == "7200"
        assert node.find("w:pPr/w:tabs/w:tab", NS).get(w("val")) == "right"
        assert node.find("w:r/w:rPr/w:b", NS) is not None
        assert node.find("w:r/w:rPr/w:rFonts", NS).get(w("ascii")) == "Arial"
    assert source.read_bytes() == original and plan == original_plan


@pytest.mark.parametrize("kind", ["direct", "inherited", "literal", "label"])
@pytest.mark.parametrize("table", [False, True])
def test_hidden_details_remove_list_labels_and_empty_rows(tmp_path, kind, table):
    """直接编号、继承列表、手输圆点和中英标签随空字段整体移除；恢复后可再次显示"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    if table:
        grid = doc.add_table(rows=3, cols=1)
        paragraphs = [row.cells[0].paragraphs[0] for row in grid.rows]
    else:
        paragraphs = [doc.add_paragraph() for _ in range(3)]
    paragraphs[0].text = "Old title"
    literal = {"literal": "• Old details", "label": "Details: Old details"}.get(kind, "Old details")
    paragraphs[1].text = literal
    paragraphs[2].text = "Fixed note"
    if kind == "inherited":
        style = doc.styles.add_style("Inherited bullet", WD_STYLE_TYPE.PARAGRAPH)
        ppr = style.element.get_or_add_pPr()
        number = etree.SubElement(ppr, w("numPr"))
        etree.SubElement(number, w("numId")).set(w("val"), "1")
        paragraphs[1].style = style
    elif kind == "direct":
        number = etree.SubElement(paragraphs[1]._p.get_or_add_pPr(), w("numPr"))
        etree.SubElement(number, w("numId")).set(w("val"), "1")
    doc.save(source)
    plan = record_plan(
        source, [("Old title", "Old title", "title"), (literal, "Old details", "details")], table
    )
    content = record_content(
        [
            {
                "id": "one",
                "title": "Shown title",
                "details": "Restored details",
                "hidden_fields": ["details"],
            }
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    filled = Document(output)
    assert [paragraph_text(node) for node in filled.element.body.iter(w("p"))] == [
        "Shown title",
        "Fixed note",
    ]
    if table:
        assert len(filled.tables[0].rows) == 2
    content.sections[0].entries[0].hidden_fields = []
    fill_template(source, output, plan, content.model_dump(), [])
    text = "\n".join(Document(output).element.body.xpath(".//w:t/text()"))
    assert "Restored details" in text and "Fixed note" in text


def test_added_metadata_share_one_indent_without_inheriting_list_numbering(tmp_path):
    """只有列表正文的陌生栏目补出三种元信息时；使用同一左边界并显式禁用继承编号"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    style = doc.styles.add_style("Body bullet", WD_STYLE_TYPE.PARAGRAPH)
    style.paragraph_format.left_indent = Pt(31)
    style.paragraph_format.first_line_indent = Pt(-9)
    number = etree.SubElement(style.element.get_or_add_pPr(), w("numPr"))
    etree.SubElement(number, w("numId")).set(w("val"), "1")
    doc.add_paragraph("Old body", style)
    doc.save(source)
    plan = record_plan(source, [("Old body", "Old body", "details")])
    content = record_content(
        [
            {
                "id": "one",
                "title": "New title",
                "subtitle": "New subtitle",
                "period": "2026",
                "details": "Hidden body",
                "hidden_fields": ["details"],
            }
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    package = TemplatePackage(output)
    styles = ParagraphStyles(package)
    paragraphs = [node for node in package.nodes.values() if node.tag == w("p")]
    assert [paragraph_text(node) for node in paragraphs] == ["New title", "New subtitle", "2026"]
    assert len({tuple(styles.indent(node).items()) for node in paragraphs}) == 1
    assert all(not styles.numbered(node) for node in paragraphs)
    assert all(node.find("w:pPr/w:ind", NS).get(w("hanging")) == "0" for node in paragraphs)


def test_empty_field_cleanup_keeps_other_text_and_right_aligned_metadata(tmp_path):
    """同段有其他内容时不能整段删除；右对齐日期也不能被左侧正文布局覆盖"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("Old title: Old body")
    period = doc.add_paragraph("Old period")
    period.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    period.paragraph_format.left_indent = Pt(50)
    doc.save(source)
    plan = record_plan(
        source,
        [
            ("Old title: Old body", "Old title", "title"),
            ("Old title: Old body", "Old body", "details"),
            ("Old period", "Old period", "period"),
        ],
    )
    content = record_content(
        [
            {
                "id": "one",
                "title": "Kept title",
                "period": "2026",
                "details": "Hidden",
                "hidden_fields": ["details"],
            }
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    paragraphs = Document(output).paragraphs
    assert paragraphs[0].text.startswith("Kept title")
    assert paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraphs[1].paragraph_format.left_indent == Pt(50)


@pytest.mark.parametrize("layout", ["columns", "cells", "frame", "rtl"])
def test_independent_columns_and_positioned_titles_keep_their_geometry(tmp_path, layout):
    """独立单元格、多栏、固定框及从右向左的标题不能被另一列正文拉到同一坐标"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    if layout == "cells":
        row = doc.add_table(rows=1, cols=2).rows[0]
        title, body = [cell.paragraphs[0] for cell in row.cells]
    else:
        title, body = doc.add_paragraph(), doc.add_paragraph()
    title.text, body.text = "Old title", "Old body"
    title.paragraph_format.left_indent = Pt(4)
    body.paragraph_format.left_indent = Pt(29)
    if layout == "columns":
        doc.sections[-1]._sectPr.find(w("cols")).set(w("num"), "2")
    elif layout in {"frame", "rtl"}:
        etree.SubElement(title._p.get_or_add_pPr(), w("framePr" if layout == "frame" else "bidi"))
    doc.save(source)
    plan = record_plan(
        source,
        [("Old title", "Old title", "title"), ("Old body", "Old body", "details")],
        layout == "cells",
    )
    content = record_content([{"id": "one", "title": "Kept title", "details": "Kept body"}])
    fill_template(source, output, plan, content.model_dump(), [])
    package = TemplatePackage(output)
    for node in package.nodes.values():
        if node.tag == w("p") and paragraph_text(node) in {"Kept title", "Kept body"}:
            expected = 80 if paragraph_text(node) == "Kept title" else 580
            assert node.find("w:pPr/w:ind", NS).get(w("left")) == str(expected)


def test_hiding_a_field_keeps_its_attached_decoration(tmp_path):
    """空字段段落还有内嵌装饰时只清空文字且不能删除有其他用途的图形"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("Old title")
    doc.add_paragraph("Old body").add_run().add_picture(BytesIO(photo_bytes(80)))
    doc.save(source)
    plan = record_plan(
        source, [("Old title", "Old title", "title"), ("Old body", "Old body", "details")]
    )
    package = TemplatePackage(source)
    plan.keep.extend(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "image")
    content = record_content(
        [
            {
                "id": "one",
                "title": "Kept title",
                "details": "Hidden body",
                "hidden_fields": ["details"],
            }
        ]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    assert len(Document(output).inline_shapes) == 1
    assert "Hidden body" not in "".join(Document(output).element.body.xpath(".//w:t/text()"))


def test_inherited_character_indent_isolated_without_changing_source_styles(tmp_path):
    """字符单位缩进通过独立样式展开后不再覆盖距离缩进；原样式、字体和多个制表位完整保留"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    base = doc.styles.add_style("Character base", WD_STYLE_TYPE.PARAGRAPH)
    base.font.name, base.font.size, base.font.bold = "Arial", Pt(12), True
    base.paragraph_format.left_indent = Pt(0)
    base.element.pPr.ind.set(w("leftChars"), "200")
    base.paragraph_format.tab_stops.add_tab_stop(Pt(160))
    base.paragraph_format.tab_stops.add_tab_stop(Pt(360), WD_TAB_ALIGNMENT.RIGHT)
    derived = doc.styles.add_style("Character derived", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    derived.font.italic = True
    doc.add_paragraph("Old title", derived)
    doc.add_paragraph("Old body").paragraph_format.left_indent = Pt(37)
    doc.save(source)
    original = source.read_bytes()
    plan = record_plan(
        source, [("Old title", "Old title", "title"), ("Old body", "Old body", "details")]
    )
    content = record_content(
        [{"id": str(i), "title": f"Title {i}", "details": "Body"} for i in range(3)]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    before, after = TemplatePackage(source), TemplatePackage(output)
    original_styles = etree.fromstring(before.files["word/styles.xml"])
    styles = ParagraphStyles(after)
    for style in original_styles.findall(w("style")):
        assert etree.tostring(styles.styles[style.get(w("styleId"))]) == etree.tostring(style)
    for node in after.nodes.values():
        if node.tag != w("p") or not paragraph_text(node).startswith("Title"):
            continue
        assert styles.indent(node) == {w("left"): "740"}
        identifier = node.find("w:pPr/w:pStyle", NS).get(w("val"))
        style = styles.styles[identifier]
        assert style.find("w:rPr/w:rFonts", NS).get(w("ascii")) == "Arial"
        assert style.find("w:rPr/w:sz", NS).get(w("val")) == "24"
        assert style.find("w:rPr/w:b", NS) is not None
        assert style.find("w:rPr/w:i", NS) is not None
        assert len(style.findall("w:pPr/w:tabs/w:tab", NS)) == 2
    assert source.read_bytes() == original


def test_hidden_field_keeps_column_boundary_without_its_list_marker(tmp_path):
    """隐藏带分栏符的编号字段时保留列边界；下一列不能错位；也不能留下空编号"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    doc = Document()
    doc.sections[-1]._sectPr.find(w("cols")).set(w("num"), "2")
    doc.add_paragraph("Old title")
    detail = doc.add_paragraph()
    detail.add_run().add_break(WD_BREAK.COLUMN)
    detail.add_run("Old body")
    number = etree.SubElement(detail._p.get_or_add_pPr(), w("numPr"))
    etree.SubElement(number, w("numId")).set(w("val"), "1")
    doc.add_paragraph("Fixed note")
    doc.save(source)
    plan = record_plan(
        source, [("Old title", "Old title", "title"), ("Old body", "Old body", "details")]
    )
    content = record_content(
        [{"id": "a", "title": "Kept title", "details": "Hidden", "hidden_fields": ["details"]}]
    )
    fill_template(source, output, plan, content.model_dump(), [])
    body = Document(output).element.body
    assert len(body.xpath(".//w:br[@w:type='column']")) == 1
    assert not body.xpath(".//w:numPr")
    assert body.xpath(".//w:t/text()") == ["Kept title", "Fixed note"]
