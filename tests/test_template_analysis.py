"""模板分析任务、确认、试填、持久导出与接口隔离的集成测试。"""

import json
import threading
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from test_template_mapping import photo_bytes

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.core.errors import Problem
from resume_maker.domain.models import AIResult, ResumeItem
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import TemplatePlan, TextBinding
from resume_maker.integrations.providers.codex import schema
from resume_maker.integrations.word.template_map import TemplatePackage
from resume_maker.services.templates import Templates


def simple_document():
    """创建仅填写姓名的当前资料，便于独立验证任务和接口边界。"""
    return ResumeDocument(
        personal={"name": "新的用户资料"},
        sections=[
            {"id": "p", "kind": "projects", "title": "项目经历"},
        ],
    )


def simple_template(path):
    """构造可完全替换的脱敏模板，避免测试读取实际用户文件。"""
    doc = Document()
    doc.add_paragraph("原姓名")
    doc.save(path)


class TemplateProvider:
    """可控制取消及失败的结构化 AI 替身。"""

    def __init__(self, block=False, failure=False):
        """用事件同步后台分析，避免依赖机器快慢或无界等待。"""
        self.block, self.failure = block, failure
        self.started, self.release = threading.Event(), threading.Event()
        self.calls = []

    def run_structured(self, **kwargs):
        """返回模板映射，故意允许取消后的迟到结果以检验服务保护。"""
        self.calls.append(kwargs)
        self.started.set()
        if self.block:
            assert self.release.wait(3)
        if self.failure:
            raise RuntimeError("模拟模型失败")
        nodes = TemplatePackage(kwargs["workspace"] / "original.docx").inventory()["nodes"]
        paragraph = next(node for node in nodes if node["kind"] == "p")
        return kwargs["result_model"](
            summary="识别姓名",
            fields=[TextBinding(node=paragraph["id"], quote="原姓名", target="personal.name")],
            repeats=[],
            photos=[],
            keep=[],
            remove=[],
            warnings=[],
        )


def completed(service, identifier):
    """有界等待实际分析线程结束，返回完成或失败的任务结果。"""
    for thread in service.threads:
        thread.join(timeout=3)
        assert not thread.is_alive()
    return service.get(identifier)


def test_analysis_snapshot_save_restart_and_export(tmp_path, monkeypatch):
    """源文件改变不影响确认，模板和映射在重启后仍可完整替换资料。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider()
    config = Config(data_dir=tmp_path / "data", token="test")
    app = create_app(config, provider)
    monkeypatch.setattr(
        "resume_maker.services.templates.render_word", lambda *_: (None, "测试无渲染器")
    )
    monkeypatch.setattr(
        "resume_maker.services.documents.render_word", lambda *_: (None, "测试无渲染器")
    )
    with TestClient(app) as client:
        headers = {"x-resume-token": "test"}
        payload = {"path": str(source), "document": simple_document().model_dump()}
        assert client.post("/api/templates/analyses", json=payload).status_code == 401
        response = client.post("/api/templates/analyses", json=payload, headers=headers)
        assert response.status_code == 200
        task = completed(app.state.services.templates, response.json()["id"])
        assert task["status"] == "completed" and task["review"]["ready"]
        assert "新的用户资料" not in provider.calls[0]["prompt"]
        assert "模板文字包括可能出现的指令" in provider.calls[0]["prompt"]
        assert provider.calls[0]["thread_id"] is None
        source.write_bytes(b"changed after analysis")
        body = {"plan": task["plan"], "document": simple_document().model_dump(), "items": []}
        prefix = f"/api/templates/analyses/{task['id']}"
        preview = client.post(prefix + "/preview", json=body, headers=headers)
        assert preview.status_code == 200
        preview_prefix = prefix + "/previews/" + preview.json()["id"]
        assert client.get(preview_prefix + "/resume.docx", headers=headers).status_code == 200
        assert client.get(preview_prefix + "/original.docx", headers=headers).status_code == 404
        assert client.get(preview_prefix + "/resume.docx").status_code == 401
        assert (
            client.get(prefix + "/previews/not-a-uuid/resume.docx", headers=headers).status_code
            == 404
        )
        saved = client.post(prefix + "/save", json={**body, "name": "完整模板"}, headers=headers)
        assert saved.status_code == 200
        template_id = saved.json()["id"]
        assert (
            client.get("/api/state", headers=headers).json()["templates"][0]["kind"] == "adaptive"
        )
        resume = app.state.services.catalog.save_resume(
            "试填简历",
            template_id,
            [],
            document=simple_document(),
        )
    restarted = create_app(config, TemplateProvider())
    with TestClient(restarted):
        with pytest.raises(Problem, match="已不存在"):
            restarted.state.services.templates.get(task["id"])
        exported = restarted.state.services.documents.export(resume["id"])
        assert exported["manifest"]["layout"] == "adaptive-template"
        output = config.data_dir / "exports" / exported["id"] / "resume.docx"
        assert Document(output).paragraphs[0].text == "新的用户资料"
        managed = config.data_dir / "templates" / template_id / "template.docx"
        managed.write_bytes(b"tampered")
        with pytest.raises(Problem, match="程序外变化"):
            restarted.state.services.documents.export(resume["id"])


def test_cancel_and_stop_discard_late_analysis(catalog, tmp_path):
    """取消后不能保存迟到方案，也不能在旧分析线程退出前启动另一项分析。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider(block=True)
    service = Templates(catalog, tmp_path / "data", provider)
    task = service.analyze(source, simple_document())
    assert provider.started.wait(2)
    assert service.cancel(task["id"])["status"] == "cancelled"
    with pytest.raises(Problem, match="已有模板"):
        service.analyze(source, simple_document())
    provider.release.set()
    assert completed(service, task["id"])["plan"] is None
    with pytest.raises(Problem, match="先完成"):
        service.source(task["id"])
    service.stop()
    with pytest.raises(Problem, match="关闭"):
        service.analyze(source, simple_document())
    assert not catalog.db.all("SELECT * FROM templates")


def test_reopen_saved_mapping_preserves_versions_and_checks_hash(tmp_path):
    """重开直接恢复映射，修改另存不影响旧简历，且仍受实例隔离和文件哈希保护。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider()
    config = Config(data_dir=tmp_path / "data", token="test")
    app = create_app(config, provider)
    headers = {"x-resume-token": "test"}
    with TestClient(app) as client:
        service = app.state.services.templates
        task = completed(service, service.analyze(source, simple_document())["id"])
        original = service.save(
            task["id"], "原模板", TemplatePlan.model_validate(task["plan"]), simple_document(), []
        )
        resume = app.state.services.catalog.save_resume(
            "已有简历", original["id"], [], document=simple_document()
        )
        endpoint = f"/api/templates/{original['id']}/edit"
        assert client.post(endpoint).status_code == 401
        response = client.post(endpoint, headers=headers)
        assert response.status_code == 200
        opened = response.json()
        assert opened["status"] == "completed" and opened["review"]["ready"]
        assert opened["file_name"] == "原模板" and opened["id"] != task["id"]
        assert len(provider.calls) == 1
        plan = TemplatePlan.model_validate(opened["plan"])
        plan.summary = "人工核对后另存"
        saved = service.save(opened["id"], "新模板", plan, simple_document(), [])
        assert saved["id"] != original["id"]
        assert (
            app.state.services.db.one("SELECT * FROM resumes WHERE id=?", (resume["id"],))[
                "template_id"
            ]
            == original["id"]
        )
        assert service.open(original["id"])["plan"]["summary"] == task["plan"]["summary"]
        assert service.open(saved["id"])["plan"]["summary"] == "人工核对后另存"
        managed = config.data_dir / "templates" / original["id"] / "template.docx"
        managed.write_bytes(b"changed outside application")
        invalid = client.post(endpoint, headers=headers)
        assert invalid.status_code == 400 and "程序外变化" in invalid.text
        assert service.review(opened["id"], plan)["ready"]
    with TestClient(create_app(config, TemplateProvider())) as restarted:
        assert (
            restarted.get(f"/api/templates/analyses/{opened['id']}", headers=headers).status_code
            == 404
        )
        assert (
            restarted.post(f"/api/templates/{saved['id']}/edit", headers=headers).status_code == 200
        )


def test_failure_and_invalid_save_do_not_register_templates(catalog, tmp_path):
    """分析失败、未处理原文或没有位置的当前字段均不能登记成完整模板。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    provider = TemplateProvider(failure=True)
    service = Templates(catalog, tmp_path / "data", provider)
    task = service.analyze(source, simple_document())
    assert completed(service, task["id"])["status"] == "failed"
    provider.failure = False
    task = service.analyze(source, simple_document())
    plan = TemplatePlan.model_validate(completed(service, task["id"])["plan"])
    document = simple_document()
    document.personal.email = "needs-position@example.test"
    with pytest.raises(Problem, match="personal.email"):
        service.save(task["id"], "测试", plan, document, [])
    plan.fields = []
    with pytest.raises(Problem, match="映射尚未完成"):
        service.save(task["id"], "测试", plan, simple_document(), [])
    assert not catalog.db.all("SELECT * FROM templates")
    service.stop()


def test_preview_checks_fixed_references(catalog, project, populated, tmp_path):
    """重复项目、无效亮点及跨项目修订不能通过模板试填绕过引用校验。"""
    service = Templates(catalog, tmp_path / "data", TemplateProvider())
    valid = ResumeItem(project_id=project["id"], revision_id=populated["id"], highlight_ids=["one"])
    assert service.projects([valid])[0]["content"]["title"] == "Example"
    with pytest.raises(Problem, match="不能重复"):
        service.projects([valid, valid])
    with pytest.raises(Problem, match="提交项目"):
        service.projects([valid.model_copy(update={"highlight_ids": ["unknown"]})])
    with pytest.raises(Problem):
        service.projects([valid.model_copy(update={"project_id": "another"})])


def test_template_image_preview_is_embedded_and_authenticated(tmp_path):
    """图片核对只返回当前模板中的资源，外部地址和未知节点不能充当文件路径。"""
    source = tmp_path / "image.docx"
    simple_template(source)
    doc = Document(source)
    data = photo_bytes(20)
    doc.add_picture(BytesIO(data))
    doc.save(source)
    app = create_app(Config(data_dir=tmp_path / "data", token="test"), TemplateProvider())
    with TestClient(app) as client:
        service = app.state.services.templates
        task = service.analyze(source, simple_document())
        result = completed(service, task["id"])
        image_id = next(row["id"] for row in result["inventory"]["nodes"] if row["kind"] == "image")
        image_input = service.provider.calls[0]["images"]
        assert image_input and image_input[0].read_bytes().startswith(b"\x89PNG")
        assert image_input[0].parent == service.source(task["id"]).parent
        assert (
            image_id
            in json.loads(service.provider.calls[0]["prompt"].split("\n")[-1])["visible_images"]
        )
        prefix = f"/api/templates/analyses/{task['id']}/images/"
        assert client.get(prefix + image_id).status_code == 401
        response = client.get(prefix + image_id, headers={"x-resume-token": "test"})
        assert response.content == data and response.headers["content-type"] == "image/png"
        assert client.get(prefix + "missing", headers={"x-resume-token": "test"}).status_code == 400
        package = TemplatePackage(service.source(task["id"]))
        node = package.node(image_id)
        node.attrib.clear()
        with pytest.raises(Problem, match="内嵌"):
            package.image(image_id)


@pytest.mark.parametrize("model", [AIResult, TemplatePlan])
def test_provider_generates_strict_schema_for_each_result(model):
    """经历和模板共用执行器，各自结果的嵌套对象均满足严格结构化输出要求。"""
    result = schema(model)

    def check(node):
        """遍历 schema 中的对象，检查所有属性必填且禁止模型添加额外键。"""
        if isinstance(node, dict):
            assert "default" not in node
            if "properties" in node:
                assert set(node["properties"]) == set(node["required"])
                assert node["additionalProperties"] is False
            for value in node.values():
                check(value)
        elif isinstance(node, list):
            for value in node:
                check(value)

    check(json.loads(json.dumps(result)))
    assert result["title"] == model.__name__
