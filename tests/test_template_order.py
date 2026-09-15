"""栏目编排必须同时控制 Word 标题、记录、层级与显隐，不能只重排前端列表。"""

from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from test_template_mapping import make_template, project_content, resume_content

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text


def body_text(path):
    """读取文档正文顺序，包含表格与图形内的文字。"""
    return "\n".join(Document(path).element.body.xpath(".//w:t/text()"))


def simple_template(path, table=False):
    """生成三个独立栏目，使用精确标题映射和可重复的条目样本。"""
    doc = Document()
    doc.add_paragraph("固定开头")
    container = doc.add_table(rows=0, cols=1) if table else doc
    for title in ("教育背景", "主修课程", "项目经历"):
        if table:
            container.add_row().cells[0].text = title
            container.add_row().cells[0].text = "旧" + title
        else:
            container.add_paragraph(title, "Heading 1")
            container.add_paragraph("旧" + title)
    doc.add_paragraph("固定结尾")
    doc.save(path)
    package = TemplatePackage(path)
    nodes = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    fields, repeats = [], []
    for title in ("教育背景", "主修课程", "项目经历"):
        fields.append(TextBinding(node=nodes[title], quote=title, target=f"section-title:{title}"))
        node = package.node(nodes["旧" + title])
        root = next(node.iterancestors(w("tr"))) if table else node
        identifier = package.ids[root]
        repeats.append(
            RepeatBinding(
                section="projects" if title == "项目经历" else title,
                start=identifier,
                end=identifier,
                sample_start=identifier,
                sample_end=identifier,
                fields=[TextBinding(node=package.ids[node], quote="旧" + title, target="title")],
            )
        )
    return TemplatePlan(
        summary="栏目顺序测试",
        fields=fields,
        repeats=repeats,
        photos=[],
        keep=[nodes["固定开头"], nodes["固定结尾"]],
        remove=[],
        warnings=[],
    )


def layout_content():
    """生成大栏目和子栏目的资料，刻意将子栏目放在数组前面检验层级展开。"""
    return ResumeDocument.model_validate(
        {
            "sections": [
                {
                    "id": "courses",
                    "title": "主修课程",
                    "kind": "text",
                    "parent_id": "education",
                    "entries": [{"id": "course", "title": "课程条目"}],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
                {
                    "id": "education",
                    "title": "教育背景",
                    "kind": "education",
                    "entries": [{"id": "school", "title": "教育条目"}],
                },
            ]
        }
    )


def minimal_projects():
    """提供两个只有名称的项目，避免测试额外要求未声明的模板字段。"""
    return [
        {
            "highlight_ids": [],
            "content": {
                "title": title,
                "period": "",
                "role": "",
                "stack": [],
                "description": "",
                "highlights": [],
            },
        }
        for title in ("项目甲", "项目乙")
    ]


def test_move_complete_sections_preserves_templates_and_fixed_content(tmp_path):
    """把项目移到教育前面时，其标题和全部项目一起移动，模板资源及首尾资料保留。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    _, plan = make_template(source)
    content = resume_content()
    content.sections.reverse()
    original, mapping = source.read_bytes(), plan.model_dump()
    fill_template(source, output, plan, content.model_dump(), project_content())
    text = body_text(output)
    assert text.index("项目经历") < text.index("新文档项目") < text.index("教育背景")
    assert text.index("教育背景") < text.index("新大学一") < text.index("新大学二")
    assert text.index("测试新姓名") < text.index("项目经历")
    assert text.index("语言：中文") > text.index("教育背景")
    assert source.read_bytes() == original and plan.model_dump() == mapping
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")
        assert b"resume-section-" not in after.read("word/document.xml")


def test_children_follow_parent_order_and_can_be_promoted(tmp_path):
    """子栏目跟随父栏目，提升为大栏目后遵循新的独立顺序。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    plan = simple_template(source)
    content = layout_content()
    fill_template(source, output, plan, content.model_dump(), minimal_projects())
    assert body_text(output).splitlines() == [
        "固定开头",
        "项目经历",
        "项目甲",
        "项目乙",
        "教育背景",
        "教育条目",
        "主修课程",
        "课程条目",
        "固定结尾",
    ]
    content.sections[0].parent_id = None
    fill_template(source, output, plan, content.model_dump(), minimal_projects())
    assert body_text(output).index("课程条目") < body_text(output).index("项目经历")


def test_hidden_or_empty_sections_remove_headings_with_their_records(tmp_path):
    """隐藏父栏目同时隐藏子栏目；项目为空时不留下项目标题。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    plan = simple_template(source)
    content = layout_content()
    content.sections[-1].visible = False
    fill_template(source, output, plan, content.model_dump(), [])
    assert body_text(output).splitlines() == ["固定开头", "固定结尾"]


def test_sections_inside_one_table_move_as_complete_rows(tmp_path):
    """整张表格中的栏目按行移动，表格结构、记录数量和首尾固定内容保持完整。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    plan = simple_template(source, table=True)
    fill_template(source, output, plan, layout_content().model_dump(), minimal_projects())
    text = body_text(output)
    assert text.index("项目经历") < text.index("项目乙") < text.index("教育背景")
    assert text.index("教育条目") < text.index("主修课程") < text.index("课程条目")
    assert len(Document(output).tables[0].rows) == 7


def test_floating_heading_moves_inline_without_losing_its_graphic(tmp_path):
    """浮动标题以嵌入图形占据排版高度，标题不会被下一条正文覆盖。"""
    from resume_maker.integrations.word.template_layout import WP, inline_heading

    paragraph = etree.Element(w("p"))
    drawing = etree.SubElement(etree.SubElement(paragraph, w("r")), w("drawing"))
    anchor = etree.SubElement(drawing, f"{{{WP}}}anchor", behindDoc="1")
    etree.SubElement(anchor, f"{{{WP}}}positionV", relativeFrom="page")
    extent = etree.SubElement(anchor, f"{{{WP}}}extent", cx="1000", cy="2000")
    graphic = etree.SubElement(
        anchor, "{http://schemas.openxmlformats.org/drawingml/2006/main}graphic"
    )
    inline_heading(paragraph)
    assert anchor.tag == f"{{{WP}}}inline"
    assert list(anchor) == [extent, graphic]
    assert "behindDoc" not in anchor.attrib
    assert paragraph.find("w:pPr/w:keepNext", NS) is not None


def shared_heading_template(path):
    """模拟课程标题与 GPA 共用段落、课程正文先于标题且位于文本框的模板。"""
    doc = Document()
    opening = doc.add_paragraph("固定开头")
    section = etree.SubElement(opening._p.get_or_add_pPr(), w("sectPr"))
    etree.SubElement(section, w("type")).set(w("val"), "continuous")
    etree.SubElement(section, w("cols")).set(w("num"), "3")
    doc.add_paragraph("教育背景")
    doc.add_paragraph("学校名称：")
    for label, text in (("", "旧课程"), ("主修课程：", "GPA：旧绩点")):
        paragraph = doc.add_paragraph(label)
        box = etree.fromstring(
            f'<w:r xmlns:w="{NS["w"]}" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            '<v:shape style="width:220pt;height:22pt"><v:textbox><w:txbxContent>'
            f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
        )
        paragraph._p.append(box)
    doc.add_paragraph("旧学校")
    doc.add_paragraph("项目经历")
    doc.add_paragraph("旧项目")
    # 正文的分栏属性会被文本框的查找命中，但不能在文本框内生成分节符。
    final = doc.sections[-1]._sectPr
    etree.SubElement(final, w("type")).set(w("val"), "continuous")
    final.find(w("cols")).set(w("num"), "2")
    doc.save(path)
    package = TemplatePackage(path)
    nodes = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    return TemplatePlan(
        summary="共享段落与分栏测试",
        fields=[
            TextBinding(node=nodes[label], quote=quote, target=target)
            for label, quote, target in (
                ("教育背景", "教育背景", "section-title:教育背景"),
                ("主修课程：", "主修课程", "section-title:主修课程"),
                ("项目经历", "项目经历", "section-title:项目经历"),
                ("GPA：旧绩点", "旧绩点", "personal.gpa"),
            )
        ],
        repeats=[
            RepeatBinding(
                section=title,
                start=nodes[text],
                end=nodes[text],
                sample_start=nodes[text],
                sample_end=nodes[text],
                fields=[TextBinding(node=nodes[text], quote=text, target="title")],
            )
            for title, text in (
                ("教育背景", "旧学校"),
                ("主修课程", "旧课程"),
                ("projects", "旧项目"),
            )
        ],
        photos=[],
        keep=[nodes["固定开头"], nodes["学校名称："]],
        remove=[],
        warnings=[],
    )


@pytest.mark.parametrize("courses", [True, False])
@pytest.mark.parametrize("school", [True, False])
def test_shared_course_heading_keeps_gpa_and_valid_section_boundaries(tmp_path, courses, school):
    """父子栏目重排或清空时 GPA 保留，课程标题在正文前，文本框不产生非法分节符。"""
    source, output = tmp_path / "source.docx", tmp_path / "result.docx"
    plan = shared_heading_template(source)
    content = layout_content()
    content.personal.gpa = "4.2/5.0"
    if not courses:
        content.sections[0].entries = []
    if not school:
        content.sections[-1].entries = []
    fill_template(source, output, plan, content.model_dump(), minimal_projects())
    text = body_text(output)
    assert text.index("项目经历") < text.index("项目乙") < text.index("教育背景")
    assert text.index("教育背景") < text.index("GPA：4.2/5.0")
    assert ("学校名称：" in text) == school
    assert ("教育条目" in text) == school
    if courses:
        assert text.index("GPA：4.2/5.0") < text.index("主修课程") < text.index("课程条目")
        assert text.count("主修课程") == 1
    else:
        assert "主修课程" not in text and "课程条目" not in text
    body = Document(output).element.body
    assert not body.xpath(".//w:txbxContent//w:sectPr | .//w:tc//w:sectPr")
    assert body[0].find("w:pPr/w:sectPr/w:cols", NS).get(w("num")) == "3"
