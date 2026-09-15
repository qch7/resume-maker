"""验证识别阶段自动补齐无空位字段，保留原文件、映射身份和真实试填内容。"""

import pytest
from docx import Document
from docx.shared import Pt
from test_template_analysis import TemplateProvider, completed, simple_document, simple_template
from test_template_mapping import make_template, project_content, resume_content

from resume_maker.domain.resume import CustomInfoField
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text
from resume_maker.integrations.word.template_supplement import supplement_personal_fields
from resume_maker.services.templates import Templates


@pytest.mark.parametrize("prebound", [True, False])
@pytest.mark.parametrize("explicit_break", [True, False])
def test_location_uses_contact_column_and_leaves_photo_spacing_empty(
    tmp_path, prebound, explicit_break
):
    """城市缺少位置或被误填进页首留白时，识别阶段即改到邮箱右列且原字体与缩进保留。"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph()
    contact = doc.add_paragraph()
    contact.paragraph_format.left_indent = Pt(12)
    contact.paragraph_format.tab_stops.add_tab_stop(Pt(270))
    contact.add_run("姓名：旧姓名\t电话：旧电话")
    if explicit_break:
        contact.add_run().add_break()
    contact.add_run("邮箱：旧邮箱").font.name = "宋体"
    doc.save(source)
    original = source.read_bytes()
    package = TemplatePackage(source)
    rows = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    fields = [
        TextBinding(node=rows[1]["id"], quote=quote, target=target)
        for quote, target in (
            ("旧姓名", "personal.name"),
            ("旧电话", "personal.phone"),
            ("旧邮箱", "personal.email"),
        )
    ]
    if prebound:
        fields.append(TextBinding(node=rows[0]["id"], quote="", target="personal.location"))
    plan = TemplatePlan(
        summary="联系信息", fields=fields, repeats=[], photos=[], keep=[], remove=[], warnings=[]
    )
    document = simple_document()
    document.personal.phone = "123456"
    document.personal.email = "new@example.test"
    document.personal.location = "测试城市"
    snapshot = tmp_path / "snapshot.docx"
    updated, plan, notices = supplement_personal_fields(package, plan, document, [], snapshot)
    assert notices and updated.review(plan)["ready"]
    location = next(field for field in plan.fields if field.target == "personal.location")
    email = next(field for field in plan.fields if field.target == "personal.email")
    assert (location.node == email.node) == explicit_break
    assert len(Document(snapshot).paragraphs) == (2 if explicit_break else 3)
    output = tmp_path / "filled.docx"
    fill_template(snapshot, output, plan, document.model_dump(), [])
    result = Document(output)
    assert not result.paragraphs[0].text
    if explicit_break:
        assert (
            result.paragraphs[1].text
            == "姓名：新的用户资料\t电话：123456\n邮箱：new@example.test\t所在地：测试城市"
        )
    else:
        # 原文没有显式换行时，不能靠第三项字段是邮箱而改写原段落的行结构。
        assert result.paragraphs[1].text == "姓名：新的用户资料\t电话：123456邮箱：new@example.test"
        assert result.paragraphs[2].text == "所在地：测试城市"
    assert result.paragraphs[1].paragraph_format.left_indent == Pt(12)
    assert result.paragraphs[1].paragraph_format.tab_stops[0].position == Pt(270)
    previous = snapshot.read_bytes()
    _, same, notices = supplement_personal_fields(
        TemplatePackage(snapshot), plan, document, [], snapshot
    )
    assert not notices and same == plan and snapshot.read_bytes() == previous
    assert source.read_bytes() == original


def test_explicit_location_field_is_not_moved(tmp_path):
    """模板已有带标签城市示例时沿用原位，不能将用户明确的布局误判成留白。"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("现居地：旧城市")
    doc.add_paragraph("邮箱：旧邮箱")
    doc.save(source)
    package = TemplatePackage(source)
    rows = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="明确的城市位置",
        fields=[
            TextBinding(node=rows[0]["id"], quote="旧城市", target="personal.location"),
            TextBinding(node=rows[1]["id"], quote="旧邮箱", target="personal.email"),
        ],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    document = simple_document()
    document.personal.name = ""
    document.personal.location = "测试城市"
    document.personal.email = "new@example.test"
    before = source.read_bytes()
    _, same, notices = supplement_personal_fields(package, plan, document, [], source)
    assert same == plan and not notices and source.read_bytes() == before


def test_first_analysis_completes_fields_without_blank_slots(catalog, tmp_path):
    """模型只识别原姓名时，一轮即可补齐所在地和自定义字段并真正写入试填。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    doc = Document(source)
    doc.paragraphs[0].runs[0].font.size = Pt(12)
    doc.save(source)
    original = source.read_bytes()
    document = simple_document()
    document.personal.location = "试填城市"
    document.personal.email = "hidden@example.test"
    document.personal.hidden_fields = ["email"]
    document.personal.custom_fields = [CustomInfoField(id="language", label="语言", value="中文")]
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    task = completed(service, service.analyze(source, document)["id"])
    assert task["review"]["ready"] and task["attempts"] == 1 and len(provider.calls) == 1
    assert source.read_bytes() == original
    assert not list((tmp_path / "template-cache").glob("*.json"))
    plan = TemplatePlan.model_validate(task["plan"])
    targets = {field.target for field in plan.fields}
    assert targets == {"personal.name", "personal.location", "personal.custom:语言"}
    snapshot = service.source(task["id"])
    package = TemplatePackage(snapshot)
    assert task["inventory"]["nodes"] == package.inventory()["nodes"]
    assert package.review(plan)["ready"]
    slot = package.node(
        next(field.node for field in plan.fields if field.target == "personal.location")
    )
    assert slot.find(".//" + w("sz")).get(w("val")) == "24"
    assert "试填城市" not in "".join(node["text"] for node in task["inventory"]["nodes"])
    output = tmp_path / "filled.docx"
    fill_template(snapshot, output, plan, document.model_dump(), [])
    text = "\n".join(p.text for p in Document(output).paragraphs)
    assert "所在地：试填城市" in text and "语言：中文" in text
    assert "原姓名" not in text and "hidden@example.test" not in text and "〔待填写〕" not in text


def test_repair_of_missing_location_needs_no_more_model_calls(catalog, tmp_path):
    """已经有正确映射时立即补出新字段位置，修复副本不修改原分析结果。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    document = simple_document()
    provider = TemplateProvider()
    service = Templates(catalog, tmp_path, provider)
    original = completed(service, service.analyze(source, document)["id"])
    original_bytes = service.source(original["id"]).read_bytes()
    plan = TemplatePlan.model_validate(original["plan"])
    plan.warnings = ["personal.location 没有合法位置", "请核对字体"]
    document.personal.location = "新增城市"
    repaired = completed(service, service.repair(original["id"], plan, document, [])["id"])
    assert repaired["review"]["ready"] and repaired["attempts"] == 0
    assert len(provider.calls) == 1
    assert repaired["plan"]["warnings"] == ["请核对字体"]
    assert service.get(original["id"])["plan"] == original["plan"]
    assert service.source(original["id"]).read_bytes() == original_bytes


def test_new_nodes_preserve_all_existing_references_and_are_idempotent(tmp_path):
    """新增段落后的字段、栏目、照片、页眉页脚和删除引用在重载后仍指向同一原文。"""
    source = tmp_path / "source.docx"
    package, plan = make_template(source)
    old_text = {identifier: paragraph_text(node) for identifier, node in package.nodes.items()}
    document = resume_content()
    document.personal.age = "22"
    projects = project_content()
    updated_package, updated, notices = supplement_personal_fields(
        package, plan, document, projects, source
    )
    assert notices and updated_package.review(updated)["ready"]
    reloaded = TemplatePackage(source)
    assert reloaded.review(updated)["ready"]
    for old, new in zip(plan.fields, updated.fields[: len(plan.fields)], strict=True):
        assert paragraph_text(reloaded.node(new.node)) == old_text[old.node]
        assert reloaded.locations[new.node] == package.locations[old.node]
    for old, new in zip(plan.repeats, updated.repeats, strict=True):
        for field, new_field in zip(old.fields, new.fields, strict=True):
            assert paragraph_text(reloaded.node(new_field.node)) == old_text[field.node]
    assert reloaded.image(updated.photos[0]) == package.image(plan.photos[0])
    output = tmp_path / "filled.docx"
    fill_template(source, output, updated, document.model_dump(), projects)
    visible = "\n".join(row["text"] for row in TemplatePackage(output).inventory()["nodes"])
    assert "年龄：22" in visible and "新文档项目" in visible and "新大学一" in visible
    previous_bytes = source.read_bytes()
    _, again, notices = supplement_personal_fields(reloaded, updated, document, projects, source)
    assert not notices and again == updated and source.read_bytes() == previous_bytes


def test_shared_table_does_not_trap_added_fields_in_repeated_rows(tmp_path):
    """姓名与经历共享表格时，新增资料放在表格前，不能随重复区删除或被固定行高裁切。"""
    source = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "原姓名"
    table.cell(1, 0).text = "旧学校"
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    rows = [row for row in package.inventory()["nodes"] if row["kind"] == "tr"]
    plan = TemplatePlan(
        summary="表格",
        fields=[TextBinding(node=paragraphs[0]["id"], quote="原姓名", target="personal.name")],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=rows[1]["id"],
                end=rows[1]["id"],
                sample_start=rows[1]["id"],
                sample_end=rows[1]["id"],
                fields=[TextBinding(node=paragraphs[1]["id"], quote="旧学校", target="title")],
            )
        ],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    document = simple_document()
    document.personal.location = "测试城市"
    package, plan, notices = supplement_personal_fields(package, plan, document, [], source)
    assert notices and package.review(plan)["ready"]
    fill_template(source, tmp_path / "filled.docx", plan, document.model_dump(), [])
    output = Document(tmp_path / "filled.docx")
    assert output.paragraphs[0].text == "所在地：测试城市"
    assert output.tables[0].cell(0, 0).text == document.personal.name


def test_unresolved_original_content_cannot_be_hidden_by_supplementation(tmp_path):
    """不能通过另加姓名把原姓名或未知经历绕过校验，照片也不能当文字字段补出。"""
    source = tmp_path / "source.docx"
    simple_template(source)
    package = TemplatePackage(source)
    plan = TemplatePlan(
        summary="尚未识别", fields=[], repeats=[], photos=[], keep=[], remove=[], warnings=[]
    )
    before = source.read_bytes()
    _, updated, notices = supplement_personal_fields(package, plan, simple_document(), [], source)
    assert not notices and updated == plan and source.read_bytes() == before


@pytest.mark.parametrize("slots", [1, 2, 5])
@pytest.mark.parametrize("flow", [False, True])
def test_project_details_and_multiple_highlight_slots_never_duplicate_content(
    tmp_path, slots, flow
):
    """亮点数量多于或少于样本位置时，每条仅填一次；综合正文不重复独立技术栈。"""
    source = tmp_path / "project.docx"
    doc = Document()
    if flow:
        doc.sections[0]._sectPr.find(w("cols")).set(w("num"), "2")
    for text in ["旧项目", "旧技术栈", "旧正文", *[f"旧亮点{i}" for i in range(slots)]]:
        doc.add_paragraph(text)
    doc.save(source)
    package = TemplatePackage(source)
    nodes = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    fields = [
        TextBinding(node=row["id"], quote=row["text"], target=target)
        for row, target in zip(
            nodes, ["title", "stack", "details", *["highlights"] * slots], strict=True
        )
    ]
    # 模型可乱序列出绑定，亮点仍按模板中的阅读顺序填入。
    fields.reverse()
    plan = TemplatePlan(
        summary="项目",
        fields=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
        repeats=[
            RepeatBinding(
                section="projects",
                start=nodes[0]["id"],
                end=nodes[-1]["id"],
                sample_start=nodes[0]["id"],
                sample_end=nodes[-1]["id"],
                fields=fields,
            )
        ],
    )
    document = simple_document()
    document.personal.name = ""
    projects = project_content()
    projects[0]["content"]["period"] = ""
    projects[0]["content"]["highlights"] = [
        {"id": str(i), "title": f"新亮点{i}", "text": f"成果{i}\n补充{i}"} for i in range(3)
    ]
    projects[0]["highlight_ids"] = ["0", "1", "2"]
    output = tmp_path / "filled.docx"
    fill_template(source, output, plan, document.model_dump(), projects)
    text = "\n".join(
        row["text"] for row in TemplatePackage(output).inventory()["nodes"] if row["kind"] == "p"
    )
    assert text.count("Python") == 1 and text.count("生成可编辑文档") == 1
    for i in range(3):
        assert text.count(f"新亮点{i}") == 1 and text.count(f"补充{i}") == 1
    assert text.index("新亮点0") < text.index("新亮点1") < text.index("新亮点2")
    assert "旧" not in text
