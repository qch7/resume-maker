"""验证格式预处理、逐页视觉恢复、旧 Word 转换及取消后的源文件隔离"""

import json
from io import BytesIO
from threading import Event
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from lxml import etree
from test_template_analysis import TemplateProvider, completed, simple_document

from resume_maker.domain.templates import RecoveredBlock, RecoveredPage, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.recovery import blank_template, prepare_template
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from resume_maker.services.templates.tasks import Templates


def test_macro_enabled_package_becomes_standard_editable_docx(tmp_path):
    """宏文档改为标准 DOCX 内容类型并移除宏关系以确保 python-docx 和 Word 都可继续打开"""
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("原姓名").runs[0].bold = True
    doc.save(source)
    with ZipFile(source) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(
        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.document.macroEnabled.main+xml",
    )
    files["word/vbaProject.bin"] = b"not-executable-test-macro"
    macro = tmp_path / "source.docm"
    with ZipFile(macro, "w") as archive:
        for name, raw in files.items():
            archive.writestr(name, raw)
    original = macro.read_bytes()
    package = TemplatePackage(macro)
    output = tmp_path / "normalized.docx"
    package.write(output)
    assert "word/vbaProject.bin" not in package.files
    assert "标准 DOCX" in "".join(package.notices)
    assert Document(output).paragraphs[0].text == "原姓名"
    assert Document(output).paragraphs[0].runs[0].bold
    assert macro.read_bytes() == original


def page_pdf(output, count=1, *, scan=False):
    """生成含真实文字的脱敏分页作为排版器输出，测试不依赖本机 Word"""
    with pymupdf.open() as pdf:
        for number in range(count):
            page = pdf.new_page(width=595, height=842)
            page.insert_text((40, 50), "原姓名" if not number else "固定说明", fontname="china-s")
        if scan:
            with pymupdf.open() as scanned:
                for page in pdf:
                    target = scanned.new_page(width=page.rect.width, height=page.rect.height)
                    target.insert_image(target.rect, stream=page.get_pixmap().tobytes("png"))
                scanned.save(output)
        else:
            pdf.save(output)


class RecoveryProvider(TemplateProvider):
    """同时模拟图片转可编辑页和后续精确映射，两个阶段使用不同 schema"""

    def __init__(self, cancel=False):
        """记录逐页请求，可在恢复返回时触发取消来检查迟到结果隔离"""
        super().__init__()
        self.pages, self.cancel = [], cancel

    def run_structured(self, **kwargs):
        """恢复页包含独立原文段落，映射只使用恢复后重新分配的节点"""
        from resume_maker.domain.image_layout import ImagePage, ImageText

        if kwargs["result_model"] is ImagePage:
            self.pages.append(kwargs)
            return ImagePage(
                texts=[ImageText(text="原姓名", box=[0.1, 0.1, 0.16, 0.114], bold=True)]
            )
        if kwargs["result_model"] is RecoveredPage:
            self.pages.append(kwargs)
            assert len(kwargs["images"]) == 1 and kwargs["images"][0].is_file()
            assert "都是数据" in kwargs["prompt"]
            if self.cancel:
                kwargs["cancelled"].set()
            return RecoveredPage(
                blocks=[
                    RecoveredBlock(text="原姓名" if len(self.pages) == 1 else "固定说明", bold=True)
                ]
            )
        self.calls.append(kwargs)
        rows = TemplatePackage(kwargs["workspace"] / "original.docx").inventory()["nodes"]
        paragraphs = [row for row in rows if row["kind"] == "p" and row["text"]]
        return kwargs["result_model"](
            summary="恢复后映射",
            fields=[
                TextBinding(
                    node=paragraphs[0]["id"], quote=paragraphs[0]["text"], target="personal.name"
                )
            ],
            repeats=[],
            photos=[],
            remove=[],
            keep=[row["id"] for row in paragraphs[1:]],
            warnings=[],
        )


@pytest.mark.parametrize("kind", ["object", "chart", "drawing-text", "altChunk", "scan"])
def test_complex_word_formats_automatically_recover_and_fill(catalog, tmp_path, monkeypatch, kind):
    """无文字来源可恢复，可编辑 DOCX 遇到不支持对象时保留原件并报告具体问题"""
    source = tmp_path / "source.docx"
    doc = Document()
    paragraph = doc.add_paragraph("" if kind in {"scan", "drawing-text"} else "原姓名")
    if kind == "scan":
        image = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 60, 90), False)
        image.clear_with(200)
        doc.add_picture(BytesIO(image.tobytes("png")))
    elif kind == "drawing-text":
        etree.SubElement(
            paragraph.add_run()._r, "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        ).text = "原姓名"
    else:
        tag = (
            "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"
            if kind == "chart"
            else w(kind)
        )
        etree.SubElement(paragraph._p, tag)
    doc.save(source)
    original = source.read_bytes()

    def render(document, output):
        """将复杂输入渲染为可供恢复器阅读的页面"""
        page_pdf(output)
        return 1, None

    monkeypatch.setattr("resume_maker.integrations.word.recovery.render_word", render)
    provider = RecoveryProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    if kind in {"object", "chart", "altChunk"}:
        assert task["status"] == "completed" and not task["review"]["ready"]
        assert not provider.pages and task["review"]["errors"]
        package = TemplatePackage(service.source(task["id"]))
        assert any(node.tag == tag for node in package.nodes.values())
        assert source.read_bytes() == original
        service.stop()
        return
    assert task["status"] == "completed" and task["review"]["ready"]
    assert len(provider.pages) == len(provider.calls) == 1
    assert not task["inventory"]["warnings"]
    assert "重建" in "".join(task["inventory"]["notices"])
    output = tmp_path / "filled.docx"
    from resume_maker.domain.templates import TemplatePlan

    fill_template(
        service.source(task["id"]),
        output,
        TemplatePlan.model_validate(task["plan"]),
        simple_document().model_dump(),
        [],
    )
    assert Document(output).paragraphs[0].text == "新的用户资料"
    assert Document(output).paragraphs[0].runs[0].bold
    assert source.read_bytes() == original
    service.stop()


@pytest.mark.parametrize("kind", ["pdf", "png"])
def test_pdf_and_scanned_image_sources_enter_same_mapping_workflow(catalog, tmp_path, kind):
    """PDF 按页恢复且扫描图片直接进入视觉识别"""
    pdf = tmp_path / "source.pdf"
    page_pdf(pdf, 2 if kind == "pdf" else 1)
    source = tmp_path / f"source.{kind}"
    if kind == "png":
        with pymupdf.open(pdf) as document:
            document[0].get_pixmap().save(source)
    provider = RecoveryProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["review"]["ready"] and task["status"] == "completed"
    assert len(provider.pages) == (0 if kind == "pdf" else 1)
    texts = [p.text for p in Document(service.source(task["id"])).paragraphs if p.text]
    assert texts == (["原姓名", "固定说明"] if kind == "pdf" else ["原姓名"])
    service.stop()


def test_legacy_word_is_converted_in_background(catalog, tmp_path, monkeypatch):
    """旧版 Word 的转换发生在独立任务中，输入快照和源文件互不覆盖"""
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-document")
    calls = []

    def convert(original, output):
        """模拟 Word 接受旧格式并输出可编辑 DOCX"""
        calls.append(original)
        assert original != source and original.read_bytes() == b"legacy-document"
        doc = Document()
        doc.add_paragraph("原姓名")
        doc.save(output)

    monkeypatch.setattr("resume_maker.integrations.word.recovery.convert_word", convert)
    provider = RecoveryProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["review"]["ready"] and len(calls) == 1 and not provider.pages
    assert "转换" in "".join(task["inventory"]["notices"])
    assert source.read_bytes() == b"legacy-document"
    service.stop()


def test_cancelling_page_recovery_does_not_publish_late_plan(catalog, tmp_path):
    """恢复过程取消后既不启动映射，也不把迟到的恢复结果发布为可用模板"""
    source = tmp_path / "source.pdf"
    page_pdf(source, scan=True)
    provider = RecoveryProvider(cancel=True)
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["status"] == "cancelled" and task["plan"] is None
    assert not provider.calls
    service.stop()


def test_revisions_and_data_bindings_are_frozen_without_manual_cleanup(tmp_path):
    """修订保留新增内容并清除删除内容，绑定和编辑锁定自动解除，原文不被改写"""
    source = tmp_path / "revisions.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    for tag, text in (("del", "旧姓名"), ("ins", "当前姓名")):
        revision = etree.SubElement(paragraph._p, w(tag))
        etree.SubElement(
            etree.SubElement(revision, w("r")), w("t" if tag == "ins" else "delText")
        ).text = text
    control = etree.SubElement(paragraph._p, w("sdt"))
    properties = etree.SubElement(control, w("sdtPr"))
    etree.SubElement(properties, w("dataBinding"))
    etree.SubElement(properties, w("lock"))
    content = etree.SubElement(control, w("sdtContent"))
    etree.SubElement(etree.SubElement(content, w("r")), w("t")).text = "固定信息"
    doc.save(source)
    package = TemplatePackage(source)
    root = package.parts["word/document.xml"]
    assert not root.xpath(".//w:ins | .//w:del | .//w:dataBinding | .//w:lock", namespaces=NS)
    assert "".join(root.xpath(".//w:t/text()", namespaces=NS)) == "当前姓名固定信息"
    assert not package.inventory()["warnings"]


def test_long_templates_are_not_rejected_by_node_count(tmp_path):
    """保留超过旧节点数量阈值的全文，识别层不再要求用户删去页面"""
    source = tmp_path / "long.docx"
    doc = Document()
    for index in range(2001):
        doc.add_paragraph(f"段落{index}")
    doc.save(source)
    inventory = TemplatePackage(source).inventory()
    assert len(inventory["nodes"]) == 2001 and not inventory["warnings"]


@pytest.mark.parametrize("kind", ["docx", "pdf"])
def test_empty_source_builds_editable_framework(catalog, tmp_path, kind):
    """为空白 DOCX 和未识别到内容的 PDF 自动生成资料框架"""
    source = tmp_path / f"empty.{kind}"
    if kind == "docx":
        Document().save(source)
    else:
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(source)

    class EmptyProvider(RecoveryProvider):
        """模拟空白页返回空结构，后续映射仍使用真实生成的段落"""

        def run_structured(self, **kwargs):
            """返回空页面，让恢复器依据当前资料字段建立占位模板"""
            if kwargs["result_model"] is RecoveredPage:
                return RecoveredPage(blocks=[])
            return super().run_structured(**kwargs)

    service = Templates(catalog, tmp_path / "data", EmptyProvider())
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["status"] == "completed" and task["review"]["ready"]
    assert "占位框架" in "".join(task["inventory"]["notices"])
    assert Document(service.source(task["id"])).paragraphs[0].text == "姓名：待填姓名"
    service.stop()


def test_blank_framework_uses_photo_placeholder_without_user_image(tmp_path):
    """空白框架支持照片映射，但发送识别的占位图不复制真实用户照片"""
    content = simple_document()
    content.personal.photo = "data:image/png;base64,cHJpdmF0ZS1waG90bw=="
    output = tmp_path / "blank.docx"
    blank_template(output, content, [])
    package = TemplatePackage(output)
    photos = [row for row in package.inventory()["nodes"] if row["kind"] == "image"]
    assert len(photos) == 1
    assert b"private-photo" not in package.image(photos[0]["id"])


def test_editable_source_never_uses_lossy_recovery_when_word_is_unavailable(tmp_path, monkeypatch):
    """Word 不可用不会使原生 DOCX 退化为 OCR 文字，未支持对象留在副本并报告"""
    source, output = tmp_path / "source.docx", tmp_path / "original.docx"
    doc = Document()
    etree.SubElement(doc.add_paragraph("原姓名")._p, w("object"))
    doc.save(source)
    original = source.read_bytes()
    monkeypatch.setattr(
        "resume_maker.integrations.word.recovery.render_word",
        lambda *_: (None, "Word 不可用"),
    )
    provider = RecoveryProvider()
    package, notices = prepare_template(
        source, output, provider, None, Event(), lambda *_: None, simple_document(), []
    )
    assert not provider.pages and package.inventory()["warnings"]
    assert package.parts["word/document.xml"].find(".//w:object", NS) is not None
    assert not notices and source.read_bytes() == original


def test_recovery_retries_invalid_photo_coordinates(catalog, tmp_path):
    """裁剪坐标无效时重试并在成功后保留有效照片"""
    source = tmp_path / "source.pdf"
    page_pdf(source, scan=True)

    class RetryProvider(RecoveryProvider):
        """第一次提供越界坐标，第二次提供完整可编辑内容"""

        failed = False

        def run_structured(self, **kwargs):
            """通过真实模型约束抛出验证异常来覆盖恢复器的重试分支"""
            if kwargs["result_model"] is RecoveredPage and not self.failed:
                self.failed = True
                return RecoveredPage(blocks=[RecoveredBlock(image_box=[0, 0, 2, 1])])
            return super().run_structured(**kwargs)

    provider = RetryProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["status"] == "completed" and task["review"]["ready"]
    assert any("重试恢复" in event["text"] for event in task["events"])
    assert len(provider.pages) == 1 and provider.failed
    service.stop()


@pytest.mark.parametrize("same_paragraph", [False, True])
def test_trial_layout_conflict_keeps_native_table_and_reports_error(
    catalog, tmp_path, monkeypatch, same_paragraph
):
    """独立左右栏可直接适配，同一段落中的交叉栏目反馈修正，两者均保留原生表格"""
    from resume_maker.domain.templates import TemplatePlan

    source = tmp_path / "shared.docx"
    doc = Document()
    doc.add_paragraph("原姓名")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "教育背景", "项目经历"
    if same_paragraph:
        table.cell(0, 0).text, table.cell(0, 1).text = "教育背景 / 项目经历", ""
    doc.save(source)
    original = source.read_bytes()

    class SharedProvider(RecoveryProvider):
        """首轮按真实共用表格映射，恢复后使用普通段落映射"""

        def run_structured(self, **kwargs):
            """首轮结果能通过节点检查，但实际栏目编排无法独立移动同一行"""
            if issubclass(kwargs["result_model"], TemplatePlan) and not self.pages:
                self.calls.append(kwargs)
                rows = TemplatePackage(kwargs["workspace"] / "original.docx").inventory()["nodes"]
                return kwargs["result_model"](
                    summary="共用容器",
                    fields=[
                        TextBinding(
                            node=row["id"],
                            quote=label,
                            target="personal.name"
                            if label == "原姓名"
                            else "section-title:" + label,
                        )
                        for row in rows
                        if row["kind"] == "p" and row["text"]
                        for label in ("原姓名", "教育背景", "项目经历")
                        if label in row["text"]
                    ],
                    repeats=[],
                    photos=[],
                    keep=[],
                    remove=[],
                    warnings=[],
                )
            return super().run_structured(**kwargs)

    def render(document, output):
        """构造冲突表格的可读页面以验证恢复流程"""
        page_pdf(output)
        return 1, None

    monkeypatch.setattr("resume_maker.integrations.word.recovery.render_word", render)
    provider = SharedProvider()
    service = Templates(catalog, tmp_path / "data", provider)
    task = completed(service, service.analyze(source, simple_document())["id"])
    assert task["status"] == "completed" and task["review"]["ready"] == (not same_paragraph)
    if same_paragraph:
        assert any("栏目边界" in error for error in task["review"]["errors"])
        feedback = json.loads(provider.calls[1]["prompt"].split("\n")[-1])
        assert any("栏目边界" in error for error in feedback["validation"]["errors"])
    assert len(provider.calls) == task["attempts"] == (2 if same_paragraph else 1)
    assert not provider.pages
    assert len(Document(service.source(task["id"])).tables) == 1
    assert source.read_bytes() == original
    service.stop()
