"""项目正文排序的版本保存、显隐和文档输出"""

from copy import deepcopy

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError

from resume_maker.api import create_app
from resume_maker.core.config import Config
from resume_maker.core.errors import Problem
from resume_maker.domain.experience import displayed_experience, replace_field
from resume_maker.domain.models import Experience, ProjectVisibility
from resume_maker.domain.project_layout import project_body_entries, project_body_order
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text
from tests.support.data import body_text, project_info
from tests.support.layouts import (
    ORDER,
    order_template,
    plan_for,
    project_document,
)


def ordered_settings(hidden=False):
    """隐藏和恢复不改变排序，亮点由外层组合独立选择"""
    return ProjectVisibility(
        order=ORDER,
        fields={"role": not hidden, "stack": True},
        custom_fields={"team": True, "link": not hidden},
    )


def test_order_normalizes_available_fields_without_mutating_settings():
    """丢失字段忽略、新增字段追加，切回旧版本仍能使用原位置"""
    content = project_info()
    settings = ProjectVisibility(order=["custom:missing", "highlights", "custom:team"])
    original = settings.model_dump()
    assert project_body_order(content, settings) == [
        "highlights",
        "custom:team",
        "role",
        "stack",
        "description",
        "custom:link",
        "custom:blank",
    ]
    content["custom_fields"].append(
        {"id": "missing", "label": "新条目", "value": "内容", "visible": True}
    )
    assert project_body_order(content, settings)[0] == "custom:missing"
    assert settings.model_dump() == original
    visible = displayed_experience(content, ordered_settings())
    entries = project_body_entries(visible, ["two", "one"], ordered_settings())
    assert [entry["key"] for entry in entries] == [
        *ORDER[:2],
        "highlights",
        "highlights",
        *ORDER[3:],
        "custom:missing",
    ]
    assert [entry["text"] for entry in entries if entry["key"] == "highlights"] == [
        "Export Word",
        "Parse documents",
    ]


@pytest.mark.parametrize(
    "order", [["title"], ["period"], ["role", "role"], ["custom:"], ["unknown"]]
)
def test_fixed_or_invalid_order_is_rejected(order):
    """标题时间不能混入正文，重复和无效字段不能进入持久化数据"""
    with pytest.raises(ValidationError):
        ProjectVisibility(order=order)


@pytest.mark.parametrize(
    "layout",
    ["builtin", "body", "cell", "row", "details", "shared_header", "split_rows", "highlight_slots"],
)
@pytest.mark.parametrize("hidden", [False, True])
def test_project_body_order_across_layouts(tmp_path, layout, hidden):
    """自定义字段可夹在元信息和亮点之间，标题时间固定顶部，每份记录独立且无内容重复"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    value = project_info()
    projects = [
        {
            "project_id": "first",
            "content": value,
            "highlight_ids": ["two"] if hidden else ["two", "one"],
        }
    ]
    second = deepcopy(projects[0])
    second["project_id"] = "second"
    second["content"]["title"] = "另一个项目"
    projects.append(second)
    original = deepcopy(projects)
    document = project_document()
    document.project_visibility = {
        "first": ordered_settings(hidden),
        "second": ProjectVisibility(
            order=["stack", "role", "description", "custom:link", "highlights"]
        ),
    }
    if layout == "builtin":
        write_full_resume(output, document.model_dump(), projects)
    else:
        plan = order_template(source, layout)
        before, plan_before = source.read_bytes(), plan.model_dump()
        fill_template(source, output, plan, document.model_dump(), projects)
        assert source.read_bytes() == before and plan.model_dump() == plan_before
    text = body_text(output)
    first, second = text.split("另一个项目")
    expected = [value["title"], value["period"]]
    expected += [] if hidden else ["https://example.test/project"]
    expected += [value["description"], "Export Word"]
    expected += [] if hidden else ["Parse documents", "角色原值"]
    expected += ["团队：隐藏团队原值", "Python"]
    positions = [first.index(term) for term in expected]
    assert positions == sorted(positions)
    assert all(first.count(term) == 1 for term in expected)
    assert "Old " not in text and "〔待填写〕" not in text and "空白条目" not in text
    assert (
        second.index(value["description"])
        < second.index("https://example.test/project")
        < second.index("Export Word")
    )
    assert "角色原值" not in second and "Python" not in second
    if hidden:
        assert all(
            term not in first
            for term in ("https://example.test", "角色原值", "担任角色", "Parse documents")
        )
    if layout in {"body", "cell", "row"}:
        paragraphs = Document(output).element.body.iter(w("p"))
        stack = next(node for node in paragraphs if paragraph_text(node) == "技术栈：Python")
        assert stack.find("w:pPr/w:spacing", NS).get(w("line")) == "300"
        assert stack.find("w:pPr/w:spacing", NS).get(w("before")) == "60"
        runs = stack.findall(w("r"))
        assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
        assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"
    assert projects == original


def test_ordered_highlights_keep_body_weight_and_separate_label_cleanup(tmp_path):
    """亮点正文不继承粗体标题，原独立角色标签随值一起迁移并且没有重复"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    for text in ("Title", "Period", "Role:", "Role value"):
        doc.add_paragraph(text)
    point = doc.add_paragraph()
    point.add_run("旧亮点：").bold = True
    point.add_run("旧正文").bold = False
    doc.save(source)
    plan = plan_for(
        source,
        [
            ("Title", "Title", "title"),
            ("Period", "Period", "period"),
            ("Role value", "Role value", "role"),
            ("旧亮点：旧正文", "旧亮点：旧正文", "details"),
        ],
        keep=["Role:"],
    )
    document = project_document()
    document.project_visibility["project"] = ordered_settings()
    fill_template(
        source,
        output,
        plan,
        document.model_dump(),
        [{"project_id": "project", "content": project_info(), "highlight_ids": ["one"]}],
    )
    text = body_text(output)
    assert "Role:" not in text
    assert text.count("担任角色") == 1
    package = TemplatePackage(output)
    point = next(
        node
        for node in package.nodes.values()
        if node.tag == w("p") and paragraph_text(node) == "Parser：Parse documents"
    )
    runs = point.findall(w("r"))
    assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
    assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"


@pytest.mark.parametrize("layout", ["builtin", "body", "row", "details"])
def test_body_order_publishes_new_revision_and_preserves_old_export(
    catalog, project, tmp_path, layout
):
    """调整进入草稿，旧版不变，提交一次产生新版本并覆盖旧简历中的排序设置"""
    identifier, initial = project["id"], project["head_revision"]
    catalog.put_draft(identifier, initial, "experience", project_info(), 0)
    base = catalog.save_revision(identifier, initial, initial)
    original = deepcopy(base["content"])
    meta = {key: value for key, value in original.items() if key != "highlights"}
    working = catalog.put_draft(identifier, base["id"], "meta", {**meta, "body_order": ORDER}, 0)
    assert working["content"]["body_order"] == ORDER
    assert catalog.revision(base["id"])["content"] == original
    count = len(catalog.db.all("SELECT * FROM revisions"))
    saved = catalog.save_revision(identifier, base["id"], base["id"])
    assert saved["number"] == base["number"] + 1
    assert len(catalog.db.all("SELECT * FROM revisions")) == count + 1
    assert saved["content"]["body_order"] == ORDER
    assert saved["content"]["highlights"] == original["highlights"]
    assert not catalog.working(identifier, saved["id"])["drafts"]
    settings = ProjectVisibility(
        order=["stack", "role", "highlights", "description", "custom:link"],
        fields={"role": True, "stack": True},
        custom_fields={"team": True},
    )
    assert project_body_order(saved["content"], settings)[:6] == ORDER
    assert project_body_order(original, settings)[:5] == settings.order
    document = project_document()
    document.project_visibility[identifier] = settings
    output = tmp_path / "result.docx"
    source = tmp_path / "source.docx"
    plan = None if layout == "builtin" else order_template(source, layout)
    for content, expected in [
        (
            saved["content"],
            [
                "example.test",
                "Document processing",
                "Export Word",
                "Parse documents",
                "角色原值",
                "隐藏团队原值",
                "Python",
            ],
        ),
        (
            original,
            [
                "Python",
                "角色原值",
                "Export Word",
                "Parse documents",
                "Document processing",
                "example.test",
            ],
        ),
    ]:
        projects = [{"project_id": identifier, "content": content, "highlight_ids": ["two", "one"]}]
        if plan is None:
            write_full_resume(output, document.model_dump(), projects)
        else:
            fill_template(source, output, plan, document.model_dump(), projects)
        text = body_text(output)
        positions = [text.index(term) for term in expected]
        assert positions == sorted(positions)
    assert (
        replace_field(saved["content"], "experience", {"title": "旧建议", "highlights": []})[
            "body_order"
        ]
        == ORDER
    )


@pytest.mark.parametrize("order", [["title"], ["period"], ["role", "role"], ["custom:"]])
def test_version_order_validated(order):
    """版本和旧布局使用相同合法字段集合，禁止重复及固定顶部字段"""
    with pytest.raises(ValidationError):
        Experience.model_validate({**project_info(), "body_order": order})


def test_atomic_discard_clears_only_confirmed_revision(catalog, project, populated):
    """撤销当前基线的元信息、排序和亮点增删并保留其他版本及简历显隐"""
    identifier, revision = project["id"], populated["id"]
    other = project["head_revision"]
    catalog.put_draft(identifier, other, "meta", {"role": "其他基线"}, 0)
    for field, value in [
        ("meta", {"role": "修改角色", "body_order": ORDER}),
        ("order", ["two", "one"]),
        ("highlight:one", None),
        ("highlight:added", {"title": "新增", "text": "内容", "evidence": []}),
    ]:
        catalog.put_draft(identifier, revision, field, value, 0)
    versions = {
        draft["field"]: draft["version"]
        for draft in catalog.working(identifier, revision)["drafts"]
    }
    before = catalog.db.all("SELECT * FROM revisions")
    catalog.discard_drafts(identifier, revision, versions)
    assert catalog.working(identifier, revision) == {"content": populated["content"], "drafts": []}
    assert catalog.working(identifier, other)["content"]["role"] == "其他基线"
    assert catalog.db.all("SELECT * FROM revisions") == before


@pytest.mark.parametrize("concurrent", ["update", "add", "remove"])
def test_discard_conflict_keeps_entire_working_copy(catalog, project, populated, concurrent):
    """确认期间草稿集合发生变化时整体拒绝撤销"""
    identifier, revision = project["id"], populated["id"]
    catalog.put_draft(identifier, revision, "meta", {"body_order": ORDER}, 0)
    catalog.put_draft(identifier, revision, "order", ["two", "one"], 0)
    versions = {"meta": 1, "order": 1}
    if concurrent == "update":
        catalog.put_draft(identifier, revision, "meta", {"role": "其他窗口"}, 1)
    elif concurrent == "add":
        catalog.put_draft(identifier, revision, "highlight:one", None, 0)
    else:
        catalog.discard_draft(identifier, revision, "order", 1)
    before = catalog.working(identifier, revision)
    with pytest.raises(Problem, match="重新确认"):
        catalog.discard_drafts(identifier, revision, versions)
    assert catalog.working(identifier, revision) == before


def test_discard_api_requires_matching_snapshot(tmp_path):
    """HTTP 撤销需要已确认的草稿版本，冲突时返回 409，成功不产生新版本"""
    app = create_app(Config(data_dir=tmp_path / "data", token="qa"))
    catalog = app.state.services.catalog
    project = catalog.create_project("测试", [str(tmp_path)])
    identifier, revision = project["id"], project["head_revision"]
    catalog.put_draft(identifier, revision, "meta", {"body_order": ORDER}, 0)
    with TestClient(app) as client:
        url = f"/api/projects/{identifier}/drafts/discard"
        assert (
            client.post(
                url,
                json={"base_revision": revision, "versions": {}},
                headers={"x-resume-token": "qa"},
            ).status_code
            == 409
        )
        response = client.post(
            url,
            json={"base_revision": revision, "versions": {"meta": 1}},
            headers={"x-resume-token": "qa"},
        )
        assert response.status_code == 200
        assert response.json()["drafts"] == []
        assert len(catalog.db.all("SELECT * FROM revisions")) == 1
