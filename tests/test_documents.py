from zipfile import ZipFile

from docx import Document
from lxml import etree

from resume_maker.documents import NS, Documents, inspect_template
from resume_maker.models import ResumeItem


def test_export_replaces_only_selected_region(catalog, project, populated, tmp_path, monkeypatch):
    source = tmp_path / "reference.docx"
    document = Document()
    document.add_paragraph("保留个人信息")
    document.add_paragraph("项目经历")
    document.add_paragraph("旧项目标题")
    document.add_paragraph("旧项目正文")
    document.add_paragraph("技能证书")
    document.add_paragraph("保留证书信息")
    document.save(source)
    inspection = inspect_template(source)
    assert (inspection["suggested_start"], inspection["suggested_end"]) == (2, 4)
    exporter = Documents(catalog, tmp_path / "data")
    template = exporter.import_template(source, "Reference", 2, 4)
    resume = catalog.save_resume(
        "My resume",
        template["id"],
        [ResumeItem(project_id=project["id"], revision_id=populated["id"], highlight_ids=["two"])],
    )
    monkeypatch.setattr("resume_maker.documents.render_word", lambda *_: (None, "No renderer"))
    result = exporter.export(resume["id"])
    output = tmp_path / "data" / "exports" / result["id"] / "resume.docx"
    with ZipFile(source) as old, ZipFile(output) as new:
        root = etree.fromstring(new.read("word/document.xml"))
        text = "".join(root.xpath("//w:t/text()", namespaces=NS))
        assert "保留个人信息" in text and "保留证书信息" in text
        assert "旧项目" not in text
        assert "Export Word" in text and "Parse documents" not in text
        for entry in old.namelist():
            if entry != "word/document.xml":
                assert old.read(entry) == new.read(entry)
    assert result["manifest"]["items"][0]["revision_id"] == populated["id"]
