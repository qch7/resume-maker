"""验证已保存模板的基本信息随当前资料自动扩展，无需重新识别。"""

from copy import deepcopy

import pytest
from docx import Document
from docx.shared import Pt
from test_template_mapping import resume_content

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import CustomInfoField, ResumeDocument, ResumeSection, SectionEntry
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.template_fill import fill_template
from resume_maker.integrations.word.template_map import TemplatePackage, paragraph_text


def personal_template(path, table):
    """制作有已满联系方式行和教育栏目的模板，覆盖正文与可扩展单元格。"""
    doc = Document()
    contact = (
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if table else doc.add_paragraph()
    )
    contact.paragraph_format.left_indent = Pt(16)
    contact.paragraph_format.first_line_indent = Pt(0)
    contact.paragraph_format.space_after = Pt(3)
    contact.paragraph_format.line_spacing = 1.25
    contact.paragraph_format.tab_stops.add_tab_stop(Pt(230))
    run = contact.add_run("邮箱：旧邮箱\t电话：旧电话")
    run.font.name = "宋体"
    run.font.size = Pt(10)
    doc.add_paragraph("教育背景")
    doc.add_paragraph("旧学校")
    doc.save(path)
    package = TemplatePackage(path)
    paragraphs = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="可扩展的个人资料区",
        fields=[
            TextBinding(node=paragraphs[0]["id"], quote="旧邮箱", target="personal.email"),
            TextBinding(node=paragraphs[0]["id"], quote="旧电话", target="personal.phone"),
        ],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=paragraphs[2]["id"],
                end=paragraphs[2]["id"],
                sample_start=paragraphs[2]["id"],
                sample_end=paragraphs[2]["id"],
                fields=[TextBinding(node=paragraphs[2]["id"], quote="旧学校", target="title")],
            )
        ],
        photos=[],
        keep=[paragraphs[1]["id"]],
        remove=[],
        warnings=[],
    )
    return plan


@pytest.mark.parametrize("table", [False, True])
def test_new_personal_fields_extend_existing_area_without_changing_template(tmp_path, table):
    """新条目继承同区字体、缩进和行距；后续隐藏、清空、改名都不留下旧行。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = personal_template(source, table)
    original, original_plan = source.read_bytes(), deepcopy(plan)
    document = ResumeDocument.model_validate(
        {
            "personal": {
                "email": "new@example.test",
                "phone": "10000000000",
                "website": "https://example.test/profile",
                "location": "不展示的城市",
                "hidden_fields": ["location"],
                "custom_fields": [
                    {"id": "language", "label": "语言", "value": "中文"},
                    {"id": "blank", "label": "空信息", "value": ""},
                    {"id": "secret", "label": "隐藏信息", "value": "不展示", "visible": False},
                ],
            },
            "sections": [
                {
                    "id": "education",
                    "title": "教育背景",
                    "entries": [{"id": "school", "title": "新学校"}],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
            ],
        }
    )
    fill_template(source, output, plan, document.model_dump(), [])
    filled = Document(output)
    paragraphs = filled.tables[0].cell(0, 0).paragraphs if table else filled.paragraphs
    assert [paragraph.text for paragraph in paragraphs[:3]] == [
        "邮箱：new@example.test\t电话：10000000000",
        "个人主页：https://example.test/profile",
        "语言：中文",
    ]
    for paragraph in paragraphs[1:3]:
        assert paragraph.paragraph_format.left_indent == Pt(16)
        assert paragraph.paragraph_format.space_after == Pt(3)
        assert paragraph.paragraph_format.line_spacing == 1.25
        assert paragraph.runs[0].font.name == "宋体"
        assert paragraph.runs[0].font.size == Pt(10)
    assert [paragraph.text for paragraph in filled.paragraphs[-2:]] == ["教育背景", "新学校"]

    # 每次以当前资料生成，绝不把补出的行持久化到原模板中。
    document.personal.hidden_fields.append("website")
    document.personal.custom_fields[0].label = "工作方式"
    document.personal.custom_fields[0].value = "远程"
    fill_template(source, output, plan, document.model_dump(), [])
    text = "\n".join(
        paragraph_text(node)
        for node in TemplatePackage(output).nodes.values()
        if node.tag == w("p")
    )
    assert "工作方式：远程" in text
    assert not any(
        value in text
        for value in ("个人主页", "语言", "空信息", "隐藏信息", "不展示", "〔待填写〕")
    )
    document.personal.custom_fields = []
    fill_template(source, output, plan, document.model_dump(), [])
    filled = Document(output)
    paragraphs = filled.tables[0].cell(0, 0).paragraphs if table else filled.paragraphs
    assert len(paragraphs) == (1 if table else 3)
    assert source.read_bytes() == original and plan == original_plan


def test_automatic_rows_do_not_mask_missing_photo_mapping(tmp_path):
    """个人区和栏目可扩展不会绕过未识别照片校验，也不写回模板或不完整输出。"""
    source, output = tmp_path / "source.docx", tmp_path / "filled.docx"
    plan = personal_template(source, False)
    original = source.read_bytes()
    document = ResumeDocument(
        sections=[
            ResumeSection(
                id="new", title="新增栏目", entries=[SectionEntry(id="entry", title="新记录")]
            ),
            ResumeSection(id="projects", title="项目经历", kind="projects"),
        ]
    )
    document.personal.website = "https://example.test"
    document.personal.custom_fields = [CustomInfoField(id="new", label="语言", value="中文")]
    document.personal.photo = resume_content().personal.photo
    with pytest.raises(Problem, match="照片"):
        fill_template(source, output, plan, document.model_dump(), [])
    assert source.read_bytes() == original and not output.exists()


@pytest.mark.parametrize(
    "hidden, expected",
    [
        (
            ["email"],
            "专业成绩：GPA 4.2/5.0\t电话：10000000000\n个人主页：https://example.test\t所在地：杭州",
        ),
        (
            ["phone"],
            "专业成绩：GPA 4.2/5.0\t所在地：杭州\n邮箱：sample@example.test\n个人主页：https://example.test",
        ),
        (["gpa", "phone", "email"], "个人主页：https://example.test\t所在地：杭州"),
        (["gpa", "phone", "email", "location"], "个人主页：https://example.test"),
        (["gpa", "phone", "email", "location", "website"], ""),
    ],
)
def test_hidden_contacts_remove_labels_and_fill_grid_gaps(tmp_path, hidden, expected):
    """两列混合段落隐藏整条信息，后续字段与自动新增主页连续补位且保留标签和值的样式。"""
    source, output = tmp_path / "grid.docx", tmp_path / "filled.docx"
    doc = Document()
    contact = doc.add_paragraph()
    contact.paragraph_format.tab_stops.add_tab_stop(Pt(230))
    contact.paragraph_format.left_indent = Pt(12)
    contact.paragraph_format.line_spacing = 1.5
    slots = [
        ("专业成绩：GPA ", "gpa"),
        ("电话：", "phone"),
        ("邮箱：", "email"),
        ("所在地：", "location"),
    ]
    for index, (label, key) in enumerate(slots):
        if index:
            contact.add_run("\n" if index == 2 else "\t")
        contact.add_run(label).bold = True
        contact.add_run(f"old-{key}").italic = True
    doc.add_paragraph("后续内容")
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="两列基本信息",
        fields=[
            TextBinding(node=paragraphs[0]["id"], quote=f"old-{key}", target=f"personal.{key}")
            for _, key in slots
        ],
        repeats=[],
        photos=[],
        keep=[paragraphs[1]["id"]],
        remove=[],
        warnings=[],
    )
    document = ResumeDocument(
        personal={
            "gpa": "4.2/5.0",
            "phone": "10000000000",
            "email": "sample@example.test",
            "location": "杭州",
            "website": "https://example.test",
            "hidden_fields": hidden,
        },
        sections=[ResumeSection(id="projects", title="项目经历", kind="projects")],
    )
    # AI 的映射顺序可以乱序，补位顺序仍须来自 Word 中的实际列位置。
    plan.fields.reverse()
    original, original_plan = source.read_bytes(), deepcopy(plan)
    fill_template(source, output, plan, document.model_dump(), [])
    result = Document(output)
    assert [p.text for p in result.paragraphs] == ([expected] if expected else []) + ["后续内容"]
    if expected:
        assert result.paragraphs[0].paragraph_format.left_indent == Pt(12)
        assert result.paragraphs[0].paragraph_format.tab_stops[0].position == Pt(230)
        assert result.paragraphs[0].paragraph_format.line_spacing == 1.5
    if "email" not in hidden:
        email_runs = [
            run for p in result.paragraphs for run in p.runs if "sample@example.test" in run.text
        ]
        assert email_runs and all(run.italic for run in email_runs)
    assert source.read_bytes() == original and plan == original_plan
    document.personal.hidden_fields = []
    fill_template(source, output, plan, document.model_dump(), [])
    restored = "\n".join(p.text for p in Document(output).paragraphs)
    assert all(label in restored for label, _ in slots)
    assert restored.count("https://example.test") == 1


def test_hidden_standalone_field_removes_its_paragraph(tmp_path):
    """单列标签和值一起隐藏，后面的姓名段落自然上移，固定标题不被当作个人条目删除。"""
    source, output = tmp_path / "single.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("邮箱：old-email")
    doc.add_paragraph("旧姓名 · 个人简历")
    doc.add_paragraph("男 21岁")
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="单列资料及固定标题",
        fields=[
            TextBinding(node=paragraphs[0]["id"], quote="old-email", target="personal.email"),
            TextBinding(node=paragraphs[1]["id"], quote="旧姓名", target="personal.name"),
            TextBinding(node=paragraphs[2]["id"], quote="男", target="personal.gender"),
            TextBinding(node=paragraphs[2]["id"], quote="21", target="personal.age"),
        ],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    document = ResumeDocument(
        personal={
            "name": "新姓名",
            "gender": "男",
            "email": "sample@example.test",
            "age": "21",
            "hidden_fields": ["email", "age"],
        },
        sections=[ResumeSection(id="projects", title="项目经历", kind="projects")],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    assert [p.text for p in Document(output).paragraphs] == ["新姓名 · 个人简历", "男 "]


def test_hidden_table_contacts_compact_each_column_and_remove_empty_rows(tmp_path):
    """表格个人区每列向上补齐，删去空行；隐藏全部信息后整张空表移除。"""
    source, output = tmp_path / "table.docx", tmp_path / "filled.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    slots = [
        ("gpa", "成绩"),
        ("phone", "电话"),
        ("email", "邮箱"),
        ("location", "所在地"),
        ("website", "主页"),
    ]
    for index, (key, label) in enumerate(slots):
        table.cell(index // 2, index % 2).text = f"{label}：old-{key}"
    doc.add_paragraph("后续内容")
    doc.save(source)
    package = TemplatePackage(source)
    paragraphs = [row for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = TemplatePlan(
        summary="表格资料",
        fields=[
            TextBinding(node=paragraphs[index]["id"], quote=f"old-{key}", target=f"personal.{key}")
            for index, (key, _) in enumerate(slots)
        ],
        repeats=[],
        photos=[],
        keep=[paragraphs[-1]["id"]],
        remove=[],
        warnings=[],
    )
    document = ResumeDocument(
        personal={
            "gpa": "4.2",
            "phone": "10000000000",
            "email": "sample@example.test",
            "location": "杭州",
            "website": "https://example.test",
            "hidden_fields": ["email"],
        },
        sections=[ResumeSection(id="projects", title="项目经历", kind="projects")],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    actual = Document(output)
    assert [[cell.text for cell in row.cells] for row in actual.tables[0].rows] == [
        ["成绩：4.2", "电话：10000000000"],
        ["主页：https://example.test", "所在地：杭州"],
    ]
    assert [p.text for p in actual.paragraphs] == ["后续内容"]
    document.personal.hidden_fields = [key for key, _ in slots]
    fill_template(source, output, plan, document.model_dump(), [])
    assert not Document(output).tables


def test_compaction_keeps_repeated_quote_label_separate_from_value(tmp_path):
    """标签与旧值同字时仍按正确出现次数替换，重排不能把标签替换为邮箱地址。"""
    source, output = tmp_path / "repeated.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("邮箱：邮箱\t所在地：城市")
    doc.save(source)
    package = TemplatePackage(source)
    identifier = next(row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p")
    plan = TemplatePlan(
        summary="重复引文",
        fields=[
            TextBinding(node=identifier, quote="邮箱", occurrence=2, target="personal.email"),
            TextBinding(node=identifier, quote="城市", target="personal.location"),
        ],
        repeats=[],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    document = ResumeDocument(
        personal={
            "email": "sample@example.test",
            "location": "杭州",
            "hidden_fields": ["location"],
        },
        sections=[ResumeSection(id="projects", title="项目经历", kind="projects")],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    assert Document(output).paragraphs[0].text == "邮箱：sample@example.test"


def test_grid_inside_single_cell_keeps_its_column_positions(tmp_path):
    """外层单元格不改变内层制表位布局，隐藏左列邮箱后右列电话仍在原列。"""
    source, output = tmp_path / "cell.docx", tmp_path / "filled.docx"
    plan = personal_template(source, True)
    document = ResumeDocument(
        personal={
            "email": "sample@example.test",
            "phone": "10000000000",
            "hidden_fields": ["email"],
        },
        sections=[ResumeSection(id="projects", title="项目经历", kind="projects")],
    )
    fill_template(source, output, plan, document.model_dump(), [])
    assert Document(output).tables[0].cell(0, 0).text == "\t电话：10000000000"
