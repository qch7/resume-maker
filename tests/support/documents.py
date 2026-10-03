"""合成 DOCX、简历资料和记录映射"""

import base64
from io import BytesIO

import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt
from lxml import etree

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.mapping import NS, TemplatePackage, paragraph_text


def header_content(**personal):
    """构造有新增字段、长网址和照片的真实填充资料"""
    return ResumeDocument.model_validate(
        {
            "personal": {
                "name": "NEW NAME",
                "phone": "10000000000",
                "email": "long.address@example.test",
                "job_title": "Software engineer",
                "gender": "女",
                "age": "28",
                "gpa": "4.2/5.0",
                "location": "Example city",
                "website": "https://example.test/portfolio",
                "photo": resume_content().personal.photo,
                **personal,
            },
            "sections": [{"id": "projects", "title": "项目经历", "kind": "projects"}],
        }
    )


def add_annotation_part(doc, kind, content):
    """为脱敏文档添加带内容类型和关系的标准注释部件"""
    part = Part(
        PackURI(f"/word/{kind}s.xml"),
        f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
        f'<w:{kind}s xmlns:w="{NS["w"]}">{content}</w:{kind}s>'.encode(),
        doc.part.package,
    )
    doc.part.relate_to(part, getattr(RELATIONSHIP_TYPE, kind.upper() + "S"))


def record_plan(path, targets, table=False):
    """按当前原文建立单条重复范围，任意空段落和样式变化都不会依赖节点编号"""
    package = TemplatePackage(path)
    nodes = {
        paragraph_text(node): key
        for key, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    fields = [
        TextBinding(node=nodes[text], quote=quote, target=target) for text, quote, target in targets
    ]
    paragraphs = list(dict.fromkeys(package.node(field.node) for field in fields))
    roots = list(
        dict.fromkeys(next(node.iterancestors(w("tr"))) if table else node for node in paragraphs)
    )
    return TemplatePlan(
        summary="独立条目布局测试",
        fields=[],
        repeats=[
            RepeatBinding(
                section="独立栏目",
                start=package.ids[roots[0]],
                end=package.ids[roots[-1]],
                sample_start=package.ids[roots[0]],
                sample_end=package.ids[roots[-1]],
                fields=fields,
            )
        ],
        keep=[key for text, key in nodes.items() if text not in {item[0] for item in targets}],
        photos=[],
        remove=[],
        warnings=[],
    )


def record_content(entries):
    """构造任意普通栏目和必需的项目区，隐藏规则和实际个人信息界面相同"""
    return ResumeDocument.model_validate(
        {
            "sections": [
                {"id": "records", "title": "独立栏目", "entries": entries},
                {"id": "projects", "kind": "projects", "title": "项目经历"},
            ]
        }
    )


def photo_bytes(color):
    """生成纯色照片用于核验包内图片替换"""
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 30, 40), False)
    pixmap.clear_with(color)
    return pixmap.tobytes("png")


def resume_content():
    """构造包含个人资料、教育、项目和隐藏字段的完整脱敏资料"""
    return ResumeDocument.model_validate(
        {
            "personal": {
                "name": "测试新姓名",
                "phone": "10000000000",
                "email": "new@example.test",
                "job_title": "文档开发工程师",
                "location": "测试城市",
                "photo": "data:image/png;base64," + base64.b64encode(photo_bytes(180)).decode(),
                "custom_fields": [{"id": "custom", "label": "语言", "value": "中文"}],
            },
            "sections": [
                {
                    "id": "education",
                    "title": "教育背景",
                    "kind": "education",
                    "entries": [
                        {
                            "id": "school1",
                            "title": "新大学一",
                            "subtitle": "计算机 · 本科",
                            "period": "2020–2024",
                            "details": "研究文档结构\n完成排版系统",
                        },
                        {
                            "id": "school2",
                            "title": "新大学二",
                            "subtitle": "计算机 · 硕士",
                            "period": "2024–2026",
                            "details": "构建自动填充工具",
                        },
                    ],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
            ],
        }
    )


def project_content():
    """只选择一条固定项目亮点以防把未选择的说明写入 Word"""
    return [
        {
            "highlight_ids": ["chosen"],
            "content": {
                "title": "新文档项目",
                "period": "2025",
                "role": "开发",
                "stack": ["Python"],
                "description": "生成可编辑文档",
                "highlights": [
                    {"id": "chosen", "title": "自动映射", "text": "覆盖姓名和教育"},
                    {"id": "skipped", "title": "不应输出", "text": "未选中的亮点"},
                ],
            },
        }
    ]


def make_template(path, *, decoration=False):
    """建立含表格、跨样式文字、页眉页脚、文本框和照片的陌生 DOCX"""
    doc = Document()
    doc.sections[0].page_width = Cm(21)
    doc.sections[0].page_height = Cm(29.7)
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(10)
    name = doc.add_paragraph()
    name.add_run("旧").bold = True
    name.add_run("姓名").italic = True
    name.add_run(" · 个人简历")
    doc.add_picture(BytesIO(photo_bytes(20)), width=Cm(1.5), height=Cm(2))
    doc.sections[0].header.paragraphs[0].text = "电话：旧电话 | 邮箱：旧邮箱"
    doc.sections[0].footer.paragraphs[0].text = "旧城市"
    box = etree.fromstring(
        f'<w:p xmlns:w="{NS["w"]}" xmlns:v="{NS["v"]}"><w:r><w:pict>'
        '<v:shape id="TextBox1" style="width:220pt;height:22pt" stroked="f">'
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>旧意向</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    doc._element.body.insert(2, box)
    doc.add_paragraph("教育背景", "Heading 1")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for row in table.rows:
        for cell, text in zip(row.cells, ["旧大学", "旧专业", "旧时间"], strict=True):
            cell.text = text
        row.cells[0].add_paragraph("旧教育正文")
    if decoration:
        table.rows[0].cells[0].paragraphs[0].add_run().add_picture(
            BytesIO(photo_bytes(20)), width=Cm(0.3)
        )
    doc.add_paragraph("项目经历", "Heading 1")
    doc.add_paragraph("旧项目 · 旧项目时间").runs[0].bold = True
    doc.add_paragraph("旧项目正文\n旧第二行")
    doc.add_paragraph("旧项目二")
    doc.add_paragraph("旧多余说明")
    doc.add_paragraph("")
    doc.save(path)
    package = TemplatePackage(path)

    def locate(text):
        """按测试原文定位首次出现的段落"""
        return next(
            key
            for key, node in package.nodes.items()
            if node.tag == w("p") and paragraph_text(node) == text
        )

    def bind(text, target, quote=None):
        """从明确样本位置创建可审查映射"""
        return TextBinding(node=locate(text), quote=text if quote is None else quote, target=target)

    rows = [key for key, node in package.nodes.items() if node.tag == w("tr")]
    empty = [
        key
        for key, node in package.nodes.items()
        if node.tag == w("p") and not len(node) and package.locations[key] == "word/document.xml"
    ][-1]
    plan = TemplatePlan(
        summary="测试完整模板",
        warnings=[],
        fields=[
            bind("旧姓名 · 个人简历", "personal.name", "旧姓名"),
            bind("电话：旧电话 | 邮箱：旧邮箱", "personal.phone", "旧电话"),
            bind("电话：旧电话 | 邮箱：旧邮箱", "personal.email", "旧邮箱"),
            bind("旧城市", "personal.location"),
            bind("旧意向", "personal.job_title"),
            TextBinding(node=empty, quote="", target="personal.custom_fields"),
        ],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=rows[0],
                end=rows[1],
                sample_start=rows[0],
                sample_end=rows[0],
                fields=[
                    bind("旧大学", "title"),
                    bind("旧专业", "subtitle"),
                    bind("旧时间", "period"),
                    bind("旧教育正文", "details"),
                ],
            ),
            RepeatBinding(
                section="projects",
                start=locate("旧项目 · 旧项目时间"),
                end=locate("旧项目二"),
                sample_start=locate("旧项目 · 旧项目时间"),
                sample_end=locate("旧项目正文旧第二行"),
                fields=[
                    bind("旧项目 · 旧项目时间", "title", "旧项目"),
                    bind("旧项目 · 旧项目时间", "period", "旧项目时间"),
                    bind("旧项目正文旧第二行", "details"),
                ],
            ),
        ],
        photos=[
            node["id"]
            for node in package.inventory()["nodes"]
            if node["kind"] == "image" and not list(package.node(node["id"]).iterancestors(w("tr")))
        ],
        keep=[locate("教育背景"), locate("项目经历")],
        remove=[locate("旧多余说明")],
    )
    return package, plan
