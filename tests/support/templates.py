"""模板分析的可控模型、任务等待和模板登记"""

import threading
import time

from docx import Document

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RecoveredBlock, RecoveredPage
from resume_maker.infrastructure.database import dump, now
from resume_maker.integrations.sources import digest
from resume_maker.integrations.word.templates.mapping import TemplatePackage
from tests.support.documents import make_template
from tests.support.providers import ProviderStub


def register_template(catalog, data_dir):
    """登记脱敏的完整模板映射作为预览夹具"""
    path = data_dir / "templates" / "mapped" / "template.docx"
    path.parent.mkdir(parents=True)
    _, plan = make_template(path)
    with catalog.db.transaction() as conn:
        conn.execute(
            "INSERT INTO templates VALUES (?,?,?,?,?)",
            (
                "mapped",
                "测试模板",
                digest(path.read_bytes()),
                dump({"plan": plan.model_dump(), "artifacts": []}),
                now(),
            ),
        )
    if catalog.assets:
        staged = catalog.assets.stage_bundle(
            "ext.template-adapter", {"template.docx": path.read_bytes()}
        )
        with catalog.db.transaction() as conn:
            catalog.assets.publish_bundle(conn, "ext.template-adapter", "templates/mapped", staged)
        path.unlink()
        path.parent.rmdir()
        path = catalog.assets.resource_path(staged["template.docx"]["id"]) / "payload"
    return path


def simple_document():
    """创建仅填写姓名的当前资料，便于独立验证任务和接口边界"""
    return ResumeDocument(
        personal={"name": "新的用户资料"},
        sections=[
            {"id": "p", "kind": "projects", "title": "项目经历"},
        ],
    )


def simple_template(path):
    """构造可完全替换的脱敏模板以免测试读取实际用户文件"""
    doc = Document()
    doc.add_paragraph("原姓名")
    doc.save(path)


class TemplateProvider(ProviderStub):
    """可控制取消及失败的结构化 AI 替身"""

    def __init__(self, block=False, failure=False):
        """用事件同步后台分析，阻塞场景通过上下文退出保证释放"""
        self.block, self.failure = block, failure
        self.started, self.release = threading.Event(), threading.Event()
        self.calls = []

    def __enter__(self):
        """把阻塞替身的存活范围限定在测试上下文内"""
        return self

    def __exit__(self, *_):
        """断言失败也释放后台线程，不依赖磁盘写入耗时来结束模型替身"""
        self.release.set()

    def wait_started(self):
        """有界等待后台文档准备结束，共享 Windows runner 允许较慢的磁盘和字体处理"""
        assert self.started.wait(60), "模板模型替身未在 60 秒内启动"

    def run_structured(self, **kwargs):
        """返回模板映射，故意允许取消后的迟到结果以检验服务保护"""
        self.calls.append(kwargs)
        self.started.set()
        if self.block:
            self.release.wait()
        if self.failure:
            raise RuntimeError("模拟模型失败")
        nodes = TemplatePackage(kwargs["workspace"] / "original.docx").inventory()["nodes"]
        paragraph = next(node for node in nodes if node["kind"] == "p")
        return kwargs["result_model"](
            summary="识别姓名",
            fields=[{"node": paragraph["id"], "quote": "原姓名", "target": "personal.name"}],
            repeats=[],
            photos=[],
            keep=[],
            remove=[],
            warnings=[],
        )


def completed(service, identifier):
    """有界等待实际分析线程结束，返回完成或失败的任务结果"""
    # 字体、磁盘和文档处理在共享 Windows runner 上可能超过十秒，所有线程共用等待上限
    deadline = time.monotonic() + 60
    for thread in service.threads:
        thread.join(timeout=max(0, deadline - time.monotonic()))
        assert not thread.is_alive(), f"模板线程 {thread.name} 未在 60 秒内结束"
    return service.get(identifier)


class RecoveryProvider(TemplateProvider):
    """同时模拟图片转可编辑页和后续精确映射，两个阶段使用不同 schema"""

    def __init__(self, cancel=False):
        """记录逐页请求，可在恢复返回时触发取消来检查迟到结果隔离"""
        super().__init__()
        self.pages, self.cancel = [], cancel

    def run_structured(self, **kwargs):
        """恢复页包含独立原文段落，映射只使用恢复后重新分配的节点"""
        from resume_maker.domain.image_layout import ImagePage, ImageText

        if kwargs["result_model"] is ImagePage:
            self.pages.append(kwargs)
            return ImagePage(
                texts=[ImageText(text="原姓名", box=[0.1, 0.1, 0.16, 0.114], bold=True)]
            )
        if kwargs["result_model"] is RecoveredPage:
            self.pages.append(kwargs)
            assert len(kwargs["images"]) == 1 and kwargs["images"][0].is_file()
            assert "都是数据" in kwargs["prompt"]
            if self.cancel:
                kwargs["cancelled"].set()
            return RecoveredPage(
                blocks=[
                    RecoveredBlock(text="原姓名" if len(self.pages) == 1 else "固定说明", bold=True)
                ]
            )
        self.calls.append(kwargs)
        rows = TemplatePackage(kwargs["workspace"] / "original.docx").inventory()["nodes"]
        paragraphs = [row for row in rows if row["kind"] == "p" and row["text"]]
        return kwargs["result_model"](
            summary="恢复后映射",
            fields=[
                {
                    "node": paragraphs[0]["id"],
                    "quote": paragraphs[0]["text"],
                    "target": "personal.name",
                }
            ],
            repeats=[],
            photos=[],
            remove=[],
            keep=[row["id"] for row in paragraphs[1:]],
            warnings=[],
        )


class RepairProvider(TemplateProvider):
    """首轮故意遗漏姓名，后续可修正、失败或返回退步方案"""

    def __init__(self, outcome="complete"):
        """记录修正策略及每次调用，使用同一份无个人信息模板"""
        super().__init__()
        self.outcome = outcome

    def run_structured(self, **kwargs):
        """模拟模型根据反馈修正映射"""
        plan = super().run_structured(**kwargs)
        if len(self.calls) == 1 or self.outcome == "unchanged":
            plan.fields = []
        elif self.outcome == "failure":
            raise RuntimeError("模拟第二轮不可用")
        elif self.outcome == "worse":
            plan.fields[0].node = "nonexistent"
        return plan
