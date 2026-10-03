"""栏目、表格及项目元信息的版式样本"""

from docx import Document
from docx.shared import Pt

from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text

ORDER = ["custom:link", "description", "highlights", "role", "custom:team", "stack"]


def order_template(path, layout):
    """构造独立正文、表格、综合字段、共享标题和多亮点槽位模板"""
    if layout in {"body", "cell", "row"}:
        return metadata_template(path, layout, 7)
    doc = Document()
    texts = ["Title", "Period", "Details"]
    if layout == "split_rows":
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Title"
        table.cell(0, 1).text = "Period"
        table.cell(1, 0).merge(table.cell(1, 1)).text = "Details"
        table.cell(2, 0).merge(table.cell(2, 1)).text = "Role"
        texts.append("Role")
    else:
        if layout == "shared_header":
            texts[0] = "Title | 担任角色：Role"
        if layout == "highlight_slots":
            texts.extend(["Point One", "Point Two", "Point Three"])
        for text in texts:
            paragraph = doc.add_paragraph(text)
            if text == "Details":
                paragraph.runs[0].bold = False
    doc.save(path)
    bindings = [
        (texts[0], "Title", "title"),
        ("Period", "Period", "period"),
        ("Details", "Details", "details"),
    ]
    if layout == "split_rows":
        bindings.append(("Role", "Role", "role"))
    if layout == "highlight_slots":
        bindings.extend((text, text, "highlights") for text in texts[3:])
    plan = plan_for(path, bindings, row=layout == "split_rows")
    if layout == "shared_header":
        node = next(field.node for field in plan.repeats[0].fields if field.target == "title")
        plan.repeats[0].fields.append(TextBinding(node=node, quote="Role", target="role"))
    return plan


def generic_template(path, layout="body", padding=0):
    """构造正文、表格行或双栏模板并按原文定位节点"""
    document = Document()
    for _ in range(padding):
        document.add_paragraph()
    document.add_paragraph("Seed Person")
    labels = ["Portfolio Ω", "Recognition · β"]
    samples = [["Old project", "Old period", "Old project body"], ["Old entry body"]]
    table = (
        document.add_table(rows=0 if layout == "rows" else 1, cols=2) if layout != "body" else None
    )
    for index, (label, content) in enumerate(zip(labels, samples, strict=True)):
        if layout == "body":
            document.add_paragraph(label)
            for literal in content:
                document.add_paragraph(literal)
        elif layout == "rows":
            table.add_row().cells[0].text = label
            cell = table.add_row().cells[0]
            cell.text = content[0]
            for literal in content[1:]:
                cell.add_paragraph(literal)
        else:
            cell = table.cell(0, index)
            cell.text = label
            for literal in content:
                cell.add_paragraph(literal)
    document.add_paragraph("Fixed ending")
    document.save(path)
    package = TemplatePackage(path)
    ids = {paragraph_text(node): key for key, node in package.nodes.items() if node.tag == w("p")}

    def binding(literal, target):
        """以清单中的完整原文建立字段绑定"""
        return TextBinding(node=ids[literal], quote=literal, target=target)

    def block(literal):
        """表格行重复整行，其他版式只重复所属容器中的段落"""
        node = package.node(ids[literal])
        return package.ids[next(node.iterancestors(w("tr")))] if layout == "rows" else ids[literal]

    plan = TemplatePlan(
        summary="Independent generic fixture",
        photos=[],
        remove=[],
        warnings=[],
        keep=[ids["Fixed ending"]],
        fields=[
            binding("Seed Person", "personal.name"),
            *[binding(label, "section-title:" + label) for label in labels],
        ],
        repeats=[
            RepeatBinding(
                section="projects",
                start=block(samples[0][0]),
                end=block(samples[0][-1]),
                sample_start=block(samples[0][0]),
                sample_end=block(samples[0][-1]),
                fields=[
                    binding(literal, target)
                    for literal, target in zip(
                        samples[0], ["title", "period", "description"], strict=True
                    )
                ],
            ),
            RepeatBinding(
                section=labels[1],
                start=block(samples[1][0]),
                end=block(samples[1][0]),
                sample_start=block(samples[1][0]),
                sample_end=block(samples[1][0]),
                fields=[binding(samples[1][0], "details")],
            ),
        ],
    )
    assert package.review(plan)["ready"]
    return package, plan


def generic_content():
    """构造两个不同项目和两条普通经历，包含未来新增、可隐藏的自定义信息"""
    document = ResumeDocument.model_validate(
        {
            "personal": {"name": "Current Person", "website": "https://profile.example.test"},
            "sections": [
                {"id": "portfolio", "kind": "projects", "title": "Portfolio Ω"},
                {
                    "id": "recognition",
                    "title": "Recognition · β",
                    "entries": [
                        {
                            "id": f"entry-{i}",
                            "title": f"Award {i}",
                            "period": f"Date {i}",
                            "details": f"Entry body {i}",
                            "custom_fields": [
                                {"id": "issuer", "label": "Issuer", "value": f"Institute {i}"}
                            ],
                        }
                        for i in range(2)
                    ],
                },
            ],
        }
    )
    projects = [
        {
            "project_id": f"project-{i}",
            "highlight_ids": [],
            "content": {
                "title": f"Project {i}",
                "period": f"Project date {i}",
                "role": "",
                "stack": [],
                "description": f"Project body {i}",
                "highlights": [],
                "custom_fields": [
                    {"id": "link", "label": "Repository", "value": f"https://repo{i}.example.test"}
                ],
            },
        }
        for i in range(2)
    ]
    return document, projects


def visible_text(path):
    """提取所有容器的段落文字，检查重复记录、隐藏资料和占位符"""
    return "\n".join(
        row["text"] for row in TemplatePackage(path).inventory()["nodes"] if row["kind"] == "p"
    )


def template_body_text(path):
    """读取文档正文顺序，包含表格和图形内的文字"""
    return "\n".join(Document(path).element.body.xpath(".//w:t/text()"))


def section_template(path, table=False):
    """生成三个独立栏目，使用精确标题映射和可重复的条目样本"""
    doc = Document()
    doc.add_paragraph("固定开头")
    container = doc.add_table(rows=0, cols=1) if table else doc
    for title in ("教育背景", "主修课程", "项目经历"):
        if table:
            container.add_row().cells[0].text = title
            container.add_row().cells[0].text = "旧" + title
        else:
            container.add_paragraph(title, "Heading 1")
            container.add_paragraph("旧" + title)
    doc.add_paragraph("固定结尾")
    doc.save(path)
    package = TemplatePackage(path)
    nodes = {
        paragraph_text(node): identifier
        for identifier, node in package.nodes.items()
        if node.tag == w("p") and paragraph_text(node)
    }
    fields, repeats = [], []
    for title in ("教育背景", "主修课程", "项目经历"):
        fields.append(TextBinding(node=nodes[title], quote=title, target=f"section-title:{title}"))
        node = package.node(nodes["旧" + title])
        root = next(node.iterancestors(w("tr"))) if table else node
        identifier = package.ids[root]
        repeats.append(
            RepeatBinding(
                section="projects" if title == "项目经历" else title,
                start=identifier,
                end=identifier,
                sample_start=identifier,
                sample_end=identifier,
                fields=[TextBinding(node=package.ids[node], quote="旧" + title, target="title")],
            )
        )
    return TemplatePlan(
        summary="栏目顺序测试",
        fields=fields,
        repeats=repeats,
        photos=[],
        keep=[nodes["固定开头"], nodes["固定结尾"]],
        remove=[],
        warnings=[],
    )


def layout_content():
    """生成大栏目和子栏目的资料，刻意将子栏目放在数组前面检验层级展开"""
    return ResumeDocument.model_validate(
        {
            "sections": [
                {
                    "id": "courses",
                    "title": "主修课程",
                    "kind": "text",
                    "parent_id": "education",
                    "entries": [{"id": "course", "title": "课程条目"}],
                },
                {"id": "projects", "title": "项目经历", "kind": "projects"},
                {
                    "id": "education",
                    "title": "教育背景",
                    "kind": "education",
                    "entries": [{"id": "school", "title": "教育条目"}],
                },
            ]
        }
    )


def project_document():
    """创建仅含项目栏目的匿名资料以免个人信息影响字段覆盖校验"""
    return ResumeDocument(sections=[{"id": "projects", "title": "Projects", "kind": "projects"}])


def plan_for(path, bindings, keep=(), row=False):
    """按段落原文和空位出现顺序建立映射，允许测试不同容器和节点编号"""
    package = TemplatePackage(path)
    paragraphs = [node for node in package.nodes.values() if node.tag == w("p")]
    fields = []
    used = set()
    for text, quote, target in bindings:
        node = next(p for p in paragraphs if paragraph_text(p) == text and p not in used)
        used.add(node)
        fields.append(TextBinding(node=package.ids[node], quote=quote, target=target))
    roots = (
        [node for node in package.nodes.values() if node.tag == w("tr")]
        if row
        else [p for p in paragraphs if p in used or paragraph_text(p) in keep]
    )
    return TemplatePlan(
        summary="独立项目空位测试",
        fields=[],
        repeats=[
            RepeatBinding(
                section="projects",
                start=package.ids[roots[0]],
                end=package.ids[roots[-1]],
                sample_start=package.ids[roots[0]],
                sample_end=package.ids[roots[-1]],
                fields=list(reversed(fields)),
            )
        ],
        keep=[package.ids[p] for p in paragraphs if paragraph_text(p) in keep],
        photos=[],
        remove=[],
        warnings=[],
    )


def metadata_template(path, container, indent, role=""):
    """构造正文或表格样本，末尾角色的原缩进故意和正文不同"""
    doc = Document()
    if container == "body":
        area = doc
    else:
        area = doc.add_table(rows=1, cols=1).cell(0, 0)
    area.add_paragraph("Old Project")
    area.add_paragraph("Old Period")
    stack = area.add_paragraph()
    stack.paragraph_format.left_indent = Pt(indent)
    stack.paragraph_format.space_before = Pt(3)
    stack.paragraph_format.line_spacing = Pt(15)
    label = stack.add_run("Stack: ")
    label.bold = True
    label.font.name, label.font.size = "Arial", Pt(11)
    value = stack.add_run("Old Stack")
    value.bold = False
    value.font.name, value.font.size = "Arial", Pt(11)
    area.add_paragraph("Old Description")
    area.add_paragraph("Old Highlight")
    area.add_paragraph(role).paragraph_format.left_indent = Pt(0)
    if container != "body":
        area._tc.remove(area.paragraphs[0]._p)
    doc.save(path)
    return plan_for(
        path,
        [
            ("Old Project", "Old Project", "title"),
            ("Old Period", "Old Period", "period"),
            ("Stack: Old Stack", "Old Stack", "stack"),
            ("Old Description", "Old Description", "description"),
            ("Old Highlight", "Old Highlight", "highlights"),
            (role, role, "role"),
        ],
        row=container == "row",
    )
