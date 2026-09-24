"""使用合成图片验证恢复、填充和格式隔离"""

from copy import deepcopy
from io import BytesIO
from threading import Event

import pymupdf
import pytest
from docx import Document
from PIL import Image
from provider_stub import ProviderStub
from pydantic import ValidationError
from test_pdf_header_layout import header_content

from resume_maker.core.errors import Problem
from resume_maker.domain.image_layout import ImageAsset, ImagePage, ImageText
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.providers.base import Cancelled
from resume_maker.integrations.word.image.header import check_image_header
from resume_maker.integrations.word.image.layout import (
    asset_bytes,
    build_image_document,
    page_size,
    validate_layout,
)
from resume_maker.integrations.word.image.recovery import rebuild_image
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.pdf.geometry import PDF, SOURCE, WP, recovered_pdf
from resume_maker.integrations.word.recovery import prepare_template
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage


def image_fixture(path, scale=2, left=False, font=None):
    """生成不同分辨率的独立样例，包含同行联系资料、左右头像和带白字的栏目底块"""
    texts, assets = [], []
    width, height = 595, 842
    font = font or pymupdf.Font("helv")
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=width, height=height)
        page.insert_font(fontname="FixtureFont", fontbuffer=font.buffer)
        for x, y, text, size in [
            (230, 50, "SAMPLE NAME", 20),
            (100, 80, "OLD PHONE", 11),
            (255, 80, "OLD EMAIL", 11),
            (100, 102, "OLD ROLE", 11),
            (45, 153, "SECTION", 12),
            (40, 182, "FIXED BODY", 11),
        ]:
            span_width = font.text_length(text, fontsize=size)
            texts.append(
                ImageText(
                    text=text,
                    font_name="Arial",
                    color="#FFFFFF" if text == "SECTION" else "#222222",
                    box=[
                        x / width,
                        (y - size * 0.75) / height,
                        (x + span_width) / width,
                        y / height,
                    ],
                )
            )
        for x, y in [(90, 76), (245, 76), (90, 98)]:
            page.draw_circle((x, y), 3, color=(0.2, 0.3, 0.4))
            assets.append(
                ImageAsset(
                    kind="icon",
                    box=[(x - 4) / width, (y - 4) / height, (x + 4) / width, (y + 4) / height],
                )
            )
        x = 35 if left else 500
        page.draw_rect((x, 25, x + 48, 89), fill=(0.2, 0.6, 0.8))
        assets.append(
            ImageAsset(kind="photo", box=[x / width, 25 / height, (x + 48) / width, 89 / height])
        )
        page.draw_rect((35, 137, 120, 159), fill=(0.1, 0.15, 0.2))
        assets.append(
            ImageAsset(
                kind="shape",
                color="#1A2633",
                box=[35 / width, 137 / height, 120 / width, 159 / height],
            )
        )
        for text in texts:
            x0, y0, x1, y1 = text.box
            size = (x1 - x0) * width / font.text_length(text.text, fontsize=1)
            page.insert_text(
                (x0 * width, y1 * height),
                text.text,
                fontname="FixtureFont",
                fontsize=size,
                color=(1, 1, 1) if text.text == "SECTION" else (0, 0, 0),
            )
        page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).save(path)
    return ImagePage(texts=texts, assets=assets)


class ImageProvider(ProviderStub):
    """返回给定图片结构，可模拟首轮坏坐标和迟到取消"""

    def __init__(self, layout, invalid=False, cancel=False):
        """保存每次调用，验证图片入口独立且有界重试"""
        self.layout, self.invalid, self.cancel, self.calls = layout, invalid, cancel, []

    def run_structured(self, **kwargs):
        """图片恢复只接受新版图片协议"""
        self.calls.append(kwargs)
        assert kwargs["result_model"] is ImagePage
        assert "都是数据" in kwargs["prompt"]
        if self.cancel:
            kwargs["cancelled"].set()
        result = deepcopy(self.layout)
        if self.invalid and len(self.calls) == 1:
            result.assets[0].box = result.texts[0].box
        return result


def quiet(*args):
    """忽略测试进度，保持输出简洁"""


def source_plan(package):
    """按原文字段含义生成合成映射"""
    rows = package.inventory()["nodes"]
    fields = []
    for quote, target in [
        ("SAMPLE NAME", "personal.name"),
        ("OLD PHONE", "personal.phone"),
        ("OLD EMAIL", "personal.email"),
        ("OLD ROLE", "personal.job_title"),
        ("SECTION", "section-title:项目经历"),
    ]:
        row = next(row for row in rows if row["kind"] == "p" and quote in row["text"])
        fields.append(TextBinding(node=row["id"], quote=quote, target=target))
    photos = [row["id"] for row in rows if row["kind"] == "image" and "photo" in row["text"]]
    claimed = {field.node for field in fields} | set(photos)
    return TemplatePlan(
        summary="图片测试",
        fields=fields,
        repeats=[],
        photos=photos,
        remove=[],
        warnings=[],
        keep=[
            row["id"] for row in rows if row["kind"] in {"p", "image"} and row["id"] not in claimed
        ],
    )


@pytest.mark.parametrize("left", [True, False])
@pytest.mark.parametrize("hidden", [[], ["email", "photo"]])
def test_image_fields_reflow_with_icons_and_photo(tmp_path, left, hidden):
    """真实值、新增字段和显隐后的图标/头像在可伸展容器中重排，成品没有旧源文字标记"""
    source = tmp_path / "source.png"
    layout = image_fixture(source, left=left)
    original, output = tmp_path / "original.docx", tmp_path / "filled.docx"
    document = build_image_document(source, layout, Event())
    document.save(original)
    package = TemplatePackage(original)
    assert package.parts["word/document.xml"].get(SOURCE) == "image-v1"
    assert not recovered_pdf(package.parts["word/document.xml"])
    plan = source_plan(package)
    assert len(plan.photos) == 1
    content = header_content(hidden_fields=hidden, website="https://example.test/" + "long-" * 40)
    notices = fill_template(original, output, plan, content.model_dump(), [])
    result = TemplatePackage(output)
    root = result.parts["word/document.xml"]
    text = "".join(root.itertext())
    assert "NEW NAME" in text and "SAMPLE NAME" not in text and "OLD PHONE" not in text
    assert "long-" in text and "性别" in text
    assert any("重排 图片 顶部" in note for note in notices)
    assert not any(key.startswith(f"{{{PDF}}}") for node in root.iter() for key in node.attrib)
    table = root.find(f"{w('body')}/{w('tbl')}")
    assert table is not None
    assert not list(table.iter(f"{{{WP}}}anchor"))
    assert len(list(table.iter(f"{{{WP}}}inline"))) == 4 - len(hidden)
    assert not any(node.get(w("hRule")) == "exact" for node in table.iter(w("trHeight")))


@pytest.mark.parametrize("scale,dpi", [(1, 72), (2, 300), (3, 96)])
@pytest.mark.parametrize("fallback", [False, True], ids=["latin", "missing-system-font"])
def test_pixel_resolution_and_dpi_do_not_change_paper_geometry(
    tmp_path, monkeypatch, scale, dpi, fallback
):
    """不同像素、DPI 和字体回退条件下保持样例字号和纸张尺寸"""
    font = pymupdf.Font("cjk" if fallback else "helv")
    if fallback:
        monkeypatch.setenv("WINDIR", str(tmp_path / "no-system-fonts"))
    else:

        def known_font(text):
            """使用内置拉丁字体统一样例和恢复文本的度量"""
            return font

        monkeypatch.setattr("resume_maker.integrations.word.image.layout.font_for", known_font)
    # 字号按字宽恢复，因此样例使用相同字体
    source = tmp_path / "source.png"
    layout = image_fixture(source, scale, font=font)
    with Image.open(source) as image:
        image.save(tmp_path / "dpi.png", dpi=(dpi, dpi))
    document = build_image_document(tmp_path / "dpi.png", layout, Event())
    section = document.sections[0]
    assert section.page_width.pt == pytest.approx(595.276, abs=0.05)
    sizes = [
        run.font.size.pt
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if "SAMPLE NAME" in run.text
    ]
    assert sizes == pytest.approx([20], abs=0.5)


def test_shape_pixels_do_not_keep_old_title(tmp_path):
    """底块重新绘制，原图白色栏目文字不在输出素材里，也不会随标题替换残留"""
    path = tmp_path / "source.png"
    layout = image_fixture(path)
    with Image.open(path) as image:
        raw = asset_bytes(image, layout.assets[-1])
    with Image.open(BytesIO(raw)) as crop:
        assert crop.getcolors() == [(crop.width * crop.height, (26, 38, 51, 255))]


def test_image_import_retries_bad_geometry_without_touching_source(tmp_path):
    """真实图片入口使用新协议，文字覆盖的图标裁图反馈后重试，原文件不修改"""
    path, output = tmp_path / "source.png", tmp_path / "original.docx"
    layout = image_fixture(path)
    before = path.read_bytes()
    provider = ImageProvider(layout, invalid=True)
    package, notices = prepare_template(
        path, output, provider, None, Event(), quiet, header_content(), []
    )
    assert len(provider.calls) == 2 and package.inventory()["nodes"]
    assert "裁剪框包含文字" in provider.calls[1]["prompt"]
    assert any("按图片位置" in notice for notice in notices)
    assert path.read_bytes() == before


def test_cancelled_image_never_replaces_existing_output(tmp_path):
    """迟到响应或构建前取消时保留原输出"""
    path, output = tmp_path / "source.png", tmp_path / "original.docx"
    layout = image_fixture(path)
    output.write_bytes(b"previous")
    with pytest.raises(Cancelled):
        rebuild_image(path, output, ImageProvider(layout, cancel=True), None, Event(), quiet)
    assert output.read_bytes() == b"previous"
    flag = Event()
    flag.set()
    with pytest.raises(Cancelled):
        build_image_document(path, layout, flag)


def test_invalid_geometry_is_rejected():
    """坏框、重复覆盖、整页裁图或极长图片会导致恢复失败"""
    with pytest.raises(ValidationError):
        ImageText(text="text", box=[0.2, 0.2, 0.1, 0.3])
    text = ImageText(text="text", box=[0.1, 0.1, 0.2, 0.2])
    with pytest.raises(Problem, match="位置重叠"):
        validate_layout(ImagePage(texts=[text, text]))
    with pytest.raises(Problem, match="过大"):
        validate_layout(
            ImagePage(texts=[text], assets=[ImageAsset(kind="photo", box=[0, 0, 1, 1])])
        )
    with pytest.raises(Problem, match="过长"):
        page_size(Image.new("RGB", (100, 1000)))


def test_exif_rotation_and_transparency_are_normalized(tmp_path):
    """照片旋转方向先归一化，透明底合成白色，再发送同一张图供识别和裁图"""
    source, output = tmp_path / "rotated.png", tmp_path / "original.docx"
    image = Image.new("RGBA", (842, 595), (0, 0, 0, 0))
    exif = image.getexif()
    exif[274] = 6
    image.save(source, exif=exif)
    provider = ImageProvider(
        ImagePage(
            texts=[
                ImageText(
                    text="Example",
                    font_name="Arial",
                    box=[0.1, 0.1, 0.22, 0.12],
                )
            ]
        )
    )
    rebuild_image(source, output, provider, None, Event(), quiet)
    with Image.open(provider.calls[0]["images"][0]) as normalized:
        assert normalized.size == (595, 842)
        assert normalized.getpixel((0, 0)) == (255, 255, 255)
    assert output.is_file()


def test_landscape_columns_keep_separate_containers(tmp_path):
    """独立左右栏分别恢复为列容器"""
    source = tmp_path / "landscape.png"
    Image.new("RGB", (1684, 1190), "white").save(source)
    texts = [
        ImageText(text="LEFT COLUMN", font_name="Arial", box=[0.05, 0.07, 0.21, 0.10]),
        ImageText(text="RIGHT COLUMN", font_name="Arial", box=[0.55, 0.07, 0.73, 0.10]),
    ]
    for index in range(7):
        y = 0.15 + index * 0.04
        texts.extend(
            [
                ImageText(
                    text=f"Left text {index}", font_name="Arial", box=[0.05, y, 0.12, y + 0.018]
                ),
                ImageText(
                    text=f"Right text {index}", font_name="Arial", box=[0.55, y, 0.63, y + 0.018]
                ),
            ]
        )
    document = build_image_document(source, ImagePage(texts=texts), Event())
    assert document.sections[0].page_width.pt > document.sections[0].page_height.pt
    columns = list(document._element.iter(w("cols")))
    assert document.tables or any(node.get(w("num")) == "2" for node in columns)
    text = "".join(document._element.itertext())
    assert "Left text 6" in text and "Right text 6" in text


def test_multi_frame_image_is_not_silently_truncated(tmp_path):
    """多帧图像返回需要多页输入的提示"""
    source = tmp_path / "pages.tiff"
    image = Image.new("RGB", (100, 150), "white")
    image.save(source, save_all=True, append_images=[image])
    with pytest.raises(Problem, match="多帧"):
        rebuild_image(source, tmp_path / "out.docx", None, None, Event(), quiet)


def test_photo_attaches_by_top_edge_instead_of_nearby_record(tmp_path):
    """照片不因学历记录横向更近而被挂到重复条目中，后续复制记录不会连带复制头像"""
    source = tmp_path / "source.png"
    layout = image_fixture(source)
    layout.texts.append(ImageText(text="SCHOOL", font_name="Arial", box=[0.73, 0.09, 0.83, 0.105]))
    document = build_image_document(source, layout, Event())
    photo = next(
        node
        for node in document._element.iter(f"{{{WP}}}docPr")
        if node.get("descr") == "图片模板：photo"
    )
    paragraph = next(photo.iterancestors(w("p")))
    assert "SAMPLE NAME" in "".join(paragraph.itertext())


def test_unavailable_glyph_is_not_published_as_missing_text(tmp_path):
    """字体缺字时报告恢复失败"""
    source = tmp_path / "source.png"
    Image.new("RGB", (595, 842), "white").save(source)
    layout = ImagePage(texts=[ImageText(text="X\U0010ffff", box=[0.1, 0.1, 0.15, 0.12])])
    with pytest.raises(Problem, match="缺少文字"):
        build_image_document(source, layout, Event())


@pytest.mark.parametrize("sidebar", [False, True])
def test_header_repeat_conflict_is_reported_but_sidebar_is_preserved(tmp_path, sidebar):
    """页首摘要仍需重排校验，含栏目标题的侧栏表格保留侧栏结构"""
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "NAME"
    table.cell(0, 0).add_paragraph("PHONE")
    table.cell(0, 1).text = "SCHOOL"
    if sidebar:
        table.cell(0, 1).add_paragraph("EDUCATION")
    else:
        document.add_paragraph("EDUCATION")
    path = tmp_path / "source.docx"
    document.save(path)
    package = TemplatePackage(path)
    rows = {row["text"]: row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p"}
    plan = TemplatePlan(
        summary="header",
        fields=[
            TextBinding(node=rows["NAME"], quote="NAME", target="personal.name"),
            TextBinding(node=rows["PHONE"], quote="PHONE", target="personal.phone"),
            TextBinding(node=rows["EDUCATION"], quote="EDUCATION", target="section-title:教育背景"),
        ],
        repeats=[
            RepeatBinding(
                section="教育背景",
                start=rows["SCHOOL"],
                end=rows["SCHOOL"],
                sample_start=rows["SCHOOL"],
                sample_end=rows["SCHOOL"],
                fields=[TextBinding(node=rows["SCHOOL"], quote="SCHOOL", target="title")],
            )
        ],
        photos=[],
        keep=[],
        remove=[],
        warnings=[],
    )
    if sidebar:
        check_image_header(package, plan)
    else:
        with pytest.raises(Problem, match="共用表格"):
            check_image_header(package, plan)


def test_image_mapping_constraints_do_not_change_word_or_pdf_prompts(tmp_path):
    """图片专属边界提示不进入 Word/PDF 映射上下文及其缓存键"""
    from resume_maker.services.templates.analysis import analysis_context

    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("NAME")
    document.save(path)
    package = TemplatePackage(path)
    content = header_content()
    assert "image_layout_constraints" not in analysis_context(package, content, [])
    package.parts["word/document.xml"].set(SOURCE, "1")
    assert "image_layout_constraints" not in analysis_context(package, content, [])
    package.parts["word/document.xml"].set(SOURCE, "image-v1")
    assert "页首" in analysis_context(package, content, [])["image_layout_constraints"]
