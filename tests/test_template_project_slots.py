"""使用合成模板验证项目空位的位置、样式和重复填充"""

from copy import deepcopy
from io import BytesIO
from itertools import permutations

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Pt
from lxml import etree
from test_template_mapping import photo_bytes, project_content

from resume_maker.core.errors import Problem
from resume_maker.domain.resume import ResumeDocument
from resume_maker.domain.templates import RepeatBinding, TemplatePlan, TextBinding
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text


def project_document():
    """创建仅含项目栏目的匿名资料以免个人信息影响字段覆盖校验"""
    return ResumeDocument(sections=[{"id": "projects", "title": "Projects", "kind": "projects"}])


def plan_for(path, bindings, keep=(), row=False):
    """按段落原文和空位出现顺序建立映射，允许测试不同容器和节点编号"""
    package = TemplatePackage(path)
    paragraphs = [node for node in package.nodes.values() if node.tag == w("p")]
    fields = []
    used = set()
    for text, quote, target in bindings:
        node = next(p for p in paragraphs if paragraph_text(p) == text and p not in used)
        used.add(node)
        fields.append(TextBinding(node=package.ids[node], quote=quote, target=target))
    roots = (
        [node for node in package.nodes.values() if node.tag == w("tr")]
        if row
        else [p for p in paragraphs if p in used or paragraph_text(p) in keep]
    )
    return TemplatePlan(
        summary="独立项目空位测试",
        fields=[],
        repeats=[
            RepeatBinding(
                section="projects",
                start=package.ids[roots[0]],
                end=package.ids[roots[-1]],
                sample_start=package.ids[roots[0]],
                sample_end=package.ids[roots[-1]],
                fields=list(reversed(fields)),
            )
        ],
        keep=[package.ids[p] for p in paragraphs if paragraph_text(p) in keep],
        photos=[],
        remove=[],
        warnings=[],
    )


def metadata_template(path, container, indent, role=""):
    """构造正文或表格样本，末尾角色的原缩进故意和正文不同"""
    doc = Document()
    if container == "body":
        area = doc
    else:
        area = doc.add_table(rows=1, cols=1).cell(0, 0)
    area.add_paragraph("Old Project")
    area.add_paragraph("Old Period")
    stack = area.add_paragraph()
    stack.paragraph_format.left_indent = Pt(indent)
    stack.paragraph_format.space_before = Pt(3)
    stack.paragraph_format.line_spacing = Pt(15)
    label = stack.add_run("Stack: ")
    label.bold = True
    label.font.name, label.font.size = "Arial", Pt(11)
    value = stack.add_run("Old Stack")
    value.bold = False
    value.font.name, value.font.size = "Arial", Pt(11)
    area.add_paragraph("Old Description")
    area.add_paragraph("Old Highlight")
    area.add_paragraph(role).paragraph_format.left_indent = Pt(0)
    if container != "body":
        area._tc.remove(area.paragraphs[0]._p)
    doc.save(path)
    return plan_for(
        path,
        [
            ("Old Project", "Old Project", "title"),
            ("Old Period", "Old Period", "period"),
            ("Stack: Old Stack", "Old Stack", "stack"),
            ("Old Description", "Old Description", "description"),
            ("Old Highlight", "Old Highlight", "highlights"),
            (role, role, "role"),
        ],
        row=container == "row",
    )


@pytest.mark.parametrize("container", ["body", "cell", "row"])
@pytest.mark.parametrize("indent", [7, 23])
def test_unlabelled_role_precedes_body_and_inherits_metadata_style(tmp_path, container, indent):
    """正文、单元格和整行重复都把角色放到基本信息区并保留邻段缩进和标签字重"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    plan = metadata_template(source, container, indent)
    before, plan_before = source.read_bytes(), plan.model_dump()
    projects = project_content()
    projects.extend(deepcopy(projects[0]) for _ in range(2))
    for index, project in enumerate(projects):
        project["content"]["title"] = f"Project {index}"
        project["content"]["role"] = "" if index == 1 else f"Engineer {index}"
    for _ in range(2):
        fill_template(source, output, plan, project_document().model_dump(), projects)
        package = TemplatePackage(output)
        paragraphs = [p for p in package.nodes.values() if p.tag == w("p")]
        text = "\n".join(paragraph_text(p) for p in paragraphs)
        assert text.count("担任角色：") == 2
        for index in (0, 2):
            role = f"担任角色：Engineer {index}"
            assert text.count(role) == 1
            start = text.index(f"Project {index}")
            assert start < text.index(role) < text.index("Stack: Python", start)
            node = next(p for p in paragraphs if paragraph_text(p) == role)
            assert node.find("w:pPr/w:ind", NS).get(w("left")) == str(indent * 20)
            assert node.find("w:pPr/w:spacing", NS).get(w("before")) == "60"
            assert node.find("w:pPr/w:spacing", NS).get(w("line")) == "300"
            runs = node.findall(w("r"))
            assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
            assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"
            assert all(run.find("w:rPr/w:sz", NS).get(w("val")) == "22" for run in runs)
            assert all(run.find("w:rPr/w:rFonts", NS).get(w("ascii")) == "Arial" for run in runs)
        assert "Old " not in text and "〔待填写〕" not in text
    assert source.read_bytes() == before and plan.model_dump() == plan_before


def test_explicit_trailing_role_keeps_its_original_position(tmp_path):
    """模板已有角色示例时沿用原位置"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    plan = metadata_template(source, "body", 17, role="Old Role")
    fill_template(source, output, plan, project_document().model_dump(), project_content())
    paragraphs = Document(output).paragraphs
    assert paragraphs[-1].text == "开发"
    assert paragraphs[-2].text == "自动映射：覆盖姓名和教育"
    assert paragraphs[-1].paragraph_format.left_indent == Pt(0)


@pytest.mark.parametrize("label", ["Role:", "担任角色：", "Role"])
def test_fixed_label_preserves_native_blank_slot(tmp_path, label):
    """标签和值分段时角色仍保留在对应标签之后"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    for text in ("Old Project", "Old Details", label, ""):
        doc.add_paragraph(text)
    doc.save(source)
    plan = plan_for(
        source,
        [
            ("Old Project", "Old Project", "title"),
            ("Old Details", "Old Details", "details"),
            ("", "", "role"),
        ],
        keep=[label],
    )
    projects = project_content()
    projects[0]["content"]["period"] = ""
    fill_template(source, output, plan, project_document().model_dump(), projects)
    assert [p.text for p in Document(output).paragraphs][-2:] == [label, "开发"]


@pytest.mark.parametrize("decoration", ["icon", "pBdr", "shd", "framePr"])
def test_decorated_blank_role_retains_its_native_location(tmp_path, decoration):
    """保留带图标或段落装饰的字段空位"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("Old Project")
    doc.add_paragraph("Old Details")
    if decoration == "icon":
        doc.add_picture(BytesIO(photo_bytes(60)))
    slot = doc.add_paragraph()
    if decoration != "icon":
        etree.SubElement(slot._p.get_or_add_pPr(), w(decoration))
    doc.save(source)
    plan = plan_for(
        source, [("Old Project", "Old Project", "title"), ("Old Details", "Old Details", "details")]
    )
    package = TemplatePackage(source)
    paragraphs = [node for node in package.nodes.values() if node.tag == w("p")]
    region = plan.repeats[0]
    region.end = region.sample_end = package.ids[paragraphs[-1]]
    region.fields.append(TextBinding(node=region.end, quote="", target="role"))
    plan.keep = [row["id"] for row in package.inventory()["nodes"] if row["kind"] == "image"]
    projects = project_content()
    projects[0]["content"]["period"] = ""
    fill_template(source, output, plan, project_document().model_dump(), projects)
    filled = Document(output)
    assert filled.paragraphs[-1].text == "开发"
    if decoration == "icon":
        assert filled.paragraphs[-2]._p.xpath(".//w:drawing")
    else:
        assert filled.paragraphs[-1]._p.find("w:pPr/w:" + decoration, NS) is not None


@pytest.mark.parametrize("targets", list(permutations(("period", "role", "stack"))))
def test_multiple_missing_metadata_fields_have_stable_semantic_order(tmp_path, targets):
    """只有项目标题示例时，无论空位和映射顺序如何，新增时间、角色、技术栈顺序一致"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("Old Project")
    for _ in targets:
        doc.add_paragraph()
    doc.save(source)
    plan = plan_for(
        source, [("Old Project", "Old Project", "title"), *[("", "", target) for target in targets]]
    )
    projects = project_content()
    projects[0]["content"]["description"] = ""
    projects[0]["highlight_ids"] = []
    fill_template(source, output, plan, project_document().model_dump(), projects)
    assert [p.text for p in Document(output).paragraphs] == [
        "新文档项目",
        "参与时间：2025",
        "担任角色：开发",
        "技术栈：Python",
    ]


def test_role_in_separate_table_cell_is_not_moved_across_columns(tmp_path):
    """表格右列的角色空位保留在本列"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Old Project"
    table.cell(0, 0).add_paragraph("Old Details")
    doc.save(source)
    plan = plan_for(
        source,
        [
            ("Old Project", "Old Project", "title"),
            ("Old Details", "Old Details", "details"),
            ("", "", "role"),
        ],
        row=True,
    )
    projects = project_content()
    projects[0]["content"]["period"] = ""
    fill_template(source, output, plan, project_document().model_dump(), projects)
    cells = Document(output).tables[0].rows[0].cells
    assert cells[1].text == "开发" and "开发" not in cells[0].text


def test_unlabelled_role_cannot_cross_a_section_boundary(tmp_path):
    """跨节补位缺少本节元信息时返回定位错误供识别修正"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph("Old Project")
    doc.add_paragraph("Old Details")
    doc.add_section(WD_SECTION_START.CONTINUOUS)
    doc.sections[0].start_type = WD_SECTION_START.CONTINUOUS
    doc.add_paragraph()
    doc.save(source)
    package = TemplatePackage(source)
    ids = [row["id"] for row in package.inventory()["nodes"] if row["kind"] == "p"]
    plan = plan_for(
        source, [("Old Project", "Old Project", "title"), ("Old Details", "Old Details", "details")]
    )
    region = plan.repeats[0]
    region.end = region.sample_end = ids[-1]
    region.fields.append(TextBinding(node=ids[-1], quote="", target="role"))
    projects = project_content()
    projects[0]["content"]["period"] = ""
    with pytest.raises(Problem, match="担任角色.*跨越分节"):
        fill_template(source, output, plan, project_document().model_dump(), projects)


def test_section_ending_after_metadata_remains_after_new_role(tmp_path):
    """原标题位于节末时新增角色沿用该节及连续分节设置"""
    source, output = tmp_path / "template.docx", tmp_path / "filled.docx"
    doc = Document()
    doc.add_paragraph()
    title = doc.add_paragraph("Old Project")
    section = deepcopy(doc.sections[0]._sectPr)
    etree.SubElement(section, w("type")).set(w("val"), "continuous")
    section.find(w("cols")).set(w("num"), "2")
    title._p.get_or_add_pPr().append(section)
    doc.save(source)
    plan = plan_for(source, [("Old Project", "Old Project", "title"), ("", "", "role")])
    projects = project_content()
    for key in ("period", "description"):
        projects[0]["content"][key] = ""
    projects[0]["content"]["stack"] = []
    projects[0]["highlight_ids"] = []
    fill_template(source, output, plan, project_document().model_dump(), projects)
    paragraphs = Document(output).paragraphs
    role = next(p for p in paragraphs if p.text == "担任角色：开发")
    assert role._p.find("w:pPr/w:sectPr/w:cols", NS).get(w("num")) == "2"
