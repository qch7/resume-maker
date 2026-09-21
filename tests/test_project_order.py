"""项目正文编排在内置及映射模板中的顺序、样式、显隐和固定标题回归"""

from copy import deepcopy

import pytest
from docx import Document
from pydantic import ValidationError
from test_project_info import body_text, project_info
from test_template_project_slots import metadata_template, plan_for, project_document

from resume_maker.domain.experience import displayed_experience
from resume_maker.domain.models import ProjectVisibility
from resume_maker.domain.project_layout import project_body_entries, project_body_order
from resume_maker.domain.templates import TextBinding
from resume_maker.integrations.word.full_resume import write_full_resume
from resume_maker.integrations.word.ooxml import NS, w
from resume_maker.integrations.word.templates.fill import fill_template
from resume_maker.integrations.word.templates.mapping import TemplatePackage, paragraph_text

ORDER = ["custom:link", "description", "highlights", "role", "custom:team", "stack"]


def ordered_settings(hidden=False):
    """隐藏和恢复不改变排序，亮点由外层组合独立选择"""
    return ProjectVisibility(
        order=ORDER,
        fields={"role": not hidden, "stack": True},
        custom_fields={"team": True, "link": not hidden},
    )


def order_template(path, layout):
    """构造独立正文、表格、综合字段、共享标题和多亮点槽位模板"""
    if layout in {"body", "cell", "row"}:
        return metadata_template(path, layout, 7)
    doc = Document()
    texts = ["Title", "Period", "Details"]
    if layout == "split_rows":
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Title"
        table.cell(0, 1).text = "Period"
        table.cell(1, 0).merge(table.cell(1, 1)).text = "Details"
        table.cell(2, 0).merge(table.cell(2, 1)).text = "Role"
        texts.append("Role")
    else:
        if layout == "shared_header":
            texts[0] = "Title | 担任角色：Role"
        if layout == "highlight_slots":
            texts.extend(["Point One", "Point Two", "Point Three"])
        for text in texts:
            paragraph = doc.add_paragraph(text)
            if text == "Details":
                paragraph.runs[0].bold = False
    doc.save(path)
    bindings = [
        (texts[0], "Title", "title"),
        ("Period", "Period", "period"),
        ("Details", "Details", "details"),
    ]
    if layout == "split_rows":
        bindings.append(("Role", "Role", "role"))
    if layout == "highlight_slots":
        bindings.extend((text, text, "highlights") for text in texts[3:])
    plan = plan_for(path, bindings, row=layout == "split_rows")
    if layout == "shared_header":
        node = next(field.node for field in plan.repeats[0].fields if field.target == "title")
        plan.repeats[0].fields.append(TextBinding(node=node, quote="Role", target="role"))
    return plan


def test_order_normalizes_available_fields_without_mutating_settings():
    """丢失字段忽略、新增字段追加，切回旧版本仍能使用原位置"""
    content = project_info()
    settings = ProjectVisibility(order=["custom:missing", "highlights", "custom:team"])
    original = settings.model_dump()
    assert project_body_order(content, settings) == [
        "highlights",
        "custom:team",
        "role",
        "stack",
        "description",
        "custom:link",
        "custom:blank",
    ]
    content["custom_fields"].append(
        {"id": "missing", "label": "新条目", "value": "内容", "visible": True}
    )
    assert project_body_order(content, settings)[0] == "custom:missing"
    assert settings.model_dump() == original
    visible = displayed_experience(content, ordered_settings())
    entries = project_body_entries(visible, ["two", "one"], ordered_settings())
    assert [entry["key"] for entry in entries] == [
        *ORDER[:2],
        "highlights",
        "highlights",
        *ORDER[3:],
        "custom:missing",
    ]
    assert [entry["text"] for entry in entries if entry["key"] == "highlights"] == [
        "Export Word",
        "Parse documents",
    ]


@pytest.mark.parametrize(
    "order", [["title"], ["period"], ["role", "role"], ["custom:"], ["unknown"]]
)
def test_fixed_or_invalid_order_is_rejected(order):
    """标题时间不能混入正文，重复和无效字段不能进入持久化数据"""
    with pytest.raises(ValidationError):
        ProjectVisibility(order=order)


@pytest.mark.parametrize(
    "layout",
    ["builtin", "body", "cell", "row", "details", "shared_header", "split_rows", "highlight_slots"],
)
@pytest.mark.parametrize("hidden", [False, True])
def test_project_body_order_across_layouts(tmp_path, layout, hidden):
    """自定义字段可夹在元信息和亮点之间，标题时间固定顶部，每份记录独立且无内容重复"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    value = project_info()
    projects = [
        {
            "project_id": "first",
            "content": value,
            "highlight_ids": ["two"] if hidden else ["two", "one"],
        }
    ]
    second = deepcopy(projects[0])
    second["project_id"] = "second"
    second["content"]["title"] = "另一个项目"
    projects.append(second)
    original = deepcopy(projects)
    document = project_document()
    document.project_visibility = {
        "first": ordered_settings(hidden),
        "second": ProjectVisibility(
            order=["stack", "role", "description", "custom:link", "highlights"]
        ),
    }
    if layout == "builtin":
        write_full_resume(output, document.model_dump(), projects)
    else:
        plan = order_template(source, layout)
        before, plan_before = source.read_bytes(), plan.model_dump()
        fill_template(source, output, plan, document.model_dump(), projects)
        assert source.read_bytes() == before and plan.model_dump() == plan_before
    text = body_text(output)
    first, second = text.split("另一个项目")
    expected = [value["title"], value["period"]]
    expected += [] if hidden else ["https://example.test/project"]
    expected += [value["description"], "Export Word"]
    expected += [] if hidden else ["Parse documents", "角色原值"]
    expected += ["团队：隐藏团队原值", "Python"]
    positions = [first.index(term) for term in expected]
    assert positions == sorted(positions)
    assert all(first.count(term) == 1 for term in expected)
    assert "Old " not in text and "〔待填写〕" not in text and "空白条目" not in text
    assert (
        second.index(value["description"])
        < second.index("https://example.test/project")
        < second.index("Export Word")
    )
    assert "角色原值" not in second and "Python" not in second
    if hidden:
        assert all(
            term not in first
            for term in ("https://example.test", "角色原值", "担任角色", "Parse documents")
        )
    if layout in {"body", "cell", "row"}:
        paragraphs = Document(output).element.body.iter(w("p"))
        stack = next(node for node in paragraphs if paragraph_text(node) == "技术栈：Python")
        assert stack.find("w:pPr/w:spacing", NS).get(w("line")) == "300"
        assert stack.find("w:pPr/w:spacing", NS).get(w("before")) == "60"
        runs = stack.findall(w("r"))
        assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
        assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"
    assert projects == original


def test_ordered_highlights_keep_body_weight_and_separate_label_cleanup(tmp_path):
    """亮点正文不继承粗体标题，原独立角色标签随值一起迁移并且没有重复"""
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    for text in ("Title", "Period", "Role:", "Role value"):
        doc.add_paragraph(text)
    point = doc.add_paragraph()
    point.add_run("旧亮点：").bold = True
    point.add_run("旧正文").bold = False
    doc.save(source)
    plan = plan_for(
        source,
        [
            ("Title", "Title", "title"),
            ("Period", "Period", "period"),
            ("Role value", "Role value", "role"),
            ("旧亮点：旧正文", "旧亮点：旧正文", "details"),
        ],
        keep=["Role:"],
    )
    document = project_document()
    document.project_visibility["project"] = ordered_settings()
    fill_template(
        source,
        output,
        plan,
        document.model_dump(),
        [{"project_id": "project", "content": project_info(), "highlight_ids": ["one"]}],
    )
    text = body_text(output)
    assert "Role:" not in text
    assert text.count("担任角色") == 1
    package = TemplatePackage(output)
    point = next(
        node
        for node in package.nodes.values()
        if node.tag == w("p") and paragraph_text(node) == "Parser：Parse documents"
    )
    runs = point.findall(w("r"))
    assert runs[0].find("w:rPr/w:b", NS).get(w("val"), "1") == "1"
    assert runs[1].find("w:rPr/w:b", NS).get(w("val")) == "0"
