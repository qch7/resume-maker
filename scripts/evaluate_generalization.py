"""独立构造陌生模板并跨供应商测试；不读取用户简历，不使用预先编写的模型映射。"""

import argparse
import base64
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from evaluate_templates import evaluate, save
from PIL import Image, ImageDraw

from resume_maker.domain.resume import ResumeDocument
from resume_maker.integrations.word.ooxml import w
from resume_maker.integrations.word.rendering import render_word
from resume_maker.integrations.word.templates.mapping import TemplatePackage


def portrait(color):
    """用简单头像图验证照片替换；不使用真人照片或外部素材。"""
    image = Image.new("RGB", (90, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.ellipse((26, 12, 64, 50), fill=color)
    draw.rounded_rectangle((12, 58, 78, 118), radius=20, fill=color)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def headings(chinese):
    """栏目名与原测试简历不同；验证结构处理不依赖固定中文名称。"""
    return (
        ["学习轨迹", "实践记录", "竞赛与认证", "能力清单", "补充学习"]
        if chinese
        else ["Academic History", "Selected Work", "Distinctions", "Toolbox", "Relevant Modules"]
    )


def trial_profile(chinese, with_photo=True):
    """生成与样本完全不同的资料；缺失字段、子栏目和大栏目均需通用补齐。"""
    education, projects_title, awards, skills, courses = headings(chinese)
    document = ResumeDocument(
        personal={
            "name": "周澄" if chinese else "Ada Cross",
            "phone": "+44 20 7000 4200",
            "email": "ada.cross@example.test",
            "location": "Bristol",
            "website": "https://example.test/portfolio/ada",
            "job_title": "Research Software Engineer",
            "photo": "data:image/png;base64," + base64.b64encode(portrait("#235b9b")).decode()
            if with_photo
            else "",
        },
        sections=[
            {
                "id": "academic",
                "title": education,
                "kind": "education",
                "entries": [
                    {
                        "id": "degree",
                        "title": "Northbridge Institute",
                        "subtitle": "MSc Computing",
                        "period": "2021-2024",
                    }
                ],
            },
            {
                "id": "modules",
                "parent_id": "academic",
                "title": courses,
                "entries": [
                    {"id": "course", "details": "Distributed systems and numerical methods"}
                ],
            },
            {"id": "projects", "title": projects_title, "kind": "projects"},
            {
                "id": "distinctions",
                "title": awards,
                "entries": [
                    {"id": "a", "title": "Open Science Fellowship", "period": "2025-11"},
                    {"id": "b", "title": "Regional Research Software Award", "period": "2024-03"},
                ],
            },
            {
                "id": "toolbox",
                "title": skills,
                "entries": [
                    {"id": "s1", "details": "Python, Rust and SQL"},
                    {"id": "s2", "details": "Testing and reproducible research"},
                ],
            },
        ],
    )
    projects = [
        {
            "project_id": f"project-{i}",
            "revision_id": f"revision-{i}",
            "highlight_ids": ["outcome"],
            "content": {
                "title": title,
                "period": f"202{4 + i}.01-202{4 + i}.09",
                "role": "Lead developer" if i else "Research engineer",
                "stack": ["Python", "SQLite"],
                "description": f"Built a reproducible research pipeline for dataset {i}.",
                "highlights": [
                    {
                        "id": "outcome",
                        "title": f"Result {i}",
                        "text": f"Validated {i + 3} independent datasets.",
                        "evidence": [],
                    }
                ],
                "custom_fields": [
                    {
                        "id": "repo",
                        "label": "Repository",
                        "value": f"https://example.test/project-{i}",
                    }
                ],
            },
        }
        for i, title in enumerate(["Cedar Research Archive", "Quartz Data Observatory"])
    ]
    return document, projects


def add_text(area, text, *, heading=False, fragment=False):
    """用不同运行切分同一段文字，改变节点编号但保持视觉和字段语义。"""
    paragraph = area.add_paragraph()
    pieces = [text[i : i + 3] for i in range(0, len(text), 3)] if fragment else [text]
    for piece in pieces:
        run = paragraph.add_run(piece)
        run.bold = heading
        run.font.size = Pt(12 if heading else 10)
    paragraph.paragraph_format.space_after = Pt(5)
    if heading:
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string("204E70")
    return paragraph


def sample_blocks(title, kind, sparse):
    """源样本与试填资料使用不同标记，支持检查旧内容是否残留。"""
    if kind == "education":
        return (
            [["SAMPLE-University"]]
            if sparse
            else [["SAMPLE-University", "SAMPLE-Degree", "2017-2020"]]
        )
    if kind == "skills":
        return [["SAMPLE-Analysis"], ["SAMPLE-Documentation"]]
    return [
        [f"SAMPLE-Project-{i}"]
        if sparse
        else [f"SAMPLE-Project-{i}\t2020.01-2020.06", f"SAMPLE-Description-{i}"]
        for i in range(2)
    ]


def build_source(path, layout, chinese, sparse=False):
    """分别建立段落、整行表格和左右独立容器；不借用用户模板。"""
    doc = Document()
    doc.sections[0].top_margin = Inches(0.45)
    doc.sections[0].bottom_margin = Inches(0.45)
    doc.sections[0].left_margin = Inches(0.55)
    doc.sections[0].right_margin = Inches(0.55)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for _ in range(5 if layout == "rows" else 1):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_text(doc, "Morgan Example", heading=True, fragment=True)
    if not sparse:
        doc.add_picture(BytesIO(portrait("#777777")), width=Inches(0.45))
    add_text(doc, "Email: sample@example.test", fragment=True)
    add_text(doc, "Phone: +1 555 0100")
    education, project_title, _, skills, _ = headings(chinese)
    # 样本栏目顺序故意不同于输出顺序。
    regions = [(skills, "skills"), (education, "education"), (project_title, "projects")]
    table = doc.add_table(rows=0, cols=3) if layout == "rows" else None
    sidebar = doc.add_table(rows=1, cols=2) if layout == "sidebar" else None
    if sidebar is not None:
        sidebar.autofit = False
        sidebar.columns[0].width, sidebar.columns[1].width = Inches(2.3), Inches(4.7)
        sidebar.cell(0, 0).width, sidebar.cell(0, 1).width = Inches(2.3), Inches(4.7)
    for title, kind in regions:
        blocks = sample_blocks(title, kind, sparse)
        if table is not None:
            cell = table.add_row().cells[0].merge(table.rows[-1].cells[-1])
            cell.text = title
            cell.paragraphs[0].runs[0].bold = True
            for block in blocks:
                cells = table.add_row().cells
                for index, value in enumerate(block):
                    cells[index].text = value
        else:
            area = sidebar.cell(0, 1 if kind == "projects" else 0) if sidebar is not None else doc
            add_text(area, title, heading=True)
            for block in blocks:
                for value in block:
                    paragraph = add_text(area, value, fragment=layout != "sidebar")
                    if "\t" in value:
                        paragraph.paragraph_format.tab_stops.add_tab_stop(
                            Inches(4.1 if sidebar is not None else 6.0), WD_TAB_ALIGNMENT.RIGHT
                        )
    doc.save(path)


def make_cases(directory):
    """生成六个独立版式；PDF 经过真实 Word 导出后从 PDF 重新识别。"""
    directory.mkdir(parents=True, exist_ok=True)
    specs = [
        ("english-paragraphs", "body", False, False, False),
        ("chinese-table-rows", "rows", True, False, False),
        ("english-sidebar", "sidebar", False, False, False),
        ("chinese-sparse", "body", True, True, False),
        ("chinese-vector-pdf", "body", True, False, True),
        ("english-sidebar-pdf", "sidebar", False, False, True),
    ]
    cases = []
    for name, layout, chinese, sparse, pdf in specs:
        source = directory / f"{name}.docx"
        build_source(source, layout, chinese, sparse)
        document, projects = trial_profile(chinese, with_photo=not sparse)
        if pdf:
            output = source.with_suffix(".pdf")
            _, error = render_word(source, output)
            if error:
                raise RuntimeError(error)
            source = output
        cases.append((source, document, projects))
    save(
        directory / "cases.json",
        [{"source": str(p), "document": d.model_dump(), "projects": r} for p, d, r in cases],
    )
    return cases


def verify_extra(directory, document):
    """独立检查旧样本、姓名重复及照片替换；不读取模型给自己的解释。"""
    package = TemplatePackage(directory / "resume.docx")
    text = "".join(node.text or "" for root in package.parts.values() for node in root.iter(w("t")))
    photo = document.personal.photo.partition(",")[2]
    return {
        "old_sample_remaining": "SAMPLE-" in text
        or "Morgan Example" in text
        or "sample@example.test" in text,
        "name_count": text.count(document.personal.name),
        "photo_replaced": not photo or base64.b64decode(photo) in package.files.values(),
    }


def main():
    """按同一六模板集合运行各供应商；所有失败都写报告并返回非零状态。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--cc-switch-db", type=Path)
    parser.add_argument("--providers", nargs="+", default=[])
    parser.add_argument("--workers", type=int, default=3)
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--cases", nargs="+", help="仅运行这些夹具名称；默认运行全部六种版式。")
    args = parser.parse_args()
    args.output = args.output.resolve()
    cases = make_cases(args.output / "fixtures")
    if args.cases:
        unknown = set(args.cases) - {p.stem for p, _, _ in cases}
        if unknown:
            parser.error("不存在的夹具：" + ", ".join(sorted(unknown)))
        cases = [case for case in cases if case[0].stem in args.cases]
    if args.generate_only:
        return
    if not args.providers or not args.cc_switch_db:
        parser.error("实际评测需要指定供应商及 CC Switch 数据库。")

    def run(supplier, case):
        """调用同一生产识别流程，并补充与模型映射无关的验证。"""
        source, document, projects = case
        report = evaluate(args, supplier, source, document, projects)
        directory = args.output / supplier / source.stem
        if report["passed"]:
            extra = verify_extra(directory, document)
            report["extra_check"] = extra
            report["passed"] = (
                not extra["old_sample_remaining"]
                and extra["name_count"] == 1
                and extra["photo_replaced"]
            )
            save(directory / "report.json", report)
        return report

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        jobs = [pool.submit(run, supplier, case) for case in cases for supplier in args.providers]
        reports = [job.result() for job in jobs]
    save(args.output / "report.json", reports)
    print(
        json.dumps(
            {"cases": len(reports), "passed": sum(r["passed"] for r in reports)}, ensure_ascii=False
        )
    )
    if not all(r["passed"] for r in reports):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
