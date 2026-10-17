"""将顶部个人信息、分层栏目、固定项目版本排成完整的蓝色主题简历"""

import base64
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from resume_maker.core.errors import Problem
from resume_maker.domain.experience import displayed_experience
from resume_maker.domain.project_layout import project_body_entries
from resume_maker.domain.resume import CustomInfoField, ResumeDocument, ResumeSection
from resume_maker.integrations.word.item_layout import mark_run, reference
from resume_maker.integrations.word.page_layout import apply_page_layout

BLUE = "718DB5"


def shade(element, color):
    """向段落或单元格属性追加背景色，形成蓝色抬头、灰色栏目条"""
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    element.append(fill)


def add_text(paragraph, text, *, bold=False, size=10.5, color="343637", ref=None):
    """以统一的中英文字体写入文本并保留用户输入的换行"""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑" if bold else "宋体")
    mark_run(run._element, ref)
    return run


def heading(document, title, child=False, ref=None):
    """大栏目使用色带，小栏目使用蓝色标题并、首条内容保持同页"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(9 if not child else 5)
    paragraph.paragraph_format.space_after = Pt(5)
    if not child:
        shade(paragraph._p.get_or_add_pPr(), "ECECEC")
        run = add_text(paragraph, f" {title}  ", bold=True, size=13, color="FFFFFF", ref=ref)
        shade(run._element.get_or_add_rPr(), BLUE)
    else:
        add_text(paragraph, title, bold=True, size=10.5, color=BLUE, ref=ref)


def labeled(document, label, text, ref=None):
    """写入带加粗标签的正文，空内容自动省略"""
    if not text:
        return
    paragraph = document.add_paragraph()
    if label:
        add_text(paragraph, label + "：", bold=True, ref=ref)
    add_text(paragraph, text, ref=ref)


def visible_custom_fields(fields: list[CustomInfoField]):
    """过滤隐藏或未填写完整的信息且只在排版副本中去掉首尾空白"""
    return [
        field.model_copy(update={"label": field.label.strip(), "value": field.value.strip()})
        for field in fields
        if field.visible and field.label.strip() and field.value.strip()
    ]


def write_header(document, personal):
    """用蓝色横幅展示照片、姓名、意向、年龄性别，在下方排列联系方式"""
    personal = personal.model_copy(update={field: "" for field in personal.hidden_fields})
    table = document.add_table(rows=1, cols=3 if personal.photo else 2)
    table.autofit = False
    widths = [2.5, 8.1, 6.8] if personal.photo else [10.6, 6.8]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
    for cell, width in zip(table.rows[0].cells, widths, strict=True):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell._tc.get_or_add_tcPr(), BLUE)
    cells = table.rows[0].cells
    if personal.photo:
        try:
            data = BytesIO(base64.b64decode(personal.photo.partition(",")[2]))
            cells[0].paragraphs[0].add_run().add_picture(data, width=Cm(2.2), height=Cm(2.93))
        except Exception as exc:
            raise Problem("照片无法读取，请重新上传 PNG 或 JPEG 图片。") from exc
    main = cells[-2].paragraphs[0]
    main.paragraph_format.space_before = Pt(14)
    if "name" not in personal.hidden_fields:
        add_text(
            main,
            personal.name or "个人简历",
            bold=True,
            size=23,
            color="FFFFFF",
            ref=reference(("personal", "name"), "name"),
        )
    age = personal.age + ("岁" if personal.age.isdigit() else "")
    demographics = [
        (field, value) for field, value in (("gender", personal.gender), ("age", age)) if value
    ]
    for index, (field, value) in enumerate(demographics):
        add_text(
            main,
            (" " if index else "  ") + value,
            size=11,
            color="FFFFFF",
            ref=reference(("personal", field), "personal"),
        )
    if personal.job_title:
        add_text(
            cells[-2].add_paragraph(),
            personal.job_title,
            size=13,
            color="FFFFFF",
            ref=reference(("personal", "job_title"), "job_title"),
        )
    title = cells[-1].paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(title, "个人简历", bold=True, size=24, color="FFFFFF")
    subtitle = cells[-1].add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(subtitle, "PERSONAL RESUME", size=12, color="FFFFFF")
    fields = [
        ("专业成绩", personal.gpa, reference(("personal", "gpa"), "personal")),
        ("电话", personal.phone, reference(("personal", "phone"), "personal")),
        ("邮箱", personal.email, reference(("personal", "email"), "personal")),
        ("所在地", personal.location, reference(("personal", "location"), "personal")),
        ("个人主页", personal.website, reference(("personal", "website"), "personal")),
        *[
            (field.label, field.value, reference(("personal", "custom", field.id), "custom"))
            for field in visible_custom_fields(personal.custom_fields)
        ],
    ]
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    visible = [(label, value, ref) for label, value, ref in fields if value]
    for start in range(0, len(visible), 2):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(8.7))
        for index, (label, value, ref) in enumerate(visible[start : start + 2]):
            add_text(paragraph, ("\t" if index else "") + label + "：" + value, ref=ref)


def displayed_entries(section: ResumeSection):
    """仅排版可见资料副本，移除名称等于文本栏目名的条目标题"""
    entries = []
    for original in section.entries:
        if not original.visible:
            continue
        entry = original.model_copy(update={field: "" for field in original.hidden_fields})
        entry.custom_fields = visible_custom_fields(entry.custom_fields)
        if section.kind == "text" and entry.title.strip() == section.title.strip():
            entry.title = ""
        if (
            any(
                value.strip()
                for value in [entry.title, entry.subtitle, entry.period, entry.details]
            )
            or entry.custom_fields
        ):
            entries.append(entry)
    return entries


def write_full_resume(output: Path, content: dict, projects: list[dict]):
    """按保存顺序生成 A4 Word 简历，隐藏父栏目时同时隐藏子栏目"""
    resume = ResumeDocument.model_validate(content)
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin, page.bottom_margin = Cm(1.2), Cm(1.2)
    page.left_margin, page.right_margin = Cm(1.8), Cm(1.8)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after = Pt(3)
    normal.line_spacing = 1.2
    normal.widow_control = True
    write_header(document, resume.personal)
    for section in resume.sections:
        if section.parent_id or not section.visible:
            continue
        group = [
            section,
            *[
                child
                for child in resume.sections
                if child.parent_id == section.id and child.visible
            ],
        ]
        # 空白栏目不会在最终文档中留下孤立标题
        if not (
            (section.kind == "projects" and projects)
            or any(displayed_entries(member) for member in group)
        ):
            continue
        heading(
            document,
            section.title,
            ref=reference(("section", section.id, "title"), "section_title"),
        )
        if section.kind == "projects":
            for item in projects:
                project_ref = ("project", item.get("project_id", ""))
                settings = resume.project_visibility.get(item.get("project_id", ""))
                value = displayed_experience(item["content"], settings, section.field_definitions)
                if value["title"] or value["period"]:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.paragraph_format.tab_stops.add_tab_stop(
                        Cm(17.3), WD_TAB_ALIGNMENT.RIGHT
                    )
                    add_text(
                        paragraph,
                        value["title"],
                        bold=True,
                        size=11,
                        ref=reference((*project_ref, "title")),
                    )
                    if value["period"]:
                        add_text(
                            paragraph,
                            "\t" + value["period"],
                            bold=True,
                            size=11,
                            ref=reference((*project_ref, "period")),
                        )
                for entry in project_body_entries(value, item["highlight_ids"], settings):
                    key = entry["key"]
                    parts = (
                        ("highlight", entry["id"])
                        if key == "highlights"
                        else ("custom", key[7:])
                        if key.startswith("custom:")
                        else (key,)
                    )
                    labeled(
                        document,
                        entry["label"],
                        entry["text"],
                        reference(
                            (*project_ref, *parts),
                            "highlight"
                            if key == "highlights"
                            else "custom"
                            if key.startswith("custom:")
                            else None,
                        ),
                    )
        for member in group:
            entries = displayed_entries(member)
            if member.parent_id and entries:
                heading(
                    document,
                    member.title,
                    child=True,
                    ref=reference(("section", member.id, "title"), "section_title"),
                )
            for entry in entries:
                entry_ref = ("entry", member.id, entry.id)
                if entry.title or entry.subtitle or entry.period:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.keep_with_next = bool(
                        entry.details or entry.custom_fields
                    )
                    if member.kind == "education":
                        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5.7))
                        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(11.7))
                        for index, field in enumerate(("period", "title", "subtitle")):
                            add_text(
                                paragraph,
                                ("\t" if index else "") + getattr(entry, field),
                                bold=True,
                                ref=reference((*entry_ref, field)),
                            )
                    else:
                        for index, field in enumerate(
                            field
                            for field in ("title", "subtitle", "period")
                            if getattr(entry, field)
                        ):
                            add_text(
                                paragraph,
                                ("   " if index else "") + getattr(entry, field),
                                bold=True,
                                ref=reference((*entry_ref, field)),
                            )
                labeled(document, "", entry.details, reference((*entry_ref, "details")))
                for field in entry.custom_fields:
                    labeled(
                        document,
                        field.label,
                        field.value,
                        reference((*entry_ref, "custom", field.id), "custom"),
                    )
    document.save(output)
    apply_page_layout(output, resume.extensions)
