"""仅用于 PDF 输入的逐页版面恢复，Word 和图片来源拥有独立的识别入口"""

from threading import RLock

import pymupdf
from docx import Document
from docx.shared import Pt

from resume_maker.integrations.providers.base import Cancelled
from resume_maker.integrations.word.pdf.assets import extract_assets, place_assets, separate_bullets
from resume_maker.integrations.word.pdf.flow import append_document, convert_flow
from resume_maker.integrations.word.pdf.symbols import font_symbols

# pdf2docx 的底层坐标矩阵是进程级状态，同进程多个工作台实例必须串行调用
LAYOUT_LOCK = RLock()


def native_text(page):
    """只接受可见且可靠的文字层，大面积扫描图片或损坏编码交给视觉识别"""
    text = page.get_text().strip()
    for _, symbol in font_symbols(page):
        text = text.replace(symbol, "", 1)
    meaningful = [char for char in text if not char.isspace()]
    if not meaningful:
        return False
    bad = sum(char == "\ufffd" or 0xE000 <= ord(char) <= 0xF8FF for char in meaningful)
    if bad > max(1, len(meaningful) * 0.02):
        return False
    if any(
        (pymupdf.Rect(item["bbox"]) & page.rect).get_area() > page.rect.get_area() * 0.55
        for item in page.get_image_info()
    ):
        return False
    traces = page.get_texttrace()
    return any(item["type"] != 3 and item.get("opacity", 1) > 0 for item in traces)


def normalized_page(document, number):
    """在独立单页副本里统一旋转和裁切坐标，保留源文件的实际可见页面"""
    result = pymupdf.open()
    source = document[number]
    page = result.new_page(width=source.rect.width, height=source.rect.height)
    if source.get_contents():
        with pymupdf.open() as copy:
            copy.insert_pdf(document, from_page=number, to_page=number)
            copy[0].set_rotation(0)
            # show_pdf_page 的角度方向和 PDF 页 rotation 相反，先复位避免裁切框被旋转两次
            page.show_pdf_page(page.rect, copy, 0, rotate=-source.rotation)
    for link in source.get_links():
        if link["kind"] == pymupdf.LINK_URI:
            page.insert_link({"kind": pymupdf.LINK_URI, "from": link["from"], "uri": link["uri"]})
    return result


def rebuild_pdf(pdf, output, flag, emit, fallback):
    """逐页恢复结构并在单页失败时改用视觉识别，全部成功后发布模板"""
    result = Document()
    notices = []
    native_count = 0
    with pymupdf.open(pdf) as pages:
        for number in range(len(pages)):
            if flag.is_set():
                raise Cancelled("模板自动整理已取消。")
            emit(
                "activity",
                {"type": "prepare", "text": f"正在恢复 PDF 版面 · 第 {number + 1}/{len(pages)} 页"},
            )
            with normalized_page(pages, number) as normalized:
                page = normalized[0]
                restored = None
                if native_text(page):
                    while not LAYOUT_LOCK.acquire(timeout=0.1):
                        if flag.is_set():
                            raise Cancelled("模板自动整理已取消。")
                    try:
                        assets, bullets = separate_bullets(page, extract_assets(page))
                        symbols = font_symbols(page)
                        restored, paragraphs = convert_flow(page, bullets, symbols)
                        for box, _ in symbols:
                            raw = page.get_pixmap(clip=box, matrix=pymupdf.Matrix(6, 6), alpha=True)
                            assets.append((box, False, raw.tobytes("png")))
                        if flag.is_set():
                            raise Cancelled("模板自动整理已取消。")
                        place_assets(assets, paragraphs)
                        native_count += 1
                    except Cancelled:
                        raise
                    except Exception as exc:
                        restored = None
                        notices.append(
                            f"第 {number + 1} 页版面转换未完成，已改用视觉恢复：{str(exc)[:180]}"
                        )
                    finally:
                        LAYOUT_LOCK.release()
                if restored is None:
                    restored = Document()
                    section = restored.sections[0]
                    section.page_width, section.page_height = (
                        Pt(page.rect.width),
                        Pt(page.rect.height),
                    )
                    section.top_margin = section.bottom_margin = Pt(36)
                    section.left_margin = section.right_margin = Pt(36)
                    notices.extend(fallback(restored, page, number + 1))
                    notices.append(
                        f"第 {number + 1} 页使用视觉识别，复杂版式可能调整，请核对试填预览。"
                    )
                if len(restored._element.body) == 1:
                    restored.add_paragraph()
                if flag.is_set():
                    raise Cancelled("模板自动整理已取消。")
                append_document(result, restored)
    temporary = output.with_name("recovered-pdf.docx")
    try:
        result.save(temporary)
        if flag.is_set():
            raise Cancelled("模板自动整理已取消。")
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return [
        f"已重建 PDF 可编辑模板：{native_count} 页保留文字布局、局部图形和照片；"
        "装饰随对应段落移动。",
        *notices,
    ]
